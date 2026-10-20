#!/usr/bin/env python3
"""CIS SUSE Linux Enterprise 16 Benchmark v1.0.0: every recommendation, from the benchmark's own workbook.

    python3 tools/cis_generators/gen_suse_16.py CIS_SUSE_Linux_Enterprise_16_Benchmark_v1.0.0.xlsx [fields.json]

Needs openpyxl. Reads id, title, Assessment Status and profiles (Level 1/2 - Server/Workstation sheets)
out of the CIS workbook (CIS-licensed, never committed), checks the ids against the hand-authored logic
table below (a recommendation without logic, or logic without a recommendation, aborts) and emits
benchmarks/cis/suse_16/: one package per benchmark section (cis_suse16.<section>), the roll-up
package cis_suse16 (controls table, profile selector, section coverage) and the bridge
cis_suse16.main (compliance_report, same keys as cis_rhel9.main), plus tests.

Layout and input contract follow benchmarks/cis/rhel_9 (the modules the shared Linux mapper
ansible/tasks/cis_linux_build_opa_input.yml feeds): wherever the control is the same, the rhel_9 field
name is reused. New fields are listed per control in the optional fields.json output.

Rules of the generated code:
  * every Automated recommendation evaluates real facts; facts not collected =>
    "<id>: evidence missing — <field> not collected" (fail closed, never a pass);
  * every Manual recommendation is an attestation boolean input.attest.<name> (not attested => violation);
  * messages are our words; only the id and the short CIS title are carried from the workbook;
  * the tests give every recommendation a pass and a fail case, plus an evidence-missing case for
    every Automated one.
"""
import json
import pathlib
import re
import sys

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE.parent.parent / "benchmarks" / "cis" / "suse_16"
BENCH = "CIS SUSE Linux Enterprise 16 Benchmark v1.0.0"
EXPECTED_TOTAL = 336
EXPECTED_AUTOMATED = 315
EXPECTED_MANUAL = 21
AS_OF = "v2.4.2"
SKIP_SHEETS = {"License", "Overview", "Combined Profiles"}

# module name, title, id prefixes (SUSE 16 section numbering)
SECTIONS = [
    ("filesystem", "Section 1.1: Filesystem", ["1.1."]),
    ("package_management", "Section 1.2: Package Management (zypper)", ["1.2."]),
    ("selinux", "Section 1.3: Mandatory Access Control (SELinux)", ["1.3."]),
    ("bootloader", "Section 1.4: Bootloader", ["1.4."]),
    ("process_hardening", "Section 1.5: Additional Process Hardening", ["1.5."]),
    ("crypto_policy", "Section 1.6: System-wide Crypto Policy", ["1.6."]),
    ("banners", "Section 1.7: Command Line Warning Banners", ["1.7."]),
    ("gdm", "Section 1.8: GNOME Display Manager", ["1.8."]),
    ("services", "Sections 2.1-2.2: Server and Client Services", ["2.1.", "2.2."]),
    ("time_sync", "Section 2.3: Time Synchronization", ["2.3."]),
    ("job_schedulers", "Section 2.4: Job Schedulers", ["2.4."]),
    ("network", "Section 3: Network", ["3."]),
    ("firewall", "Section 4: Host Based Firewall (firewalld)", ["4."]),
    ("ssh", "Section 5.1: SSH Server", ["5.1."]),
    ("sudo", "Section 5.2: Privilege Escalation", ["5.2."]),
    ("pam", "Section 5.3: Pluggable Authentication Modules", ["5.3."]),
    ("accounts", "Section 5.4: User Accounts and Environment", ["5.4."]),
    ("logging", "Section 6.1: System Logging", ["6.1."]),
    ("auditd", "Section 6.2: System Auditing (auditd)", ["6.2."]),
    ("integrity", "Section 6.3: Integrity Checking (AIDE)", ["6.3."]),
    ("file_permissions", "Section 7.1: System File and Directory Access", ["7.1."]),
    ("user_group", "Section 7.2: Local User and Group Settings", ["7.2."]),
]


def section_of(rid):
    for name, _, prefixes in SECTIONS:
        if any(rid.startswith(p) for p in prefixes):
            return name
    raise SystemExit(f"no section for {rid}")


# ── input paths ────────────────────────────────────────────────────────────────
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def ref(path):
    """['a', 'b-c'] -> input.a["b-c"]"""
    out = "input"
    for p in path:
        out += f".{p}" if IDENT.match(p) else f"[{json.dumps(p)}]"
    return out


def dotted(path):
    return ".".join(path)


def nest(path, value):
    d = value
    for p in reversed(path):
        d = {p: d}
    return d


def merge(*ds):
    out = {}
    for d in ds:
        for k, v in d.items():
            if k in out and isinstance(out[k], dict) and isinstance(v, dict):
                out[k] = merge(out[k], v)
            else:
                out[k] = v
    return out


def P(s):
    """'a.b.c' -> ['a','b','c'] (keys containing dots are passed as lists by callers)"""
    return s.split(".")


# ── a control's logic ──────────────────────────────────────────────────────────
class Spec:
    """ev: rego lines (AND) proving the facts were collected (alternatives: list of lists in ev_alts);
    passes: list of alternatives, each a list of rego lines (AND); field: evidence-missing message field;
    tp/tf/tm: test inputs that pass / fail / lack evidence; reads: input paths read (for the report)."""

    def __init__(self, ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
        self.ev_alts = ev_alts if ev_alts is not None else [ev]
        self.passes = passes
        self.field = field
        self.tp, self.tf = tp, tf
        self.tm = {} if tm is None else tm
        self.reads = reads or []
        self.manual = False
        # True when an all-empty input (strings "", lists []) is a genuine pass: offender lists, where an
        # empty list collected by the collector means "none found". Otherwise a hollow-input test is generated.
        self.hollow_ok = False


class Attest:
    def __init__(self, name):
        self.name = name
        self.manual = True
        self.reads = [["attest", name]]


# ── builders ───────────────────────────────────────────────────────────────────
def kmod(listkey, name):
    lst = ref([listkey, "modules"])
    ev = [f'some e in {lst}', f'e.module == {json.dumps(name)}', 'kmod_evidence(e)']
    passes = [[f'count([e | some e in {lst}; e.module == {json.dumps(name)}]) > 0',
               f'every e in [e | some e in {lst}; e.module == {json.dumps(name)}] {{ kmod_unavailable(e) }}']]
    ok = {"module": name, "disabled": True, "loaded": False}
    return Spec(ev, passes, f"{listkey}.modules[{name}]",
                nest([listkey, "modules"], [ok]),
                nest([listkey, "modules"], [{"module": name, "disabled": False, "loaded": False}]),
                nest([listkey, "modules"], [{"module": "other", "disabled": True, "loaded": False}]),
                reads=[[listkey, "modules"]])


MOUNTS = ["ansible_facts", "ansible_mounts"]


def mnt_sep(path, extra_tmp=False):
    passes = [[f'mounted({json.dumps(path)})']]
    if extra_tmp:
        passes = [[f'mounted({json.dumps(path)})', 'not tmp_mount_unit_off']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "tmpfs", "options": "rw,nodev,nosuid,noexec"}]),
                nest(MOUNTS, [other]), {}, reads=[MOUNTS] + ([["services", "tmp.mount"]] if extra_tmp else []))


def mnt_opt(path, opt):
    # The option applies to the partition when it exists; the mount table is the collected fact
    # (the separate-partition recommendation fails on its own when the partition is missing).
    passes = [[f'not mounted({json.dumps(path)})'],
              [f'every m in mounts_at({json.dumps(path)}) {{ {json.dumps(opt)} in mount_opts(m) }}']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    good = "rw,nodev,nosuid,noexec"
    bad = ",".join(o for o in ["rw", "nodev", "nosuid", "noexec"] if o != opt)
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": good}]),
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": bad}]),
                {}, reads=[MOUNTS])


def sysctl(param, allowed, ipv6=False):
    vals = json.dumps(allowed)
    run, per = ref(["sysctl", param]), ref(["sysctl_persisted", param])
    ev = [f'{run} != null', f'{per} != null']
    passes = [[f'sprintf("%v", [{run}]) in {vals}', f'sprintf("%v", [{per}]) in {vals}']]
    good = {"sysctl": {param: allowed[0]}, "sysctl_persisted": {param: allowed[0]}}
    bad_v = "99"
    tf = {"sysctl": {param: allowed[0]}, "sysctl_persisted": {param: bad_v}}
    ev_alts = [ev]
    reads = [["sysctl", param], ["sysctl_persisted", param]]
    if ipv6:  # "IF IPv6 is enabled": a host with IPv6 disabled passes on that collected fact
        ev_alts = [ev, ['input.ipv6.disabled == true']]
        passes = passes + [['input.ipv6.disabled == true']]
        reads.append(["ipv6", "disabled"])
    return Spec(None, passes, f"sysctl/sysctl_persisted[{param}]", good, tf, {}, reads=reads, ev_alts=ev_alts)


def pkg_present(pkg, path=None):
    p = path or ["packages", pkg]
    return Spec([f'is_boolean({ref(p)})'], [[f'{ref(p)} == true']], dotted(p),
                nest(p, True), nest(p, False), {}, reads=[p])


def pkg_absent(pkgs, path=None):
    if path:  # a single rhel_9-style boolean
        return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == false']], dotted(path),
                    nest(path, False), nest(path, True), {}, reads=[path])
    ev = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    passes = [[f'{ref(["packages", p])} == false' for p in pkgs]]
    tp = {"packages": {p: False for p in pkgs}}
    tf = {"packages": {p: (i == 0) for i, p in enumerate(pkgs)}}
    return Spec(ev, passes, "packages[" + ",".join(pkgs) + "]", tp, tf, {"packages": {"audit": True}},
                reads=[["packages", p] for p in pkgs])


def not_in_use(units_by_pkg):
    """CIS 'not in use': every package not installed, OR (installed for a dependency) each unit the audit
    names for it neither enabled nor active. Ansible's service_facts lists .service units only, so a
    unit missing from input.services proves nothing: for an installed package every named unit must be
    reported (LoadState not-found is fine), or the evidence is missing."""
    pkgs = list(units_by_pkg)
    pk = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    none_installed = [f'{ref(["packages", p])} == false' for p in pkgs]
    m = json.dumps(units_by_pkg)
    reported = f'every p, us in {m} {{ unit_set_reported(p, us) }}'
    ev_alts = [pk + none_installed, pk + ['services_collected', reported]]
    passes = [none_installed,
              ['services_collected', reported, f'every p, us in {m} {{ unit_set_off(p, us) }}']]
    tp = {"packages": {p: False for p in pkgs}}
    allu = {u: {"state": "inactive", "status": "disabled"} for us in units_by_pkg.values() for u in us}
    first = units_by_pkg[pkgs[0]][0]
    tf = {"packages": {p: True for p in pkgs},
          "services": dict(allu, **{first: {"state": "running", "status": "enabled"}},
                           **{"sshd.service": {"state": "running", "status": "enabled"}})}
    tm = {"packages": {p: True for p in pkgs},
          "services": {"sshd.service": {"state": "running", "status": "enabled"}}}
    return Spec(None, passes, "packages[" + ",".join(pkgs) + "] / services[" + ",".join(
        u for us in units_by_pkg.values() for u in us) + "]", tp, tf, tm,
                reads=[["packages", p] for p in pkgs] + [["services"]], ev_alts=ev_alts)


def units_off(units):
    """The named units must each be reported (not-found is fine) and be neither enabled nor active."""
    passes = [[f'every u in {json.dumps(units)} {{ not unit_on(u) }}']]
    other = {"sshd.service": {"state": "running", "status": "enabled"}}
    nf = {u: {"state": "inactive", "status": "not-found"} for u in units}
    return Spec(['services_collected', f'every u in {json.dumps(units)} {{ unit_reported(u) }}'], passes,
                "services[" + ",".join(units) + "]",
                {"services": merge(other, nf)},
                {"services": merge(other, nf, {units[0]: {"state": "running", "status": "enabled"}})},
                {"services": other}, reads=[["services"]])


def unit_running(unit):
    passes = [[f'unit_enabled_active({json.dumps(unit)})']]
    other = {"sshd.service": {"state": "running", "status": "enabled"}}
    return Spec(['services_collected'], passes, "services",
                {"services": merge(other, {unit: {"state": "running", "status": "enabled"}})},
                {"services": merge(other, {unit: {"state": "stopped", "status": "disabled"}})},
                {}, reads=[["services"]])


def boolean(path, want):
    return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == {json.dumps(want)}']], dotted(path),
                nest(path, want), nest(path, not want), {}, reads=[path])


def access(mode, owner, group, mx, owners, groups, exists=None, optional=False):
    """One file: mode within mx, owner in owners, group in groups. exists: path of a boolean; when
    optional, a collected 'does not exist' passes; when required, it fails."""
    ev = [f'{ref(mode)} != null', f'{ref(owner)} != null', f'{ref(group)} != null']
    ok = [f'mode_within({ref(mode)}, {json.dumps(mx)})', f'{ref(owner)} in {json.dumps(owners)}',
          f'{ref(group)} in {json.dumps(groups)}']
    tp = merge(nest(mode, mx), nest(owner, owners[0]), nest(group, groups[0]))
    tf = merge(nest(mode, "0777"), nest(owner, owners[0]), nest(group, groups[0]))
    ev_alts, passes, reads = [ev], [ok], [mode, owner, group]
    if exists:
        reads.append(exists)
        ev_alts = [[f'{ref(exists)} == true'] + ev, [f'{ref(exists)} == false']]
        passes = [[f'{ref(exists)} == true'] + ok]
        tp = merge(tp, nest(exists, True))
        tf = merge(tf, nest(exists, True))
        if optional:
            passes.append([f'{ref(exists)} == false'])
        else:
            tf = nest(exists, False)
    return Spec(None, passes, "/".join(dotted(p) for p in reads), tp, tf, {}, reads=reads, ev_alts=ev_alts)


def access_list(path, mx, owners, groups, mode="mode", owner="owner", group="group", sample="/x", nonempty=True):
    """An inventory of files that always exist on a host must be non-empty to count as evidence."""
    lst = ref(path)
    passes = [[f'every f in {lst} {{ list_item_ok(f, {json.dumps(mode)}, {json.dumps(owner)}, '
               f'{json.dumps(group)}, {json.dumps(mx)}, {json.dumps(owners)}, {json.dumps(groups)}) }}']]
    good = {"path": sample, mode: mx, owner: owners[0], group: groups[0]}
    bad = dict(good, **{mode: "0777"})
    ev = [f'is_array({lst})'] + ([f'count({lst}) > 0'] if nonempty else [])
    return Spec(ev, passes, dotted(path), nest(path, [good]), nest(path, [good, bad]), {}, reads=[path])


def crit(path, mx, groups, optional=False):
    """critical_files.analysis entry {file, exists, actual_mode, actual_owner, actual_group} (rhel_9 shape);
    the expected_* / *_compliant fields the rhel_9 collector pre-computes are not trusted here."""
    p = json.dumps(path)
    ent = f'[f | some f in object.get(input, ["critical_files", "analysis"], []); f.file == {p}]'
    ok = (f'every f in {ent} {{ f.exists == true; mode_within(f.actual_mode, {json.dumps(mx)}); '
          f'f.actual_owner == "root"; f.actual_group in {json.dumps(groups)} }}')
    ev = [f'count({ent}) > 0', f'every f in {ent} {{ is_boolean(f.exists) }}']
    passes = [[f'count({ent}) > 0', ok]]
    if optional:
        passes.append([f'count({ent}) > 0', f'every f in {ent} {{ f.exists == false }}'])
    good = {"file": path, "exists": True, "actual_mode": mx, "actual_owner": "root", "actual_group": groups[0]}
    bad = dict(good, actual_mode="0666")
    other = {"file": "/etc/other", "exists": True, "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"}
    return Spec(ev, passes, f"critical_files.analysis[{path}]",
                {"critical_files": {"analysis": [other, good]}},
                {"critical_files": {"analysis": [other, bad]}},
                {"critical_files": {"analysis": [other]}}, reads=[["critical_files", "analysis"]])


def attest(name):
    return Attest(name)


def custom(ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
    return Spec(ev, passes, field, tp, tf, tm, reads=reads, ev_alts=ev_alts)


# ── sshd (effective config; `sshd -T` output preferred, raw sshd_config accepted) ──────────────
SSH_RAW = ["ssh", "sshd_config_raw"]


def ssh_in(lines):
    return {"ssh": {"sshd_config_raw": "\n".join(lines)}}


def sshd(key, cond, good, bad, explicit=False):
    """cond: rego expression over `v` (the effective lowercase value). explicit: the directive must be
    present (no reliable compiled default), otherwise its absence is missing evidence."""
    ev = ['ssh_collected', f'sshd_first({json.dumps(key)})']
    passes = [[f'v := lower(sshd_value({json.dumps(key)}))', cond]]
    tm = ssh_in(["port 22"])
    return Spec(ev, passes, f"ssh.sshd_config_raw (effective {key})",
                ssh_in(["port 22", f"{key} {good}"]), ssh_in(["port 22", f"{key} {bad}"]), tm, reads=[SSH_RAW])


def text_has(path, regex, want, good, bad):
    """want=True: some active line of the text matches regex; False: none does."""
    t = ref(path)
    cond = f'regex_in_text({t}, {json.dumps(regex)})'
    return Spec([f'is_string({t})'], [[cond if want else f'not {cond}']], dotted(path),
                nest(path, good), nest(path, bad), {}, reads=[path])


def conf(path, key_re, cond, good, bad, default=None):
    """key = value configuration text; last occurrence wins. cond is an expression over `v`.
    default: the documented compiled-in value used when the key is not set (None: unset fails)."""
    t = ref(path)
    passes = [[f'v := conf_last({t}, {json.dumps(key_re)})', cond]]
    if default is not None:
        passes.append([f'not conf_last({t}, {json.dumps(key_re)})', f'v := {json.dumps(default)}', cond])
    return Spec([f'is_string({t})'], passes, dotted(path), nest(path, good), nest(path, bad), {}, reads=[path])


# ── Rego helpers copied into every module (no file may call a function defined in another) ────
HELPERS = r'''
# Parse an octal permission string ("644", "0644", "0o644") to an integer; undefined if not octal.
octal(s) := n if {
	t := trim_prefix(lower(trim_space(sprintf("%v", [s]))), "0o")
	regex.match(`^[0-7]{1,4}$`, t)
	p := concat("", [substring("0000", 0, 4 - count(t)), t])
	ds := [to_number(substring(p, i, 1)) | some i in numbers.range(0, 3)]
	n := ((ds[0] * 512) + (ds[1] * 64)) + ((ds[2] * 8) + ds[3])
}

# true when mode grants nothing beyond max (both octal strings)
mode_within(mode, max) if bits.and(octal(mode), bits.and(bits.negate(octal(max)), 4095)) == 0

# Active (non-blank, non-comment) lines of a configuration text.
active_lines(text) := [l |
	some raw in split(text, "\n")
	l := trim_space(raw)
	l != ""
	not startswith(l, "#")
]

# a collected text: a non-blank string ("" is not evidence)
text_collected(t) if {
	is_string(t)
	trim_space(t) != ""
}

regex_in_text(text, re) if {
	some l in active_lines(text)
	regex.match(re, l)
}

# Values of a `key = value` / `key value` directive (key is a case-insensitive regex); last one wins.
conf_values(text, key_re) := [trim(m[1], "\"' ") |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(concat("", ["(?i)^", key_re, `(?:\s*=\s*|\s+)(.*?)\s*$`]), l, 1)
]

conf_last(text, key_re) := vs[count(vs) - 1] if {
	vs := conf_values(text, key_re)
	count(vs) > 0
}

num(v) := to_number(trim_space(sprintf("%v", [v]))) if regex.match(`^\s*-?[0-9]+\s*$`, sprintf("%v", [v]))

# ── packages / services ──
services_collected if {
	is_object(input.services)
	count(input.services) > 0
}

unit_on(u) if input.services[u].status == "enabled"

unit_on(u) if lower(input.services[u].state) in {"running", "active"}

# every unit named for an installed package is reported in the inventory
unit_set_reported(p, _) if input.packages[p] == false

unit_set_reported(p, us) if {
	input.packages[p] == true
	every u in us { unit_reported(u) }
}

unit_reported(u) if {
	is_string(input.services[u].status)
	input.services[u].status != ""
}

unit_set_off(p, _) if input.packages[p] == false

unit_set_off(p, us) if {
	input.packages[p] == true
	every u in us { not unit_on(u) }
}

unit_enabled_active(u) if {
	input.services[u].status == "enabled"
	lower(input.services[u].state) in {"running", "active"}
}

# tmp.mount masked or disabled means systemd will not mount /tmp at boot
tmp_mount_unit_off if input.services["tmp.mount"].status in {"masked", "disabled"}

# ── kernel modules: {module, disabled, loaded, available?}; "unknown" proves nothing ──
kmod_evidence(e) if is_boolean(e.available)

kmod_evidence(e) if {
	is_boolean(e.disabled)
	is_boolean(e.loaded)
}

kmod_unavailable(e) if e.available == false

kmod_unavailable(e) if {
	e.disabled == true
	e.loaded == false
}

# ── mounts ──
mounts_collected if {
	is_array(input.ansible_facts.ansible_mounts)
	count(input.ansible_facts.ansible_mounts) > 0
}

mounts_at(path) := [m | some m in input.ansible_facts.ansible_mounts; m.mount == path]

mounted(path) if count(mounts_at(path)) > 0

mount_opts(m) := split(m.options, ",") if is_string(m.options)

mount_opts(m) := m.options if is_array(m.options)

# ── file lists ──
list_item_ok(f, mk, ok, gk, max, owners, groups) if {
	mode_within(f[mk], max)
	f[ok] in owners
	f[gk] in groups
}

# ── sshd: first occurrence wins (sshd semantics); `sshd -T` prints effective values ──
ssh_collected if {
	is_string(input.ssh.sshd_config_raw)
	trim_space(input.ssh.sshd_config_raw) != ""
}

sshd_pairs := [[lower(m[1]), trim_space(m[2])] |
	some l in active_lines(input.ssh.sshd_config_raw)
	some m in regex.find_all_string_submatch_n(`^([A-Za-z0-9]+)(?:\s*=\s*|\s+)(.*)$`, l, 1)
]

sshd_first(key) := vs[0] if {
	vs := [p[1] | some p in sshd_pairs; p[0] == key]
	count(vs) > 0
}

# effective value: `sshd -T` prints every directive, so a missing directive means the effective
# configuration was not collected (a raw sshd_config misses /usr/etc/ssh and sshd_config.d)
sshd_value(key) := sshd_first(key)

sshd_list(key) := [trim_space(x) | some x in split(lower(sshd_value(key)), ",")]

# seconds from an sshd time value ("60", "60s", "1m")
sshd_seconds(v) := to_number(v) if regex.match(`^[0-9]+$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "s")) if regex.match(`^[0-9]+s$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "m")) * 60 if regex.match(`^[0-9]+m$`, v)

# ── /etc/passwd, /etc/group (rhel_9 fields user_group.passwd_content / group_content) ──
rows(text) := [f |
	some l in active_lines(text)
	f := split(l, ":")
]

passwd_rows := [f | some f in rows(input.user_group.passwd_content); count(f) >= 7]

group_rows := [f | some f in rows(input.user_group.group_content); count(f) >= 3]

valid_shells := {s |
	some s in active_lines(input.user_group.shells_content)
	not regex.match(`/nologin$`, s)
}
'''


# ════════════════════════════════════════════════════════════════════════════════════════════
# LOGIC TABLE — one entry per recommendation, written from each audit procedure.
# ════════════════════════════════════════════════════════════════════════════════════════════
L = {}
FSM = "disabled_filesystem_modules"   # mapper field (rhel_9 shape)
NETM = "disabled_network_modules"     # same shape, network modules (new)

# 1.1.1 filesystem kernel modules
for rid, mod in [("1.1.1.1", "cramfs"), ("1.1.1.2", "firewire-core"), ("1.1.1.3", "freevxfs"), ("1.1.1.4", "hfs"),
                 ("1.1.1.5", "hfsplus"), ("1.1.1.6", "jffs2"), ("1.1.1.7", "overlay"), ("1.1.1.8", "squashfs"),
                 ("1.1.1.9", "udf")]:
    L[rid] = kmod(FSM, mod)
L["1.1.1.10"] = attest("unused_filesystem_modules_reviewed")

# 1.1.2 partitions
L["1.1.2.1.1"] = mnt_sep("/tmp", extra_tmp=True)
L["1.1.2.2.1"] = mnt_sep("/dev/shm")
for base, part in [("1.1.2.3", "/home"), ("1.1.2.4", "/var"), ("1.1.2.5", "/var/tmp"), ("1.1.2.6", "/var/log"),
                   ("1.1.2.7", "/var/log/audit")]:
    L[base + ".1"] = mnt_sep(part)
OPTS = {"1.1.2.1": ("/tmp", ["nodev", "nosuid", "noexec"]), "1.1.2.2": ("/dev/shm", ["nodev", "nosuid", "noexec"]),
        "1.1.2.3": ("/home", ["nodev", "nosuid"]), "1.1.2.4": ("/var", ["nodev", "nosuid"]),
        "1.1.2.5": ("/var/tmp", ["nodev", "nosuid", "noexec"]), "1.1.2.6": ("/var/log", ["nodev", "nosuid", "noexec"]),
        "1.1.2.7": ("/var/log/audit", ["nodev", "nosuid", "noexec"])}
for base, (part, opts) in OPTS.items():
    for i, o in enumerate(opts, start=2):
        L[f"{base}.{i}"] = mnt_opt(part, o)

# 1.2 package management (zypper)
L["1.2.1.1"] = attest("gpg_keys_reviewed")
L["1.2.1.2"] = custom(
    ['is_boolean(input.initial_setup.global_gpgcheck_enabled)', 'is_array(input.initial_setup.repos_without_gpgcheck)'],
    [['input.initial_setup.global_gpgcheck_enabled == true', 'count(input.initial_setup.repos_without_gpgcheck) == 0']],
    "initial_setup.global_gpgcheck_enabled/repos_without_gpgcheck",
    {"initial_setup": {"global_gpgcheck_enabled": True, "repos_without_gpgcheck": []}},
    {"initial_setup": {"global_gpgcheck_enabled": True, "repos_without_gpgcheck": ["repo-oss"]}},
    reads=[["initial_setup", "global_gpgcheck_enabled"], ["initial_setup", "repos_without_gpgcheck"]])
L["1.2.1.3"] = attest("repo_gpgcheck_reviewed")
L["1.2.1.4"] = attest("repositories_reviewed")
L["1.2.1.5"] = boolean(["initial_setup", "zypper_solver_only_requires"], True)
L["1.2.2.1"] = attest("updates_installed")

# 1.3 SELinux (SUSE 16 ships SELinux, not AppArmor)
L["1.3.1.1"] = boolean(["selinux", "installed"], True)
L["1.3.1.2"] = custom(
    ['is_string(input.selinux.grub_cmdline)'],
    [['not regex.match(`(^|[\\s"])(selinux|enforcing)=0\\b`, input.selinux.grub_cmdline)']],
    "selinux.grub_cmdline",
    {"selinux": {"grub_cmdline": 'GRUB_CMDLINE_LINUX="quiet security=selinux selinux=1"'}},
    {"selinux": {"grub_cmdline": 'GRUB_CMDLINE_LINUX="quiet selinux=0"'}}, reads=[["selinux", "grub_cmdline"]])
L["1.3.1.3"] = custom(['is_string(input.selinux.policy)'], [['lower(input.selinux.policy) in {"targeted", "mls"}']],
                      "selinux.policy", {"selinux": {"policy": "targeted"}}, {"selinux": {"policy": "minimum"}},
                      reads=[["selinux", "policy"]])
L["1.3.1.4"] = pkg_absent(None, ["selinux", "mcstrans_installed"])
L["1.3.1.5"] = custom(['is_string(input.selinux.mode)'], [['lower(input.selinux.mode) in {"enforcing", "permissive"}']],
                      "selinux.mode", {"selinux": {"mode": "permissive"}}, {"selinux": {"mode": "disabled"}},
                      reads=[["selinux", "mode"]])
L["1.3.1.6"] = custom(['is_string(input.selinux.mode)'], [['lower(input.selinux.mode) == "enforcing"']],
                      "selinux.mode", {"selinux": {"mode": "enforcing"}}, {"selinux": {"mode": "permissive"}},
                      reads=[["selinux", "mode"]])
L["1.3.1.7"] = pkg_absent(None, ["selinux", "setroubleshoot_installed"])
L["1.3.1.8"] = attest("no_unconfined_services")

# 1.4 bootloader (/boot/grub2/grub.cfg — not the mapper's /etc/default/grub text)
L["1.4.1"] = custom(
    ['is_boolean(input.boot_security.grub_password_set)', 'is_boolean(input.boot_security.grub_password_hashed)'],
    [['input.boot_security.grub_password_set == true', 'input.boot_security.grub_password_hashed == true']],
    "boot_security.grub_password_set/grub_password_hashed",
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": True}},
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": False}},
    reads=[["boot_security", "grub_password_set"], ["boot_security", "grub_password_hashed"]])
L["1.4.2"] = access(["boot_security", "grub_cfg_mode"], ["boot_security", "grub_cfg_owner"],
                    ["boot_security", "grub_cfg_group"], "0600", ["root"], ["root"])

# 1.5 process hardening
for rid, p, vals in [("1.5.1", "fs.protected_hardlinks", ["1"]), ("1.5.2", "fs.protected_symlinks", ["1"]),
                     ("1.5.3", "fs.suid_dumpable", ["0"]), ("1.5.4", "kernel.dmesg_restrict", ["1"]),
                     ("1.5.5", "kernel.kptr_restrict", ["1", "2"]), ("1.5.6", "kernel.randomize_va_space", ["2"]),
                     ("1.5.7", "kernel.yama.ptrace_scope", ["1", "2", "3"])]:
    L[rid] = sysctl(p, vals)
L["1.5.8"] = custom(['input.boot_security.systemd_coredump_process_size_max != null'],
                    [['sprintf("%v", [input.boot_security.systemd_coredump_process_size_max]) == "0"']],
                    "boot_security.systemd_coredump_process_size_max",
                    {"boot_security": {"systemd_coredump_process_size_max": "0"}},
                    {"boot_security": {"systemd_coredump_process_size_max": "2G"}},
                    reads=[["boot_security", "systemd_coredump_process_size_max"]])
L["1.5.9"] = custom(['is_string(input.boot_security.systemd_coredump_storage)'],
                    [['lower(input.boot_security.systemd_coredump_storage) == "none"']],
                    "boot_security.systemd_coredump_storage",
                    {"boot_security": {"systemd_coredump_storage": "none"}},
                    {"boot_security": {"systemd_coredump_storage": "external"}},
                    reads=[["boot_security", "systemd_coredump_storage"]])

# 1.6 crypto policy
POL = ["crypto_policy", "current_pol_raw"]
L["1.6.1"] = pkg_present("crypto-policies-scripts")
L["1.6.2"] = custom(['is_string(input.initial_setup.crypto_policy)'],
                    [['not regex.match(`(?i)^\\s*LEGACY\\b`, input.initial_setup.crypto_policy)']],
                    "initial_setup.crypto_policy", {"initial_setup": {"crypto_policy": "DEFAULT"}},
                    {"initial_setup": {"crypto_policy": "LEGACY"}}, reads=[["initial_setup", "crypto_policy"]])
L["1.6.3"] = text_has(["crypto_policy", "sysconfig_ssh_raw"], "(?i)^CRYPTO_POLICY\\s*=", False,
                      "# CRYPTO_POLICY=\nSSHD_OPTS=", "CRYPTO_POLICY=")
L["1.6.3"].ev_alts.append(['input.crypto_policy.sysconfig_ssh_exists == false'])
L["1.6.3"].passes.append(['input.crypto_policy.sysconfig_ssh_exists == false'])
L["1.6.3"].reads.append(["crypto_policy", "sysconfig_ssh_exists"])
L["1.6.3"].field = "crypto_policy.sysconfig_ssh_raw/sysconfig_ssh_exists"
L["1.6.4"] = custom(
    [f'is_string({ref(POL)})'],
    [[f'not sha1_hash_or_sign_enabled', f'conf_last({ref(POL)}, "sha1_in_certs") == "0"']],
    dotted(POL),
    nest(POL, "hash = SHA2-256 SHA2-384\nsign = RSA-PSS-SHA2-256\nsha1_in_certs = 0"),
    nest(POL, "hash = SHA2-256 SHA1\nsign = RSA-PSS-SHA2-256\nsha1_in_certs = 0"), reads=[POL])
L["1.6.5"] = text_has(POL, "(?i)(^mac\\s*=\\s*([^#]+)?-128|hmac-sha1)", False,
                      "mac = HMAC-SHA2-256 HMAC-SHA2-512", "mac = HMAC-SHA2-256 UMAC-128")
L["1.6.6"] = custom([f'is_string({ref(POL)})'], [['not ssh_cipher_has("-cbc")']], dotted(POL),
                    nest(POL, "cipher = AES-256-GCM AES-256-CBC\ncipher@openssh = AES-256-GCM AES-128-GCM"),
                    nest(POL, "cipher = AES-256-GCM AES-256-CBC"), reads=[POL])
L["1.6.7"] = custom([f'is_string({ref(POL)})'], [['not ssh_cipher_has("chacha20-poly1305")']], dotted(POL),
                    nest(POL, "cipher = AES-256-GCM CHACHA20-POLY1305\ncipher@openssh = AES-256-GCM"),
                    nest(POL, "cipher = AES-256-GCM CHACHA20-POLY1305"), reads=[POL])

# 1.7 banners — no OS/version escapes (\m \r \s \v) or the os-release ID ("sles")
SYSINFO = "(?i)(\\\\[mrsv]|\\bsles\\b)"
for rid, f in [("1.7.1", "motd"), ("1.7.2", "issue"), ("1.7.3", "issue_net")]:
    ex, ct = ["initial_setup", f"{f}_exists"], ["initial_setup", f"{f}_content"]
    alts = [[f'{ref(ex)} == false'], [f'{ref(ex)} == true', f'is_string({ref(ct)})']]
    ok = [f'{ref(ex)} == true', f'not regex.match({json.dumps(SYSINFO)}, {ref(ct)})']
    passes = [ok] + ([[f'{ref(ex)} == false']] if f == "motd" else [])
    if f != "motd":  # /etc/issue and /etc/issue.net: the banner is expected to exist
        pass
    L[rid] = custom(None, passes, f"initial_setup.{f}_exists/{f}_content",
                    merge(nest(ex, True), nest(ct, "Authorized uses only. All activity may be monitored.")),
                    merge(nest(ex, True), nest(ct, "Welcome to SLES \\r on \\m")),
                    {}, reads=[ex, ct], ev_alts=alts)
L["1.7.4"] = custom(['is_string(input.pam.pam_d_raw)'],
                    [['every l in active_lines(input.pam.pam_d_raw) { pam_motd_ok(l) }']], "pam.pam_d_raw",
                    {"pam": {"pam_d_raw": "session optional pam_motd.so motd=/etc/motd\nauth include common-auth"}},
                    {"pam": {"pam_d_raw": "session optional pam_motd.so"}}, reads=[["pam", "pam_d_raw"]])
L["1.7.5"] = custom(
    None,
    [['v := sshd_value("banner")', 'lower(v) != "none"', 'is_string(input.ssh.banner_content)',
      f'not regex.match({json.dumps(SYSINFO)}, input.ssh.banner_content)']],
    "ssh.sshd_config_raw/ssh.banner_content",
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_content": "Authorized uses only."}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_content": "SLES 16 \\r"}},
    {}, reads=[SSH_RAW, ["ssh", "banner_content"]],
    ev_alts=[['ssh_collected', 'lower(sshd_value("banner")) == "none"'],
             ['ssh_collected', 'sshd_first("banner")', 'is_string(input.ssh.banner_content)']])
L["1.7.6"] = access(["initial_setup", "motd_mode"], ["initial_setup", "motd_owner"], ["initial_setup", "motd_group"],
                    "0644", ["root"], ["root"], exists=["initial_setup", "motd_exists"], optional=True)
L["1.7.7"] = access(["initial_setup", "issue_mode"], ["initial_setup", "issue_owner"], ["initial_setup", "issue_group"],
                    "0644", ["root"], ["root"], exists=["initial_setup", "issue_exists"])
L["1.7.8"] = access(["initial_setup", "issue_net_mode"], ["initial_setup", "issue_net_owner"],
                    ["initial_setup", "issue_net_group"], "0644", ["root"], ["root"],
                    exists=["initial_setup", "issue_net_exists"])
BP = ["ssh", "banner_permissions"]
L["1.7.9"] = custom(
    None,
    [['lower(sshd_value("banner")) != "none"', 'mode_within(input.ssh.banner_permissions.mode, "0644")',
      'input.ssh.banner_permissions.pw_name == "root"', 'input.ssh.banner_permissions.gr_name == "root"']],
    "ssh.banner_permissions",
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_permissions": {"mode": "0644", "pw_name": "root", "gr_name": "root"}}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_permissions": {"mode": "0666", "pw_name": "root", "gr_name": "root"}}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net"}}, reads=[SSH_RAW, BP],
    ev_alts=[['ssh_collected', 'lower(sshd_value("banner")) == "none"'],
             ['ssh_collected', 'sshd_first("banner")', 'input.ssh.banner_permissions.mode != null', 'input.ssh.banner_permissions.pw_name != null',
              'input.ssh.banner_permissions.gr_name != null']])


# 1.8 GDM — not applicable (passes) when GDM is not installed, a collected boolean
def gdm(fields, conds, good, bad):
    inst = ["initial_setup", "gdm_installed"]
    fp = [["initial_setup", f] for f in fields]
    have = [f'{ref(p)} != null' for p in fp]
    return custom(None, [['input.initial_setup.gdm_installed == false'],
                         ['input.initial_setup.gdm_installed == true'] + conds],
                  "initial_setup." + "/".join(["gdm_installed"] + fields),
                  {"initial_setup": dict(good, gdm_installed=True)},
                  {"initial_setup": dict(bad, gdm_installed=True)},
                  {"initial_setup": {"gdm_installed": True}}, reads=[inst] + fp,
                  ev_alts=[['input.initial_setup.gdm_installed == false'],
                           ['input.initial_setup.gdm_installed == true'] + have])


I = "input.initial_setup."
L["1.8.1"] = gdm(["gdm_banner_configured", "gdm_banner_override_allowed"],
                 [I + "gdm_banner_configured == true", I + "gdm_banner_override_allowed == false"],
                 {"gdm_banner_configured": True, "gdm_banner_override_allowed": False},
                 {"gdm_banner_configured": True, "gdm_banner_override_allowed": True})
L["1.8.2"] = gdm(["gdm_disable_user_list", "gdm_disable_user_list_override_allowed"],
                 [I + "gdm_disable_user_list == true", I + "gdm_disable_user_list_override_allowed == false"],
                 {"gdm_disable_user_list": True, "gdm_disable_user_list_override_allowed": False},
                 {"gdm_disable_user_list": False, "gdm_disable_user_list_override_allowed": False})
L["1.8.3"] = gdm(["gdm_idle_delay", "gdm_lock_delay", "gdm_lock_override_allowed"],
                 [f"num({I}gdm_idle_delay) > 0", f"num({I}gdm_idle_delay) <= 900",
                  f"num({I}gdm_lock_delay) >= 0", f"num({I}gdm_lock_delay) <= 5", I + "gdm_lock_override_allowed == false"],
                 {"gdm_idle_delay": 900, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False},
                 {"gdm_idle_delay": 0, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False})
L["1.8.4"] = gdm(["gdm_automount_disabled", "gdm_automount_open_disabled", "gdm_automount_override_allowed"],
                 [I + "gdm_automount_disabled == true", I + "gdm_automount_open_disabled == true",
                  I + "gdm_automount_override_allowed == false"],
                 {"gdm_automount_disabled": True, "gdm_automount_open_disabled": True, "gdm_automount_override_allowed": False},
                 {"gdm_automount_disabled": True, "gdm_automount_open_disabled": False, "gdm_automount_override_allowed": False})
L["1.8.5"] = gdm(["gdm_autorun_never_enabled", "gdm_autorun_override_allowed"],
                 [I + "gdm_autorun_never_enabled == true", I + "gdm_autorun_override_allowed == false"],
                 {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": False},
                 {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": True})
L["1.8.6"] = gdm(["gdm_xdmcp_enabled"], [I + "gdm_xdmcp_enabled == false"],
                 {"gdm_xdmcp_enabled": False}, {"gdm_xdmcp_enabled": True})
L["1.8.7"] = gdm(["gdm_wayland_enabled"], [I + "gdm_wayland_enabled == false"],
                 {"gdm_wayland_enabled": False}, {"gdm_wayland_enabled": True})

# 2.1 server services: package not installed, or every unit neither enabled nor active
for rid, units_by_pkg in [
    ("2.1.1", {"autofs": ["autofs.service"]}),
    ("2.1.2", {"avahi": ["avahi-daemon.socket", "avahi-daemon.service"]}),
    ("2.1.3", {"kea": ["kea-dhcp-ddns.service", "kea-dhcp4.service", "kea-dhcp6.service"]}),
    ("2.1.4", {"bind": ["named.service"]}),
    ("2.1.5", {"dnsmasq": ["dnsmasq.service"]}),
    ("2.1.6", {"samba": ["smb.service"]}),
    ("2.1.7", {"openldap2": ["slapd.service"], "openldap2_6": ["slapd.service"]}),
    ("2.1.8", {"vsftpd": ["vsftpd.service"]}),
    ("2.1.9", {"dovecot": ["dovecot.socket", "dovecot.service"], "cyrus-imapd": ["cyrus-imapd.service"]}),
    ("2.1.10", {"nfs-kernel-server": ["nfs-server.service"]}),
    ("2.1.11", {"cups": ["cups.socket", "cups.service"]}),
    ("2.1.12", {"rpcbind": ["rpcbind.socket", "rpcbind.service"]}),
    ("2.1.13", {"rsync": ["rsyncd.socket", "rsyncd.service"]}),
    ("2.1.14", {"net-snmp": ["snmpd.service"]}),
    ("2.1.15", {"telnet-server": ["telnet.socket"]}),
    ("2.1.16", {"tftp-server": ["tftp.socket", "tftp.service"]}),
    ("2.1.17", {"squid": ["squid.service"]}),
    # the audit names httpd.* units; SUSE's apache2 unit is apache2.service, so it is checked as well
    ("2.1.18", {"apache2": ["httpd.socket", "httpd.service", "apache2.service"], "nginx": ["nginx.service"]}),
    ("2.1.19", {"xinetd": ["xinetd.service"]}),
]:
    L[rid] = not_in_use(units_by_pkg)
L["2.1.20"] = pkg_absent(["xwayland"])
L["2.1.21"] = custom(
    ['is_array(input.mta.non_loopback_listeners)'],
    [['count(input.mta.non_loopback_listeners) == 0', 'not mta_interfaces_open']],
    "mta.non_loopback_listeners",
    {"mta": {"non_loopback_listeners": [], "inet_interfaces": "localhost"}},
    {"mta": {"non_loopback_listeners": ["0.0.0.0:25"], "inet_interfaces": "all"}},
    reads=[["mta", "non_loopback_listeners"], ["mta", "inet_interfaces"]])
L["2.1.22"] = attest("listening_services_approved")
# 2.2 client packages
L["2.2.1"] = pkg_absent(["ftp"])
L["2.2.2"] = pkg_absent(["openldap2_6-client"])
L["2.2.3"] = pkg_absent(["ypbind"])
L["2.2.4"] = pkg_absent(["telnet"])
L["2.2.5"] = pkg_absent(["tftp"])

# 2.3 chrony
L["2.3.1.1"] = custom(
    ['input.chronyd.servers_configured != null', 'input.chronyd.pools_configured != null'],
    [['num(input.chronyd.servers_configured) > 0'], ['num(input.chronyd.pools_configured) > 0']],
    "chronyd.servers_configured/pools_configured",
    {"chronyd": {"servers_configured": 0, "pools_configured": 2}},
    {"chronyd": {"servers_configured": 0, "pools_configured": 0}},
    reads=[["chronyd", "servers_configured"], ["chronyd", "pools_configured"]])
L["2.3.1.2"] = custom(
    ['is_boolean(input.chronyd.service_enabled)', 'is_boolean(input.chronyd.service_active)'],
    [['input.chronyd.service_enabled == true', 'input.chronyd.service_active == true']],
    "chronyd.service_enabled/service_active",
    {"chronyd": {"service_enabled": True, "service_active": True}},
    {"chronyd": {"service_enabled": True, "service_active": False}},
    reads=[["chronyd", "service_enabled"], ["chronyd", "service_active"]])

# 2.4 job schedulers (rhel_9 cron.* field names; group fields added)
L["2.4.1.1"] = custom(
    ['is_boolean(input.cron.crond_enabled)', 'is_boolean(input.cron.crond_running)'],
    [['input.cron.crond_enabled == true', 'input.cron.crond_running == true']],
    "cron.crond_enabled/crond_running",
    {"cron": {"crond_enabled": True, "crond_running": True}},
    {"cron": {"crond_enabled": True, "crond_running": False}},
    reads=[["cron", "crond_enabled"], ["cron", "crond_running"]])
L["2.4.1.2"] = access(["cron", "crontab_mode"], ["cron", "crontab_owner"], ["cron", "crontab_group"], "0600", ["root"], ["root"])
for rid, d in [("2.4.1.3", "cron_hourly"), ("2.4.1.4", "cron_daily"), ("2.4.1.5", "cron_weekly"),
               ("2.4.1.6", "cron_monthly"), ("2.4.1.7", "cron_yearly"), ("2.4.1.8", "cron_d")]:
    L[rid] = access(["cron", f"{d}_mode"], ["cron", f"{d}_owner"], ["cron", f"{d}_group"], "0700", ["root"], ["root"])


def allow_deny(svc, groups):
    """<svc>.allow must exist (mode 0640, root, groups); <svc>.deny, if present, the same."""
    c = "input.cron."
    a = [f"{svc}_allow_{k}" for k in ("exists", "mode", "owner", "group")]
    d = [f"{svc}_deny_{k}" for k in ("exists", "mode", "owner", "group")]
    g = json.dumps(groups)
    allow_ok = [f'{c}{a[0]} == true', f'mode_within({c}{a[1]}, "0640")', f'{c}{a[2]} == "root"', f'{c}{a[3]} in {g}']
    deny_ok_absent = [f'{c}{d[0]} == false']
    deny_ok_present = [f'{c}{d[0]} == true', f'mode_within({c}{d[1]}, "0640")', f'{c}{d[2]} == "root"', f'{c}{d[3]} in {g}']
    good = {a[0]: True, a[1]: "0640", a[2]: "root", a[3]: groups[0], d[0]: False}
    bad = {a[0]: False, d[0]: False}
    ev_allow = [[f'{c}{a[0]} == false'], [f'{c}{a[0]} == true', f'{c}{a[1]} != null', f'{c}{a[2]} != null', f'{c}{a[3]} != null']]
    ev_deny = [[f'{c}{d[0]} == false'], [f'{c}{d[0]} == true', f'{c}{d[1]} != null', f'{c}{d[2]} != null', f'{c}{d[3]} != null']]
    return custom(None, [allow_ok + deny_ok_absent, allow_ok + deny_ok_present],
                  f"cron.{svc}_allow_*/{svc}_deny_*", {"cron": good}, {"cron": bad}, {},
                  reads=[["cron", k] for k in a + d],
                  ev_alts=[x + y for x in ev_allow for y in ev_deny])


L["2.4.1.9"] = allow_deny("cron", ["root", "crontab"])
L["2.4.2.1"] = allow_deny("at", ["root", "daemon"])
L["2.4.3.1"] = access_list(["systemd", "timer_unit_files"], "0644", ["root"], ["root"], sample="/etc/systemd/system/x.timer")
L["2.4.3.2"] = access_list(["systemd", "service_unit_files"], "0644", ["root"], ["root"], sample="/etc/systemd/system/x.service")
# drop-in directories exist only when an override was made: an empty collected list means none
L["2.4.3.3"] = access_list(["systemd", "dropin_dirs"], "0750", ["root"], ["root"], sample="/etc/systemd/system/x.timer.d",
                           nonempty=False)
L["2.4.3.3"].hollow_ok = True
L["2.4.3.4"] = custom(['is_array(input.systemd.lingering_users)'],
                      [['count([u | some u in input.systemd.lingering_users; u != "root"]) == 0']],
                      "systemd.lingering_users", {"systemd": {"lingering_users": []}},
                      {"systemd": {"lingering_users": ["alice"]}}, reads=[["systemd", "lingering_users"]])

# 3 network
L["3.1.1"] = attest("ipv6_status_identified")
L["3.1.2"] = custom(
    None,
    [['input.wireless_interfaces.has_wireless == false'],
     ['input.wireless_interfaces.has_wireless == true', 'is_array(input.wireless_interfaces.modules)',
      'count(input.wireless_interfaces.modules) > 0',
      'every e in input.wireless_interfaces.modules { kmod_unavailable(e) }']],
    "wireless_interfaces.has_wireless/modules",
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": True, "loaded": False}]}},
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": False, "loaded": True}]}},
    {"wireless_interfaces": {"has_wireless": True}},
    reads=[["wireless_interfaces", "has_wireless"], ["wireless_interfaces", "modules"]],
    ev_alts=[['input.wireless_interfaces.has_wireless == false'],
             ['input.wireless_interfaces.has_wireless == true', 'is_array(input.wireless_interfaces.modules)']])
L["3.1.3"] = not_in_use({"bluez": ["bluetooth.service"]})
for rid, mod in [("3.2.1", "atm"), ("3.2.2", "can"), ("3.2.3", "dccp"), ("3.2.4", "tipc"), ("3.2.5", "rds"), ("3.2.6", "sctp")]:
    L[rid] = kmod(NETM, mod)
for i, (p, v) in enumerate([
        ("net.ipv4.ip_forward", "0"), ("net.ipv4.conf.all.forwarding", "0"), ("net.ipv4.conf.default.forwarding", "0"),
        ("net.ipv4.conf.all.send_redirects", "0"), ("net.ipv4.conf.default.send_redirects", "0"),
        ("net.ipv4.icmp_ignore_bogus_error_responses", "1"), ("net.ipv4.icmp_echo_ignore_broadcasts", "1"),
        ("net.ipv4.conf.all.accept_redirects", "0"), ("net.ipv4.conf.default.accept_redirects", "0"),
        ("net.ipv4.conf.all.secure_redirects", "0"), ("net.ipv4.conf.default.secure_redirects", "0"),
        ("net.ipv4.conf.all.rp_filter", "1"), ("net.ipv4.conf.default.rp_filter", "1"),
        ("net.ipv4.conf.all.accept_source_route", "0"), ("net.ipv4.conf.default.accept_source_route", "0"),
        ("net.ipv4.conf.all.log_martians", "1"), ("net.ipv4.conf.default.log_martians", "1"),
        ("net.ipv4.tcp_syncookies", "1"), ("net.ipv4.conf.lo.route_localnet", "0")], start=1):
    L[f"3.3.1.{i}"] = sysctl(p, [v])
for i, p in enumerate(["net.ipv6.conf.all.forwarding", "net.ipv6.conf.default.forwarding",
                       "net.ipv6.conf.all.accept_redirects", "net.ipv6.conf.default.accept_redirects",
                       "net.ipv6.conf.all.accept_source_route", "net.ipv6.conf.default.accept_source_route",
                       "net.ipv6.conf.all.accept_ra", "net.ipv6.conf.default.accept_ra"], start=1):
    L[f"3.3.2.{i}"] = sysctl(p, ["0"], ipv6=True)

# 4 firewalld
L["4.1.1"] = pkg_present("firewalld")
L["4.1.2"] = custom(['is_string(input.firewalld.backend)'], [['lower(input.firewalld.backend) == "nftables"']],
                    "firewalld.backend", {"firewalld": {"backend": "nftables"}}, {"firewalld": {"backend": "iptables"}},
                    reads=[["firewalld", "backend"]])
L["4.1.3"] = unit_running("firewalld.service")
L["4.1.4"] = custom(
    ['is_array(input.firewalld.zones)', 'count(input.firewalld.zones) > 0'],
    [['every z in input.firewalld.zones { zone_target_ok(z) }']], "firewalld.zones",
    {"firewalld": {"zones": [{"name": "public", "interfaces": ["eth0"], "target": "default", "permanent_target": "default"},
                             {"name": "trusted", "interfaces": ["lo"], "target": "ACCEPT", "permanent_target": "ACCEPT"}]}},
    {"firewalld": {"zones": [{"name": "public", "interfaces": ["eth0"], "target": "ACCEPT", "permanent_target": "ACCEPT"}]}},
    reads=[["firewalld", "zones"]])
L["4.1.5"] = attest("firewalld_loopback_traffic_configured")
L["4.1.6"] = attest("firewalld_loopback_source_traffic_configured")
L["4.1.7"] = attest("firewalld_services_ports_reviewed")
L["4.1.8"] = attest("firewall_default_drop_configured")

# 5.1 SSH server (sshd_config in /etc/ssh and the /usr/etc/ssh vendor copy on SUSE)
L["5.1.1"] = access_list(["ssh", "config_files"], "0600", ["root"], ["root"], mode="mode", owner="pw_name",
                         group="gr_name", sample="/etc/ssh/sshd_config")
L["5.1.2"] = custom(
    ['is_array(input.ssh.private_key_permissions)', 'count(input.ssh.private_key_permissions) > 0'],
    [['every k in input.ssh.private_key_permissions { mode_within(k.stat.mode, "0600"); k.stat.pw_name == "root"; k.stat.gr_name == "root" }']],
    "ssh.private_key_permissions",
    {"ssh": {"private_key_permissions": [{"stat": {"path": "/etc/ssh/ssh_host_ed25519_key", "mode": "0600", "pw_name": "root", "gr_name": "root"}}]}},
    {"ssh": {"private_key_permissions": [{"stat": {"path": "/etc/ssh/ssh_host_ed25519_key", "mode": "0640", "pw_name": "root", "gr_name": "root"}}]}},
    reads=[["ssh", "private_key_permissions"]])
L["5.1.3"] = access_list(["ssh", "public_key_permissions"], "0644", ["root"], ["root"], mode="mode", owner="pw_name",
                         group="gr_name", sample="/etc/ssh/ssh_host_ed25519_key.pub")
L["5.1.4"] = custom(['ssh_collected'], [['some k in ["allowusers", "allowgroups", "denyusers", "denygroups"]', 'sshd_first(k) != ""']],
                    "ssh.sshd_config_raw", ssh_in(["port 22", "allowgroups sshusers"]), ssh_in(["port 22"]), {},
                    reads=[SSH_RAW])
L["5.1.5"] = sshd("banner", 'startswith(v, "/")', "/etc/issue.net", "none")
WEAK_CIPH = ["3des-cbc", "aes128-cbc", "aes192-cbc", "aes256-cbc", "arcfour", "arcfour128", "arcfour256",
             "blowfish-cbc", "cast128-cbc", "rijndael-cbc@lysator.liu.se"]
WEAK_MAC = ["hmac-md5", "hmac-md5-96", "hmac-ripemd160", "hmac-sha1-96", "umac-64@openssh.com",
            "hmac-md5-etm@openssh.com", "hmac-md5-96-etm@openssh.com", "hmac-ripemd160-etm@openssh.com",
            "hmac-sha1-96-etm@openssh.com", "umac-64-etm@openssh.com", "umac-128-etm@openssh.com"]
WEAK_KEX = ["diffie-hellman-group1-sha1", "diffie-hellman-group14-sha1", "diffie-hellman-group-exchange-sha1"]


def sshd_list_none(key, weak, good, bad):
    s = sshd(key, f'count([x | some x in sshd_list({json.dumps(key)}); x in {json.dumps(weak)}]) == 0', good, bad,
             explicit=True)
    s.passes = [[f'count([x | some x in sshd_list({json.dumps(key)}); x in {json.dumps(weak)}]) == 0']]
    return s


L["5.1.6"] = sshd_list_none("ciphers", WEAK_CIPH, "aes256-gcm@openssh.com,aes256-ctr", "aes256-ctr,aes128-cbc")
L["5.1.7"] = custom(
    ['ssh_collected', 'sshd_first("clientaliveinterval")', 'sshd_first("clientalivecountmax")'],
    [['num(sshd_value("clientaliveinterval")) > 0', 'num(sshd_value("clientalivecountmax")) > 0']],
    "ssh.sshd_config_raw (effective clientaliveinterval/clientalivecountmax)", ssh_in(["clientaliveinterval 15", "clientalivecountmax 3"]),
    ssh_in(["clientaliveinterval 0", "clientalivecountmax 3"]), {}, reads=[SSH_RAW])
L["5.1.8"] = sshd("disableforwarding", 'v == "yes"', "yes", "no")
L["5.1.9"] = sshd("gssapiauthentication", 'v == "no"', "no", "yes")
L["5.1.10"] = sshd("hostbasedauthentication", 'v == "no"', "no", "yes")
L["5.1.11"] = sshd("ignorerhosts", 'v == "yes"', "yes", "no")
L["5.1.12"] = sshd_list_none("kexalgorithms", WEAK_KEX, "curve25519-sha256,mlkem768x25519-sha256",
                             "curve25519-sha256,diffie-hellman-group14-sha1")
L["5.1.13"] = sshd("logingracetime", 'sshd_seconds(v) >= 1', "60", "120")
L["5.1.13"].passes = [['v := lower(sshd_value("logingracetime"))', 'sshd_seconds(v) >= 1', 'sshd_seconds(v) <= 60']]
L["5.1.14"] = sshd("loglevel", 'v in {"info", "verbose"}', "VERBOSE", "QUIET")
L["5.1.15"] = sshd_list_none("macs", WEAK_MAC, "hmac-sha2-512-etm@openssh.com,hmac-sha2-256",
                             "hmac-sha2-256,umac-64-etm@openssh.com")
L["5.1.16"] = sshd("maxauthtries", 'num(v) <= 4', "4", "6")
L["5.1.17"] = sshd("maxstartups", 'maxstartups_ok(v)', "10:30:60", "10:30:100")
L["5.1.18"] = sshd("maxsessions", 'num(v) <= 10', "10", "20")
L["5.1.19"] = sshd("permitemptypasswords", 'v == "no"', "no", "yes")
L["5.1.20"] = sshd("permitrootlogin", 'v == "no"', "no", "prohibit-password")
L["5.1.21"] = sshd("permituserenvironment", 'v == "no"', "no", "yes")
L["5.1.22"] = sshd("usepam", 'v == "yes"', "yes", "no")
L["5.1.23"] = custom(
    ['ssh_collected', 'sshd_first("kexalgorithms")', 'is_string(input.ssh.openssh_version)'],
    [['"sntrup761x25519-sha512" in sshd_list("kexalgorithms")', 'not openssh_needs_mlkem'],
     ['"sntrup761x25519-sha512" in sshd_list("kexalgorithms")', '"mlkem768x25519-sha256" in sshd_list("kexalgorithms")']],
    "ssh.sshd_config_raw (effective kexalgorithms)/ssh.openssh_version",
    {"ssh": {"sshd_config_raw": "kexalgorithms mlkem768x25519-sha256,sntrup761x25519-sha512,curve25519-sha256", "openssh_version": "9.9"}},
    {"ssh": {"sshd_config_raw": "kexalgorithms sntrup761x25519-sha512,curve25519-sha256", "openssh_version": "10.0"}},
    {"ssh": {"sshd_config_raw": "kexalgorithms sntrup761x25519-sha512"}}, reads=[SSH_RAW, ["ssh", "openssh_version"]])
L["5.1.24"] = sshd("listenaddress", 'not listen_all(v)', "192.0.2.10", "0.0.0.0")

# 5.2 privilege escalation (sudoers: /etc/sudoers*, /usr/etc/sudoers* on SUSE)
SU = ["sudo", "sudoers_content"]
L["5.2.1"] = boolean(["sudo", "installed"], True)
L["5.2.2"] = custom([f'is_string({ref(SU)})'],
                    [['sudo_defaults_has(`(^|[\\s,])use_pty\\b`)', 'not sudo_defaults_has(`!\\s*use_pty\\b`)']],
                    dotted(SU), nest(SU, "Defaults use_pty\nroot ALL=(ALL:ALL) ALL"),
                    nest(SU, "Defaults !use_pty\nroot ALL=(ALL:ALL) ALL"), reads=[SU])
L["5.2.3"] = custom([f'is_string({ref(SU)})'], [['sudo_logfile']], dotted(SU),
                    nest(SU, 'Defaults logfile="/var/log/sudo.log"'), nest(SU, "# Defaults logfile=/var/log/sudo.log"),
                    reads=[SU])
L["5.2.4"] = text_has(SU, "NOPASSWD", False, "%wheel ALL=(ALL:ALL) ALL", "%wheel ALL=(ALL:ALL) NOPASSWD: ALL")
L["5.2.5"] = text_has(SU, "!authenticate", False, "Defaults use_pty", "Defaults !authenticate")
L["5.2.6"] = custom(
    None,
    [['count(sudo_timeouts) > 0', 'every t in sudo_timeouts { t >= 0; t <= 15 }'],
     ['count(sudo_timeouts) == 0', 'num(input.sudo.default_timestamp_timeout) >= 0',
      'num(input.sudo.default_timestamp_timeout) <= 15']],
    "sudo.sudoers_content/sudo.default_timestamp_timeout",
    {"sudo": {"sudoers_content": "Defaults timestamp_timeout=15"}},
    {"sudo": {"sudoers_content": "Defaults timestamp_timeout=-1"}},
    {"sudo": {"sudoers_content": "Defaults use_pty"}}, reads=[SU, ["sudo", "default_timestamp_timeout"]],
    ev_alts=[[f'is_string({ref(SU)})', 'count(sudo_timeouts) > 0'],
             [f'is_string({ref(SU)})', 'input.sudo.default_timestamp_timeout != null']])
L["5.2.7"] = custom(
    ['is_string(input.pam.su)', 'is_string(input.user_group.group_content)'],
    [['some g in su_wheel_groups', 'su_group_empty(g)']],
    "pam.su/user_group.group_content",
    {"pam": {"su": "auth required pam_wheel.so use_uid group=sugroup"},
     "user_group": {"group_content": "root:x:0:\nsugroup:x:1001:"}},
    {"pam": {"su": "auth required pam_wheel.so use_uid group=sugroup"},
     "user_group": {"group_content": "root:x:0:\nsugroup:x:1001:alice"}},
    reads=[["pam", "su"], ["user_group", "group_content"]])
L["5.1.24"].passes = [['vals := [p[1] | some p in sshd_pairs; p[0] == "listenaddress"]', 'count(vals) > 0',
                       'every x in vals { not listen_all(lower(x)) }']]

# 5.3 PAM (SUSE: pam-config managed /etc/pam.d/common-*; pwquality and faillock vendor files in /usr)
L["5.3.1.1"] = custom(
    ['is_boolean(input.packages.pam)', 'is_array(input.pending_updates)'],
    [['input.packages.pam == true', 'not "pam" in input.pending_updates']], "packages.pam/pending_updates",
    {"packages": {"pam": True}, "pending_updates": ["vim"]}, {"packages": {"pam": True}, "pending_updates": ["pam"]},
    reads=[["packages", "pam"], ["pending_updates"]])
FL = ["pam", "faillock_conf_raw"]
L["5.3.2.1.1"] = conf(FL, "deny", 'num(v) >= 1', "deny = 3", "deny = 10")
L["5.3.2.1.1"].passes = [p + ['num(v) <= 5'] for p in L["5.3.2.1.1"].passes]
L["5.3.2.1.2"] = conf(FL, "unlock_time", 'unlock_ok(v)', "unlock_time = 900", "unlock_time = 600")
L["5.3.2.1.3"] = custom(
    [f'is_string({ref(FL)})'],
    [['regex_in_text(input.pam.faillock_conf_raw, `(?i)^even_deny_root\\b`)', 'not root_unlock_short'],
     ['conf_last(input.pam.faillock_conf_raw, "root_unlock_time")', 'not root_unlock_short']],
    dotted(FL), nest(FL, "deny = 3\neven_deny_root\nroot_unlock_time = 60"), nest(FL, "deny = 3\nroot_unlock_time = 30"),
    reads=[FL])
PQ = ["pam", "pwquality_conf_raw"]
L["5.3.2.2.1"] = conf(PQ, "dictcheck", 'num(v) != 0', "dictcheck = 1", "dictcheck = 0")
L["5.3.2.2.2"] = conf(PQ, "difok", 'num(v) >= 2', "difok = 2", "difok = 1")
L["5.3.2.2.3"] = conf(PQ, "minlen", 'num(v) >= 14', "minlen = 14", "minlen = 8")
L["5.3.2.2.4"] = custom(
    [f'is_string({ref(PQ)})'],
    [['num(conf_last(input.pam.pwquality_conf_raw, "minclass")) >= 4'],
     ['every c in ["dcredit", "ucredit", "lcredit", "ocredit"] { num(conf_last(input.pam.pwquality_conf_raw, c)) <= -1 }']],
    dotted(PQ), nest(PQ, "minclass = 4"), nest(PQ, "minclass = 2\ndcredit = -1"), reads=[PQ])
L["5.3.2.2.5"] = conf(PQ, "maxrepeat", 'num(v) >= 1', "maxrepeat = 3", "maxrepeat = 0")
L["5.3.2.2.5"].passes = [p + ['num(v) <= 3'] for p in L["5.3.2.2.5"].passes]
L["5.3.2.2.6"] = conf(PQ, "maxsequence", 'num(v) >= 1', "maxsequence = 3", "maxsequence = 0")
L["5.3.2.2.6"].passes = [p + ['num(v) <= 3'] for p in L["5.3.2.2.6"].passes]
L["5.3.2.2.7"] = text_has(PQ, "(?i)^enforce_for_root\\b", True, "minlen = 14\nenforce_for_root", "minlen = 14")
PS = ["pam", "system_auth"]   # mapper: PAM stack text (common-password on SUSE)
PWH = "password requisite pam_pwhistory.so remember=24 enforce_for_root use_authtok"
UNIX = "password required pam_unix.so sha512 shadow try_first_pass use_authtok"
L["5.3.2.3.1"] = custom([f'is_string({ref(PS)})'], [['count(pwhistory_lines) > 0', 'every l in pwhistory_lines { pam_arg_num(l, "remember") >= 24 }']],
                        dotted(PS), nest(PS, PWH), nest(PS, PWH.replace("24", "5")), reads=[PS])
L["5.3.2.3.2"] = custom([f'is_string({ref(PS)})'], [['count(pwhistory_lines) > 0', 'every l in pwhistory_lines { regex.match(`\\senforce_for_root\\b`, l) }']],
                        dotted(PS), nest(PS, PWH), nest(PS, PWH.replace(" enforce_for_root", "")), reads=[PS])
L["5.3.2.3.3"] = custom([f'is_string({ref(PS)})'], [['count(pwhistory_lines) > 0', 'every l in pwhistory_lines { regex.match(`\\suse_authtok\\b`, l) }']],
                        dotted(PS), nest(PS, PWH), nest(PS, PWH.replace(" use_authtok", "")), reads=[PS])
L["5.3.2.4.1"] = custom([f'is_string({ref(PS)})', 'count(pam_unix_auth_lines) > 0'], [['count([l | some l in pam_unix_lines; regex.match(`\\snullok\\b`, l)]) == 0']],
                        dotted(PS), nest(PS, "auth required pam_unix.so try_first_pass\n" + UNIX),
                        nest(PS, "auth required pam_unix.so nullok try_first_pass\n" + UNIX), reads=[PS])
L["5.3.2.4.2"] = custom([f'is_string({ref(PS)})', 'count(pam_unix_password_lines) > 0'], [['count([l | some l in pam_unix_password_lines; regex.match(`\\sremember=`, l)]) == 0']],
                        dotted(PS), nest(PS, UNIX), nest(PS, UNIX + " remember=5"), reads=[PS])
L["5.3.2.4.3"] = custom([f'is_string({ref(PS)})', 'count(pam_unix_password_lines) > 0'], [['count(pam_unix_password_lines) > 0', 'every l in pam_unix_password_lines { regex.match(`\\s(sha512|yescrypt)\\b`, l) }']],
                        dotted(PS), nest(PS, UNIX), nest(PS, UNIX.replace("sha512", "md5")), reads=[PS])
L["5.3.2.4.4"] = custom([f'is_string({ref(PS)})', 'count(pam_unix_password_lines) > 0'], [['count(pam_unix_password_lines) > 0', 'every l in pam_unix_password_lines { regex.match(`\\suse_authtok\\b`, l) }']],
                        dotted(PS), nest(PS, UNIX), nest(PS, UNIX.replace(" use_authtok", "")), reads=[PS])

# 5.4 accounts — login.defs values (mapper fields pam.pass_*), per-user shadow facts (no hashes)
SH = ["user_group", "shadow_entries"]


def shadow_rule(defs_path, defs_cond, user_cond, good_def, bad_def, good_user, field):
    """login.defs default AND every user that has a password."""
    dp = ref(defs_path)
    return custom(
        [f'{dp} != null', f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
        [[f'v := num({dp})'] + defs_cond + [
          f'every u in [u | some u in {ref(SH)}; u.has_password == true] {{ {user_cond} }}']],
        f"{dotted(defs_path)}/user_group.shadow_entries",
        merge(nest(defs_path, good_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        merge(nest(defs_path, bad_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        nest(defs_path, good_def), reads=[defs_path, SH])


L["5.4.1.1"] = shadow_rule(["pam", "pass_max_days"], ['v >= 1', 'v <= 365'], 'u.max_days >= 1; u.max_days <= 365',
                           365, 99999, {"max_days": 365}, "max_days")
L["5.4.1.2"] = attest("minimum_password_days_reviewed")
L["5.4.1.3"] = shadow_rule(["pam", "pass_warn_age"], ['v >= 7'], 'u.warn_days >= 7', 7, 3, {"warn_days": 7}, "warn_days")
L["5.4.1.4"] = custom(['is_string(input.pam.encrypt_method)'], [['upper(input.pam.encrypt_method) in {"SHA512", "YESCRYPT"}']],
                      "pam.encrypt_method", {"pam": {"encrypt_method": "SHA512"}}, {"pam": {"encrypt_method": "MD5"}},
                      reads=[["pam", "encrypt_method"]])
L["5.4.1.5"] = shadow_rule(["pam", "inactive_days"], ['v >= 0', 'v <= 45'], 'u.inactive >= 0; u.inactive <= 45',
                           45, -1, {"inactive": 30}, "inactive")
L["5.4.1.6"] = custom(
    [f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [[f'every u in {ref(SH)} {{ u.last_change <= today_days }}']], "user_group.shadow_entries",
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 19000}]),
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 9999999}]), reads=[SH])
PW = ["user_group", "passwd_content"]
GR = ["user_group", "group_content"]
PASSWD = "root:x:0:0:root:/root:/bin/bash\nbin:x:1:1:bin:/bin:/usr/sbin/nologin\nalice:x:1000:100::/home/alice:/bin/bash"
GROUP = "root:x:0:\nbin:x:1:\nusers:x:100:"
L["5.4.2.1"] = custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'],
                      [['count([f | some f in passwd_rows; f[2] == "0"; f[0] != "root"]) == 0']], dotted(PW),
                      nest(PW, PASSWD), nest(PW, PASSWD + "\ntoor:x:0:0::/root:/bin/bash"), reads=[PW])
L["5.4.2.2"] = custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'],
                      [['count([f | some f in passwd_rows; f[3] == "0"; not f[0] in {"root", "sync", "shutdown", "halt", "operator"}]) == 0']],
                      dotted(PW), nest(PW, PASSWD), nest(PW, PASSWD + "\nbob:x:1001:0::/home/bob:/bin/bash"), reads=[PW])
L["5.4.2.3"] = custom([f'is_string({ref(GR)})', 'count(group_rows) > 0'],
                      [['count([f | some f in group_rows; f[2] == "0"; f[0] != "root"]) == 0']], dotted(GR),
                      nest(GR, GROUP), nest(GR, GROUP + "\nwheel2:x:0:"), reads=[GR])
L["5.4.2.4"] = custom(['is_string(input.user_group.root_password_status)'],
                      [['regex.match(`^(P|L)`, input.user_group.root_password_status)']], "user_group.root_password_status",
                      {"user_group": {"root_password_status": "P"}}, {"user_group": {"root_password_status": "NP"}},
                      reads=[["user_group", "root_password_status"]])
L["5.4.2.5"] = custom(
    ['is_string(input.user_group.root_path)', 'is_array(input.user_group.root_path_entries)', 'count(input.user_group.root_path_entries) > 0'],
    [['root_path_text_ok', 'every d in input.user_group.root_path_entries { d.is_dir == true; d.owner == "root"; mode_within(d.mode, "0755") }']],
    "user_group.root_path/root_path_entries",
    {"user_group": {"root_path": "/usr/sbin:/usr/bin", "root_path_entries": [{"path": "/usr/bin", "is_dir": True, "owner": "root", "mode": "0755"}]}},
    {"user_group": {"root_path": "/usr/sbin:/usr/bin:.", "root_path_entries": [{"path": "/usr/bin", "is_dir": True, "owner": "root", "mode": "0755"}]}},
    reads=[["user_group", "root_path"], ["user_group", "root_path_entries"]])
L["5.4.2.6"] = custom(['is_string(input.user_group.root_shell_init_raw)'],
                      [['every u in umask_values(input.user_group.root_shell_init_raw) { umask_ok(u) }']],
                      "user_group.root_shell_init_raw", {"user_group": {"root_shell_init_raw": "umask 0027"}},
                      {"user_group": {"root_shell_init_raw": "umask 0022"}}, reads=[["user_group", "root_shell_init_raw"]])
# neither /root/.bash_profile nor /root/.bashrc present: no umask is set there (collected as root_shell_init_exists)
L["5.4.2.6"].ev_alts.append(['input.user_group.root_shell_init_exists == false'])
L["5.4.2.6"].passes.append(['input.user_group.root_shell_init_exists == false'])
L["5.4.2.6"].reads.append(["user_group", "root_shell_init_exists"])
L["5.4.2.6"].field = "user_group.root_shell_init_raw/root_shell_init_exists"
SHELLS = ["user_group", "shells_content"]
UMIN = ["user_group", "uid_min"]
SHELLS_TXT = "/bin/bash\n/bin/sh\n/usr/sbin/nologin"
L["5.4.2.7"] = custom(
    [f'is_string({ref(PW)})', 'count(passwd_rows) > 0', f'is_string({ref(SHELLS)})', f'is_number({ref(UMIN)})'],
    [['count(system_accounts_with_shell) == 0']], "user_group.passwd_content/shells_content/uid_min",
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "uid_min": 1000}},
    {"user_group": {"passwd_content": PASSWD + "\ndaemon:x:2:2::/:/bin/bash", "shells_content": SHELLS_TXT, "uid_min": 1000}},
    reads=[PW, SHELLS, UMIN])
L["5.4.2.8"] = custom(
    [f'is_string({ref(PW)})', 'count(passwd_rows) > 0', f'is_string({ref(SHELLS)})', f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [['count(unlocked_no_shell_accounts) == 0']], "user_group.passwd_content/shells_content/shadow_entries",
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "shadow_entries": [{"user": "bin", "locked": True}]}},
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "shadow_entries": [{"user": "bin", "locked": False}]}},
    reads=[PW, SHELLS, SH])
L["5.4.3.1"] = text_has(SHELLS, "/nologin\\b", False, "/bin/bash\n/bin/sh", SHELLS_TXT)
L["5.4.3.2"] = custom(['input.pam.tmout != null'], [['num(input.pam.tmout) > 0', 'num(input.pam.tmout) <= 900']], "pam.tmout",
                      {"pam": {"tmout": 900}}, {"pam": {"tmout": 0}}, reads=[["pam", "tmout"]])
L["5.4.3.3"] = custom(['input.pam.default_umask != null'], [['umask_ok(sprintf("%v", [input.pam.default_umask]))']],
                      "pam.default_umask", {"pam": {"default_umask": "027"}}, {"pam": {"default_umask": "022"}},
                      reads=[["pam", "default_umask"]])
L["5.4.3.4"] = text_has(["pam", "profile_d_raw"], "(?i)^umask\\s+0?027\\b", True, "umask 027", "umask 022")

# 6.1 logging (mapper journald / rsyslog fields)
L["6.1.1.1"] = boolean(["journald", "service_active"], True)
L["6.1.1.2"] = units_off(["systemd-journal-remote.socket", "systemd-journal-remote.service"])
L["6.1.1.3"] = custom(['is_string(input.journald.compress_enabled)'], [['lower(input.journald.compress_enabled) == "yes"']],
                      "journald.compress_enabled", {"journald": {"compress_enabled": "yes"}},
                      {"journald": {"compress_enabled": "no"}}, reads=[["journald", "compress_enabled"]])
L["6.1.1.4"] = custom(['is_string(input.journald.storage_mode)'], [['lower(input.journald.storage_mode) == "persistent"']],
                      "journald.storage_mode", {"journald": {"storage_mode": "persistent"}},
                      {"journald": {"storage_mode": "auto"}}, reads=[["journald", "storage_mode"]])
for rid in ("6.1.1.5", "6.1.2.3"):
    L[rid] = custom(['is_string(input.journald.forward_to_syslog)'], [['lower(input.journald.forward_to_syslog) == "yes"']],
                    "journald.forward_to_syslog", {"journald": {"forward_to_syslog": "yes"}},
                    {"journald": {"forward_to_syslog": "no"}}, reads=[["journald", "forward_to_syslog"]])
L["6.1.1.6"] = attest("journald_rotation_reviewed")
L["6.1.1.7"] = attest("journald_file_access_reviewed")
L["6.1.2.1"] = boolean(["rsyslog", "installed"], True)
L["6.1.2.2"] = custom(
    ['is_boolean(input.rsyslog.service_enabled)', 'is_boolean(input.rsyslog.service_active)'],
    [['input.rsyslog.service_enabled == true', 'input.rsyslog.service_active == true']],
    "rsyslog.service_enabled/service_active", {"rsyslog": {"service_enabled": True, "service_active": True}},
    {"rsyslog": {"service_enabled": False, "service_active": True}},
    reads=[["rsyslog", "service_enabled"], ["rsyslog", "service_active"]])
# $FileCreateMode unset => rsyslog's 0644 default, which fails; the mapper omits the key when unset
L["6.1.2.4"] = custom(['is_boolean(input.rsyslog.installed)'],
                      [['mode_within(input.rsyslog.file_create_mode, "0640")']], "rsyslog",
                      {"rsyslog": {"installed": True, "file_create_mode": "0640"}},
                      {"rsyslog": {"installed": True, "file_create_mode": "0644"}},
                      reads=[["rsyslog", "installed"], ["rsyslog", "file_create_mode"]])
L["6.1.2.5"] = attest("rsyslog_logging_reviewed")
L["6.1.2.6"] = attest("rsyslog_remote_host_configured")
RC = ["rsyslog", "config_raw"]
L["6.1.2.7"] = text_has(RC, "(?i)(^module\\(load=\"?imtcp\"?\\)|^input\\(type=\"?imtcp\"?|^\\$ModLoad\\s+imtcp\\b|^\\$InputTCPServerRun\\b)",
                        False, "*.* /var/log/messages", 'module(load="imtcp")')
L["6.1.2.8"] = attest("logrotate_reviewed")
L["6.1.2.9"] = pkg_present("rsyslog-module-gtls")
L["6.1.2.10"] = text_has(RC, "(?i)StreamDriver=\"gtls\"", True,
                         '*.* action(type="omfwd" target="loghost" StreamDriver="gtls" StreamDriverMode="1")',
                         '*.* action(type="omfwd" target="loghost" protocol="tcp")')
L["6.1.2.11"] = attest("rsyslog_ca_certificates_configured")
LF = ["log_file_permissions", "analysis"]
L["6.1.3.1"] = custom(
    [f'is_array({ref(LF)})', f'count({ref(LF)}) > 0'], [[f'every f in {ref(LF)} {{ log_file_ok(f) }}']], dotted(LF),
    nest(LF, [{"file": "/var/log/messages", "actual_mode": "0640", "actual_owner": "root", "actual_group": "root"},
              {"file": "/var/log/wtmp", "actual_mode": "0664", "actual_owner": "root", "actual_group": "utmp"}]),
    nest(LF, [{"file": "/var/log/messages", "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"}]),
    reads=[LF])

# 6.2 auditd
L["6.2.1.1"] = pkg_present("audit")
L["6.2.1.2"] = unit_running("auditd.service")
GRUB = ["bootloader", "grub_config"]  # mapper: /etc/default/grub
L["6.2.1.3"] = custom([f'is_string({ref(GRUB)})'], [['grub_cmdline_has(`^audit=1$`)']], dotted(GRUB),
                      nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit=1"'), nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit=10"'),
                      reads=[GRUB])
L["6.2.1.4"] = custom([f'is_string({ref(GRUB)})'], [['some t in grub_cmdline_tokens', 'startswith(t, "audit_backlog_limit=")',
                                                     'num(substring(t, 20, -1)) >= 8192']], dotted(GRUB),
                      nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit=1 audit_backlog_limit=8192"'),
                      nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit_backlog_limit=64"'), reads=[GRUB])
AC = ["auditd", "config"]
L["6.2.2.1"] = conf(AC, "max_log_file", 'num(v) > 0', "max_log_file = 8", "max_log_file = 0")
L["6.2.2.2"] = conf(AC, "max_log_file_action", 'lower(v) == "keep_logs"', "max_log_file_action = keep_logs",
                    "max_log_file_action = rotate")
L["6.2.2.3"] = custom([f'is_string({ref(AC)})'],
                      [['lower(conf_last(input.auditd.config, "disk_full_action")) in {"halt", "single"}',
                        'lower(conf_last(input.auditd.config, "disk_error_action")) in {"syslog", "single", "halt"}']],
                      dotted(AC), nest(AC, "disk_full_action = halt\ndisk_error_action = syslog"),
                      nest(AC, "disk_full_action = suspend\ndisk_error_action = syslog"), reads=[AC])
L["6.2.2.4"] = custom([f'is_string({ref(AC)})'],
                      [['lower(conf_last(input.auditd.config, "space_left_action")) in {"email", "exec", "single", "halt"}',
                        'lower(conf_last(input.auditd.config, "admin_space_left_action")) in {"single", "halt"}']],
                      dotted(AC), nest(AC, "space_left_action = email\nadmin_space_left_action = single"),
                      nest(AC, "space_left_action = ignore\nadmin_space_left_action = single"), reads=[AC])
AR = ["auditd", "active_rules"]
W = "-a always,exit -F arch=b64 -F path={p} -F perm={perm} -F key=k"
S64 = "-a always,exit -F arch=b64 -S {s} {extra}-F key=k"
AUID = "-F auid>=1000 -F auid!=unset "


def audit(watches=(), syscalls=(), good_extra=(), raw_rules=None):
    """watches: [(path, perm)]; syscalls: [(names, [field regexes])]. Rules are matched on tokens."""
    conds, good = [], []
    for p, perm in watches:
        conds.append(f'audit_watch({json.dumps(p)}, {json.dumps(perm)})')
        good.append(W.format(p=p, perm=perm))
    for names, fields in syscalls:
        conds.append(f'every s in {json.dumps(names)} {{ audit_syscall(s, {json.dumps(fields)}) }}')
        extra = ""
        for f in fields:
            extra += {"auid": AUID, "euid": "-C euid!=uid -F auid!=unset ", "eacces": "-F exit=-EACCES " + AUID,
                      "eperm": "-F exit=-EPERM " + AUID}[f]
        good.append(S64.format(s=",".join(names), extra=extra))
    good += list(good_extra)
    bad = good[:-1]
    return custom([f'is_array({ref(AR)})'], [conds], dotted(AR), nest(AR, good), nest(AR, bad), reads=[AR])


L["6.2.3.1"] = audit([("/etc/sudoers", "wa"), ("/etc/sudoers.d", "wa")])
L["6.2.3.2"] = audit(syscalls=[(["execve"], ["euid"])])
L["6.2.3.3"] = custom(
    [f'is_array({ref(AR)})', f'is_string({ref(SU)})'], [['some f in sudo_logfiles', 'audit_watch(f, "wa")']],
    "auditd.active_rules/sudo.sudoers_content",
    merge(nest(AR, [W.format(p="/var/log/sudo.log", perm="wa")]), nest(SU, 'Defaults logfile="/var/log/sudo.log"')),
    merge(nest(AR, [W.format(p="/var/log/other.log", perm="wa")]), nest(SU, 'Defaults logfile="/var/log/sudo.log"')),
    reads=[AR, SU])
L["6.2.3.4"] = audit([("/etc/localtime", "wa")], [(["adjtimex", "settimeofday", "clock_settime"], [])])
L["6.2.3.5"] = audit(syscalls=[(["sethostname", "setdomainname"], [])])
L["6.2.3.6"] = audit([("/etc/issue", "wa"), ("/etc/issue.net", "wa")])
L["6.2.3.7"] = audit([("/etc/hosts", "wa"), ("/etc/hostname", "wa")])
L["6.2.3.8"] = audit([("/etc/NetworkManager", "wa")])
L["6.2.3.9"] = custom(
    [f'is_array({ref(AR)})', 'is_array(input.auditd.privileged_commands)', 'count(input.auditd.privileged_commands) > 0'],
    [['every c in input.auditd.privileged_commands { audit_watch(c, "x") }']], "auditd.active_rules/privileged_commands",
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su"]}},
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su", "/usr/bin/passwd"]}},
    reads=[AR, ["auditd", "privileged_commands"]])
L["6.2.3.10"] = audit(syscalls=[(["creat", "open", "openat", "truncate", "ftruncate"], ["eacces"]),
                                (["creat", "open", "openat", "truncate", "ftruncate"], ["eperm"])])
L["6.2.3.11"] = audit([("/etc/group", "wa")])
L["6.2.3.12"] = audit([("/etc/passwd", "wa")])
L["6.2.3.13"] = audit([("/etc/gshadow", "wa"), ("/etc/shadow", "wa")])
L["6.2.3.14"] = audit([("/etc/security/opasswd", "wa")])
L["6.2.3.15"] = audit([("/etc/pam.d", "wa")])
L["6.2.3.16"] = audit(syscalls=[(["chmod", "fchmod", "fchmodat", "fchmodat2"], ["auid"])])
L["6.2.3.17"] = audit(syscalls=[(["chown", "fchown", "lchown", "fchownat"], ["auid"])])
L["6.2.3.18"] = audit(syscalls=[(["setxattr", "lsetxattr", "fsetxattr", "removexattr", "lremovexattr", "fremovexattr"], ["auid"])])
L["6.2.3.19"] = audit(syscalls=[(["mount"], ["auid"])])
L["6.2.3.20"] = audit([("/var/run/utmp", "wa"), ("/var/log/wtmp", "wa"), ("/var/log/btmp", "wa")])
L["6.2.3.21"] = audit([("/var/log/lastlog", "wa"), ("/var/run/faillock", "wa")])
L["6.2.3.22"] = audit(syscalls=[(["unlink", "unlinkat"], ["auid"])])
L["6.2.3.23"] = audit(syscalls=[(["rename", "renameat", "renameat2"], ["auid"])])
L["6.2.3.24"] = audit([("/etc/selinux", "wa"), ("/usr/share/selinux", "wa"), ("/var/lib/selinux", "wa")])
L["6.2.3.25"] = audit([("/usr/bin/chcon", "x")])
L["6.2.3.26"] = audit([("/usr/bin/setfacl", "x")])
L["6.2.3.27"] = audit([("/usr/bin/chacl", "x")])
L["6.2.3.28"] = audit([("/usr/sbin/usermod", "x")])
L["6.2.3.29"] = audit([("/usr/bin/kmod", "x")])
L["6.2.3.30"] = audit(syscalls=[(["init_module", "finit_module"], ["auid"])])
L["6.2.3.31"] = audit(syscalls=[(["delete_module"], ["auid"])])
L["6.2.3.32"] = audit(syscalls=[(["create_module", "query_module"], ["auid"])])
ARD = ["auditd", "rules"]   # on-disk rules (/etc/audit/rules.d)
L["6.2.3.33"] = custom([f'is_array({ref(ARD)})'], [['some l in input.auditd.rules', 'regex.match(`^\\s*-c\\s*$`, l)']],
                       dotted(ARD), nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 2"]), reads=[ARD])
L["6.2.3.34"] = custom([f'is_array({ref(ARD)})'], [['audit_last_e == "2"']], dotted(ARD),
                       nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 1"]), reads=[ARD])
L["6.2.3.35"] = attest("audit_running_and_disk_rules_match")
L["6.2.4.1"] = custom(['input.auditd.log_dir_permissions.mode != null'],
                      [['mode_within(input.auditd.log_dir_permissions.mode, "0750")']], "auditd.log_dir_permissions.mode",
                      {"auditd": {"log_dir_permissions": {"mode": "0750"}}}, {"auditd": {"log_dir_permissions": {"mode": "0755"}}},
                      reads=[["auditd", "log_dir_permissions", "mode"]])
ALF = ["auditd", "log_files"]
ALF_GOOD = [{"path": "/var/log/audit/audit.log", "mode": "0600", "owner": "root", "group": "root"}]
L["6.2.4.2"] = custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0'], [[f'every f in {ref(ALF)} {{ mode_within(f.mode, "0640") }}']], dotted(ALF),
                      nest(ALF, ALF_GOOD), nest(ALF, [dict(ALF_GOOD[0], mode="0644")]), reads=[ALF])
L["6.2.4.3"] = custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0'], [[f'every f in {ref(ALF)} {{ f.owner == "root" }}']], dotted(ALF),
                      nest(ALF, ALF_GOOD), nest(ALF, [dict(ALF_GOOD[0], owner="adm")]), reads=[ALF])
L["6.2.4.4"] = custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0', f'is_string({ref(AC)})'],
                      [['lower(conf_last(input.auditd.config, "log_group")) in {"root", "audit"}',
                        f'every f in {ref(ALF)} {{ f.group in {{"root", "audit"}} }}']],
                      "auditd.log_files/auditd.config",
                      merge(nest(ALF, ALF_GOOD), nest(AC, "log_group = root")),
                      merge(nest(ALF, ALF_GOOD), nest(AC, "log_group = adm")), reads=[ALF, AC])
for rid, path, key, cond, badv in [
        ("6.2.4.5", ["auditd", "config_files"], "mode", 'mode_within(f.mode, "0640")', "0644"),
        ("6.2.4.6", ["auditd", "config_files"], "owner", 'f.owner == "root"', "adm"),
        ("6.2.4.7", ["auditd", "config_files"], "group", 'f.group == "root"', "adm"),
        ("6.2.4.8", ["auditd", "tools"], "mode", 'mode_within(f.mode, "0755")', "0777"),
        ("6.2.4.9", ["auditd", "tools"], "owner", 'f.owner == "root"', "adm"),
        ("6.2.4.10", ["auditd", "tools"], "group", 'f.group == "root"', "adm")]:
    g = {"path": "/etc/audit/auditd.conf", "mode": "0640", "owner": "root", "group": "root"}
    L[rid] = custom([f'is_array({ref(path)})', f'count({ref(path)}) > 0'], [[f'every f in {ref(path)} {{ {cond} }}']],
                    dotted(path), nest(path, [g]), nest(path, [dict(g, **{key: badv})]), reads=[path])

# 6.3 integrity (rhel_9 boot_security.aide_* fields; SUSE schedules aidecheck.timer)
L["6.3.1"] = boolean(["boot_security", "aide_installed"], True)
L["6.3.2"] = custom(
    ['is_boolean(input.boot_security.aide_cron_configured)', 'is_boolean(input.boot_security.aide_timer_configured)'],
    [['input.boot_security.aide_cron_configured == true'], ['input.boot_security.aide_timer_configured == true']],
    "boot_security.aide_cron_configured/aide_timer_configured",
    {"boot_security": {"aide_cron_configured": False, "aide_timer_configured": True}},
    {"boot_security": {"aide_cron_configured": False, "aide_timer_configured": False}},
    reads=[["boot_security", "aide_cron_configured"], ["boot_security", "aide_timer_configured"]])
AIDE_OPTS = "p+i+n+u+g+s+b+acl+xattrs+sha512"
L["6.3.3"] = custom(
    ['is_object(input.aide.audit_tool_options)'],
    [['every t in ["auditctl", "auditd", "ausearch", "aureport", "augenrules"] { aide_opts_ok(input.aide.audit_tool_options[t]) }']],
    "aide.audit_tool_options",
    {"aide": {"audit_tool_options": {t: AIDE_OPTS for t in ["auditctl", "auditd", "ausearch", "aureport", "augenrules"]}}},
    {"aide": {"audit_tool_options": dict({t: AIDE_OPTS for t in ["auditctl", "auditd", "ausearch", "aureport"]}, augenrules="p+i+n")}},
    reads=[["aide", "audit_tool_options"]])

# 7.1 system files (rhel_9 critical_files.analysis entries, judged here from actual_* values)
L["7.1.1"] = crit("/etc/passwd", "0644", ["root"])
L["7.1.2"] = crit("/etc/passwd-", "0644", ["root"])
L["7.1.3"] = crit("/etc/group", "0644", ["root"])
L["7.1.4"] = crit("/etc/group-", "0644", ["root"])
L["7.1.5"] = crit("/etc/shadow", "0640", ["root", "shadow"])
L["7.1.6"] = crit("/etc/shadow-", "0640", ["root", "shadow"])
L["7.1.7"] = crit("/etc/gshadow", "0640", ["root", "shadow"])
L["7.1.8"] = crit("/etc/gshadow-", "0640", ["root", "shadow"], optional=True)
L["7.1.9"] = crit("/etc/shells", "0644", ["root"])
# opasswd and opasswd.old: each checked only if it exists
_op = crit("/etc/security/opasswd", "0600", ["root"], optional=True)
_op2 = crit("/etc/security/opasswd.old", "0600", ["root"], optional=True)
_op.ev_alts = [_op.ev_alts[0] + _op2.ev_alts[0]]
_op.passes = [a + b for a in _op.passes for b in _op2.passes]
_op.field = "critical_files.analysis[/etc/security/opasswd, /etc/security/opasswd.old]"
_old = {"file": "/etc/security/opasswd.old", "exists": False}
_op.tp = {"critical_files": {"analysis": _op.tp["critical_files"]["analysis"] + [_old]}}
_op.tf = {"critical_files": {"analysis": _op.tf["critical_files"]["analysis"] + [_old]}}
L["7.1.10"] = _op
L["7.1.11"] = custom(
    ['is_array(input.world_writable.files)', 'is_array(input.sticky_bit.world_writable_without_sticky)'],
    [['count(input.world_writable.files) == 0', 'count(input.sticky_bit.world_writable_without_sticky) == 0']],
    "world_writable.files/sticky_bit.world_writable_without_sticky",
    {"world_writable": {"files": []}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": ["/opt/app/data.txt"]}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": []}},
    reads=[["world_writable", "files"], ["sticky_bit", "world_writable_without_sticky"]])
L["7.1.12"] = custom(
    ['is_array(input.unowned_files.files)', 'is_array(input.ungrouped_files.files)'],
    [['count(input.unowned_files.files) == 0', 'count(input.ungrouped_files.files) == 0']],
    "unowned_files.files/ungrouped_files.files",
    {"unowned_files": {"files": []}, "ungrouped_files": {"files": []}},
    {"unowned_files": {"files": []}, "ungrouped_files": {"files": ["/srv/orphan"]}},
    reads=[["unowned_files", "files"], ["ungrouped_files", "files"]])
L["7.1.13"] = attest("suid_sgid_files_reviewed")

# 7.2 local users and groups
L["7.2.1"] = custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'], [['count([f | some f in passwd_rows; f[1] != "x"]) == 0']],
                    dotted(PW), nest(PW, PASSWD), nest(PW, PASSWD + "\nbob::1001:100::/home/bob:/bin/bash"), reads=[PW])
L["7.2.2"] = custom([f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
                    [[f'count([u | some u in {ref(SH)}; u.password_empty != false]) == 0']], dotted(SH),
                    nest(SH, [{"user": "root", "password_empty": False}]), nest(SH, [{"user": "bob", "password_empty": True}]),
                    reads=[SH])
L["7.2.3"] = custom([f'is_string({ref(PW)})', f'is_string({ref(GR)})', 'count(passwd_rows) > 0', 'count(group_rows) > 0'],
                    [['count([f | some f in passwd_rows; not f[3] in {g[2] | some g in group_rows}]) == 0']],
                    "user_group.passwd_content/group_content", {"user_group": {"passwd_content": PASSWD, "group_content": GROUP}},
                    {"user_group": {"passwd_content": PASSWD + "\nbob:x:1001:4242::/home/bob:/bin/bash", "group_content": GROUP}},
                    reads=[PW, GR])
for rid, rows_name, col, path, txt, dup in [
        ("7.2.4", "passwd_rows", 2, PW, PASSWD, "\ntoor:x:1000:100::/home/t:/bin/bash"),
        ("7.2.5", "group_rows", 2, GR, GROUP, "\nstaff:x:100:"),
        ("7.2.6", "passwd_rows", 0, PW, PASSWD, "\nalice:x:1005:100::/home/a2:/bin/bash"),
        ("7.2.7", "group_rows", 0, GR, GROUP, "\nusers:x:1200:")]:
    L[rid] = custom([f'is_string({ref(path)})', f'count({rows_name}) > 0'],
                    [[f'count([f | some f in {rows_name}]) == count({{f[{col}] | some f in {rows_name}}})']],
                    dotted(path), nest(path, txt), nest(path, txt + dup), reads=[path])
HOMES = ["user_group", "interactive_homes"]
L["7.2.8"] = custom(
    [f'is_array({ref(HOMES)})'],
    [[f'every h in {ref(HOMES)} {{ h.exists == true; h.owner == h.user; mode_within(h.mode, "0750") }}']], dotted(HOMES),
    nest(HOMES, [{"user": "alice", "home": "/home/alice", "exists": True, "owner": "alice", "mode": "0750"}]),
    nest(HOMES, [{"user": "alice", "home": "/home/alice", "exists": True, "owner": "alice", "mode": "0755"}]), reads=[HOMES])
DOT = ["user_group", "dotfiles"]
L["7.2.9"] = custom(
    [f'is_array({ref(DOT)})'], [[f'every d in {ref(DOT)} {{ dotfile_ok(d) }}']], dotted(DOT),
    nest(DOT, [{"user": "alice", "path": "/home/alice/.bashrc", "mode": "0644", "owner": "alice", "group": "users", "primary_group": "users"},
               {"user": "alice", "path": "/home/alice/.bash_history", "mode": "0600", "owner": "alice", "group": "users", "primary_group": "users"}]),
    nest(DOT, [{"user": "alice", "path": "/home/alice/.rhosts", "mode": "0600", "owner": "alice", "group": "users", "primary_group": "users"}]),
    reads=[DOT])
L["7.2.10"] = custom([f'is_string({ref(GR)})', 'count(group_rows) > 0'],
                     [['count([f | some f in group_rows; not f[1] in {"x", "!", ""}]) == 0']], dotted(GR),
                     nest(GR, GROUP), nest(GR, GROUP + "\nold:secret:1300:"), reads=[GR])


# An empty collected list or blank body is a genuine pass here (offender lists; empty banners; no interactive
# users / dotfiles; no drop-in directories). Collector must always emit these keys.
for _rid in ["5.3.1.1", "1.2.1.2", "1.7.1", "1.7.2", "1.7.3", "2.4.3.3", "2.4.3.4", "7.1.11", "7.1.12", "7.2.8", "7.2.9"]:
    L[_rid].hollow_ok = True

# ── helpers used by one module only ──────────────────────────────────────────────────────────
EXTRA = {
    "crypto_policy": r'''
# CURRENT.pol: a hash/sign line that still enables SHA1 (a leading "-" removes it)
sha1_hash_or_sign_enabled if {
	some l in active_lines(input.crypto_policy.current_pol_raw)
	some m in regex.find_all_string_submatch_n(`^([^=]*)=(.*)$`, l, 1)
	regex.match(`(hash|sign)`, m[1])
	contains(m[2], "SHA1")
	not regex.match(`^\s*-\s*([^#]+)?SHA1`, m[2])
}

# the cipher list SSH uses: the cipher@openssh / cipher@libssh override when present, else the global one
ssh_cipher_lines := [m[3] |
	some l in active_lines(input.crypto_policy.current_pol_raw)
	some m in regex.find_all_string_submatch_n(`(?i)^cipher@(lib|open)ssh(-server|-client)?\s*=\s*(.*)$`, l, 1)
]

ssh_cipher_has(x) if {
	count(ssh_cipher_lines) > 0
	some v in ssh_cipher_lines
	contains(lower(v), x)
}

ssh_cipher_has(x) if {
	count(ssh_cipher_lines) == 0
	contains(lower(conf_last(input.crypto_policy.current_pol_raw, "cipher")), x)
}
''',
    "banners": r'''
pam_motd_ok(l) if not contains(l, "pam_motd.so")

pam_motd_ok(l) if {
	contains(l, "pam_motd.so")
	regex.match(`\smotd=\S+`, l)
}
''',
    "services": r'''
# postfix inet_interfaces (when collected): "all" or a non-loopback value listens off-host
mta_interfaces_open if regex.match(`\ball\b`, lower(input.mta.inet_interfaces))

mta_interfaces_open if {
	is_string(input.mta.inet_interfaces)
	not regex.match(`(0\.0\.0\.0|::1|loopback-only|localhost)`, lower(input.mta.inet_interfaces))
}
''',
    "firewall": r'''
virtual_iface(i) if regex.match(`^(lo|virbr.*|docker.*|veth.*|br-.*|podman.*|cni.*)$`, i)

# zones serving only loopback/virtual interfaces are out of scope
zone_target_ok(z) if every i in object.get(z, "interfaces", []) { virtual_iface(i) }

zone_target_ok(z) if {
	some i in z.interfaces
	not virtual_iface(i)
	upper(z.target) != "ACCEPT"
	lower(z.target) == lower(z.permanent_target)
}
''',
    "ssh": r'''
maxstartups_ok(v) if {
	parts := split(v, ":")
	num(parts[0]) <= 10
	ms_le(parts, 1, 30)
	ms_le(parts, 2, 60)
}

ms_le(parts, i, _) if count(parts) <= i

ms_le(parts, i, lim) if num(parts[i]) <= lim

listen_all(v) if regex.match(`^(0\.0\.0\.0|\[?::\]?)(:[0-9]+)?$`, v)

openssh_ver := [to_number(m[1]), to_number(m[2])] if {
	some m in regex.find_all_string_submatch_n(`([0-9]+)\.([0-9]+)`, input.ssh.openssh_version, 1)
}

openssh_needs_mlkem if openssh_ver[0] > 9

openssh_needs_mlkem if {
	openssh_ver[0] == 9
	openssh_ver[1] >= 9
}
''',
    "sudo": r'''
sudo_defaults_has(re) if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	regex.match(re, l)
}

sudo_logfile if sudo_defaults_has(`(^|[\s,])logfile\s*=`)

sudo_timeouts := [num(m[1]) |
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`timestamp_timeout\s*=\s*(-?[0-9]+)`, l, -1)
]

# groups named by `auth required|requisite pam_wheel.so use_uid group=<g>`
su_wheel_groups contains g if {
	some l in active_lines(input.pam.su)
	regex.match(`^auth\s+(required|requisite)\s+pam_wheel\.so\b`, l)
	regex.match(`\suse_uid\b`, l)
	some m in regex.find_all_string_submatch_n(`\sgroup=(\S+)`, l, 1)
	g := m[1]
}

su_group_empty(g) if {
	some f in group_rows
	f[0] == g
	count(f) >= 4
	trim_space(f[3]) == ""
}
''',
    "pam": r'''
pwhistory_lines := [l |
	some l in active_lines(input.pam.system_auth)
	regex.match(`^password\s+(\[[^\]]*\]|\S+)\s+pam_pwhistory\.so\b`, l)
]

pam_unix_lines := [l | some l in active_lines(input.pam.system_auth); regex.match(`\spam_unix\.so\b`, l)]

pam_unix_password_lines := [l | some l in pam_unix_lines; regex.match(`^password\s`, l)]

pam_unix_auth_lines := [l | some l in pam_unix_lines; regex.match(`^auth\s`, l)]

pam_arg_num(l, k) := num(m[1]) if {
	some m in regex.find_all_string_submatch_n(concat("", [`\s`, k, `=([0-9]+)`]), l, 1)
}

unlock_ok(v) if num(v) == 0

unlock_ok(v) if num(v) >= 900

root_unlock_short if num(conf_last(input.pam.faillock_conf_raw, "root_unlock_time")) < 60
''',
    "accounts": r'''
today_days := time.now_ns() / 86400000000000

root_path_text_ok if {
	p := input.user_group.root_path
	not contains(p, "::")
	not endswith(trim_space(p), ":")
	not startswith(trim_space(p), ":")
	not "." in split(p, ":")
}

umask_values(text) := [m[1] |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(`(?i)^umask\s+(\S+)`, l, 1)
]

# at least as restrictive as 027 (027 octal = 23)
umask_ok(u) if bits.and(octal(u), 23) == 23

umask_ok(u) if {
	regex.match(`^([ugoa]+[=+-][rwx]*,?)+$`, u)
	not regex.match(`g=[rx]*w`, u)
	not regex.match(`o=[rwx]+`, u)
}

system_uid(uid) if uid < input.user_group.uid_min

system_uid(uid) if uid == 65534

system_accounts_with_shell := [f[0] |
	some f in passwd_rows
	not f[0] in {"root", "halt", "sync", "shutdown", "nfsnobody"}
	system_uid(to_number(f[2]))
	f[6] in valid_shells
]

shadow_locked(u) if {
	some e in input.user_group.shadow_entries
	e.user == u
	e.locked == true
}

unlocked_no_shell_accounts := [f[0] |
	some f in passwd_rows
	f[0] != "root"
	not f[6] in valid_shells
	not shadow_locked(f[0])
]
''',
    "logging": r'''
# 6.1.3.1: per-file expectations by file name, from the benchmark's audit script (max mode, owners, groups)
log_rule(b) := ["0664", {"root"}, {"root", "utmp"}] if {
	regex.match(`^((lastlog|wtmp|btmp)([.-].*)?|README)$`, b)
} else := ["0644", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(cloud-init\.log|localmessages|waagent\.log)`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(secure([.-].*)?|auth\.log|syslog|messages)$`, b)
} else := ["0660", {"root", "SSSD"}, {"root", "SSSD"}] if {
	b in {"SSSD", "sssd"}
} else := ["0660", {"root"}, {"root", "gdm", "gdm3"}] if {
	b in {"gdm", "gdm3"}
} else := ["0640", {"root"}, {"root", "systemd-journal"}] if {
	regex.match(`\.journal~?$`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}]

log_file_ok(f) if {
	r := log_rule(regex.replace(f.file, `^.*/`, ""))
	mode_within(f.actual_mode, r[0])
	f.actual_owner in r[1]
	f.actual_group in r[2]
}
''',
    "auditd": r'''
audit_rule_lines := [l |
	some r in input.auditd.active_rules
	l := trim_space(r)
	l != ""
	not startswith(l, "#")
]

audit_toks(l) := regex.split(`\s+`, l)

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some t in toks
	t in {concat("", ["path=", p]), concat("", ["dir=", p]), concat("", ["path=", p, "/"]), concat("", ["dir=", p, "/"])}
}

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-w"
	toks[i + 1] in {p, concat("", [p, "/"])}
}

audit_perm(l) := substring(t, 5, -1) if {
	some t in audit_toks(l)
	startswith(t, "perm=")
}

audit_perm(l) := toks[i + 1] if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-p"
}

audit_watch(p, perm) if {
	some l in audit_rule_lines
	audit_path_tok(l, p)
	pv := audit_perm(l)
	every i in numbers.range(0, count(perm) - 1) { contains(pv, substring(perm, i, 1)) }
}

line_syscalls(l) := {x |
	toks := audit_toks(l)
	some i, t in toks
	t == "-S"
	some x in split(toks[i + 1], ",")
}

audit_unset := `auid!=(unset|-1|4294967295)(\s|$)`

audit_field(l, "auid") if {
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "euid") if {
	regex.match(`(euid!=uid|uid!=euid)(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eacces") if {
	regex.match(`exit=-EACCES(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eperm") if {
	regex.match(`exit=-EPERM(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_syscall(s, fields) if {
	some l in audit_rule_lines
	regex.match(`^-a\s+(always,exit|exit,always)(\s|$)`, l)
	s in line_syscalls(l)
	every f in fields { audit_field(l, f) }
}

# the last `-e` line of the on-disk rules decides the enabled flag
audit_e_values := [m[1] |
	some r in input.auditd.rules
	some m in regex.find_all_string_submatch_n(`^\s*-e\s+([0-9])\b`, r, 1)
]

audit_last_e := audit_e_values[count(audit_e_values) - 1] if count(audit_e_values) > 0

sudo_logfiles contains f if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`logfile\s*=\s*("[^"]+"|'[^']+'|[^,\s]+)`, l, -1)
	f := trim(m[1], "\"'")
}

grub_cmdline_tokens := [t |
	some l in active_lines(input.bootloader.grub_config)
	some m in regex.find_all_string_submatch_n(`^GRUB_CMDLINE_LINUX="([^"]*)"`, l, 1)
	some t in regex.split(`\s+`, m[1])
	t != ""
]

grub_cmdline_has(re) if {
	some t in grub_cmdline_tokens
	regex.match(re, t)
}
''',
    "integrity": r'''
aide_opts_ok(s) if {
	opts := {trim_space(o) | some o in split(s, "+")}
	every want in ["p", "i", "n", "u", "g", "s", "b", "acl", "xattrs", "sha512"] { want in opts }
}
''',
    "user_group": r'''
dot_max(b) := "0600" if b in {".netrc", ".bash_history"} else := "0644"

dotfile_ok(d) if {
	b := regex.replace(d.path, `^.*/`, "")
	not b in {".forward", ".rhosts"}
	mode_within(d.mode, dot_max(b))
	d.owner == d.user
	d.group == d.primary_group
}
''',
}


# ════════════════════════════════════════════════════════════════════════════════════════════
# workbook
# ════════════════════════════════════════════════════════════════════════════════════════════
def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    recs, order = {}, []
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        rows = ws.iter_rows(values_only=True)
        head = next(rows)
        for r in rows:
            d = dict(zip(head, r))
            rid = d.get("Recommendation #")
            if not rid:
                continue
            rid = str(rid).strip()
            if rid not in recs:
                recs[rid] = {"title": str(d["Title"]).strip(), "assessment": str(d["Assessment Status"]).strip(),
                             "profiles": []}
                order.append(rid)
            recs[rid]["profiles"].append(ws.title)
    comb = wb["Combined Profiles"]
    rows = comb.iter_rows(values_only=True)
    head = next(rows)
    combined = {str(dict(zip(head, r))["Recommendation #"]).strip() for r in rows if dict(zip(head, r))["Recommendation #"]}
    if combined != set(recs):
        raise SystemExit(f"Combined Profiles disagrees with the level sheets: {sorted(combined ^ set(recs))}")
    return recs, order


def level(profiles, kind):
    lv = [int(m.group(1)) for p in profiles for m in [re.match(r"Level (\d) - (\w+)", p)] if m and m.group(2) == kind]
    return min(lv) if lv else None


def J(x):
    return json.dumps(x, ensure_ascii=False)


def idkey(rid):
    return rid.replace(".", "_")


def sort_key(rid):
    return [int(x) for x in rid.split(".")]


# ════════════════════════════════════════════════════════════════════════════════════════════
# emit
# ════════════════════════════════════════════════════════════════════════════════════════════
def header(title, checks, entry="none", density="sized", basis="size"):
    return f'''# METADATA
# title: {J(title)}
# custom:
#   class: security
#   framework: cis_suse_16
#   source: cis
#   domains: [linux, suse]
#   coverage:
#     checks: {checks}
#     cited: {checks}
#     density: {density}
#     basis: {basis}
#     tested: true
#     entrypoint: {entry}
#     guard: self
#     as_of: {AS_OF}
'''


def rule_block(name, alts):
    out = [f"default {name} := false", ""]
    for alt in alts:
        out.append(f"{name} if {{")
        out += [f"\t{c}" for c in alt]
        out += ["}", ""]
    return out


def profiles_txt(rec):
    return ", ".join(rec["profiles"])


EMPTY_TEXT_OK = ("motd_content)", "issue_content)", "issue_net_content)", "banner_content)")  # motd/issue/issue.net/sshd banner bodies: an empty banner is a real value


def evidence_alts(spec):
    out = []
    for alt in spec.ev_alts:
        out.append([("text_collected(" + c[len("is_string("):]) if c.startswith("is_string(") and not c.endswith(EMPTY_TEXT_OK)
                    else c for c in alt])
    return out


def hollow(x):
    """the same shape with every string blanked and every list emptied"""
    if isinstance(x, dict):
        return {k: hollow(v) for k, v in x.items()}
    if isinstance(x, list):
        return []
    if isinstance(x, str):
        return ""
    return x


def emit_control(rid, rec, spec):
    k = idkey(rid)
    out = [f"# {rid} ({rec['assessment']}; {profiles_txt(rec)})", f"# {rec['title']}"]
    if spec.manual:
        out += rule_block(f"p_{k}", [[f'input.attest.{spec.name} == true']])
        msg = f"{rid}: manual review not attested — set attest.{spec.name} to true once reviewed ({rec['title']})"
        out += [f"violations contains {J(msg)} if not p_{k}", ""]
        return out
    out += rule_block(f"e_{k}", evidence_alts(spec))
    out += rule_block(f"p_{k}", spec.passes)
    miss = f"{rid}: evidence missing — {spec.field} not collected"
    fail = f"{rid}: not met — {rec['title']}"
    out += [f"violations contains {J(miss)} if not e_{k}", "",
            f"violations contains {J(fail)} if {{", f"\te_{k}", f"\tnot p_{k}", "}", ""]
    return out


def emit_module(name, title, ids, recs):
    lines = [header(f"{BENCH} — {title}", len(ids)).rstrip("\n"), f"package cis_suse16.{name}", "",
             f"# {BENCH}, {title}.",
             "# Generated by tools/cis_generators/gen_suse_16.py from the CIS workbook; do not edit by hand.",
             "# Each recommendation: e_<id> (facts collected), p_<id> (requirement met). Automated checks fail",
             "# closed with '<id>: evidence missing — <field> not collected'; Manual ones need input.attest.<name>.",
             "", "import rego.v1", "",
             "default compliant := false", "",
             "compliant if count(violations) == 0", "",
             f"controls := {json.dumps(sorted(ids, key=sort_key))}", ""]
    lines += HELPERS.strip("\n").split("\n") + [""]
    if name in EXTRA:
        lines += EXTRA[name].strip("\n").split("\n") + [""]
    for rid in ids:
        lines += emit_control(rid, recs[rid], L[rid])
    return "\n".join(lines).rstrip("\n") + "\n"


def emit_tests(name, ids, recs):
    lines = [f"package cis_suse16.{name}_test", "", "import rego.v1", "",
             "# Generated by tools/cis_generators/gen_suse_16.py: a pass and a fail case for every recommendation,",
             "# plus an evidence-missing case for every Automated one. Assertions key on '<id>: ' (colon included).", ""]
    mod = f"data.cis_suse16.{name}.violations"
    for rid in ids:
        s, k = L[rid], idkey(rid)
        pre = json.dumps(rid + ": ")
        if s.manual:
            tp, tf = {"attest": {s.name: True}}, {"attest": {s.name: False}}
            miss = None
        else:
            tp, tf = s.tp, s.tf
            miss = J(f"{rid}: evidence missing — {s.field} not collected")
        lines += [f"test_{k}_pass if {{", f"\tvs := {mod} with input as {J(tp)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) == 0", "}", ""]
        lines += [f"test_{k}_fail if {{", f"\tvs := {mod} with input as {J(tf)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0"]
        if miss:
            lines.append(f"\tnot {miss} in vs")
        lines += ["}", ""]
        if miss:
            lines += [f"test_{k}_evidence_missing if {{", f"\tvs := {mod} with input as {J(s.tm)}",
                      f"\t{miss} in vs", "}", ""]
        if miss and not s.hollow_ok and hollow(s.tp) != s.tp:  # present-but-hollow facts (blank strings, empty lists) never pass
            lines += [f"test_{k}_hollow_evidence_not_a_pass if {{", f"\tvs := {mod} with input as {J(hollow(s.tp))}",
                      f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0", "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


PROFILES = {"level1_server": ("Server", 1), "level2_server": ("Server", 2),
            "level1_workstation": ("Workstation", 1), "level2_workstation": ("Workstation", 2)}


def emit_rollup(recs, by_mod):
    controls = {}
    for rid in sorted(recs, key=sort_key):
        r = recs[rid]
        lv = {}
        for kind in ("Server", "Workstation"):
            x = level(r["profiles"], kind)
            if x is not None:
                lv[kind.lower()] = x
        controls[rid] = {"title": r["title"], "assessment": r["assessment"], "profiles": r["profiles"],
                         "levels": lv, "section": section_of(rid)}
    sec_inputs = {}
    for name, _, _ in SECTIONS:
        keys = sorted({L[rid].reads[i][0] for rid in by_mod[name] for i in range(len(L[rid].reads))} - {"attest"})
        sec_inputs[name] = keys if keys else ["attest"]
    imports = "\n".join(f"import data.cis_suse16.{n}" for n, _, _ in SECTIONS)
    modv = ",\n".join(f'\t"{n}": {n}.violations' for n, _, _ in SECTIONS)
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    return header(f"{BENCH} — roll-up", len(recs), density="orchestrator", basis="role") + f'''package cis_suse16

# {BENCH}: roll-up of the {len(SECTIONS)} section packages.
# Generated by tools/cis_generators/gen_suse_16.py from the CIS workbook; do not edit by hand.
# {len(recs)} recommendations ({a} Automated, {len(recs) - a} Manual), every one in the benchmark. Ids, titles,
# Automated/Manual status and profiles are the workbook's. Levels differ by platform (e.g. 1.1.1.2 is
# Level 1 Server but Level 2 Workstation), so each control carries its own server/workstation level, and five
# recommendations (1.3.1.7, 2.1.11, 2.1.20, 3.1.2, 4.1.4) have no Workstation profile at all.
#
# PROFILE: input.profile selects one of level1_server (default), level2_server, level1_workstation,
# level2_workstation. A Level 2 profile includes the Level 1 recommendations of the same platform.
#
# FAIL CLOSED: every Automated recommendation emits its own evidence-missing violation when its facts were
# not collected, and every Manual one fails until attested, so a control is counted as passed only when
# real facts satisfy it. sections_not_assessed is informational (no input key a section reads is present).

import rego.v1

{imports}

standard_total := {len(recs)}

controls := {json.dumps(controls, indent=1, ensure_ascii=False)}

profiles := {json.dumps({k: {"platform": v[0].lower(), "level": v[1]} for k, v in PROFILES.items()})}

default profile := "level1_server"

profile := input.profile if input.profile in object.keys(profiles)

applies(c) if c.levels[profiles[profile].platform] <= profiles[profile].level

applicable_ids := {{id | some id, c in controls; applies(c)}}

module_violations := {{
{modv},
}}

all_violations := sort([v |
	some _, vs in module_violations
	some v in vs
	split(v, ":")[0] in applicable_ids
])

failing_ids := {{split(v, ":")[0] | some v in all_violations}}

# =============================================================================
# ASSESSMENT COVERAGE — which sections received any of the input keys they read
# =============================================================================

section_inputs := {json.dumps(sec_inputs, indent=1)}

section_assessed(name) if {{
	some k in section_inputs[name]
	input[k]
}}

missing_inputs := {{name: [k | some k in section_inputs[name]; not input[k]] | some name, _ in section_inputs}}

not_assessed_sections := sort([name | some name, _ in section_inputs; not section_assessed(name)])

assessment_coverage_percentage := round(((count(section_inputs) - count(not_assessed_sections)) * 100) / count(section_inputs))

# =============================================================================
# SECTIONS (controls applicable to the selected profile)
# =============================================================================

section_ids(name) := {{id | some id in applicable_ids; controls[id].section == name}}

section_compliance := {{name: {{
	"compliant": count(section_ids(name) & failing_ids) == 0,
	"controls_covered": count(section_ids(name)),
	"violation_count": count([v | some v in all_violations; split(v, ":")[0] in section_ids(name)]),
}} |
	some name, _ in section_inputs
	count(section_ids(name)) > 0
}}

compliant_sections := count([name | some name, s in section_compliance; s.compliant == true; section_assessed(name)])

default compliance_percentage := 0

compliance_percentage := (compliant_sections * 100) / count(section_compliance) if count(section_compliance) > 0

default compliant := false

compliant if {{
	count(object.keys(input)) > 0
	count(applicable_ids) > 0
	count(failing_ids) == 0
}}
'''


def emit_main(recs):
    return header("Bridge: expose cis_suse16 under the /v1/data/<framework>/main/compliance_report", 0,
                  entry="compliance_report", density="orchestrator", basis="role") + f'''package cis_suse16.main

# Bridge: expose cis_suse16 under the /v1/data/<framework>/main/compliance_report convention, with the
# same keys as cis_rhel9.main (plus profile and standard_total). Every field is sourced through a DEFAULTED
# helper: an undefined field inside an object literal would collapse the whole report to {{}}.
#
# total_controls is the number of recommendations applicable to the selected profile; standard_total is
# the benchmark's {len(recs)}. A control passes only when real facts (or an attestation) satisfy it, so
# failed_controls counts distinct failing recommendation ids, including those whose evidence is missing.

import data.cis_suse16 as bench
import rego.v1

default _compliant := false

_compliant := bench.compliant

default _facts_supplied := false

_facts_supplied if count(object.keys(input)) > 0

_no_facts_msg := "FAIL-CLOSED: no facts supplied for cis_suse16 — the assessment could not be evaluated. This is NOT a passing result; check that fact collection ran and produced input."

default _bench_violations := []

_bench_violations := bench.all_violations

default _not_assessed := []

_not_assessed := bench.not_assessed_sections

default _missing_inputs := {{}}

_missing_inputs := bench.missing_inputs

_not_assessed_msgs := [sprintf("NOT ASSESSED: %s — required input not collected: %s", [name, concat(", ", _missing_inputs[name])]) | some name in _not_assessed]

default _sections := {{}}

_sections := bench.section_compliance

default _section_percentage := 0

_section_percentage := bench.compliance_percentage

default _coverage := 0

_coverage := bench.assessment_coverage_percentage

default _profile := "level1_server"

_profile := bench.profile

default _total_controls := 0

_total_controls := count(bench.applicable_ids)

default _failing := 0

_failing := count(bench.failing_ids)

_violations := array.concat(_bench_violations, [_no_facts_msg]) if not _facts_supplied

_violations := array.concat(_bench_violations, _not_assessed_msgs) if _facts_supplied

# No facts => every control is unverified, which counts as failed, not passed.
_failed := _total_controls if not _facts_supplied

_failed := min([_total_controls, _failing]) if _facts_supplied

_passed := max([0, _total_controls - _failed])

default _percentage := 0

_percentage := round((_passed * 100000) / _total_controls) / 1000 if _total_controls > 0

compliance_report := {{
	"framework": "cis_suse16",
	"benchmark": "{BENCH}",
	"version": "v1.0.0",
	"profile": _profile,
	"total_controls": _total_controls,
	"standard_total": {len(recs)},
	"controls_basis": "CIS recommendations applicable to the selected profile (every one evaluated)",
	"passed_controls": _passed,
	"failed_controls": _failed,
	"compliance_percentage": _percentage,
	"section_compliance_percentage": _section_percentage,
	"compliant": _compliant,
	"violations": _violations,
	"section_results": _sections,
	"sections_not_assessed": _not_assessed,
	"missing_inputs": _missing_inputs,
	"assessment_coverage_percentage": _coverage,
}}
'''


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    recs, _ = read_workbook(sys.argv[1])
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    m = sum(1 for r in recs.values() if r["assessment"] == "Manual")
    if (len(recs), a, m) != (EXPECTED_TOTAL, EXPECTED_AUTOMATED, EXPECTED_MANUAL):
        raise SystemExit(f"workbook has {len(recs)} recommendations ({a} Automated, {m} Manual)")
    if set(L) != set(recs):
        raise SystemExit(f"logic/workbook mismatch: no logic {sorted(set(recs) - set(L), key=sort_key)}; "
                         f"no recommendation {sorted(set(L) - set(recs), key=sort_key)}")
    for rid, r in recs.items():
        if (r["assessment"] == "Manual") != L[rid].manual:
            raise SystemExit(f"{rid}: workbook says {r['assessment']}, logic table disagrees")
    by_mod = {n: [] for n, _, _ in SECTIONS}
    for rid in sorted(recs, key=sort_key):
        by_mod[section_of(rid)].append(rid)
    (OUT / "tests").mkdir(parents=True, exist_ok=True)
    for name, title, _ in SECTIONS:
        (OUT / f"{name}.rego").write_text(emit_module(name, title, by_mod[name], recs))
        (OUT / "tests" / f"test_{name}.rego").write_text(emit_tests(name, by_mod[name], recs))
    (OUT / "cis_suse_16.rego").write_text(emit_rollup(recs, by_mod))
    (OUT / "cis_suse16_main.rego").write_text(emit_main(recs))
    (OUT / "tests" / "test_cis_suse16_main.rego").write_text(emit_rollup_tests(recs))
    if len(sys.argv) > 2:  # input paths each control reads (for collector/mapper gap analysis)
        pathlib.Path(sys.argv[2]).write_text(json.dumps(
            {rid: [".".join(p) for p in L[rid].reads] for rid in sorted(recs, key=sort_key)}, indent=1))
    print(f"{len(recs)} recommendations ({a} Automated, {m} Manual) -> {OUT}")



def emit_rollup_tests(recs):
    counts = {}
    for p, (kind, lv) in PROFILES.items():
        counts[p] = sum(1 for r in recs.values() if (level(r["profiles"], kind) or 99) <= lv)
    keys = ["framework", "benchmark", "version", "profile", "total_controls", "standard_total", "controls_basis",
            "passed_controls", "failed_controls", "compliance_percentage", "section_compliance_percentage", "compliant",
            "violations", "section_results", "sections_not_assessed", "missing_inputs", "assessment_coverage_percentage"]
    lines = ["package cis_suse16.main_test", "", "import rego.v1", "",
             "# Generated by tools/cis_generators/gen_suse_16.py: roll-up and bridge smoke tests.", "",
             "report(i) := r if {", "\tr := data.cis_suse16.main.compliance_report with input as i", "}", "",
             "test_every_recommendation_in_table if {",
             f"\tcount(data.cis_suse16.controls) == {len(recs)}",
             f"\tcount([c | some c in data.cis_suse16.controls; c.assessment == \"Manual\"]) == {EXPECTED_MANUAL}",
             "}", "",
             "test_empty_input_fails_closed if {", "\tr := report({})",
             f"\tcount(object.keys(r)) == {len(keys)}",
             f"\tevery k in {json.dumps(keys)} {{ k in object.keys(r) }}",
             "\tr.compliant == false", f"\tr.total_controls == {counts['level1_server']}",
             "\tr.failed_controls == r.total_controls", "\tr.passed_controls == 0",
             "\tsome v in r.violations", "\tstartswith(v, \"FAIL-CLOSED\")", "}", ""]
    for p, n in counts.items():
        lines += [f"test_profile_{p}_count if {{", f"\tr := report({{\"profile\": \"{p}\"}})",
                  f"\tr.profile == \"{p}\"", f"\tr.total_controls == {n}", f"\tr.standard_total == {len(recs)}",
                  "\tr.failed_controls == r.total_controls", "}", ""]
    lines += ["test_unknown_profile_defaults_to_level1_server if {", "\tr := report({\"profile\": \"bogus\"})",
              "\tr.profile == \"level1_server\"", f"\tr.total_controls == {counts['level1_server']}", "}", "",
              "test_one_attested_control_passes if {",
              "\tr := report({\"attest\": {\"gpg_keys_reviewed\": true}})",
              "\tr.passed_controls == 1", "\tr.failed_controls == r.total_controls - 1", "\tr.compliant == false", "}", "",
              "test_unassessed_sections_listed if {", "\tr := report({\"ssh\": {\"sshd_config_raw\": \"permitrootlogin no\"}})",
              "\tnot \"ssh\" in r.sections_not_assessed", "\t\"auditd\" in r.sections_not_assessed",
              "\tr.assessment_coverage_percentage < 100", "}", "",
              "test_level2_violation_hidden_in_level1_profile if {",
              "\tvs := data.cis_suse16.all_violations with input as {\"profile\": \"level1_server\"}",
              "\tcount([v | some v in vs; startswith(v, \"6.2.1.1: \")]) == 0",
              "\tvs2 := data.cis_suse16.all_violations with input as {\"profile\": \"level2_server\"}",
              "\tcount([v | some v in vs2; startswith(v, \"6.2.1.1: \")]) > 0", "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


if __name__ == "__main__":
    main()
