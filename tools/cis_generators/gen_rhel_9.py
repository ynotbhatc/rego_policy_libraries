#!/usr/bin/env python3
"""CIS Red Hat Enterprise Linux 9 Benchmark v3.0.0: every recommendation, from the benchmark's own workbook.

    python3 tools/cis_generators/gen_rhel_9.py CIS_Red_Hat_Enterprise_Linux_9_Benchmark_v3.0.0.xlsx [fields.json]

Needs openpyxl. Reads id, title, Assessment Status and profiles (Level 1/2 - Server/Workstation sheets)
out of the CIS workbook (CIS-licensed, never committed), checks the ids against the hand-authored logic
table below and emits benchmarks/cis/rhel_9/. The run aborts on any count, Automated/Manual status, logic
or title-keyword mismatch with the workbook: put(id, keyword, spec) requires the keyword to occur in the
workbook title of that id, so a logic entry attached to a renumbered id cannot slip through.

Adapted from gen_suse_16.py and gen_ubuntu_26_04.py, which share this layout. Unlike them it writes into the
existing rhel_9 package paths, because callers query them: the fourteen section packages keep their names
and files (cis_rhel9.filesystem in filesystem_validation.rego, ...), the roll-up stays package cis_rhel9
(cis_rhel9_complete.rego, with section_inputs / NOT ASSESSED and every rule the v2.0.0 roll-up exposed),
the bridge stays cis_rhel9.main (compliance_report), and cis_rhel9.l2 becomes the Level 2 view of the
same controls. The input contract is the one the shared Linux mapper
(ansible/tasks/cis_linux_build_opa_input.yml) emits: wherever a v3.0.0 control checks the same thing as a
v2.0.0 control, the v2.0.0 field name and shape are reused. Fields the mapper does not emit yet are listed
per control in the optional fields.json output.

The STIG/NIST extended-hardening modules in the same directory (authorized_keys, certificate_validation,
storage_encryption) and the deprecated legacy package cis.rhel_9 are not CIS v3.0.0 recommendations and
are not generated.

Rules of the generated code:
  * every Automated recommendation evaluates real facts; facts not collected =>
    "<id>: evidence missing — <field> not collected" (fail closed, never a pass);
  * every Manual recommendation is an attestation boolean input.attest.<name> (not attested => violation);
  * messages are our words; only the id and the short CIS title are carried from the workbook;
  * the tests give every recommendation a pass and a fail case, plus an evidence-missing case for
    every Automated one.
"""
import json
import pathlib
import re
import sys

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE.parent.parent / "benchmarks" / "cis" / "rhel_9"
PKG = "cis_rhel9"
GEN = "tools/cis_generators/gen_rhel_9.py"
BENCH = "CIS Red Hat Enterprise Linux 9 Benchmark v3.0.0"
VERSION = "v3.0.0"
EXPECTED_TOTAL = 352
EXPECTED_AUTOMATED = 329
EXPECTED_MANUAL = 23
AS_OF = "v2.4.2"
SKIP_SHEETS = {"License", "Overview", "Combined Profiles"}

# module name, file, section key (the v2.0.0 keys are kept: callers and tests read them), title, id prefixes
SECTIONS = [
    ("filesystem", "filesystem_validation.rego", "section_1_1_filesystem",
     "Section 1.1: Filesystem", ["1.1."]),
    ("initial_setup", "initial_setup_validation.rego", "section_1_2_1_7_1_8_initial_setup",
     "Sections 1.2, 1.6-1.8: Package Management, Crypto Policy, Banners, GNOME Display Manager",
     ["1.2.", "1.6.", "1.7.", "1.8."]),
    ("boot_security", "boot_security_validation.rego", "section_1_3_1_4_1_5_boot_security",
     "Sections 1.4, 1.5, 6.3: Bootloader, Process Hardening, Integrity Checking (AIDE)", ["1.4.", "1.5.", "6.3."]),
    ("selinux", "selinux_validation.rego", "section_1_6_selinux",
     "Section 1.3: Mandatory Access Control (SELinux)", ["1.3."]),
    ("services", "service_validation.rego", "section_2_services",
     "Sections 2.1-2.3: Server Services, Client Services, Time Synchronization", ["2.1.", "2.2.", "2.3."]),
    ("network", "network_validation.rego", "section_3_network",
     "Sections 3-4: Network, Host Based Firewall (firewalld)", ["3.", "4."]),
    ("auditd", "auditd_validation.rego", "section_4_1_auditd",
     "Section 6.2: System Auditing (auditd)", ["6.2."]),
    ("logging", "logging_validation.rego", "section_4_2_logging",
     "Section 6.1: System Logging (journald, rsyslog)", ["6.1."]),
    ("cron", "cron_validation.rego", "section_5_1_cron",
     "Section 2.4: Job Schedulers", ["2.4."]),
    ("ssh", "ssh_validation.rego", "section_5_2_ssh",
     "Section 5.1: SSH Server", ["5.1."]),
    ("sudo", "sudo_validation.rego", "section_5_3_sudo",
     "Section 5.2: Privilege Escalation", ["5.2."]),
    ("pam", "pam_validation.rego", "section_5_4_5_5_pam",
     "Sections 5.3-5.4: Pluggable Authentication Modules, User Accounts and Environment", ["5.3.", "5.4."]),
    ("file_permissions", "file_permissions_validation.rego", "section_6_1_file_permissions",
     "Section 7.1: System File Permissions", ["7.1."]),
    ("user_group", "user_group_validation.rego", "section_6_2_user_group",
     "Section 7.2: Local User and Group Settings", ["7.2."]),
]

# NOT ASSESSED: a section is assessed when any of these input keys is present. The v2.0.0 keys are kept and
# the top-level keys new v3.0.0 controls read are added. ansible_facts, services, packages and attest are
# deliberately absent (read by many sections; an ansible_facts-only payload is not an assessment).
SECTION_INPUTS = {
    "filesystem": ["mount_options", "disabled_filesystem_modules", "separate_partitions", "usb_storage", "sticky_bit"],
    "initial_setup": ["initial_setup", "crypto_policy"],
    "boot_security": ["boot_security"],
    "selinux": ["selinux"],
    "services": ["unnecessary_services", "time_synchronization", "chronyd", "client_packages", "autofs",
                 "mta", "gui", "x11", "webserver"],
    "network": ["sysctl_parameters", "ip_forwarding", "ipv6", "firewall", "wireless_interfaces",
                "disabled_network_modules"],
    "auditd": ["auditd"],
    "logging": ["rsyslog", "journald", "log_file_permissions"],
    "cron": ["cron"],
    "ssh": ["ssh"],
    "sudo": ["sudo"],
    "pam": ["pam"],
    "file_permissions": ["suid_files", "critical_files", "world_writable", "unowned_files", "ungrouped_files"],
    "user_group": ["user_group"],
}

# v2.0.0 recommendations (generate_recommendations) kept per section: priority and caller-facing section name
RECOMMEND = {
    "ssh": ("critical", "SSH Server Configuration", "Harden SSH configuration per CIS guidelines"),
    "auditd": ("critical", "System Auditing", "Configure auditd service and implement required audit rules"),
    "selinux": ("critical", "SELinux Configuration", "Enable and configure SELinux in enforcing mode"),
    "pam": ("high", "Password Authentication", "Configure password policies and PAM modules"),
    "sudo": ("high", "Privilege Escalation", "Harden sudo configuration and restrict su access"),
    "logging": ("high", "Logging Configuration", "Configure rsyslog and journald properly"),
    "boot_security": ("high", "Boot Security & AIDE", "Set bootloader password, enable AIDE, configure process hardening"),
    "cron": ("medium", "Cron Configuration", "Set correct permissions on cron files and restrict access"),
    "user_group": ("medium", "User and Group Settings", "Fix user account issues, duplicate UIDs/GIDs, and home directory permissions"),
    "initial_setup": ("medium", "Initial Setup & Banners", "Configure warning banners, GPG keys, crypto policy and GNOME settings"),
}


def section_of(rid):
    for name, _, _, _, prefixes in SECTIONS:
        if any(rid.startswith(p) for p in prefixes):
            return name
    raise SystemExit(f"no section for {rid}")


# ── input paths ────────────────────────────────────────────────────────────────
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def ref(path):
    """['a', 'b-c'] -> input.a["b-c"]"""
    out = "input"
    for p in path:
        out += f".{p}" if IDENT.match(p) else f"[{json.dumps(p)}]"
    return out


def dotted(path):
    return ".".join(path)


def nest(path, value):
    d = value
    for p in reversed(path):
        d = {p: d}
    return d


def merge(*ds):
    out = {}
    for d in ds:
        for k, v in d.items():
            if k in out and isinstance(out[k], dict) and isinstance(v, dict):
                out[k] = merge(out[k], v)
            else:
                out[k] = v
    return out


def P(s):
    """'a.b.c' -> ['a','b','c'] (keys containing dots are passed as lists by callers)"""
    return s.split(".")


# ── a control's logic ──────────────────────────────────────────────────────────
class Spec:
    """ev: rego lines (AND) proving the facts were collected (alternatives: list of lists in ev_alts);
    passes: list of alternatives, each a list of rego lines (AND); field: evidence-missing message field;
    tp/tf/tm: test inputs that pass / fail / lack evidence; reads: input paths read (for the report)."""

    def __init__(self, ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
        self.ev_alts = ev_alts if ev_alts is not None else [ev]
        self.passes = passes
        self.field = field
        self.tp, self.tf = tp, tf
        self.tm = {} if tm is None else tm
        self.reads = reads or []
        self.manual = False
        # True when an all-empty input (strings "", lists []) is a genuine pass: offender lists, where an
        # empty list collected by the collector means "none found". Otherwise a hollow-input test is generated.
        self.hollow_ok = False


class Attest:
    def __init__(self, name):
        self.name = name
        self.manual = True
        self.reads = [["attest", name]]


# ── builders ───────────────────────────────────────────────────────────────────
def kmod(listkey, name):
    lst = ref([listkey, "modules"])
    ev = [f'some e in {lst}', f'e.module == {json.dumps(name)}', 'kmod_evidence(e)']
    passes = [[f'count([e | some e in {lst}; e.module == {json.dumps(name)}]) > 0',
               f'every e in [e | some e in {lst}; e.module == {json.dumps(name)}] {{ kmod_unavailable(e) }}']]
    ok = {"module": name, "disabled": True, "loaded": False}
    return Spec(ev, passes, f"{listkey}.modules[{name}]",
                nest([listkey, "modules"], [ok]),
                nest([listkey, "modules"], [{"module": name, "disabled": False, "loaded": False}]),
                nest([listkey, "modules"], [{"module": "other", "disabled": True, "loaded": False}]),
                reads=[[listkey, "modules"]])


MOUNTS = ["ansible_facts", "ansible_mounts"]


def mnt_sep(path, extra_tmp=False):
    passes = [[f'mounted({json.dumps(path)})']]
    if extra_tmp:
        passes = [[f'mounted({json.dumps(path)})', 'not tmp_mount_unit_off']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "tmpfs", "options": "rw,nodev,nosuid,noexec"}]),
                nest(MOUNTS, [other]), {}, reads=[MOUNTS] + ([["services", "tmp.mount"]] if extra_tmp else []))


def mnt_opt(path, opt):
    # The option applies to the partition when it exists; the mount table is the collected fact
    # (the separate-partition recommendation fails on its own when the partition is missing).
    passes = [[f'not mounted({json.dumps(path)})'],
              [f'every m in mounts_at({json.dumps(path)}) {{ {json.dumps(opt)} in mount_opts(m) }}']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    good = "rw,nodev,nosuid,noexec"
    bad = ",".join(o for o in ["rw", "nodev", "nosuid", "noexec"] if o != opt)
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": good}]),
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": bad}]),
                {}, reads=[MOUNTS])



SP = "sysctl_parameters"


def sysctl(param, allowed, ipv6=False, alias=None):
    """Running value: sysctl_parameters.analysis [{parameter, current_value}] (the rhel_9 mapper shape).
    Persisted value: sysctl_parameters.persisted[param] (the value systemd-sysctl applies at boot, last file
    wins). alias: (path, rego condition) of a v2.0.0 field that already carries the running value; it is
    used only when the analysis list has no entry for the parameter."""
    vals, p = json.dumps(allowed), json.dumps(param)
    per = ref([SP, "persisted", param])
    run_ok = f'every v in sysctl_running({p}) {{ v in {vals} }}'
    per_ok = f'sprintf("%v", [{per}]) in {vals}'
    ev_alts = [[f'count(sysctl_running({p})) > 0', f'{per} != null']]
    passes = [[f'count(sysctl_running({p})) > 0', run_ok, per_ok]]
    reads = [[SP, "analysis"], [SP, "persisted", param]]
    if alias:
        apath, acond = alias
        ev_alts.append([f'{ref(apath)} != null', f'{per} != null'])
        passes.append([f'count(sysctl_running({p})) == 0', acond, per_ok])
        reads.append(apath)
    run = [{"parameter": param, "current_value": allowed[0]}]
    good = {SP: {"analysis": run, "persisted": {param: allowed[0]}}}
    tf = {SP: {"analysis": run, "persisted": {param: "99"}}}
    tm = {SP: {"analysis": run}}  # running value only: the persisted setting was not collected
    if ipv6:  # "IF IPv6 is enabled": a host with IPv6 disabled passes on that collected fact
        ev_alts.append(['input.ipv6.disabled == true'])
        passes.append(['input.ipv6.disabled == true'])
        reads.append(["ipv6", "disabled"])
    return Spec(None, passes, f"sysctl_parameters.analysis/persisted[{param}]", good, tf, tm, reads=reads,
                ev_alts=ev_alts)


# ── /tmp and /dev/shm: the v2.0.0 mount_options.<tmp|dev_shm> shape the mapper emits ──
def mo_sep(key, tmp_unit=False):
    e = ref(["mount_options", key, "exists"])
    passes = [[f'{e} == true'] + (['not tmp_mount_unit_off'] if tmp_unit else [])]
    ok = {"exists": True, "has_nodev": True, "has_nosuid": True, "has_noexec": True}
    return Spec([f'is_boolean({e})'], passes, f"mount_options.{key}.exists",
                {"mount_options": {key: ok}}, {"mount_options": {key: {"exists": False}}}, {"mount_options": {key: {}}},
                reads=[["mount_options", key, "exists"]] + ([["services", "tmp.mount"]] if tmp_unit else []))


def mo_opt(key, opt):
    """the option applies to the partition when it exists (the separate-partition recommendation fails on
    its own when it does not)"""
    e, h = ref(["mount_options", key, "exists"]), ref(["mount_options", key, f"has_{opt}"])
    ok = {"exists": True, "has_nodev": True, "has_nosuid": True, "has_noexec": True}
    return Spec(None, [[f'{e} == false'], [f'{e} == true', f'{h} == true']], f"mount_options.{key}.has_{opt}",
                {"mount_options": {key: ok}}, {"mount_options": {key: dict(ok, **{f"has_{opt}": False})}},
                {"mount_options": {key: {"exists": True}}},
                reads=[["mount_options", key, "exists"], ["mount_options", key, f"has_{opt}"]],
                ev_alts=[[f'{e} == false'], [f'{e} == true', f'is_boolean({h})']])


def alt(spec, ev, passes, reads, field_note):
    """Add a v2.0.0 field as an alternative evidence source. It decides the control only when the primary
    evidence was not collected (emitted as `not ep_<id>`), so two sources can never disagree into a pass."""
    spec.alts = getattr(spec, "alts", []) + [(ev, passes)]
    spec.reads = spec.reads + reads
    spec.field = f"{spec.field} (or {field_note})"
    return spec




def pkg_present(pkg, path=None):
    p = path or ["packages", pkg]
    return Spec([f'is_boolean({ref(p)})'], [[f'{ref(p)} == true']], dotted(p),
                nest(p, True), nest(p, False), {}, reads=[p])


def pkg_absent(pkgs, path=None):
    if path:  # a single rhel_9-style boolean
        return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == false']], dotted(path),
                    nest(path, False), nest(path, True), {}, reads=[path])
    ev = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    passes = [[f'{ref(["packages", p])} == false' for p in pkgs]]
    tp = {"packages": {p: False for p in pkgs}}
    tf = {"packages": {p: (i == 0) for i, p in enumerate(pkgs)}}
    return Spec(ev, passes, "packages[" + ",".join(pkgs) + "]", tp, tf, {"packages": {"audit": True}},
                reads=[["packages", p] for p in pkgs])


def not_in_use(units_by_pkg):
    """CIS 'not in use': every package not installed, OR (installed for a dependency) each unit the audit
    names for it neither enabled nor active. Ansible's service_facts lists .service units only, so a
    unit missing from input.services proves nothing: for an installed package every named unit must be
    reported (LoadState not-found is fine), or the evidence is missing."""
    pkgs = list(units_by_pkg)
    pk = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    none_installed = [f'{ref(["packages", p])} == false' for p in pkgs]
    m = json.dumps(units_by_pkg)
    reported = f'every p, us in {m} {{ unit_set_reported(p, us) }}'
    ev_alts = [pk + none_installed, pk + ['services_collected', reported]]
    passes = [none_installed,
              ['services_collected', reported, f'every p, us in {m} {{ unit_set_off(p, us) }}']]
    tp = {"packages": {p: False for p in pkgs}}
    allu = {u: {"state": "inactive", "status": "disabled"} for us in units_by_pkg.values() for u in us}
    first = units_by_pkg[pkgs[0]][0]
    tf = {"packages": {p: True for p in pkgs},
          "services": dict(allu, **{first: {"state": "running", "status": "enabled"}},
                           **{"sshd.service": {"state": "running", "status": "enabled"}})}
    tm = {"packages": {p: True for p in pkgs},
          "services": {"sshd.service": {"state": "running", "status": "enabled"}}}
    return Spec(None, passes, "packages[" + ",".join(pkgs) + "] / services[" + ",".join(
        u for us in units_by_pkg.values() for u in us) + "]", tp, tf, tm,
                reads=[["packages", p] for p in pkgs] + [["services"]], ev_alts=ev_alts)


def units_off(units):
    """The named units must each be reported (not-found is fine) and be neither enabled nor active."""
    passes = [[f'every u in {json.dumps(units)} {{ not unit_on(u) }}']]
    other = {"sshd.service": {"state": "running", "status": "enabled"}}
    nf = {u: {"state": "inactive", "status": "not-found"} for u in units}
    return Spec(['services_collected', f'every u in {json.dumps(units)} {{ unit_reported(u) }}'], passes,
                "services[" + ",".join(units) + "]",
                {"services": merge(other, nf)},
                {"services": merge(other, nf, {units[0]: {"state": "running", "status": "enabled"}})},
                {"services": other}, reads=[["services"]])


def unit_running(unit):
    passes = [[f'unit_enabled_active({json.dumps(unit)})']]
    other = {"sshd.service": {"state": "running", "status": "enabled"}}
    return Spec(['services_collected'], passes, "services",
                {"services": merge(other, {unit: {"state": "running", "status": "enabled"}})},
                {"services": merge(other, {unit: {"state": "stopped", "status": "disabled"}})},
                {}, reads=[["services"]])


def boolean(path, want):
    return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == {json.dumps(want)}']], dotted(path),
                nest(path, want), nest(path, not want), {}, reads=[path])


def access(mode, owner, group, mx, owners, groups, exists=None, optional=False):
    """One file: mode within mx, owner in owners, group in groups. exists: path of a boolean; when
    optional, a collected 'does not exist' passes; when required, it fails."""
    ev = [f'{ref(mode)} != null', f'{ref(owner)} != null', f'{ref(group)} != null']
    ok = [f'mode_within({ref(mode)}, {json.dumps(mx)})', f'{ref(owner)} in {json.dumps(owners)}',
          f'{ref(group)} in {json.dumps(groups)}']
    tp = merge(nest(mode, mx), nest(owner, owners[0]), nest(group, groups[0]))
    tf = merge(nest(mode, "0777"), nest(owner, owners[0]), nest(group, groups[0]))
    ev_alts, passes, reads = [ev], [ok], [mode, owner, group]
    if exists:
        reads.append(exists)
        ev_alts = [[f'{ref(exists)} == true'] + ev, [f'{ref(exists)} == false']]
        passes = [[f'{ref(exists)} == true'] + ok]
        tp = merge(tp, nest(exists, True))
        tf = merge(tf, nest(exists, True))
        if optional:
            passes.append([f'{ref(exists)} == false'])
        else:
            tf = nest(exists, False)
    return Spec(None, passes, "/".join(dotted(p) for p in reads), tp, tf, {}, reads=reads, ev_alts=ev_alts)


def access_list(path, mx, owners, groups, mode="mode", owner="owner", group="group", sample="/x", nonempty=True):
    """An inventory of files that always exist on a host must be non-empty to count as evidence."""
    lst = ref(path)
    passes = [[f'every f in {lst} {{ list_item_ok(f, {json.dumps(mode)}, {json.dumps(owner)}, '
               f'{json.dumps(group)}, {json.dumps(mx)}, {json.dumps(owners)}, {json.dumps(groups)}) }}']]
    good = {"path": sample, mode: mx, owner: owners[0], group: groups[0]}
    bad = dict(good, **{mode: "0777"})
    ev = [f'is_array({lst})'] + ([f'count({lst}) > 0'] if nonempty else [])
    return Spec(ev, passes, dotted(path), nest(path, [good]), nest(path, [good, bad]), {}, reads=[path])


def crit(path, mx, groups, optional=False):
    """critical_files.analysis entry {file, exists, actual_mode, actual_owner, actual_group} (rhel_9 shape);
    the expected_* / *_compliant fields the rhel_9 collector pre-computes are not trusted here."""
    p = json.dumps(path)
    ent = f'[f | some f in object.get(input, ["critical_files", "analysis"], []); f.file == {p}]'
    ok = (f'every f in {ent} {{ f.exists == true; mode_within(f.actual_mode, {json.dumps(mx)}); '
          f'f.actual_owner == "root"; f.actual_group in {json.dumps(groups)} }}')
    ev = [f'count({ent}) > 0', f'every f in {ent} {{ is_boolean(f.exists) }}']
    passes = [[f'count({ent}) > 0', ok]]
    if optional:
        passes.append([f'count({ent}) > 0', f'every f in {ent} {{ f.exists == false }}'])
    good = {"file": path, "exists": True, "actual_mode": mx, "actual_owner": "root", "actual_group": groups[0]}
    bad = dict(good, actual_mode="0666")
    other = {"file": "/etc/other", "exists": True, "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"}
    return Spec(ev, passes, f"critical_files.analysis[{path}]",
                {"critical_files": {"analysis": [other, good]}},
                {"critical_files": {"analysis": [other, bad]}},
                {"critical_files": {"analysis": [other]}}, reads=[["critical_files", "analysis"]])


def attest(name):
    return Attest(name)


def custom(ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
    return Spec(ev, passes, field, tp, tf, tm, reads=reads, ev_alts=ev_alts)


# ── sshd (effective config; `sshd -T` output preferred, raw sshd_config accepted) ──────────────
SSH_RAW = ["ssh", "sshd_config_raw"]


def ssh_in(lines):
    return {"ssh": {"sshd_config_raw": "\n".join(lines)}}


def sshd(key, cond, good, bad, explicit=False):
    """cond: rego expression over `v` (the effective lowercase value). explicit: the directive must be
    present (no reliable compiled default), otherwise its absence is missing evidence."""
    ev = ['ssh_collected', f'sshd_first({json.dumps(key)})']
    passes = [[f'v := lower(sshd_value({json.dumps(key)}))', cond]]
    tm = ssh_in(["port 22"])
    return Spec(ev, passes, f"ssh.sshd_config_raw (effective {key})",
                ssh_in(["port 22", f"{key} {good}"]), ssh_in(["port 22", f"{key} {bad}"]), tm, reads=[SSH_RAW])


def text_has(path, regex, want, good, bad):
    """want=True: some active line of the text matches regex; False: none does."""
    t = ref(path)
    cond = f'regex_in_text({t}, {json.dumps(regex)})'
    return Spec([f'is_string({t})'], [[cond if want else f'not {cond}']], dotted(path),
                nest(path, good), nest(path, bad), {}, reads=[path])


def conf(path, key_re, cond, good, bad, default=None):
    """key = value configuration text; last occurrence wins. cond is an expression over `v`.
    default: the documented compiled-in value used when the key is not set (None: unset fails)."""
    t = ref(path)
    passes = [[f'v := conf_last({t}, {json.dumps(key_re)})', cond]]
    if default is not None:
        passes.append([f'not conf_last({t}, {json.dumps(key_re)})', f'v := {json.dumps(default)}', cond])
    return Spec([f'is_string({t})'], passes, dotted(path), nest(path, good), nest(path, bad), {}, reads=[path])


HELPERS = r'''
# Parse an octal permission string ("644", "0644", "0o644") to an integer; undefined if not octal.
octal(s) := n if {
	t := trim_prefix(lower(trim_space(sprintf("%v", [s]))), "0o")
	regex.match(`^[0-7]{1,4}$`, t)
	p := concat("", [substring("0000", 0, 4 - count(t)), t])
	ds := [to_number(substring(p, i, 1)) | some i in numbers.range(0, 3)]
	n := ((ds[0] * 512) + (ds[1] * 64)) + ((ds[2] * 8) + ds[3])
}

# true when mode grants nothing beyond max (both octal strings)
mode_within(mode, max) if bits.and(octal(mode), bits.and(bits.negate(octal(max)), 4095)) == 0

# Active (non-blank, non-comment) lines of a configuration text.
active_lines(text) := [l |
	some raw in split(text, "\n")
	l := trim_space(raw)
	l != ""
	not startswith(l, "#")
]

# a collected text: a non-blank string ("" is not evidence)
text_collected(t) if {
	is_string(t)
	trim_space(t) != ""
}

regex_in_text(text, re) if {
	some l in active_lines(text)
	regex.match(re, l)
}

# Values of a `key = value` / `key value` directive (key is a case-insensitive regex); last one wins.
conf_values(text, key_re) := [trim(m[1], "\"' ") |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(concat("", ["(?i)^", key_re, `(?:\s*=\s*|\s+)(.*?)\s*$`]), l, 1)
]

conf_last(text, key_re) := vs[count(vs) - 1] if {
	vs := conf_values(text, key_re)
	count(vs) > 0
}

num(v) := to_number(trim_space(sprintf("%v", [v]))) if regex.match(`^\s*-?[0-9]+\s*$`, sprintf("%v", [v]))

# ── packages / services ──
services_collected if {
	is_object(input.services)
	count(input.services) > 0
}

unit_on(u) if input.services[u].status == "enabled"

unit_on(u) if lower(input.services[u].state) in {"running", "active"}

# every unit named for an installed package is reported in the inventory
unit_set_reported(p, _) if input.packages[p] == false

unit_set_reported(p, us) if {
	input.packages[p] == true
	every u in us { unit_reported(u) }
}

unit_reported(u) if {
	is_string(input.services[u].status)
	input.services[u].status != ""
}

unit_set_off(p, _) if input.packages[p] == false

unit_set_off(p, us) if {
	input.packages[p] == true
	every u in us { not unit_on(u) }
}

unit_enabled_active(u) if {
	input.services[u].status == "enabled"
	lower(input.services[u].state) in {"running", "active"}
}

# tmp.mount masked or disabled means systemd will not mount /tmp at boot
tmp_mount_unit_off if input.services["tmp.mount"].status in {"masked", "disabled"}

# ── kernel modules: {module, disabled, loaded, available?}; "unknown" proves nothing ──
kmod_evidence(e) if is_boolean(e.available)

kmod_evidence(e) if {
	is_boolean(e.disabled)
	is_boolean(e.loaded)
}

kmod_unavailable(e) if e.available == false

kmod_unavailable(e) if {
	e.disabled == true
	e.loaded == false
}

# ── mounts ──
mounts_collected if {
	is_array(input.ansible_facts.ansible_mounts)
	count(input.ansible_facts.ansible_mounts) > 0
}

mounts_at(path) := [m | some m in input.ansible_facts.ansible_mounts; m.mount == path]

mounted(path) if count(mounts_at(path)) > 0

mount_opts(m) := split(m.options, ",") if is_string(m.options)

mount_opts(m) := m.options if is_array(m.options)

# ── file lists ──
list_item_ok(f, mk, ok, gk, max, owners, groups) if {
	mode_within(f[mk], max)
	f[ok] in owners
	f[gk] in groups
}

# ── sshd: first occurrence wins (sshd semantics); `sshd -T` prints effective values ──
ssh_collected if {
	is_string(input.ssh.sshd_config_raw)
	trim_space(input.ssh.sshd_config_raw) != ""
}

sshd_pairs := [[lower(m[1]), trim_space(m[2])] |
	some l in active_lines(input.ssh.sshd_config_raw)
	some m in regex.find_all_string_submatch_n(`^([A-Za-z0-9]+)(?:\s*=\s*|\s+)(.*)$`, l, 1)
]

sshd_first(key) := vs[0] if {
	vs := [p[1] | some p in sshd_pairs; p[0] == key]
	count(vs) > 0
}

# effective value: `sshd -T` prints every directive, so a missing directive means the effective
# configuration was not collected (a raw sshd_config misses /usr/etc/ssh and sshd_config.d)
sshd_value(key) := sshd_first(key)

sshd_list(key) := [trim_space(x) | some x in split(lower(sshd_value(key)), ",")]

# seconds from an sshd time value ("60", "60s", "1m")
sshd_seconds(v) := to_number(v) if regex.match(`^[0-9]+$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "s")) if regex.match(`^[0-9]+s$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "m")) * 60 if regex.match(`^[0-9]+m$`, v)

# ── sysctl: running values from the mapper's analysis list (rhel_9 shape) ──
sysctl_running(p) := [sprintf("%v", [e.current_value]) |
	some e in object.get(input, ["sysctl_parameters", "analysis"], [])
	e.parameter == p
	e.current_value != null
]
'''


# ════════════════════════════════════════════════════════════════════════════════════════════
# LOGIC TABLE — one entry per recommendation, written from each audit procedure of the RHEL 9 v3.0.0
# workbook. put(id, keyword, spec): the keyword must occur in the workbook title of that id
# (case-insensitive), so a logic entry attached to a renumbered id aborts the run.
# ════════════════════════════════════════════════════════════════════════════════════════════
L = {}
KW = {}


def put(rid, kw, spec):
    if rid in L:
        raise SystemExit(f"{rid} defined twice")
    L[rid] = spec
    KW[rid] = kw
    return spec


def gate(spec, path, value, note):
    """'IF <x>': a collected fact saying the condition does not hold passes (and counts as evidence)."""
    g = f'{ref(path)} == {json.dumps(value)}'
    spec.ev_alts = spec.ev_alts + [[g]]
    spec.passes = spec.passes + [[g]]
    spec.reads = spec.reads + [path]
    spec.field = f"{spec.field} (or {note})"
    return spec


def pkgs_absent_any(pkgs):
    return pkg_absent(pkgs)


FSM = "disabled_filesystem_modules"   # v2.0.0 mapper field: [{module, disabled, loaded}]
NETM = "disabled_network_modules"     # same shape, network modules (new; SUSE/Ubuntu siblings use it)

# ── 1.1.1 filesystem kernel modules ──
for rid, mod in [("1.1.1.1", "cramfs"), ("1.1.1.2", "freevxfs"), ("1.1.1.3", "hfs"), ("1.1.1.4", "hfsplus"),
                 ("1.1.1.5", "jffs2"), ("1.1.1.6", "overlay"), ("1.1.1.7", "squashfs"), ("1.1.1.8", "udf"),
                 ("1.1.1.9", "firewire-core")]:
    put(rid, f"ensure {mod} kernel module", kmod(FSM, mod))
# usb-storage: the mapper's module probe lists it; the v2.0.0 usb_storage object is accepted when it does not
_usb = kmod(FSM, "usb-storage")
alt(_usb, ['is_boolean(input.usb_storage.disabled)', 'is_boolean(input.usb_storage.loaded)'],
    ['input.usb_storage.disabled == true', 'input.usb_storage.loaded == false'],
    [["usb_storage", "disabled"], ["usb_storage", "loaded"]], "usb_storage.disabled/loaded")
put("1.1.1.10", "usb-storage kernel module", _usb)
put("1.1.1.11", "unused filesystems", attest("unused_filesystem_modules_reviewed"))

# ── 1.1.2 partitions: /tmp and /dev/shm from mount_options (v2.0.0), the rest from ansible_mounts ──
put("1.1.2.1.1", "/tmp is tmpfs", mo_sep("tmp", tmp_unit=True))
for i, o in enumerate(["nodev", "nosuid", "noexec"], start=2):
    put(f"1.1.2.1.{i}", f"{o} option set on /tmp", mo_opt("tmp", o))
put("1.1.2.2.1", "/dev/shm is tmpfs", mo_sep("dev_shm"))
for i, o in enumerate(["nodev", "nosuid", "noexec"], start=2):
    put(f"1.1.2.2.{i}", f"{o} option set on /dev/shm", mo_opt("dev_shm", o))
for base, part, opts in [("1.1.2.3", "/home", ["nodev", "nosuid"]), ("1.1.2.4", "/var", ["nodev", "nosuid"]),
                         ("1.1.2.5", "/var/tmp", ["nodev", "nosuid", "noexec"]),
                         ("1.1.2.6", "/var/log", ["nodev", "nosuid", "noexec"]),
                         ("1.1.2.7", "/var/log/audit", ["nodev", "nosuid", "noexec"])]:
    put(base + ".1", "separate partition exists for " + part + " ", mnt_sep(part))
    for i, o in enumerate(opts, start=2):
        put(f"{base}.{i}", f"{o} option set on {part} partition", mnt_opt(part, o))

# ── 1.2 package management (dnf) ──
IS = "input.initial_setup."
DNF = ["initial_setup", "dnf_conf_raw"]   # new: text of /etc/dnf/dnf.conf
put("1.2.1.1", "GPG keys", attest("gpg_keys_reviewed"))
put("1.2.1.2", "ensure gpgcheck is configured", custom(
    [f'is_boolean({IS}global_gpgcheck_enabled)', f'is_array({IS}repos_without_gpgcheck)', f'is_array({IS}enabled_repos)'],
    [[f'{IS}global_gpgcheck_enabled == true', f'count({IS}repos_without_gpgcheck) == 0']],
    "initial_setup.global_gpgcheck_enabled/repos_without_gpgcheck/enabled_repos",
    {"initial_setup": {"global_gpgcheck_enabled": True, "repos_without_gpgcheck": [], "enabled_repos": ["baseos"]}},
    {"initial_setup": {"global_gpgcheck_enabled": True, "repos_without_gpgcheck": ["epel"], "enabled_repos": ["baseos", "epel"]}},
    reads=[["initial_setup", "global_gpgcheck_enabled"], ["initial_setup", "repos_without_gpgcheck"],
           ["initial_setup", "enabled_repos"]]))
put("1.2.1.3", "repo_gpgcheck", conf(DNF, "repo_gpgcheck", 'lower(v) in {"1", "true", "yes"}',
                                     "[main]\nrepo_gpgcheck=1", "[main]\nrepo_gpgcheck=0"))
put("1.2.1.4", "repositories", attest("repositories_reviewed"))
put("1.2.1.5", "weak dependencies", custom(
    [f'text_collected({ref(DNF)})'],
    [['count([v | some v in conf_values(input.initial_setup.dnf_conf_raw, "install_weak_deps"); lower(v) in {"0", "false", "no"}]) > 0',
      'count([v | some v in conf_values(input.initial_setup.dnf_conf_raw, "install_weak_deps"); lower(v) in {"1", "true", "yes"}]) == 0']],
    dotted(DNF), nest(DNF, "[main]\ninstall_weak_deps=False"), nest(DNF, "[main]\ngpgcheck=1"), reads=[DNF]))
put("1.2.2.1", "updates, patches", attest("updates_installed"))

# ── 1.3 SELinux (v2.0.0 selinux.* fields; configured mode from selinux.config_mode or the mapper's
#    ansible_facts.ansible_selinux.config_mode) ──
SE = "input.selinux."
put("1.3.1.1", "SELinux is installed", boolean(["selinux", "installed"], True))
put("1.3.1.2", "not disabled in bootloader", custom(
    [f'is_string({SE}grub_cmdline)'], [[f'not regex.match(`(^|[\\s"])(selinux|enforcing)=0\\b`, {SE}grub_cmdline)']],
    "selinux.grub_cmdline", {"selinux": {"grub_cmdline": 'GRUB_CMDLINE_LINUX="quiet audit=1"'}},
    {"selinux": {"grub_cmdline": 'GRUB_CMDLINE_LINUX="quiet selinux=0"'}}, reads=[["selinux", "grub_cmdline"]]))
put("1.3.1.3", "SELinux policy", custom(
    [f'is_string({SE}policy)', f'{SE}policy != ""'], [[f'lower({SE}policy) in {{"targeted", "mls"}}']], "selinux.policy",
    {"selinux": {"policy": "targeted"}}, {"selinux": {"policy": "minimum"}}, reads=[["selinux", "policy"]]))
for rid, kw, allowed in [("1.3.1.4", "mode is not disabled", ["enforcing", "permissive"]),
                         ("1.3.1.5", "mode is enforcing", ["enforcing"])]:
    a = json.dumps(allowed)
    put(rid, kw, custom(
        [f'is_string({SE}mode)', 'is_string(selinux_config_mode)'],
        [[f'lower({SE}mode) in {a}', f'lower(selinux_config_mode) in {a}']],
        "selinux.mode/selinux.config_mode",
        {"selinux": {"mode": "enforcing", "config_mode": "enforcing"}},
        {"selinux": {"mode": "enforcing", "config_mode": "disabled" if len(allowed) == 2 else "permissive"}},
        {"selinux": {"mode": "enforcing"}},
        reads=[["selinux", "mode"], ["selinux", "config_mode"], ["ansible_facts", "ansible_selinux", "config_mode"]]))
put("1.3.1.6", "unconfined services", attest("no_unconfined_services"))
put("1.3.1.7", "mcstrans", pkg_absent(None, ["selinux", "mcstrans_installed"]))
put("1.3.1.8", "SETroubleshoot", pkg_absent(None, ["selinux", "setroubleshoot_installed"]))

# ── 1.4 bootloader (v2.0.0 boot_security.* fields) ──
BS = "input.boot_security."
put("1.4.1", "bootloader password", custom(
    [f'is_boolean({BS}grub_password_set)', f'is_boolean({BS}grub_password_hashed)'],
    [[f'{BS}grub_password_set == true', f'{BS}grub_password_hashed == true']],
    "boot_security.grub_password_set/grub_password_hashed",
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": True}},
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": False}},
    reads=[["boot_security", "grub_password_set"], ["boot_security", "grub_password_hashed"]]))
GF = ["boot_security", "grub_files"]   # new: every grub*/user.cfg file under /boot {path, mode, owner, group}
_gcfg = {"grub_cfg_mode": "0600", "grub_cfg_owner": "root", "grub_cfg_group": "root"}
_gfiles = [{"path": "/boot/grub2/grub.cfg", "mode": "0600", "owner": "root", "group": "root"},
           {"path": "/boot/grub2/grubenv", "mode": "0600", "owner": "root", "group": "root"}]
put("1.4.2", "access to bootloader config", custom(
    [f'{BS}grub_cfg_mode != null', f'{BS}grub_cfg_owner != null', f'{BS}grub_cfg_group != null', f'is_array({ref(GF)})',
     f'count({ref(GF)}) > 0'],
    [[f'mode_within({BS}grub_cfg_mode, "0600")', f'{BS}grub_cfg_owner == "root"', f'{BS}grub_cfg_group == "root"',
      f'every f in {ref(GF)} {{ grub_file_ok(f) }}']],
    "boot_security.grub_cfg_mode/grub_cfg_owner/grub_cfg_group/grub_files",
    {"boot_security": dict(_gcfg, grub_files=_gfiles)},
    {"boot_security": dict(_gcfg, grub_files=_gfiles + [{"path": "/boot/grub2/grubenv", "mode": "0644", "owner": "root", "group": "root"}])},
    {"boot_security": _gcfg},
    reads=[["boot_security", "grub_cfg_mode"], ["boot_security", "grub_cfg_owner"], ["boot_security", "grub_cfg_group"], GF]))
for rid, kw, unit in [("1.4.3", "rescue mode", "rescue"), ("1.4.4", "emergency mode", "emergency")]:
    f = f"{unit}_auth_required"   # new per-unit field; v2.0.0 single_user_auth_required covers both
    put(rid, kw, custom(
        None,
        [[f'{BS}{f} == true'], [f'not {f}_collected', f'{BS}single_user_auth_required == true']],
        f"boot_security.{f} (or single_user_auth_required)",
        {"boot_security": {f: True}}, {"boot_security": {f: False}}, {"boot_security": {"aide_installed": True}},
        reads=[["boot_security", f], ["boot_security", "single_user_auth_required"]],
        ev_alts=[[f'is_boolean({BS}{f})'], [f'is_boolean({BS}single_user_auth_required)']]))
KA = ["boot_security", "kernel_args"]   # new: args= of every boot entry (grubby --info=ALL)
GRUB = ["bootloader", "grub_config"]    # v2.0.0 mapper field: /etc/default/grub
put("1.4.5", "interactive boot", custom(
    None,
    [[f'is_array({ref(KA)})', f'count([a | some a in {ref(KA)}; contains(a, "systemd.confirm_spawn")]) == 0'],
     ['not kernel_args_collected', f'not regex_in_text({ref(GRUB)}, `systemd\\.confirm_spawn`)']],
    "boot_security.kernel_args (or bootloader.grub_config)",
    {"boot_security": {"kernel_args": ["ro crashkernel=1G audit=1"]}},
    {"boot_security": {"kernel_args": ["ro systemd.confirm_spawn=1"]}}, {"boot_security": {"aide_installed": True}},
    reads=[KA, GRUB],
    ev_alts=[[f'is_array({ref(KA)})', f'count({ref(KA)}) > 0'], [f'text_collected({ref(GRUB)})']]))

# ── 1.5 process hardening ──
put("1.5.1", "fs.protected_hardlinks", sysctl("fs.protected_hardlinks", ["1"]))
put("1.5.2", "fs.protected_symlinks", sysctl("fs.protected_symlinks", ["1"]))
put("1.5.3", "fs.suid_dumpable", sysctl("fs.suid_dumpable", ["0"],
                                        alias=(["boot_security", "core_dumps_disabled_sysctl"], f'{BS}core_dumps_disabled_sysctl == true')))
put("1.5.4", "kernel.dmesg_restrict", sysctl("kernel.dmesg_restrict", ["1"],
                                             alias=(["boot_security", "kernel_dmesg_restrict"], f'sprintf("%v", [{BS}kernel_dmesg_restrict]) == "1"')))
put("1.5.5", "kernel.kptr_restrict", sysctl("kernel.kptr_restrict", ["1", "2"],
                                            alias=(["boot_security", "kernel_kptr_restrict"], f'sprintf("%v", [{BS}kernel_kptr_restrict]) in {{"1", "2"}}')))
put("1.5.6", "kernel.randomize_va_space", sysctl("kernel.randomize_va_space", ["2"],
                                                 alias=(["boot_security", "aslr_value"], f'sprintf("%v", [{BS}aslr_value]) == "2"')))
put("1.5.7", "kernel.yama.ptrace_scope", sysctl("kernel.yama.ptrace_scope", ["1", "2", "3"]))
put("1.5.8", "core file size", boolean(["boot_security", "core_dumps_disabled_limits"], True))
put("1.5.9", "ProcessSizeMax", custom(
    [f'{BS}systemd_coredump_process_size_max != null'],
    [[f'sprintf("%v", [{BS}systemd_coredump_process_size_max]) == "0"']], "boot_security.systemd_coredump_process_size_max",
    {"boot_security": {"systemd_coredump_process_size_max": "0"}},
    {"boot_security": {"systemd_coredump_process_size_max": "unset (systemd default)"}},
    reads=[["boot_security", "systemd_coredump_process_size_max"]]))
put("1.5.10", "systemd-coredump Storage", custom(
    [f'is_string({BS}systemd_coredump_storage)'], [[f'lower({BS}systemd_coredump_storage) == "none"']],
    "boot_security.systemd_coredump_storage", {"boot_security": {"systemd_coredump_storage": "none"}},
    {"boot_security": {"systemd_coredump_storage": "external"}}, reads=[["boot_security", "systemd_coredump_storage"]]))

# ── 1.6 crypto policy: the active policy name (v2.0.0 initial_setup.crypto_policy, `update-crypto-policies
#    --show`) and the resolved policy /etc/crypto-policies/state/CURRENT.pol (new crypto_policy.current_pol_raw,
#    the SUSE sibling's field) ──
POL = ["crypto_policy", "current_pol_raw"]
put("1.6.1", "not legacy", custom(
    [f'text_collected({IS}crypto_policy)'], [[f'not regex.match(`(?i)^\\s*LEGACY\\b`, {IS}crypto_policy)']],
    "initial_setup.crypto_policy", {"initial_setup": {"crypto_policy": "DEFAULT"}},
    {"initial_setup": {"crypto_policy": "LEGACY"}}, reads=[["initial_setup", "crypto_policy"]]))
put("1.6.2", "sha1 hash and signature", custom(
    [f'text_collected({ref(POL)})'],
    [['not sha1_enabled', 'sha1_in_certs_off']], dotted(POL),
    nest(POL, "hash = SHA2-256 SHA2-384\nsign = RSA-PSS-SHA2-256\nhash@RPM = SHA2-256 SHA1\nsha1_in_certs = 0"),
    nest(POL, "hash = SHA2-256 SHA1\nsign = RSA-PSS-SHA2-256\nsha1_in_certs = 0"), reads=[POL]))
put("1.6.3", "crypto policy macs", text_has(POL, "(?i)^mac(@[^=\\s]+)?\\s*=\\s*([^#]+)?-64\\b", False,
                              "mac = HMAC-SHA2-256 UMAC-128", "mac = HMAC-SHA2-256 UMAC-64"))
put("1.6.4", "cbc for ssh", custom(
    [f'text_collected({ref(POL)})'],
    [[f'not regex_in_text({ref(POL)}, `(?i)^cipher\\s*=\\s*([^#]+)?-CBC\\b`)',
      f'not regex_in_text({ref(POL)}, `(?i)^cipher@(lib|open)ssh(-server|-client)?\\s*=\\s*([^#]+)?-CBC\\b`)']],
    dotted(POL), nest(POL, "cipher = AES-256-GCM AES-128-GCM\ncipher@openssh-server = AES-256-GCM"),
    nest(POL, "cipher = AES-256-GCM\ncipher@openssh-server = AES-256-GCM AES-256-CBC"), reads=[POL]))
put("1.6.5", "chacha20-poly1305", attest("crypto_policy_chacha20_ssh_reviewed"))
put("1.6.6", "EtM for ssh", attest("crypto_policy_etm_ssh_reviewed"))
PQA = ["crypto_policy", "pq_module_available"]   # new: /usr/share/crypto-policies/policies/modules/PQ.pmod exists
put("1.6.7", "post-quantum", custom(
    None,
    [[f'{ref(PQA)} == false'],
     [f'{ref(PQA)} == true', f'"PQ" in [upper(trim_space(x)) | some x in split({IS}crypto_policy, ":")]']],
    "crypto_policy.pq_module_available/initial_setup.crypto_policy",
    {"crypto_policy": {"pq_module_available": True}, "initial_setup": {"crypto_policy": "DEFAULT:PQ"}},
    {"crypto_policy": {"pq_module_available": True}, "initial_setup": {"crypto_policy": "DEFAULT"}},
    {"initial_setup": {"crypto_policy": "DEFAULT"}}, reads=[PQA, ["initial_setup", "crypto_policy"]],
    ev_alts=[[f'{ref(PQA)} == false'], [f'{ref(PQA)} == true', f'text_collected({IS}crypto_policy)']]))

# ── 1.7 banners (v2.0.0 initial_setup.<f>_exists / <f>_contains_sensitive_info / <f>_mode / <f>_owner;
#    <f>_group is new) ──
for rid, f, kw, required in [("1.7.1", "motd", "ensure /etc/motd is configured", False),
                             ("1.7.2", "issue", "ensure /etc/issue is configured", True),
                             ("1.7.3", "issue_net", "ensure /etc/issue.net is configured", True)]:
    ex, si = ["initial_setup", f"{f}_exists"], ["initial_setup", f"{f}_contains_sensitive_info"]
    passes = [[f'{ref(ex)} == true', f'{ref(si)} == false']] + ([] if required else [[f'{ref(ex)} == false']])
    put(rid, kw, custom(None, passes, f"initial_setup.{f}_exists/{f}_contains_sensitive_info",
                        merge(nest(ex, True), nest(si, False)), merge(nest(ex, True), nest(si, True)), {},
                        reads=[ex, si], ev_alts=[[f'{ref(ex)} == false'], [f'{ref(ex)} == true', f'is_boolean({ref(si)})']]))
for rid, f, kw, optional in [("1.7.4", "motd", "access to /etc/motd", True),
                             ("1.7.5", "issue", "access to /etc/issue is", False),
                             ("1.7.6", "issue_net", "access to /etc/issue.net", False)]:
    put(rid, kw, access(["initial_setup", f"{f}_mode"], ["initial_setup", f"{f}_owner"],
                                ["initial_setup", f"{f}_group"], "0644", ["root"], ["root"],
                                exists=["initial_setup", f"{f}_exists"], optional=optional))


# ── 1.8 GDM (v2.0.0 initial_setup.gdm_* fields; not applicable when gdm_installed is collected false) ──
def gdm(fields, conds, good, bad):
    inst = ["initial_setup", "gdm_installed"]
    fp = [["initial_setup", f] for f in fields]
    have = [f'{ref(p)} != null' for p in fp]
    return custom(None, [[f'{IS}gdm_installed == false'], [f'{IS}gdm_installed == true'] + conds],
                  "initial_setup." + "/".join(["gdm_installed"] + fields),
                  {"initial_setup": dict(good, gdm_installed=True)},
                  {"initial_setup": dict(bad, gdm_installed=True)},
                  {"initial_setup": {"gdm_installed": True}}, reads=[inst] + fp,
                  ev_alts=[[f'{IS}gdm_installed == false'], [f'{IS}gdm_installed == true'] + have])


put("1.8.1", "GDM login banner", gdm(["gdm_banner_configured", "gdm_banner_override_allowed"],
                                     [IS + "gdm_banner_configured == true", IS + "gdm_banner_override_allowed == false"],
                                     {"gdm_banner_configured": True, "gdm_banner_override_allowed": False},
                                     {"gdm_banner_configured": True, "gdm_banner_override_allowed": True}))
put("1.8.2", "disable-user-list", gdm(["gdm_disable_user_list", "gdm_disable_user_list_override_allowed"],
                                      [IS + "gdm_disable_user_list == true", IS + "gdm_disable_user_list_override_allowed == false"],
                                      {"gdm_disable_user_list": True, "gdm_disable_user_list_override_allowed": False},
                                      {"gdm_disable_user_list": False, "gdm_disable_user_list_override_allowed": False}))
put("1.8.3", "screen lock", gdm(["gdm_idle_delay", "gdm_lock_delay", "gdm_lock_override_allowed"],
                                [f"num({IS}gdm_idle_delay) > 0", f"num({IS}gdm_idle_delay) <= 900",
                                 f"num({IS}gdm_lock_delay) >= 0", f"num({IS}gdm_lock_delay) <= 5",
                                 IS + "gdm_lock_override_allowed == false"],
                                {"gdm_idle_delay": 900, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False},
                                {"gdm_idle_delay": 0, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False}))
put("1.8.4", "automount", gdm(["gdm_automount_disabled", "gdm_automount_open_disabled", "gdm_automount_override_allowed"],
                              [IS + "gdm_automount_disabled == true", IS + "gdm_automount_open_disabled == true",
                               IS + "gdm_automount_override_allowed == false"],
                              {"gdm_automount_disabled": True, "gdm_automount_open_disabled": True, "gdm_automount_override_allowed": False},
                              {"gdm_automount_disabled": True, "gdm_automount_open_disabled": False, "gdm_automount_override_allowed": False}))
put("1.8.5", "autorun-never", gdm(["gdm_autorun_never_enabled", "gdm_autorun_override_allowed"],
                                  [IS + "gdm_autorun_never_enabled == true", IS + "gdm_autorun_override_allowed == false"],
                                  {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": False},
                                  {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": True}))
put("1.8.6", "XDMCP", gdm(["gdm_xdmcp_enabled"], [IS + "gdm_xdmcp_enabled == false"],
                          {"gdm_xdmcp_enabled": False}, {"gdm_xdmcp_enabled": True}))
put("1.8.7", "Wayland is configured", gdm(["gdm_wayland_disabled"], [IS + "gdm_wayland_disabled == false"],
                            {"gdm_wayland_disabled": False}, {"gdm_wayland_disabled": True}))

# ── 2.1 server services: package not installed, or every unit the audit names neither enabled nor active
#    (packages.<name> booleans + ansible service_facts `services`, the v2.0.0 auditd shapes) ──
SVC = [
    ("2.1.1", "autofs", {"autofs": ["autofs.service"]}),
    ("2.1.2", "avahi", {"avahi": ["avahi-daemon.socket", "avahi-daemon.service"]}),
    ("2.1.3", "cockpit", {"cockpit-ws": ["cockpit.socket", "cockpit.service"]}),
    ("2.1.4", "dhcp server", {"dhcp-server": ["dhcpd.service", "dhcpd6.service"]}),
    ("2.1.5", "dns server", {"bind": ["named.service"]}),
    ("2.1.6", "dnsmasq", {"dnsmasq": ["dnsmasq.service"]}),
    ("2.1.7", "samba", {"samba": ["smb.service"]}),
    ("2.1.8", "ensure ftp server", {"vsftpd": ["vsftpd.service"]}),
    ("2.1.9", "message access server", {"dovecot": ["dovecot.socket", "dovecot.service"], "cyrus-imapd": ["cyrus-imapd.service"]}),
    ("2.1.10", "network file system", {"nfs-utils": ["nfs-server.service"]}),
    ("2.1.11", "print server", {"cups": ["cups.socket", "cups.service"]}),
    ("2.1.12", "rpcbind", {"rpcbind": ["rpcbind.socket", "rpcbind.service"]}),
    ("2.1.13", "rsync", {"rsync-daemon": ["rsyncd.socket", "rsyncd.service"]}),
    ("2.1.14", "snmp", {"net-snmp": ["snmpd.service"]}),
    ("2.1.15", "telnet server", {"telnet-server": ["telnet.socket"]}),
    ("2.1.16", "tftp server", {"tftp-server": ["tftp.socket", "tftp.service"]}),
    ("2.1.17", "web proxy", {"squid": ["squid.service"]}),
    ("2.1.18", "web server", {"httpd": ["httpd.socket", "httpd.service"], "nginx": ["nginx.service"]}),
]
for rid, kw, ubp in SVC:
    put(rid, kw, not_in_use(ubp))
# v2.0.0 autofs object {installed, active, enabled} (the mapper emits it from service_facts)
alt(L["2.1.1"], ['is_boolean(input.autofs.installed)', 'is_boolean(input.autofs.active)', 'is_boolean(input.autofs.enabled)'],
    ['input.autofs.active == false', 'input.autofs.enabled == false'],
    [["autofs", "installed"], ["autofs", "active"], ["autofs", "enabled"]], "autofs.installed/active/enabled")
# v2.0.0 webserver inventory {any_installed, packages}
alt(L["2.1.18"], ['is_boolean(input.webserver.any_installed)'], ['input.webserver.any_installed == false'],
    [["webserver", "any_installed"]], "webserver.any_installed")
put("2.1.19", "mail transfer agents", custom(
    None,
    [['input.mta.has_mta == false'],
     ['count(input.mta.non_loopback_listeners) == 0', 'not mta_interfaces_open']],
    "mta.has_mta (or mta.non_loopback_listeners/inet_interfaces)",
    {"mta": {"has_mta": True, "non_loopback_listeners": [], "inet_interfaces": "loopback-only"}},
    {"mta": {"has_mta": True, "non_loopback_listeners": ["0.0.0.0:25"], "inet_interfaces": "all"}},
    {"mta": {"has_mta": True}},
    reads=[["mta", "has_mta"], ["mta", "non_loopback_listeners"], ["mta", "inet_interfaces"]],
    ev_alts=[['input.mta.has_mta == false'], ['is_array(input.mta.non_loopback_listeners)']]))
put("2.1.20", "graphical user interface", custom(
    None,
    [['input.gui.environment_installed == false'],
     ['input.gui.environment_installed == true', 'lower(input.gui.default_target) != "graphical.target"']],
    "gui.environment_installed/default_target",
    {"gui": {"environment_installed": True, "default_target": "multi-user.target"}},
    {"gui": {"environment_installed": True, "default_target": "graphical.target"}},
    {"gui": {"environment_installed": True}},
    reads=[["gui", "environment_installed"], ["gui", "default_target"]],
    ev_alts=[['input.gui.environment_installed == false'],
             ['input.gui.environment_installed == true', 'text_collected(input.gui.default_target)']]))
put("2.1.21", "X window server", alt(pkg_absent(["xorg-x11-server-common"]),
                                     ['is_boolean(input.x11.any_installed)'], ['input.x11.any_installed == false'],
                                     [["x11", "any_installed"]], "x11.any_installed"))
put("2.1.22", "Xwayland", pkg_absent(["xorg-x11-server-Xwayland"]))
put("2.1.23", "approved services are listening", attest("listening_services_approved"))

# ── 2.2 client packages (packages.<name>, or the v2.0.0 client_packages.analysis [{package, installed}]) ──
for rid, kw, pkg in [("2.2.1", "ensure ftp client", "ftp"), ("2.2.2", "ldap client", "openldap-clients"),
                     ("2.2.3", "telnet client", "telnet"), ("2.2.4", "tftp client", "tftp")]:
    cp = f'[c | some c in object.get(input, ["client_packages", "analysis"], []); c["package"] == {json.dumps(pkg)}]'
    put(rid, kw, alt(pkg_absent([pkg]), [f'count({cp}) > 0', f'every c in {cp} {{ is_boolean(c.installed) }}'],
                     [f'every c in {cp} {{ c.installed == false }}'], [["client_packages", "analysis"]],
                     "client_packages.analysis"))

# ── 2.3 time synchronization (v2.0.0 chronyd.* fields) ──
CH = "input.chronyd."
put("2.3.1", "time synchronization is in use", alt(boolean(["chronyd", "installed"], True),
                                                   ['is_boolean(input.packages.chrony)'], ['input.packages.chrony == true'],
                                                   [["packages", "chrony"]], "packages.chrony"))
put("2.3.2", "chrony is configured", custom(
    [f'{CH}servers_configured != null', f'{CH}pools_configured != null'],
    [[f'num({CH}servers_configured) > 0'], [f'num({CH}pools_configured) > 0']],
    "chronyd.servers_configured/pools_configured",
    {"chronyd": {"servers_configured": 0, "pools_configured": 1}},
    {"chronyd": {"servers_configured": 0, "pools_configured": 0}},
    reads=[["chronyd", "servers_configured"], ["chronyd", "pools_configured"]]))
put("2.3.3", "not run as the root user", gate(custom(
    [f'is_string({CH}sysconfig_options)'], [[f'not regex.match(`(^|\\s)-u\\s+root\\b`, {CH}sysconfig_options)']],
    "chronyd.sysconfig_options", {"chronyd": {"sysconfig_options": "-F 2"}}, {"chronyd": {"sysconfig_options": "-F 2 -u root"}},
    reads=[["chronyd", "sysconfig_options"]]), ["chronyd", "installed"], False, "chronyd.installed"))

# ── 2.4 job schedulers (v2.0.0 cron.* fields; *_group, cron_yearly and the deny-file fields are new) ──
CR = "input.cron."
put("2.4.1.1", "cron daemon is enabled", custom(
    [f'is_boolean({CR}crond_enabled)', f'is_boolean({CR}crond_running)'],
    [[f'{CR}crond_enabled == true', f'{CR}crond_running == true']], "cron.crond_enabled/crond_running",
    {"cron": {"crond_enabled": True, "crond_running": True}}, {"cron": {"crond_enabled": True, "crond_running": False}},
    reads=[["cron", "crond_enabled"], ["cron", "crond_running"]]))
put("2.4.1.2", "/etc/crontab", access(["cron", "crontab_mode"], ["cron", "crontab_owner"], ["cron", "crontab_group"],
                                      "0600", ["root"], ["root"]))
for rid, d, kw in [("2.4.1.3", "cron_hourly", "/etc/cron.hourly"), ("2.4.1.4", "cron_daily", "/etc/cron.daily"),
                   ("2.4.1.5", "cron_weekly", "/etc/cron.weekly"), ("2.4.1.6", "cron_monthly", "/etc/cron.monthly"),
                   ("2.4.1.8", "cron_d", "access to /etc/cron.d ")]:
    put(rid, kw, access(["cron", f"{d}_mode"], ["cron", f"{d}_owner"], ["cron", f"{d}_group"], "0700", ["root"], ["root"]))
# stock RHEL 9 cronie ships no /etc/cron.yearly (seen on a RHEL 9.8 host): an absent directory has nothing to
# secure, so a collected cron_yearly_exists == false passes; a present one is checked like the others
_cy = access(["cron", "cron_yearly_mode"], ["cron", "cron_yearly_owner"], ["cron", "cron_yearly_group"], "0700", ["root"], ["root"])
_cy.ev_alts = _cy.ev_alts + [[f'{CR}cron_yearly_exists == false']]
_cy.passes = _cy.passes + [[f'{CR}cron_yearly_exists == false']]
_cy.reads = _cy.reads + [["cron", "cron_yearly_exists"]]
_cy.field = _cy.field + "/cron.cron_yearly_exists"
put("2.4.1.7", "/etc/cron.yearly", _cy)


def allow_deny(svc, groups):
    """<svc>.allow must exist (mode 0640, root, groups); <svc>.deny, if present, the same."""
    a = [f"{svc}_allow_{k}" for k in ("exists", "mode", "owner", "group")]
    d = [f"{svc}_deny_{k}" for k in ("exists", "mode", "owner", "group")]
    g = json.dumps(groups)
    allow_ok = [f'{CR}{a[0]} == true', f'mode_within({CR}{a[1]}, "0640")', f'{CR}{a[2]} == "root"', f'{CR}{a[3]} in {g}']
    deny_ok_present = [f'{CR}{d[0]} == true', f'mode_within({CR}{d[1]}, "0640")', f'{CR}{d[2]} == "root"', f'{CR}{d[3]} in {g}']
    good = {a[0]: True, a[1]: "0640", a[2]: "root", a[3]: groups[0], d[0]: False}
    bad = {a[0]: False, d[0]: False}
    ev_allow = [[f'{CR}{a[0]} == false'], [f'{CR}{a[0]} == true', f'{CR}{a[1]} != null', f'{CR}{a[2]} != null', f'{CR}{a[3]} != null']]
    ev_deny = [[f'{CR}{d[0]} == false'], [f'{CR}{d[0]} == true', f'{CR}{d[1]} != null', f'{CR}{d[2]} != null', f'{CR}{d[3]} != null']]
    return custom(None, [allow_ok + [f'{CR}{d[0]} == false'], allow_ok + deny_ok_present],
                  f"cron.{svc}_allow_*/{svc}_deny_*", {"cron": good}, {"cron": bad}, {},
                  reads=[["cron", k] for k in a + d], ev_alts=[x + y for x in ev_allow for y in ev_deny])


put("2.4.1.9", "access to crontab", gate(allow_deny("cron", ["root", "crontab"]), ["packages", "cronie"], False,
                                         "packages.cronie"))
put("2.4.2.1", "access to at", gate(allow_deny("at", ["root", "daemon"]), ["packages", "at"], False, "packages.at"))

# ── 3 network ──
put("3.1.1", "IPv6 status", attest("ipv6_status_identified"))
WI = "input.wireless_interfaces."
put("3.1.2", "wireless interfaces", custom(
    None,
    [[f'{WI}has_wireless == false'],
     [f'{WI}has_wireless == true', f'count({WI}modules) > 0', f'every e in {WI}modules {{ kmod_unavailable(e) }}']],
    "wireless_interfaces.has_wireless/modules",
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": True, "loaded": False}]}},
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": False, "loaded": True}]}},
    {"wireless_interfaces": {"has_wireless": True}},
    reads=[["wireless_interfaces", "has_wireless"], ["wireless_interfaces", "modules"]],
    ev_alts=[[f'{WI}has_wireless == false'], [f'{WI}has_wireless == true', f'is_array({WI}modules)']]))
put("3.1.3", "bluetooth", not_in_use({"bluez": ["bluetooth.service"]}))
for rid, mod in [("3.2.1", "dccp"), ("3.2.2", "tipc"), ("3.2.3", "rds"), ("3.2.4", "sctp"), ("3.2.5", "atm"), ("3.2.6", "can")]:
    put(rid, mod + " kernel module", kmod(NETM, mod))
for i, (p, v) in enumerate([
        ("net.ipv4.ip_forward", "0"), ("net.ipv4.conf.all.forwarding", "0"), ("net.ipv4.conf.default.forwarding", "0"),
        ("net.ipv4.conf.all.send_redirects", "0"), ("net.ipv4.conf.default.send_redirects", "0"),
        ("net.ipv4.icmp_ignore_bogus_error_responses", "1"), ("net.ipv4.icmp_echo_ignore_broadcasts", "1"),
        ("net.ipv4.conf.all.accept_redirects", "0"), ("net.ipv4.conf.default.accept_redirects", "0"),
        ("net.ipv4.conf.all.secure_redirects", "0"), ("net.ipv4.conf.default.secure_redirects", "0"),
        ("net.ipv4.conf.all.rp_filter", "1"), ("net.ipv4.conf.default.rp_filter", "1"),
        ("net.ipv4.conf.all.accept_source_route", "0"), ("net.ipv4.conf.default.accept_source_route", "0"),
        ("net.ipv4.conf.all.log_martians", "1"), ("net.ipv4.conf.default.log_martians", "1"),
        ("net.ipv4.tcp_syncookies", "1"), ("net.ipv4.conf.all.route_localnet", "0"),
        ("net.ipv4.conf.default.route_localnet", "0")], start=1):
    put(f"3.3.1.{i}", p, sysctl(p, [v]))
for i, p in enumerate(["net.ipv6.conf.all.forwarding", "net.ipv6.conf.default.forwarding",
                       "net.ipv6.conf.all.accept_redirects", "net.ipv6.conf.default.accept_redirects",
                       "net.ipv6.conf.all.accept_source_route", "net.ipv6.conf.default.accept_source_route",
                       "net.ipv6.conf.all.accept_ra", "net.ipv6.conf.default.accept_ra"], start=1):
    put(f"3.3.2.{i}", p, sysctl(p, ["0"], ipv6=True))

# ── 4 firewalld (v2.0.0 firewall {type, active}; backend/zones/log_denied are new firewall.* fields) ──
FW = "input.firewall."
put("4.1.1", "firewalld is installed", alt(pkg_present("firewalld"),
                                           [f'{FW}type == "firewalld"', f'{FW}active == true'], ['true'],
                                           [["firewall", "type"], ["firewall", "active"]], "firewall.type=firewalld (active)"))
put("4.1.2", "backend", custom([f'is_string({FW}backend)'], [[f'lower({FW}backend) != "iptables"']], "firewall.backend",
                               {"firewall": {"backend": "nftables"}}, {"firewall": {"backend": "iptables"}},
                               reads=[["firewall", "backend"]]))
put("4.1.3", "firewalld.service", unit_running("firewalld.service"))
_z = [{"name": "public", "interfaces": ["eth0"], "target": "default", "permanent_target": "default",
       "icmp_block_inversion": False, "icmp_blocks": ["timestamp-reply", "timestamp-request"]},
      {"name": "trusted", "interfaces": ["lo"], "target": "ACCEPT", "permanent_target": "ACCEPT",
       "icmp_block_inversion": False, "icmp_blocks": []}]
put("4.1.4", "active zone target", custom(
    [f'is_array({FW}zones)', f'count({FW}zones) > 0'],
    [[f'count([z | some z in {FW}zones; not loopback_zone(z)]) > 0', f'every z in {FW}zones {{ zone_target_ok(z) }}']],
    "firewall.zones", {"firewall": {"zones": _z}},
    {"firewall": {"zones": [dict(_z[0], target="ACCEPT", permanent_target="ACCEPT")]}}, reads=[["firewall", "zones"]]))
put("4.1.5", "loopback traffic", attest("firewalld_loopback_traffic_configured"))
put("4.1.6", "loopback source address", attest("firewalld_loopback_source_traffic_configured"))
put("4.1.7", "services and ports", attest("firewalld_services_ports_reviewed"))
put("4.1.8", "default zone is set to drop", attest("firewalld_default_zone_drop"))
_c = units_off(["nftables.service"])
put("4.1.9", "conflicting firewall", custom(
    _c.ev_alts[0] + ['is_boolean(input.packages["iptables-services"])', 'is_boolean(input.packages.ufw)'],
    [_c.passes[0] + ['input.packages["iptables-services"] == false', 'input.packages.ufw == false']],
    "services[nftables.service]/packages[iptables-services,ufw]",
    merge(_c.tp, {"packages": {"iptables-services": False, "ufw": False}}),
    merge(_c.tp, {"packages": {"iptables-services": True, "ufw": False}}),
    {"packages": {"iptables-services": False, "ufw": False}},
    reads=[["services"], ["packages", "iptables-services"], ["packages", "ufw"]]))
put("4.1.10", "logs denied packets", custom([f'text_collected({FW}log_denied)'], [[f'lower({FW}log_denied) != "off"']],
                                            "firewall.log_denied", {"firewall": {"log_denied": "unicast"}},
                                            {"firewall": {"log_denied": "off"}}, reads=[["firewall", "log_denied"]]))
put("4.1.11", "ICMP timestamp", custom(
    [f'is_array({FW}zones)', f'count({FW}zones) > 0'], [[f'every z in {FW}zones {{ zone_icmp_ok(z) }}']], "firewall.zones",
    {"firewall": {"zones": _z}}, {"firewall": {"zones": [dict(_z[0], icmp_blocks=["timestamp-request"])]}},
    reads=[["firewall", "zones"]]))

# ── 5.1 SSH server (v2.0.0 ssh.sshd_config_raw = `sshd -T`, ssh.config_permissions {mode, pw_name, gr_name},
#    ssh.private_key_permissions [{stat: {path, mode, pw_name, gr_name}}], ssh.public_key_permissions
#    [{path, mode, pw_name, gr_name}]) ──
SS = "input.ssh."
DP = ["ssh", "dropin_permissions"]   # new: /etc/ssh/sshd_config.d/*.conf [{path, mode, pw_name, gr_name}]
_cp = {"mode": "0600", "pw_name": "root", "gr_name": "root"}
put("5.1.1", "/etc/ssh/sshd_config is", custom(
    [f'{SS}config_permissions.mode != null', f'{SS}config_permissions.pw_name != null',
     f'{SS}config_permissions.gr_name != null', f'is_array({ref(DP)})'],
    [[f'perm_ok({SS}config_permissions, "0600")', f'every f in {ref(DP)} {{ perm_ok(f, "0600") }}']],
    "ssh.config_permissions/dropin_permissions",
    {"ssh": {"config_permissions": _cp, "dropin_permissions": [dict(_cp, path="/etc/ssh/sshd_config.d/50-redhat.conf")]}},
    {"ssh": {"config_permissions": _cp, "dropin_permissions": [dict(_cp, path="/etc/ssh/sshd_config.d/50-redhat.conf", mode="0644")]}},
    {"ssh": {"config_permissions": _cp}}, reads=[["ssh", "config_permissions"], DP]))
DD = ["ssh", "dropin_dir_permissions"]   # new: /etc/ssh/sshd_config.d/ {mode, pw_name, gr_name}
put("5.1.2", "sshd_config.d/ directory", custom(
    [f'{ref(DD)}.mode != null', f'{ref(DD)}.pw_name != null', f'{ref(DD)}.gr_name != null'],
    [[f'perm_ok({ref(DD)}, "0700")']], "ssh.dropin_dir_permissions",
    {"ssh": {"dropin_dir_permissions": {"mode": "0700", "pw_name": "root", "gr_name": "root"}}},
    {"ssh": {"dropin_dir_permissions": {"mode": "0755", "pw_name": "root", "gr_name": "root"}}}, reads=[DD]))
_pk = {"path": "/etc/ssh/ssh_host_ed25519_key", "mode": "0600", "pw_name": "root", "gr_name": "root"}
put("5.1.3", "private host key", custom(
    [f'is_array({SS}private_key_permissions)', f'count({SS}private_key_permissions) > 0'],
    [[f'every k in {SS}private_key_permissions {{ host_private_key_ok(k.stat) }}']], "ssh.private_key_permissions",
    {"ssh": {"private_key_permissions": [{"stat": _pk}, {"stat": dict(_pk, path="/etc/ssh/ssh_host_rsa_key", mode="0640", gr_name="ssh_keys")}]}},
    {"ssh": {"private_key_permissions": [{"stat": dict(_pk, mode="0640")}]}}, reads=[["ssh", "private_key_permissions"]]))
_pub = {"path": "/etc/ssh/ssh_host_ed25519_key.pub", "mode": "0644", "pw_name": "root", "gr_name": "root"}
put("5.1.4", "public host key", custom(
    [f'is_array({SS}public_key_permissions)', f'count({SS}public_key_permissions) > 0'],
    [[f'every k in {SS}public_key_permissions {{ host_public_key_ok(k) }}']], "ssh.public_key_permissions",
    {"ssh": {"public_key_permissions": [_pub]}}, {"ssh": {"public_key_permissions": [dict(_pub, mode="0664")]}},
    reads=[["ssh", "public_key_permissions"]]))
WEAK_CIPH = ["3des-cbc", "aes128-cbc", "aes192-cbc", "aes256-cbc", "arcfour", "arcfour128", "arcfour256",
             "blowfish-cbc", "cast128-cbc", "rijndael-cbc@lysator.liu.se"]
WEAK_MAC = ["hmac-md5", "hmac-md5-96", "hmac-ripemd160", "hmac-sha1", "hmac-sha1-96", "umac-64@openssh.com",
            "hmac-md5-etm@openssh.com", "hmac-md5-96-etm@openssh.com", "hmac-ripemd160-etm@openssh.com",
            "hmac-sha1-etm@openssh.com", "hmac-sha1-96-etm@openssh.com", "umac-64-etm@openssh.com"]
WEAK_KEX = ["diffie-hellman-group1-sha1", "diffie-hellman-group14-sha1", "diffie-hellman-group-exchange-sha1"]


def sshd_list_none(key, weak, good, bad):
    s = sshd(key, 'true', good, bad)
    s.passes = [[f'count([x | some x in sshd_list({json.dumps(key)}); x in {json.dumps(weak)}]) == 0']]
    return s


put("5.1.5", "Ciphers", sshd_list_none("ciphers", WEAK_CIPH, "aes256-gcm@openssh.com,aes256-ctr", "aes256-ctr,aes128-cbc"))
put("5.1.6", "KexAlgorithms", sshd_list_none("kexalgorithms", WEAK_KEX, "curve25519-sha256,ecdh-sha2-nistp256",
                                             "curve25519-sha256,diffie-hellman-group14-sha1"))
put("5.1.7", "sshd MACs", sshd_list_none("macs", WEAK_MAC, "hmac-sha2-512-etm@openssh.com,hmac-sha2-256",
                                    "hmac-sha2-256,umac-64-etm@openssh.com"))
put("5.1.8", "sshd access", custom(
    ['ssh_collected'], [['some k in ["allowusers", "allowgroups", "denyusers", "denygroups"]', 'sshd_first(k) != ""']],
    "ssh.sshd_config_raw", ssh_in(["port 22", "allowgroups sshusers"]), ssh_in(["port 22"]), {}, reads=[SSH_RAW]))
put("5.1.9", "sshd Banner", custom(
    None,
    [['startswith(sshd_value("banner"), "/")', 'banner_text_ok']],
    "ssh.sshd_config_raw (effective banner)/ssh.banner_contains_sensitive_info",
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_contains_sensitive_info": False}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_contains_sensitive_info": True}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net"}},
    reads=[SSH_RAW, ["ssh", "banner_contains_sensitive_info"], ["initial_setup", "issue_net_contains_sensitive_info"]],
    ev_alts=[['ssh_collected', 'lower(sshd_value("banner")) == "none"'],
             ['ssh_collected', 'sshd_first("banner")', 'banner_text_collected']]))
put("5.1.10", "ClientAliveInterval", custom(
    ['ssh_collected', 'sshd_first("clientaliveinterval")', 'sshd_first("clientalivecountmax")'],
    [['num(sshd_value("clientaliveinterval")) > 0', 'num(sshd_value("clientalivecountmax")) > 0']],
    "ssh.sshd_config_raw (effective clientaliveinterval/clientalivecountmax)",
    ssh_in(["clientaliveinterval 15", "clientalivecountmax 3"]), ssh_in(["clientaliveinterval 0", "clientalivecountmax 3"]),
    ssh_in(["port 22"]), reads=[SSH_RAW]))
put("5.1.11", "DisableForwarding", sshd("disableforwarding", 'v == "yes"', "yes", "no"))
put("5.1.12", "GSSAPIAuthentication", sshd("gssapiauthentication", 'v == "no"', "no", "yes"))
put("5.1.13", "HostbasedAuthentication", sshd("hostbasedauthentication", 'v == "no"', "no", "yes"))
put("5.1.14", "IgnoreRhosts", sshd("ignorerhosts", 'v == "yes"', "yes", "no"))
_lg = sshd("logingracetime", 'sshd_seconds(v) >= 1', "60", "120")
_lg.passes = [['v := lower(sshd_value("logingracetime"))', 'sshd_seconds(v) >= 1', 'sshd_seconds(v) <= 60']]
put("5.1.15", "LoginGraceTime", _lg)
put("5.1.16", "LogLevel", sshd("loglevel", 'v in {"info", "verbose"}', "VERBOSE", "QUIET"))
put("5.1.17", "MaxAuthTries", sshd("maxauthtries", 'num(v) <= 4', "4", "6"))
put("5.1.18", "MaxStartups", sshd("maxstartups", 'maxstartups_ok(v)', "10:30:60", "10:30:100"))
put("5.1.19", "MaxSessions", sshd("maxsessions", 'num(v) <= 10', "10", "20"))
put("5.1.20", "PermitEmptyPasswords", sshd("permitemptypasswords", 'v == "no"', "no", "yes"))
put("5.1.21", "PermitRootLogin", sshd("permitrootlogin", 'v == "no"', "no", "prohibit-password"))
put("5.1.22", "PermitUserEnvironment", sshd("permituserenvironment", 'v == "no"', "no", "yes"))
put("5.1.23", "UsePAM", sshd("usepam", 'v == "yes"', "yes", "no"))

# ── 5.2 privilege escalation (v2.0.0 sudo.* fields; sudoers.d file and directory modes are new) ──
SU = ["sudo", "sudoers_content"]
SD = "input.sudo."
put("5.2.1", "sudo is installed", boolean(["sudo", "installed"], True))
_sf = {"path": "/etc/sudoers.d/admins", "mode": "0440", "owner": "root", "group": "root"}
put("5.2.2", "sudoers configuration files", custom(
    [f'{SD}sudoers_mode != null', f'{SD}sudoers_owner != null', f'{SD}sudoers_group != null', f'is_array({SD}sudoers_d_files)'],
    [[f'mode_within({SD}sudoers_mode, "0440")', f'{SD}sudoers_owner == "root"', f'{SD}sudoers_group == "root"',
      f'every f in {SD}sudoers_d_files {{ mode_within(f.mode, "0440"); f.owner == "root"; f.group == "root" }}']],
    "sudo.sudoers_mode/sudoers_owner/sudoers_group/sudoers_d_files",
    {"sudo": {"sudoers_mode": "0440", "sudoers_owner": "root", "sudoers_group": "root", "sudoers_d_files": [_sf]}},
    {"sudo": {"sudoers_mode": "0440", "sudoers_owner": "root", "sudoers_group": "root", "sudoers_d_files": [dict(_sf, mode="0644")]}},
    {"sudo": {"sudoers_mode": "0440", "sudoers_owner": "root", "sudoers_group": "root"}},
    reads=[["sudo", "sudoers_mode"], ["sudo", "sudoers_owner"], ["sudo", "sudoers_group"], ["sudo", "sudoers_d_files"]]))
put("5.2.3", "/etc/sudoers.d directory", access(["sudo", "sudoers_d_dir", "mode"], ["sudo", "sudoers_d_dir", "owner"],
                                                ["sudo", "sudoers_d_dir", "group"], "0750", ["root"], ["root"]))
put("5.2.4", "use pty", custom([f'text_collected({ref(SU)})'],
                               [['sudo_defaults_has(`(^|[\\s,])use_pty\\b`)', 'not sudo_defaults_has(`!\\s*use_pty\\b`)']],
                               dotted(SU), nest(SU, "Defaults use_pty\nroot ALL=(ALL) ALL"),
                               nest(SU, "Defaults !use_pty\nroot ALL=(ALL) ALL"), reads=[SU]))
put("5.2.5", "ensure sudo log file exists", custom(
    [f'text_collected({ref(SU)})'], [['sudo_logfile', f'not {SD}logfile_exists == false']], "sudo.sudoers_content/logfile_exists",
    {"sudo": {"sudoers_content": 'Defaults logfile="/var/log/sudo.log"', "logfile_exists": True}},
    {"sudo": {"sudoers_content": "# Defaults logfile=/var/log/sudo.log"}}, reads=[SU, ["sudo", "logfile_exists"]]))
put("5.2.6", "password for escalation", text_has(SU, "(^|[\\s:,])NOPASSWD\\s*:", False,
                                                 "%wheel ALL=(ALL) ALL", "%wheel ALL=(ALL) NOPASSWD: ALL"))
put("5.2.7", "re-authentication", text_has(SU, "!authenticate\\b", False, "Defaults use_pty", "Defaults !authenticate"))
put("5.2.8", "pwfeedback", custom([f'text_collected({ref(SU)})'], [['not sudo_defaults_has(`(^|[\\s,])pwfeedback\\b`)']],
                                  dotted(SU), nest(SU, "Defaults use_pty"), nest(SU, "Defaults pwfeedback"), reads=[SU]))
put("5.2.9", "visiblepw", custom([f'text_collected({ref(SU)})'], [['not sudo_defaults_has(`(^|[\\s,])visiblepw\\b`)']],
                                 dotted(SU), nest(SU, "Defaults !visiblepw"), nest(SU, "Defaults visiblepw"), reads=[SU]))
put("5.2.10", "timestamp_timeout", custom(
    [f'text_collected({ref(SU)})'], [['every t in sudo_timeouts { t >= 0; t <= 15 }']], dotted(SU),
    nest(SU, "Defaults timestamp_timeout=15"), nest(SU, "Defaults timestamp_timeout=-1"), reads=[SU]))
put("5.2.11", "timestamp_type", text_has(SU, "(?i)timestamp_type\\s*=\\s*global\\b", False,
                                         "Defaults timestamp_type=tty", "Defaults timestamp_type=global"))
put("5.2.12", "su command", boolean(["sudo", "su_restricted"], True))
for _rid in ("5.2.6", "5.2.7", "5.2.11"):
    L[_rid].ev_alts = [[f'text_collected({ref(SU)})']]

# ── 5.3 PAM (v2.0.0 pam.system_auth = the mapper's collected PAM text, pam.password_auth, pam.pwquality_conf_raw;
#    faillock.conf / pwhistory.conf text, package versions and authselect state are new) ──
PA = "input.pam."
PV = ["pam", "package_versions"]
for rid, kw, pkg, mn in [("5.3.1.1", "version of pam", "pam", "1.5.1-14"),
                         ("5.3.1.2", "version of authselect", "authselect", "1.2.6-1"),
                         ("5.3.1.3", "version of libpwquality", "libpwquality", "1.4.4-8")]:
    v = f'{ref(PV)}[{json.dumps(pkg)}]'
    put(rid, kw, custom([f'text_collected({v})'], [[f'rpm_version_ge({v}, {json.dumps(mn)})']], f"pam.package_versions[{pkg}]",
                        {"pam": {"package_versions": {pkg: mn + ".el9"}}},
                        {"pam": {"package_versions": {pkg: "1.0.0-1.el9"}}}, {"pam": {"package_versions": {}}}, reads=[PV]))
put("5.3.2.1", "active authselect profile", custom(
    [f'is_string({PA}authselect_profile)', f'is_boolean({PA}authselect_check_valid)'],
    [[f'trim_space({PA}authselect_profile) != ""', f'{PA}authselect_check_valid == true']],
    "pam.authselect_profile/authselect_check_valid",
    {"pam": {"authselect_profile": "sssd", "authselect_check_valid": True}},
    {"pam": {"authselect_profile": "sssd", "authselect_check_valid": False}},
    reads=[["pam", "authselect_profile"], ["pam", "authselect_check_valid"]]))
PAMOK = ("auth        required      pam_faillock.so preauth silent\n"
         "auth        sufficient    pam_unix.so\n"
         "auth        required      pam_faillock.so authfail\n"
         "account     required      pam_faillock.so\n"
         "account     required      pam_unix.so\n"
         "password    requisite     pam_pwquality.so local_users_only\n"
         "password    required      pam_pwhistory.so use_authtok\n"
         "password    sufficient    pam_unix.so sha512 shadow use_authtok\n"
         "session     required      pam_unix.so")


def pam_stack(cond, good_text, bad_text, field_suffix=""):
    """cond: rego expression over `t` (one PAM file's text); it must hold for system-auth AND password-auth"""
    return custom(
        [f'text_collected({PA}system_auth)', f'text_collected({PA}password_auth)'],
        [[f'every t in [pam_system_auth_text, {PA}password_auth] {{ {cond} }}']],
        "pam.system_auth/password_auth" + field_suffix,
        {"pam": {"system_auth": good_text, "password_auth": good_text}},
        {"pam": {"system_auth": good_text, "password_auth": bad_text}},
        {"pam": {"system_auth": good_text}}, reads=[["pam", "system_auth"], ["pam", "password_auth"]])


def drop(text, needle):
    return "\n".join(l for l in text.split("\n") if needle not in l)


put("5.3.2.2", "pam_faillock module", pam_stack(
    'pam_line(t, "auth", "pam_faillock.so", `\\spreauth\\b`); pam_line(t, "auth", "pam_faillock.so", `\\sauthfail\\b`); pam_line(t, "account", "pam_faillock.so", ``)',
    PAMOK, drop(PAMOK, "authfail")))
put("5.3.2.3", "pam_pwquality module", pam_stack('pam_line(t, "password", "pam_pwquality.so", ``)', PAMOK, drop(PAMOK, "pam_pwquality")))
put("5.3.2.4", "pam_pwhistory module", pam_stack('pam_line(t, "password", "pam_pwhistory.so", ``)', PAMOK, drop(PAMOK, "pam_pwhistory")))
put("5.3.2.5", "pam_unix module", pam_stack(
    'every ty in ["auth", "account", "password", "session"] { pam_line(t, ty, "pam_unix.so", ``) }',
    PAMOK, drop(PAMOK, "session")))
FL = ["pam", "faillock_conf_raw"]
FLT = f'text_collected({ref(FL)})'
PAMS = [f'text_collected({PA}system_auth)', f'text_collected({PA}password_auth)']


def pam_conf(conf_path, conds, args_bad, good_conf, bad_conf, field):
    """value from a module's configuration file AND no module argument on the PAM lines that undercuts it"""
    return custom([f'text_collected({ref(conf_path)})'] + PAMS,
                  [conds + [f'count(pam_args_bad({json.dumps(args_bad[0])}, {json.dumps(args_bad[1])})) == 0']],
                  f"{dotted(conf_path)}/pam.system_auth/password_auth",
                  {"pam": {conf_path[1]: good_conf, "system_auth": PAMOK, "password_auth": PAMOK}},
                  {"pam": {conf_path[1]: bad_conf, "system_auth": PAMOK, "password_auth": PAMOK}},
                  {"pam": {"system_auth": PAMOK, "password_auth": PAMOK}},
                  reads=[conf_path, ["pam", "system_auth"], ["pam", "password_auth"]])


put("5.3.3.1.1", "failed attempts lockout is", pam_conf(
    FL, ['v := num(conf_last(input.pam.faillock_conf_raw, "deny"))', 'v >= 1', 'v <= 5'],
    ("pam_faillock.so", "deny=(0|[6-9]|[1-9][0-9]+)\\b"), "deny = 5", "deny = 10", "deny"))
put("5.3.3.1.2", "unlock time", pam_conf(
    FL, ['unlock_ok(conf_last(input.pam.faillock_conf_raw, "unlock_time"))'],
    ("pam_faillock.so", "unlock_time=([1-9]|[1-9][0-9]|[1-8][0-9][0-9])\\b"), "unlock_time = 900", "unlock_time = 600",
    "unlock_time"))
put("5.3.3.1.3", "includes root account", pam_conf(
    FL, ['root_lockout_on', 'not root_unlock_short'],
    ("pam_faillock.so", "root_unlock_time=([1-9]|[1-5][0-9])\\b"), "deny = 5\neven_deny_root\nroot_unlock_time = 60",
    "deny = 5\nroot_unlock_time = 30", "even_deny_root"))
PQ = ["pam", "pwquality_conf_raw"]
for rid, kw, key, cond, bad_arg, good, bad in [
        ("5.3.3.2.1", "changed characters", "difok", 'v >= 2', "difok=[0-1]\\b", "difok = 2", "difok = 1"),
        ("5.3.3.2.2", "password length", "minlen", 'v >= 14', "minlen=([0-9]|1[0-3])\\b", "minlen = 14", "minlen = 8"),
        ("5.3.3.2.4", "same consecutive", "maxrepeat", 'v >= 1; v <= 3', "maxrepeat=(0|[4-9]|[1-9][0-9]+)\\b",
         "maxrepeat = 3", "maxrepeat = 0"),
        ("5.3.3.2.5", "sequential characters", "maxsequence", 'v >= 1; v <= 3', "maxsequence=(0|[4-9]|[1-9][0-9]+)\\b",
         "maxsequence = 3", "maxsequence = 0"),
        ("5.3.3.2.6", "dictionary check", "dictcheck", 'v != 0', "dictcheck=0\\b", "dictcheck = 1", "dictcheck = 0")]:
    put(rid, kw, pam_conf(PQ, [f'v := num(conf_last(input.pam.pwquality_conf_raw, {json.dumps(key)}))'] + cond.split("; "),
                          ("pam_pwquality.so", bad_arg), good, bad, key))
put("5.3.3.2.3", "password complexity", attest("password_complexity_reviewed"))
put("5.3.3.2.7", "quality is enforced for the root", text_has(PQ, "(?i)^enforce_for_root\\b", True,
                                                               "minlen = 14\nenforce_for_root", "minlen = 14"))
PH = ["pam", "pwhistory_conf_raw"]
put("5.3.3.3.1", "history remember", pam_conf(
    PH, ['v := num(conf_last(input.pam.pwhistory_conf_raw, "remember"))', 'v >= 24'],
    ("pam_pwhistory.so", "remember=(1?[0-9]|2[0-3])\\b"), "remember = 24", "remember = 5", "remember"))
put("5.3.3.3.2", "history is enforced for the root", text_has(PH, "(?i)^enforce_for_root\\b", True,
                                                               "remember = 24\nenforce_for_root", "remember = 24"))
put("5.3.3.3.3", "pam_pwhistory includes use_authtok", pam_stack(
    'count(pam_lines(t, "password", "pam_pwhistory.so")) > 0; every l in pam_lines(t, "password", "pam_pwhistory.so") { regex.match(`\\suse_authtok\\b`, l) }',
    PAMOK, PAMOK.replace("pam_pwhistory.so use_authtok", "pam_pwhistory.so")))
put("5.3.3.4.1", "does not include nullok", pam_stack(
    'count([l | some ty in ["auth", "account", "password", "session"]; some l in pam_lines(t, ty, "pam_unix.so"); regex.match(`\\snullok\\b`, l)]) == 0',
    PAMOK, PAMOK.replace("auth        sufficient    pam_unix.so", "auth        sufficient    pam_unix.so nullok")))
put("5.3.3.4.2", "does not include remember", pam_stack(
    'count([l | some l in pam_lines(t, "password", "pam_unix.so"); regex.match(`\\sremember=`, l)]) == 0',
    PAMOK, PAMOK.replace("sha512 shadow", "sha512 shadow remember=5")))
put("5.3.3.4.3", "pam_unix includes a strong", pam_stack(
    'count(pam_lines(t, "password", "pam_unix.so")) > 0; every l in pam_lines(t, "password", "pam_unix.so") { regex.match(`\\s(sha512|yescrypt)\\b`, l) }',
    PAMOK, PAMOK.replace("sha512 shadow", "md5 shadow")))
put("5.3.3.4.4", "pam_unix includes use_authtok", pam_stack(
    'count(pam_lines(t, "password", "pam_unix.so")) > 0; every l in pam_lines(t, "password", "pam_unix.so") { regex.match(`\\suse_authtok\\b`, l) }',
    PAMOK, PAMOK.replace("sha512 shadow use_authtok", "sha512 shadow")))

# ── 5.4 user accounts (v2.0.0 pam.pass_* / inactive_days / tmout / default_umask / system_accounts_with_login_shell,
#    user_group.uid_zero_accounts / root_path*; per-user shadow aging and the rest are new) ──
SH = ["user_group", "shadow_entries"]   # new: [{user, has_password, max_days, warn_days, inactive, last_change}]


def shadow_rule(defs_path, defs_cond, user_cond, good_def, bad_def, good_user):
    """login.defs default AND every user that has a password."""
    dp = ref(defs_path)
    return custom(
        [f'{dp} != null', f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
        [[f'v := num({dp})'] + defs_cond + [f'every u in [u | some u in {ref(SH)}; u.has_password == true] {{ {user_cond} }}']],
        f"{dotted(defs_path)}/user_group.shadow_entries",
        merge(nest(defs_path, good_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        merge(nest(defs_path, bad_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        nest(defs_path, good_def), reads=[defs_path, SH])


put("5.4.1.1", "password expiration is", shadow_rule(["pam", "pass_max_days"], ['v >= 1', 'v <= 365'],
                                                     'num(u.max_days) >= 1; num(u.max_days) <= 365', 365, 99999, {"max_days": 365}))
put("5.4.1.2", "minimum password days", attest("minimum_password_days_reviewed"))
put("5.4.1.3", "expiration warning days", shadow_rule(["pam", "pass_warn_age"], ['v >= 7'], 'num(u.warn_days) >= 7', 7, 3,
                                                      {"warn_days": 7}))
put("5.4.1.4", "ensure strong password hashing algorithm", custom(
    [f'text_collected({PA}encrypt_method)'], [[f'upper({PA}encrypt_method) in {{"SHA512", "YESCRYPT"}}']], "pam.encrypt_method",
    {"pam": {"encrypt_method": "SHA512"}}, {"pam": {"encrypt_method": "MD5"}}, reads=[["pam", "encrypt_method"]]))
# INACTIVE: -1 (and any negative value) disables the lock and fails, like a value over 45 (decisions 2026-10-19)
put("5.4.1.5", "inactive password lock", shadow_rule(["pam", "inactive_days"], ['v >= 0', 'v <= 45'],
                                                     'num(u.inactive) >= 0; num(u.inactive) <= 45', 30, -1, {"inactive": 30}))
put("5.4.1.6", "last password change date", custom(
    [f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [[f'every u in [u | some u in {ref(SH)}; u.has_password == true] {{ num(u.last_change) <= today_days }}']], dotted(SH),
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 19000}]),
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 9999999}]), reads=[SH]))
UG = "input.user_group."
put("5.4.2.1", "only UID 0", custom(
    [f'is_array({UG}uid_zero_accounts)', f'count({UG}uid_zero_accounts) > 0'],
    [[f'count([u | some u in {UG}uid_zero_accounts; u != "root"]) == 0']], "user_group.uid_zero_accounts",
    {"user_group": {"uid_zero_accounts": ["root"]}}, {"user_group": {"uid_zero_accounts": ["root", "toor"]}},
    reads=[["user_group", "uid_zero_accounts"]]))
put("5.4.2.2", "only GID 0 account", custom(
    [f'is_array({UG}gid_zero_accounts)', f'count({UG}gid_zero_accounts) > 0'],
    [[f'count([u | some u in {UG}gid_zero_accounts; not u in {{"root", "sync", "shutdown", "halt", "operator"}}]) == 0']],
    "user_group.gid_zero_accounts", {"user_group": {"gid_zero_accounts": ["root", "sync"]}},
    {"user_group": {"gid_zero_accounts": ["root", "bob"]}}, reads=[["user_group", "gid_zero_accounts"]]))
put("5.4.2.3", "only GID 0 group", custom(
    [f'is_array({UG}gid_zero_groups)', f'count({UG}gid_zero_groups) > 0'],
    [[f'count([g | some g in {UG}gid_zero_groups; g != "root"]) == 0']], "user_group.gid_zero_groups",
    {"user_group": {"gid_zero_groups": ["root"]}}, {"user_group": {"gid_zero_groups": ["root", "wheel2"]}},
    reads=[["user_group", "gid_zero_groups"]]))
put("5.4.2.4", "root account access", custom(
    [f'text_collected({UG}root_password_status)'], [[f'regex.match(`^(P|L)`, {UG}root_password_status)']],
    "user_group.root_password_status", {"user_group": {"root_password_status": "PS"}},
    {"user_group": {"root_password_status": "NP"}}, reads=[["user_group", "root_password_status"]]))
put("5.4.2.5", "root path integrity", custom(
    [f'text_collected({UG}root_path)', f'is_array({UG}root_path_nonexistent_dirs)', f'is_array({UG}root_path_writable_dirs)',
     f'is_array({UG}root_path_entries)', f'count({UG}root_path_entries) > 0'],
    [['root_path_text_ok', f'count({UG}root_path_nonexistent_dirs) == 0', f'count({UG}root_path_writable_dirs) == 0',
      f'every d in {UG}root_path_entries {{ d.owner == "root"; mode_within(d.mode, "0755") }}']],
    "user_group.root_path/root_path_nonexistent_dirs/root_path_writable_dirs/root_path_entries",
    {"user_group": {"root_path": "/usr/sbin:/usr/bin", "root_path_nonexistent_dirs": [], "root_path_writable_dirs": [],
                    "root_path_entries": [{"path": "/usr/bin", "owner": "root", "mode": "0755"}]}},
    {"user_group": {"root_path": "/usr/sbin:/usr/bin:.", "root_path_nonexistent_dirs": [], "root_path_writable_dirs": [],
                    "root_path_entries": [{"path": "/usr/bin", "owner": "root", "mode": "0755"}]}},
    {"user_group": {"root_path": "/usr/sbin:/usr/bin", "root_path_nonexistent_dirs": [], "root_path_writable_dirs": []}},
    reads=[["user_group", "root_path"], ["user_group", "root_path_nonexistent_dirs"],
           ["user_group", "root_path_writable_dirs"], ["user_group", "root_path_entries"]]))
put("5.4.2.6", "root user umask", custom(
    [f'text_collected({UG}root_umask)'], [[f'umask_ok({UG}root_umask)']], "user_group.root_umask",
    {"user_group": {"root_umask": "0027"}}, {"user_group": {"root_umask": "0022"}}, reads=[["user_group", "root_umask"]]))
put("5.4.2.7", "system accounts do not have a valid login shell", custom(
    [f'is_array({PA}system_accounts_with_login_shell)'], [[f'count({PA}system_accounts_with_login_shell) == 0']],
    "pam.system_accounts_with_login_shell", {"pam": {"system_accounts_with_login_shell": []}},
    {"pam": {"system_accounts_with_login_shell": ["daemon"]}}, reads=[["pam", "system_accounts_with_login_shell"]]))
put("5.4.2.8", "without a valid login shell are locked", custom(
    [f'is_array({UG}unlocked_accounts_without_shell)'], [[f'count({UG}unlocked_accounts_without_shell) == 0']],
    "user_group.unlocked_accounts_without_shell", {"user_group": {"unlocked_accounts_without_shell": []}},
    {"user_group": {"unlocked_accounts_without_shell": ["bin"]}}, reads=[["user_group", "unlocked_accounts_without_shell"]]))
put("5.4.3.1", "nologin is not listed", text_has(["user_group", "shells_content"], "/nologin\\b", False,
                                                  "/bin/bash\n/bin/sh", "/bin/bash\n/usr/sbin/nologin"))
put("5.4.3.2", "shell timeout", custom(
    [f'{PA}tmout != null', f'is_boolean({PA}tmout_readonly_exported)'],
    [[f'num({PA}tmout) > 0', f'num({PA}tmout) <= 900', f'{PA}tmout_readonly_exported == true']],
    "pam.tmout/tmout_readonly_exported", {"pam": {"tmout": 900, "tmout_readonly_exported": True}},
    {"pam": {"tmout": 0, "tmout_readonly_exported": True}}, {"pam": {"tmout": 900}},
    reads=[["pam", "tmout"], ["pam", "tmout_readonly_exported"]]))
put("5.4.3.3", "umask is configured in /etc/login.defs", custom(
    [f'{PA}login_defs_umask != null'], [[f'umask_ok(sprintf("%v", [{PA}login_defs_umask]))']], "pam.login_defs_umask",
    {"pam": {"login_defs_umask": "027"}}, {"pam": {"login_defs_umask": "022"}}, reads=[["pam", "login_defs_umask"]]))
put("5.4.3.4", "default user umask is configured ", custom(
    [f'{PA}default_umask != null'], [[f'umask_ok(sprintf("%v", [{PA}default_umask]))']], "pam.default_umask",
    {"pam": {"default_umask": "027"}}, {"pam": {"default_umask": "022"}}, reads=[["pam", "default_umask"]]))

# ── 6.1 logging (v2.0.0 journald.* / rsyslog.* / log_file_permissions.analysis) ──
put("6.1.1.1", "journald service is active", boolean(["journald", "service_active"], True))
put("6.1.1.2", "systemd-journal-remote", units_off(["systemd-journal-remote.socket", "systemd-journal-remote.service"]))
put("6.1.1.3", "send logs to rsyslog", custom(
    ['is_string(input.journald.forward_to_syslog)'], [['lower(input.journald.forward_to_syslog) == "yes"']],
    "journald.forward_to_syslog", {"journald": {"forward_to_syslog": "yes"}}, {"journald": {"forward_to_syslog": "no"}},
    reads=[["journald", "forward_to_syslog"]]))
put("6.1.1.4", "log file rotation", attest("journald_rotation_reviewed"))
put("6.1.1.5", "log file access", attest("journald_file_access_reviewed"))
put("6.1.2.1", "rsyslog is installed", boolean(["rsyslog", "installed"], True))
put("6.1.2.2", "rsyslog service is enabled", custom(
    ['is_boolean(input.rsyslog.service_enabled)', 'is_boolean(input.rsyslog.service_active)'],
    [['input.rsyslog.service_enabled == true', 'input.rsyslog.service_active == true']],
    "rsyslog.service_enabled/service_active", {"rsyslog": {"service_enabled": True, "service_active": True}},
    {"rsyslog": {"service_enabled": False, "service_active": True}},
    reads=[["rsyslog", "service_enabled"], ["rsyslog", "service_active"]]))
# file_create_mode absent while rsyslog is installed: rsyslog's 0644 default applies, which fails
put("6.1.2.3", "creation mode", custom(
    ['input.rsyslog.installed == true'], [['mode_within(input.rsyslog.file_create_mode, "0640")']],
    "rsyslog.installed/file_create_mode", {"rsyslog": {"installed": True, "file_create_mode": "0640"}},
    {"rsyslog": {"installed": True, "file_create_mode": "0644"}}, reads=[["rsyslog", "installed"], ["rsyslog", "file_create_mode"]]))
RNI = (["rsyslog", "installed"], False, "rsyslog.installed")
_rc = [{"path": "/etc/rsyslog.conf", "mode": "0640", "owner": "root", "group": "root"}]
put("6.1.2.4", "rsyslog configuration files", gate(access_list(["rsyslog", "config_files"], "0640", ["root"], ["root"],
                                                               sample="/etc/rsyslog.conf"), *RNI))
put("6.1.2.5", "/etc/rsyslog.d directory", gate(access(["rsyslog", "d_dir", "mode"], ["rsyslog", "d_dir", "owner"],
                                                       ["rsyslog", "d_dir", "group"], "0755", ["root"], ["root"]), *RNI))
put("6.1.2.6", "rsyslog work directory", gate(access(["rsyslog", "work_dir", "mode"], ["rsyslog", "work_dir", "owner"],
                                                     ["rsyslog", "work_dir", "group"], "0700", ["root"], ["root"]), *RNI))
RC = ["rsyslog", "config_raw"]   # new (the mapper already reads it): effective rsyslog configuration text
put("6.1.2.7", "not configured as a log server", gate(text_has(
    RC, "(?i)(^module\\(load=\"?(imtcp|imudp|imptcp|imrelp)\"?\\)|^input\\(type=\"?(imtcp|imudp|imptcp|imrelp)\"?|"
        "^\\$ModLoad\\s+(imtcp|imudp|imptcp|imrelp)\\b|^\\$(InputTCPServerRun|UDPServerRun|InputPTCPServerRun|InputRELPServerRun)\\b)",
    False, "*.* /var/log/messages", 'module(load="imtcp")'), *RNI))
put("6.1.2.8", "logging rules", attest("rsyslog_logging_rules_reviewed"))
put("6.1.2.9", "remote log host", attest("rsyslog_remote_host_configured"))
put("6.1.2.10", "TLS network stream driver", custom(
    ['is_boolean(input.packages["rsyslog-gnutls"])', 'is_boolean(input.packages["rsyslog-openssl"])'],
    [['input.packages["rsyslog-gnutls"] == true'], ['input.packages["rsyslog-openssl"] == true']],
    "packages[rsyslog-gnutls,rsyslog-openssl]", {"packages": {"rsyslog-gnutls": True, "rsyslog-openssl": False}},
    {"packages": {"rsyslog-gnutls": False, "rsyslog-openssl": False}},
    reads=[["packages", "rsyslog-gnutls"], ["packages", "rsyslog-openssl"]]))
put("6.1.2.11", "forwarding uses TLS", custom(
    [f'text_collected({ref(RC)})'], [['some a in rsyslog_actions', 'tls_forward_ok(a)']], dotted(RC),
    nest(RC, 'action(type="omfwd" target="loghost.example.com" port="6514" protocol="tcp"\n  StreamDriver="gtls" StreamDriverMode="1" StreamDriverAuthMode="x509/name")'),
    nest(RC, 'action(type="omfwd" target="loghost.example.com" port="514" protocol="tcp")'), reads=[RC]))
put("6.1.2.12", "TLS CA certificate", attest("rsyslog_tls_ca_configured"))
put("6.1.2.13", "TLS client certificate", attest("rsyslog_tls_client_cert_configured"))
put("6.1.2.14", "logrotate", attest("logrotate_reviewed"))
LF = ["log_file_permissions", "analysis"]
put("6.1.3.1", "access to all logfiles", custom(
    [f'is_array({ref(LF)})', f'count({ref(LF)}) > 0'], [[f'every f in {ref(LF)} {{ log_file_ok(f) }}']], dotted(LF),
    nest(LF, [{"file": "/var/log/messages", "actual_mode": "0600", "actual_owner": "root", "actual_group": "root"},
              {"file": "/var/log/wtmp", "actual_mode": "0664", "actual_owner": "root", "actual_group": "utmp"}]),
    nest(LF, [{"file": "/var/log/messages", "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"}]),
    reads=[LF]))

# ── 6.2 auditd (v2.0.0 auditd.config / active_rules / rules / log_permissions / log_dir_permissions /
#    config_permissions, packages.audit, services["auditd.service"], bootloader.grub_config) ──
put("6.2.1.1", "auditd packages", custom(
    ['is_boolean(input.packages.audit)', 'is_boolean(input.packages["audit-libs"])'],
    [['input.packages.audit == true', 'input.packages["audit-libs"] == true']], "packages[audit,audit-libs]",
    {"packages": {"audit": True, "audit-libs": True}}, {"packages": {"audit": True, "audit-libs": False}},
    reads=[["packages", "audit"], ["packages", "audit-libs"]]))
put("6.2.1.2", "prior to auditd", custom([f'text_collected({ref(GRUB)})'], [['grub_cmdline_has(`^audit=1$`)']], dotted(GRUB),
                                         nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit=1"'),
                                         nest(GRUB, 'GRUB_CMDLINE_LINUX="quiet audit=10"'), reads=[GRUB]))
put("6.2.1.3", "audit_backlog_limit", custom(
    [f'text_collected({ref(GRUB)})'],
    [['some t in grub_cmdline_tokens', 'startswith(t, "audit_backlog_limit=")', 'num(substring(t, 20, -1)) >= 8192']],
    dotted(GRUB), nest(GRUB, 'GRUB_CMDLINE_LINUX="audit=1 audit_backlog_limit=8192"'),
    nest(GRUB, 'GRUB_CMDLINE_LINUX="audit=1 audit_backlog_limit=64"'), reads=[GRUB]))
put("6.2.1.4", "auditd service is enabled", unit_running("auditd.service"))
AC = ["auditd", "config"]
put("6.2.2.1", "log storage size", conf(AC, "max_log_file", 'num(v) > 0', "max_log_file = 8", "max_log_file = x"))
put("6.2.2.2", "not automatically deleted", conf(AC, "max_log_file_action", 'lower(v) == "keep_logs"',
                                                 "max_log_file_action = keep_logs", "max_log_file_action = rotate"))
put("6.2.2.3", "disabled when audit logs are full", custom(
    [f'text_collected({ref(AC)})'],
    [['lower(conf_last(input.auditd.config, "disk_full_action")) in {"halt", "single"}',
      'lower(conf_last(input.auditd.config, "disk_error_action")) in {"syslog", "single", "halt"}']],
    dotted(AC), nest(AC, "disk_full_action = halt\ndisk_error_action = syslog"),
    nest(AC, "disk_full_action = suspend\ndisk_error_action = syslog"), reads=[AC]))
put("6.2.2.4", "low on space", custom(
    [f'text_collected({ref(AC)})'],
    [['lower(conf_last(input.auditd.config, "space_left_action")) in {"email", "exec", "single", "halt"}',
      'lower(conf_last(input.auditd.config, "admin_space_left_action")) in {"single", "halt"}']],
    dotted(AC), nest(AC, "space_left_action = email\nadmin_space_left_action = single"),
    nest(AC, "space_left_action = ignore\nadmin_space_left_action = single"), reads=[AC]))
AR = ["auditd", "active_rules"]
W = "-a always,exit -S all -F path={p} -F perm={perm} -F key=k"
S64 = "-a always,exit -F arch=b64 -S {s} {extra}-F key=k"
AUID = "-F auid>=1000 -F auid!=unset "


def audit(watches=(), syscalls=()):
    """watches: [(path, perm)]; syscalls: [(names, [field kinds])]. Rules are matched on tokens."""
    conds, good = [], []
    for p, perm in watches:
        conds.append(f'audit_watch({json.dumps(p)}, {json.dumps(perm)})')
        good.append(W.format(p=p, perm=perm))
    for names, fields in syscalls:
        conds.append(f'every s in {json.dumps(names)} {{ audit_syscall(s, {json.dumps(fields)}) }}')
        extra = "".join({"auid": AUID, "euid": "-C euid!=uid -F auid!=unset ", "eacces": "-F exit=-EACCES " + AUID,
                         "eperm": "-F exit=-EPERM " + AUID}[f] for f in fields)
        good.append(S64.format(s=",".join(names), extra=extra))
    return custom([f'is_array({ref(AR)})', f'count({ref(AR)}) > 0'], [conds], dotted(AR), nest(AR, good),
                  nest(AR, good[:-1] or ["-e 2"]), reads=[AR])


SU_ = ["sudo", "sudoers_content"]
put("6.2.3.1", "/etc/sudoers file", audit([("/etc/sudoers", "wa"), ("/etc/sudoers.d", "wa")]))
put("6.2.3.2", "actions as another user", audit(syscalls=[(["execve", "execveat"], ["euid"])]))
put("6.2.3.3", "modify the sudo log file", custom(
    # sudoers collected with no `Defaults logfile` is a finding (nothing to watch), not missing evidence
    [f'is_array({ref(AR)})', f'count({ref(AR)}) > 0', f'text_collected({ref(SU_)})'],
    [['count(sudo_logfiles) > 0', 'every f in sudo_logfiles { audit_watch(f, "wa") }']],
    "auditd.active_rules/sudo.sudoers_content",
    merge(nest(AR, [W.format(p="/var/log/sudo.log", perm="wa")]), nest(SU_, 'Defaults logfile="/var/log/sudo.log"')),
    merge(nest(AR, [W.format(p="/var/log/other.log", perm="wa")]), nest(SU_, 'Defaults logfile="/var/log/sudo.log"')),
    nest(AR, [W.format(p="/var/log/sudo.log", perm="wa")]), reads=[AR, SU_]))
put("6.2.3.4", "date and time", audit([("/etc/localtime", "wa")], [(["adjtimex", "settimeofday", "clock_settime"], [])]))
put("6.2.3.5", "sethostname and setdomainname", audit(syscalls=[(["sethostname", "setdomainname"], [])]))
put("6.2.3.6", "/etc/issue and /etc/issue.net", audit([("/etc/issue", "wa"), ("/etc/issue.net", "wa")]))
put("6.2.3.7", "/etc/hosts and /etc/hostname", audit([("/etc/hosts", "wa"), ("/etc/hostname", "wa")]))
put("6.2.3.8", "/etc/sysconfig/network", audit([("/etc/sysconfig/network", "wa"),
                                                ("/etc/NetworkManager/system-connections", "wa")]))
put("6.2.3.9", "/etc/NetworkManager directory", audit([("/etc/NetworkManager", "wa")]))
PC = ["auditd", "privileged_commands"]   # new: setuid/setgid binaries on local filesystems
put("6.2.3.10", "privileged commands", custom(
    [f'is_array({ref(AR)})', f'count({ref(AR)}) > 0', f'is_array({ref(PC)})', f'count({ref(PC)}) > 0'],
    [[f'every c in {ref(PC)} {{ audit_watch(c, "x") }}']], "auditd.active_rules/privileged_commands",
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su"]}},
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su", "/usr/bin/passwd"]}},
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")]}}, reads=[AR, PC]))
put("6.2.3.11", "unsuccessful file access", audit(syscalls=[(["creat", "open", "openat", "truncate", "ftruncate"], ["eacces"]),
                                                            (["creat", "open", "openat", "truncate", "ftruncate"], ["eperm"])]))
put("6.2.3.12", "modify /etc/group", audit([("/etc/group", "wa")]))
put("6.2.3.13", "modify /etc/passwd", audit([("/etc/passwd", "wa")]))
put("6.2.3.14", "/etc/shadow and /etc/gshadow", audit([("/etc/shadow", "wa"), ("/etc/gshadow", "wa")]))
put("6.2.3.15", "modify /etc/security/opasswd", audit([("/etc/security/opasswd", "wa")]))
put("6.2.3.16", "/etc/nsswitch.conf", audit([("/etc/nsswitch.conf", "wa")]))
put("6.2.3.17", "/etc/pam.conf and /etc/pam.d/", audit([("/etc/pam.conf", "wa"), ("/etc/pam.d", "wa")]))
put("6.2.3.18", "chmod,fchmod,fchmodat", audit(syscalls=[(["chmod", "fchmod", "fchmodat"], ["auid"])]))
put("6.2.3.19", "chown,fchown,lchown,fchownat", audit(syscalls=[(["chown", "fchown", "lchown", "fchownat"], ["auid"])]))
put("6.2.3.20", "setxattr", audit(syscalls=[(["setxattr", "lsetxattr", "fsetxattr", "removexattr", "lremovexattr",
                                              "fremovexattr"], ["auid"])]))
put("6.2.3.21", "file system mounts", audit(syscalls=[(["mount"], ["auid"])]))
put("6.2.3.22", "session initiation", audit([("/var/run/utmp", "wa"), ("/var/log/wtmp", "wa"), ("/var/log/btmp", "wa")]))
put("6.2.3.23", "login and logout", audit([("/var/log/lastlog", "wa"), ("/var/run/faillock", "wa")]))
put("6.2.3.24", "unlink file deletion", audit(syscalls=[(["unlink", "unlinkat"], ["auid"])]))
put("6.2.3.25", "rename file deletion", audit(syscalls=[(["rename", "renameat", "renameat2"], ["auid"])]))
put("6.2.3.26", "Mandatory Access Controls", audit([("/etc/selinux", "wa"), ("/usr/share/selinux", "wa"),
                                                    ("/var/lib/selinux", "wa")]))
put("6.2.3.27", "chcon", audit([("/usr/bin/chcon", "x")]))
put("6.2.3.28", "setfacl", audit([("/usr/bin/setfacl", "x")]))
put("6.2.3.29", "chacl", audit([("/usr/bin/chacl", "x")]))
put("6.2.3.30", "usermod", audit([("/usr/sbin/usermod", "x")]))
put("6.2.3.31", "kernel module loading", audit([("/usr/bin/kmod", "x")]))
put("6.2.3.32", "init_module", audit(syscalls=[(["init_module", "finit_module"], ["auid"])]))
put("6.2.3.33", "delete_module", audit(syscalls=[(["delete_module"], ["auid"])]))
ARD = ["auditd", "rules"]   # on-disk rules (/etc/audit/rules.d)
put("6.2.3.34", "regardless of errors", custom(
    [f'is_array({ref(ARD)})', f'count({ref(ARD)}) > 0'], [['some l in input.auditd.rules', 'regex.match(`^\\s*-c\\s*$`, l)']],
    dotted(ARD), nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 2"]), reads=[ARD]))
put("6.2.3.35", "immutable", custom([f'is_array({ref(ARD)})', f'count({ref(ARD)}) > 0'], [['audit_last_e == "2"']], dotted(ARD),
                                    nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 1"]), reads=[ARD]))
# active_rules and rules are filled from the same collected text by the mapper, so comparing them proves
# nothing: the on-disk/running comparison is its own collected fact (`augenrules --check`)
AG = ["auditd", "augenrules_check"]
put("6.2.3.36", "running and on disk", custom(
    [f'text_collected({ref(AG)})'], [[f'regex.match(`(?i)^\\s*(/usr)?/sbin/augenrules:\\s*no change\\s*$`, {ref(AG)})']],
    dotted(AG), nest(AG, "/sbin/augenrules: No change"), nest(AG, "/usr/sbin/augenrules: Rules have changed and should be updated"),
    reads=[AG]))
put("6.2.4.1", "log file directory mode", custom(
    ['input.auditd.log_dir_permissions.mode != null'], [['mode_within(input.auditd.log_dir_permissions.mode, "0750")']],
    "auditd.log_dir_permissions.mode", {"auditd": {"log_dir_permissions": {"mode": "0750"}}},
    {"auditd": {"log_dir_permissions": {"mode": "0755"}}}, reads=[["auditd", "log_dir_permissions", "mode"]]))
LP = "input.auditd.log_permissions."
put("6.2.4.2", "log files mode", custom([f'{LP}mode != null'], [[f'mode_within({LP}mode, "0640")']], "auditd.log_permissions.mode",
                                        {"auditd": {"log_permissions": {"mode": "0600"}}},
                                        {"auditd": {"log_permissions": {"mode": "0644"}}}, reads=[["auditd", "log_permissions", "mode"]]))
put("6.2.4.3", "log files owner", custom([f'{LP}pw_name != null'], [[f'{LP}pw_name == "root"']], "auditd.log_permissions.pw_name",
                                         {"auditd": {"log_permissions": {"pw_name": "root"}}},
                                         {"auditd": {"log_permissions": {"pw_name": "adm"}}}, reads=[["auditd", "log_permissions", "pw_name"]]))
put("6.2.4.4", "log files group owner", custom(
    [f'{LP}gr_name != null', f'text_collected({ref(AC)})'],
    [[f'{LP}gr_name in {{"root", "audit"}}', 'audit_log_group_ok']], "auditd.log_permissions.gr_name/auditd.config",
    {"auditd": {"log_permissions": {"gr_name": "root"}, "config": "log_group = root"}},
    {"auditd": {"log_permissions": {"gr_name": "root"}, "config": "log_group = adm"}},
    {"auditd": {"log_permissions": {"gr_name": "root"}}}, reads=[["auditd", "log_permissions", "gr_name"], AC]))
for rid, kw, path, cond, key, badv in [
        ("6.2.4.5", "configuration files mode", ["auditd", "config_permissions"], 'mode_within(f.stat.mode, "0640")', "mode", "0644"),
        ("6.2.4.6", "configuration files owner", ["auditd", "config_permissions"], 'f.stat.pw_name == "root"', "pw_name", "adm"),
        ("6.2.4.7", "configuration files group owner", ["auditd", "config_permissions"], 'f.stat.gr_name == "root"', "gr_name", "adm"),
        ("6.2.4.8", "audit tools mode", ["auditd", "tools_permissions"], 'mode_within(f.stat.mode, "0755")', "mode", "0777"),
        ("6.2.4.9", "audit tools owner", ["auditd", "tools_permissions"], 'f.stat.pw_name == "root"', "pw_name", "adm"),
        ("6.2.4.10", "audit tools group owner", ["auditd", "tools_permissions"], 'f.stat.gr_name == "root"', "gr_name", "adm")]:
    st = {"path": "/etc/audit/auditd.conf" if "config" in path[1] else "/usr/sbin/auditctl",
          "mode": "0640" if "config" in path[1] else "0755", "pw_name": "root", "gr_name": "root"}
    k = "stat." + key
    put(rid, kw, custom([f'is_array({ref(path)})', f'count({ref(path)}) > 0',
                         f'every f in {ref(path)} {{ f.{k} != null }}'],
                        [[f'every f in {ref(path)} {{ {cond} }}']], f"{dotted(path)}[].{k}",
                        nest(path, [{"stat": st}]), nest(path, [{"stat": dict(st, **{key: badv})}]), reads=[path]))

# ── 6.3 integrity checking (v2.0.0 boot_security.aide_* fields) ──
put("6.3.1", "AIDE is installed", boolean(["boot_security", "aide_installed"], True))
put("6.3.2", "regularly checked", custom(
    [f'is_boolean({BS}aide_cron_configured)', f'is_boolean({BS}aide_timer_configured)'],
    [[f'{BS}aide_cron_configured == true'], [f'{BS}aide_timer_configured == true']],
    "boot_security.aide_cron_configured/aide_timer_configured",
    {"boot_security": {"aide_cron_configured": False, "aide_timer_configured": True}},
    {"boot_security": {"aide_cron_configured": False, "aide_timer_configured": False}},
    reads=[["boot_security", "aide_cron_configured"], ["boot_security", "aide_timer_configured"]]))
AIDE_OPTS = "p+i+n+u+g+s+b+acl+xattrs+sha512"
TOOLS = ["auditctl", "auditd", "ausearch", "aureport", "autrace", "augenrules"]
AT = ["boot_security", "aide_audit_tool_options"]   # new: tool name -> selection attributes in aide.conf
put("6.3.3", "integrity of audit tools", custom(
    [f'is_object({ref(AT)})'], [[f'every t in {json.dumps(TOOLS)} {{ aide_opts_ok({ref(AT)}[t]) }}']], dotted(AT),
    nest(AT, {t: AIDE_OPTS for t in TOOLS}), nest(AT, dict({t: AIDE_OPTS for t in TOOLS[:-1]}, augenrules="p+i+n")),
    reads=[AT]))
RU = ["boot_security", "aide_report_urls"]   # new: report_url values of the effective aide.conf
put("6.3.4", "report_url", custom([f'is_array({ref(RU)})'], [[f'some u in {ref(RU)}', 'startswith(lower(u), "syslog:")']],
                                  dotted(RU), nest(RU, ["syslog:LOG_LOCAL0"]), nest(RU, ["file:/var/log/aide.log"]), reads=[RU]))

# ── 7.1 system files (v2.0.0 critical_files.analysis {file, exists, actual_mode, actual_owner, actual_group},
#    world_writable, sticky_bit, unowned_files, ungrouped_files) ──
for rid, kw, path, mx in [("7.1.1", "/etc/passwd is", "/etc/passwd", "0644"), ("7.1.2", "/etc/passwd-", "/etc/passwd-", "0644"),
                          ("7.1.3", "/etc/group is", "/etc/group", "0644"), ("7.1.4", "/etc/group-", "/etc/group-", "0644"),
                          ("7.1.5", "/etc/shadow is", "/etc/shadow", "0000"), ("7.1.6", "/etc/shadow-", "/etc/shadow-", "0000"),
                          ("7.1.7", "/etc/gshadow is", "/etc/gshadow", "0000"), ("7.1.8", "/etc/gshadow-", "/etc/gshadow-", "0000"),
                          ("7.1.9", "access to /etc/shells", "/etc/shells", "0644")]:
    put(rid, kw, crit(path, mx, ["root"]))
_op = crit("/etc/security/opasswd", "0600", ["root"], optional=True)
_op2 = crit("/etc/security/opasswd.old", "0600", ["root"], optional=True)
_op.ev_alts = [_op.ev_alts[0] + _op2.ev_alts[0]]
_op.passes = [a + b for a in _op.passes for b in _op2.passes]
_op.field = "critical_files.analysis[/etc/security/opasswd, /etc/security/opasswd.old]"
_old = {"file": "/etc/security/opasswd.old", "exists": False}
_op.tp = {"critical_files": {"analysis": _op.tp["critical_files"]["analysis"] + [_old]}}
_op.tf = {"critical_files": {"analysis": _op.tf["critical_files"]["analysis"] + [_old]}}
put("7.1.10", "access to /etc/security/opasswd", _op)
put("7.1.11", "world writable files", custom(
    ['is_array(input.world_writable.files)', 'not input.world_writable.truncated == true',
     'is_array(input.sticky_bit.world_writable_without_sticky)'],
    [['count(input.world_writable.files) == 0', 'object.get(input.world_writable, "count", 0) == 0',
      'count(input.sticky_bit.world_writable_without_sticky) == 0']],
    "world_writable.files/sticky_bit.world_writable_without_sticky",
    {"world_writable": {"files": [], "count": 0}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": ["/opt/app/data.txt"], "count": 1}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": [], "count": 0}},
    reads=[["world_writable", "files"], ["sticky_bit", "world_writable_without_sticky"]]))
put("7.1.12", "without an owner and a group", custom(
    ['is_array(input.unowned_files.files)', 'is_array(input.ungrouped_files.files)'],
    [['count(input.unowned_files.files) == 0', 'count(input.ungrouped_files.files) == 0',
      'object.get(input.unowned_files, "count", 0) == 0', 'object.get(input.ungrouped_files, "count", 0) == 0']],
    "unowned_files.files/ungrouped_files.files",
    {"unowned_files": {"files": [], "count": 0}, "ungrouped_files": {"files": [], "count": 0}},
    {"unowned_files": {"files": [], "count": 0}, "ungrouped_files": {"files": ["/srv/orphan"], "count": 1}},
    reads=[["unowned_files", "files"], ["ungrouped_files", "files"]]))
put("7.1.13", "SUID and SGID", attest("suid_sgid_files_reviewed"))


# ── 7.2 local users and groups (v2.0.0 user_group.* offender lists / maps the mapper emits) ──
def ug_empty(rid, kw, keys, bad_key, bad_val):
    ev = [f'{UG}{k} != null' for k in keys]
    good = {k: ({} if k.startswith("duplicate_") else []) for k in keys}
    return put(rid, kw, custom(ev, [[f'count({UG}{k}) == 0' for k in keys]], "user_group." + "/".join(keys),
                               {"user_group": good}, {"user_group": dict(good, **{bad_key: bad_val})},
                               reads=[["user_group", k] for k in keys]))


ug_empty("7.2.1", "shadowed passwords", ["users_without_shadow"], "users_without_shadow", ["bob"])
ug_empty("7.2.2", "password fields are not empty", ["users_with_empty_password"], "users_with_empty_password", ["bob"])
ug_empty("7.2.3", "groups in /etc/passwd exist", ["passwd_groups_not_in_group_file"], "passwd_groups_not_in_group_file", ["bob:4242"])
ug_empty("7.2.4", "duplicate UIDs", ["duplicate_uids"], "duplicate_uids", {"1000": ["alice", "toor"]})
ug_empty("7.2.5", "duplicate GIDs", ["duplicate_gids"], "duplicate_gids", {"100": ["users", "staff"]})
ug_empty("7.2.6", "duplicate user names", ["duplicate_usernames"], "duplicate_usernames", {"alice": 2})
ug_empty("7.2.7", "duplicate group names", ["duplicate_groupnames"], "duplicate_groupnames", {"users": 2})
ug_empty("7.2.8", "home directories", ["users_without_home_directory", "users_not_owning_home", "home_dirs_too_permissive"],
         "home_dirs_too_permissive", [{"user": "alice", "home": "/home/alice", "mode": "0755"}])
ug_empty("7.2.9", "dot files", ["users_with_netrc", "users_with_forward", "users_with_rhosts", "dotfiles_group_or_world_writable"],
         "users_with_rhosts", [{"user": "alice", "path": "/home/alice/.rhosts"}])

# Offender lists and empty maps: a collected empty list/object is a genuine pass (the collector always emits them).
for _rid in ["1.2.1.2", "2.3.3", "7.1.11", "7.1.12", "5.4.2.7", "5.4.2.8", "2.1.19"] + [f"7.2.{i}" for i in range(1, 10)]:
    L[_rid].hollow_ok = True

# ── helpers used by one module only ──
EXTRA = {
    "initial_setup": r'''
# CURRENT.pol: a global hash/sign line, or a scoped one outside the legacy-interop scopes, that enables SHA1
sha1_legacy_scope(k) if regex.match(`(?i)^(hash|sign)@(rpm|dnssec|smime|pkcs12)`, k)

sha1_enabled if {
	some l in active_lines(input.crypto_policy.current_pol_raw)
	some m in regex.find_all_string_submatch_n(`^((hash|sign)(@\S+)?)\s*=\s*(.*)$`, l, 1)
	not sha1_legacy_scope(m[1])
	regex.match(`(?i)\bSHA-?1\b`, m[4])
}

# sha1_in_certs absent: the default 0 (SHA1 signatures disallowed in certificates) applies
sha1_in_certs_off if not conf_last(input.crypto_policy.current_pol_raw, "sha1_in_certs")

sha1_in_certs_off if num(conf_last(input.crypto_policy.current_pol_raw, "sha1_in_certs")) == 0
''',
    "boot_security": r'''
# helper rules: `not is_boolean(input.missing)` is undefined and never fires
rescue_auth_required_collected if is_boolean(input.boot_security.rescue_auth_required)

emergency_auth_required_collected if is_boolean(input.boot_security.emergency_auth_required)

kernel_args_collected if is_array(input.boot_security.kernel_args)

grub_max(p) := "0700" if startswith(p, "/boot/efi/")

grub_max(p) := "0600" if not startswith(p, "/boot/efi/")

# grub files under /boot/efi (vfat) may be 0700; everything else 0600
grub_file_ok(f) if {
	mode_within(f.mode, grub_max(f.path))
	f.owner == "root"
	f.group == "root"
}

aide_opts_ok(s) if {
	opts := {trim_space(o) | some o in split(s, "+")}
	every want in ["p", "i", "n", "u", "g", "s", "b", "acl", "xattrs", "sha512"] { want in opts }
}
''',
    "selinux": r'''
# configured mode (/etc/selinux/config SELINUX=): selinux.config_mode, else the mapper's ansible_selinux
# helper rule: `not is_string(input.missing)` is undefined and never fires
selinux_own_config_mode if is_string(input.selinux.config_mode)

selinux_config_mode := input.selinux.config_mode if selinux_own_config_mode

selinux_config_mode := input.ansible_facts.ansible_selinux.config_mode if {
	not selinux_own_config_mode
	is_string(input.ansible_facts.ansible_selinux.config_mode)
}
''',
    "services": r'''
# postfix inet_interfaces (when collected): "all" or a non-loopback value listens off-host
mta_interfaces_open if regex.match(`\ball\b`, lower(input.mta.inet_interfaces))

mta_interfaces_open if {
	is_string(input.mta.inet_interfaces)
	not regex.match(`(0\.0\.0\.0|::1|loopback-only|localhost)`, lower(input.mta.inet_interfaces))
}
''',
    "network": r'''
virtual_iface(i) if regex.match(`^(lo|virbr.*|docker.*|veth.*|br-.*|podman.*|cni.*)$`, i)

# a zone serving only loopback/virtual interfaces is out of scope
loopback_zone(z) if every i in object.get(z, "interfaces", []) { virtual_iface(i) }

zone_target_ok(z) if loopback_zone(z)

zone_target_ok(z) if {
	not loopback_zone(z)
	is_string(z.target)
	upper(z.target) != "ACCEPT"
	lower(z.target) == lower(z.permanent_target)
}

# ICMP timestamps: only a zone whose target is "default" needs the explicit block
zone_icmp_ok(z) if lower(object.get(z, "permanent_target", "")) in {"accept", "drop", "%%reject%%", "reject"}

zone_icmp_ok(z) if {
	lower(z.permanent_target) == "default"
	z.icmp_block_inversion == false
	every t in ["timestamp-request", "timestamp-reply"] { t in z.icmp_blocks }
}

zone_icmp_ok(z) if {
	lower(z.permanent_target) == "default"
	z.icmp_block_inversion == true
	every t in ["timestamp-request", "timestamp-reply"] { not t in z.icmp_blocks }
}
''',
    "ssh": r'''
maxstartups_ok(v) if {
	parts := split(v, ":")
	num(parts[0]) <= 10
	ms_le(parts, 1, 30)
	ms_le(parts, 2, 60)
}

ms_le(parts, i, _) if count(parts) <= i

ms_le(parts, i, lim) if num(parts[i]) <= lim


# {mode, pw_name, gr_name}: within max, root:root
perm_ok(f, max) if {
	mode_within(f.mode, max)
	f.pw_name == "root"
	f.gr_name == "root"
}

host_private_key_ok(s) if {
	s.pw_name == "root"
	s.gr_name == "root"
	mode_within(s.mode, "0600")
}

host_private_key_ok(s) if {
	s.pw_name == "root"
	s.gr_name == "ssh_keys"
	mode_within(s.mode, "0640")
}

host_public_key_ok(k) if {
	mode_within(k.mode, "0644")
	k.pw_name == "root"
	k.gr_name in {"root", "ssh_keys"}
}

# the banner file's text: its own flag, or the matching /etc/issue(.net) flag of initial_setup
# helper rule: `not is_boolean(input.missing)` is undefined and never fires
banner_flag_collected if is_boolean(input.ssh.banner_contains_sensitive_info)

banner_text_collected if banner_flag_collected

banner_text_collected if {
	sshd_value("banner") == "/etc/issue.net"
	is_boolean(input.initial_setup.issue_net_contains_sensitive_info)
}

banner_text_collected if {
	sshd_value("banner") == "/etc/issue"
	is_boolean(input.initial_setup.issue_contains_sensitive_info)
}

banner_text_ok if input.ssh.banner_contains_sensitive_info == false

banner_text_ok if {
	not banner_flag_collected
	sshd_value("banner") == "/etc/issue.net"
	input.initial_setup.issue_net_contains_sensitive_info == false
}

banner_text_ok if {
	not banner_flag_collected
	sshd_value("banner") == "/etc/issue"
	input.initial_setup.issue_contains_sensitive_info == false
}

# v2.0.0 compatibility: directive -> value(s); a directive `sshd -T` repeats (hostkey, listenaddress)
# is comma-joined rather than raising an object-key conflict
sshd_config := {k: concat(",", [p[1] | some p in sshd_pairs; p[0] == k]) | some q in sshd_pairs; k := q[0]}

# v2.0.0 compatibility: the 5.1.1 / 5.1.3 / 5.1.4 (file access) violations
permission_violations := {v | some v in violations; regex.match(`^5\.1\.[1-4]: `, v)}
''',
    "sudo": r'''
sudo_defaults_has(re) if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	regex.match(re, l)
}

sudo_logfile if sudo_defaults_has(`(^|[\s,])logfile\s*=`)

sudo_timeouts := [num(m[1]) |
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`timestamp_timeout\s*=\s*(-?[0-9]+)`, l, -1)
]

''',
    "pam": r'''
# the mapper's pam.system_auth carries every collected PAM file under "=== <name> ===" headers;
# system-auth is the part under its own header (the whole text when there are no headers)
pam_system_auth_text := t if {
	parts := split(input.pam.system_auth, "=== system-auth ===")
	count(parts) > 1
	t := split(parts[1], "\n===")[0]
}

pam_system_auth_text := input.pam.system_auth if not contains(input.pam.system_auth, "=== system-auth ===")

# active lines of a PAM file for one module type (auth/account/password/session)
pam_lines(t, ty, mod) := [l |
	some l in active_lines(t)
	regex.match(concat("", [`^-?`, ty, `\s+(\[[^\]]*\]|\S+)\s+`, replace(mod, ".", `\.`), `(\s|$)`]), l)
]

pam_line(t, ty, mod, re) if {
	some l in pam_lines(t, ty, mod)
	regex.match(re, l)
}

# module arguments in system-auth or password-auth that undercut the configuration file
pam_args_bad(mod, re) := [l |
	some t in [pam_system_auth_text, input.pam.password_auth]
	some ty in ["auth", "password"]
	some l in pam_lines(t, ty, mod)
	regex.match(concat("", [`\s`, re]), l)
]

unlock_ok(v) if num(v) == 0

unlock_ok(v) if num(v) >= 900

root_lockout_on if regex_in_text(input.pam.faillock_conf_raw, `(?i)^even_deny_root\b`)

root_lockout_on if conf_last(input.pam.faillock_conf_raw, "root_unlock_time")

root_unlock_short if num(conf_last(input.pam.faillock_conf_raw, "root_unlock_time")) < 60

# rpm version-release at least min ("1.5.1-14.el9" >= "1.5.1-14"): numeric fields compared in order
rpm_version_ge(v, min) if {
	a := [to_number(m[0]) | some m in regex.find_all_string_submatch_n(`[0-9]+`, v, -1)]
	b := [to_number(m[0]) | some m in regex.find_all_string_submatch_n(`[0-9]+`, min, -1)]
	count(a) >= count(b)
	not ver_lt(a, b)
}

# a < b on the first count(b) fields
ver_lt(a, b) if {
	some i in numbers.range(0, count(b) - 1)
	a[i] < b[i]
	every j in numbers.range(0, i - 1) { a[j] == b[j] }
}

today_days := time.now_ns() / 86400000000000

root_path_text_ok if {
	p := input.user_group.root_path
	not contains(p, "::")
	not endswith(trim_space(p), ":")
	not startswith(trim_space(p), ":")
	not "." in split(p, ":")
	not "" in split(p, ":")
}

umask_values(text) := [m[1] |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(`(?i)^umask\s+(\S+)`, l, 1)
]

# at least as restrictive as 027 (027 octal = 23)
umask_ok(u) if bits.and(octal(u), 23) == 23

umask_ok(u) if {
	regex.match(`^([ugoa]+[=+-][rwx]*,?)+$`, u)
	not regex.match(`g=[rx]*w`, u)
	not regex.match(`o=[rwx]+`, u)
}

''',
    "logging": r'''
# 6.1.3.1: per-file expectations by file name, from the benchmark's audit script (max mode, owners, groups)
log_rule(b) := ["0664", {"root"}, {"root", "utmp"}] if {
	regex.match(`^((lastlog|wtmp|btmp)([.-].*)?|README)$`, b)
} else := ["0644", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(cloud-init\.log|localmessages|waagent\.log)`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(secure([.-].*)?|auth\.log|syslog|messages)$`, b)
} else := ["0660", {"root", "SSSD"}, {"root", "SSSD"}] if {
	b in {"SSSD", "sssd"}
} else := ["0660", {"root"}, {"root", "gdm", "gdm3"}] if {
	b in {"gdm", "gdm3"}
} else := ["0640", {"root"}, {"root", "systemd-journal"}] if {
	regex.match(`\.journal~?$`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}]

log_file_ok(f) if {
	r := log_rule(regex.replace(f.file, `^.*/`, ""))
	mode_within(f.actual_mode, r[0])
	f.actual_owner in r[1]
	f.actual_group in r[2]
}

# rsyslog action(...) bodies of the collected configuration (comments dropped, lines joined)
rsyslog_actions := [m[1] |
	some m in regex.find_all_string_submatch_n(`(?is)action\s*\(([^)]*)\)`, concat(" ", active_lines(input.rsyslog.config_raw)), -1)
]

tls_forward_ok(a) if {
	regex.match(`(?i)type\s*=\s*"?omfwd"?`, a)
	regex.match(`(?i)StreamDriver(\.Name)?\s*=\s*"?(gtls|ossl)"?`, a)
	regex.match(`(?i)StreamDriverMode\s*=\s*"?1"?`, a)
	regex.match(`(?i)StreamDriverAuthMode\s*=\s*"?(x509/name|x509/fingerprint|x509/certvalid)"?`, a)
	regex.match(`(?i)port\s*=\s*"?6514"?`, a)
}
''',
    "auditd": r'''
audit_rule_lines := [l |
	some r in input.auditd.active_rules
	l := trim_space(r)
	l != ""
	not startswith(l, "#")
]

audit_toks(l) := regex.split(`\s+`, l)

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some t in toks
	t in {concat("", ["path=", p]), concat("", ["dir=", p]), concat("", ["path=", p, "/"]), concat("", ["dir=", p, "/"])}
}

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-w"
	toks[i + 1] in {p, concat("", [p, "/"])}
}

audit_perm(l) := substring(t, 5, -1) if {
	some t in audit_toks(l)
	startswith(t, "perm=")
}

audit_perm(l) := toks[i + 1] if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-p"
}

audit_watch(p, perm) if {
	some l in audit_rule_lines
	audit_path_tok(l, p)
	pv := audit_perm(l)
	every i in numbers.range(0, count(perm) - 1) { contains(pv, substring(perm, i, 1)) }
}

line_syscalls(l) := {x |
	toks := audit_toks(l)
	some i, t in toks
	t == "-S"
	some x in split(toks[i + 1], ",")
}

audit_unset := `auid!=(unset|-1|4294967295)(\s|$)`

audit_field(l, "auid") if {
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "euid") if {
	regex.match(`(euid!=uid|uid!=euid)(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eacces") if {
	regex.match(`exit=-EACCES(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eperm") if {
	regex.match(`exit=-EPERM(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_syscall(s, fields) if {
	some l in audit_rule_lines
	regex.match(`^-a\s+(always,exit|exit,always)(\s|$)`, l)
	s in line_syscalls(l)
	every f in fields { audit_field(l, f) }
}

# the last `-e` line of the on-disk rules decides the enabled flag
audit_e_values := [m[1] |
	some r in input.auditd.rules
	some m in regex.find_all_string_submatch_n(`^\s*-e\s+([0-9])\b`, r, 1)
]

audit_last_e := audit_e_values[count(audit_e_values) - 1] if count(audit_e_values) > 0

sudo_logfiles contains f if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`logfile\s*=\s*("[^"]+"|'[^']+'|[^,\s]+)`, l, -1)
	f := trim(m[1], "\"'")
}

grub_cmdline_tokens := [t |
	some l in active_lines(input.bootloader.grub_config)
	some m in regex.find_all_string_submatch_n(`^GRUB_CMDLINE_LINUX="([^"]*)"`, l, 1)
	some t in regex.split(`\s+`, m[1])
	t != ""
]

grub_cmdline_has(re) if {
	some t in grub_cmdline_tokens
	regex.match(re, t)
}

# log_group unset: auditd's default root applies
audit_log_group_ok if not conf_last(input.auditd.config, "log_group")

audit_log_group_ok if lower(conf_last(input.auditd.config, "log_group")) in {"root", "audit"}
''',
}


# ════════════════════════════════════════════════════════════════════════════════════════════
# workbook
# ════════════════════════════════════════════════════════════════════════════════════════════
def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    recs = {}
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        rows = ws.iter_rows(values_only=True)
        head = next(rows)
        for r in rows:
            d = dict(zip(head, r))
            rid = d.get("Recommendation #")
            if not rid:
                continue
            rid = str(rid).strip()
            if rid not in recs:
                recs[rid] = {"title": " ".join(str(d["Title"]).split()), "assessment": str(d["Assessment Status"]).strip(),
                             "profiles": []}
            recs[rid]["profiles"].append(ws.title)
    comb = wb["Combined Profiles"]
    rows = comb.iter_rows(values_only=True)
    head = next(rows)
    combined = {str(dict(zip(head, r))["Recommendation #"]).strip() for r in rows if dict(zip(head, r))["Recommendation #"]}
    if combined != set(recs):
        raise SystemExit(f"Combined Profiles disagrees with the level sheets: {sorted(combined ^ set(recs))}")
    return recs


def level(profiles, kind):
    lv = [int(m.group(1)) for p in profiles for m in [re.match(r"Level (\d) - (\w+)", p)] if m and m.group(2) == kind]
    return min(lv) if lv else None


def J(x):
    return json.dumps(x, ensure_ascii=False)


def idkey(rid):
    return rid.replace(".", "_")


def sort_key(rid):
    return [int(x) for x in rid.split(".")]


SEC = {n: (f, key, title, pre) for n, f, key, title, pre in SECTIONS}


# ════════════════════════════════════════════════════════════════════════════════════════════
# emit
# ════════════════════════════════════════════════════════════════════════════════════════════
def header(title, checks, entry="none", density="sized", basis="size", extra=""):
    return f'''# METADATA
# title: {J(title)}
# custom:
#   class: security
#   framework: cis_rhel_9
#   source: cis
#   domains: [linux, rhel]
#   coverage:
#     checks: {checks}
#     cited: {checks}
{extra}#     density: {density}
#     basis: {basis}
#     tested: true
#     entrypoint: {entry}
#     guard: self
#     as_of: {AS_OF}
'''


def rule_block(name, alts):
    out = [f"default {name} := false", ""]
    for a in alts:
        out.append(f"{name} if {{")
        out += [f"\t{c}" for c in a]
        out += ["}", ""]
    return out


def profiles_txt(rec):
    return ", ".join(rec["profiles"])


EMPTY_TEXT_OK = ("sysconfig_options)",)  # an empty OPTIONS= is a real value


def evidence_alts(alts):
    """is_string(x) evidence becomes text_collected(x): a blank string is not evidence"""
    return [[("text_collected(" + c[len("is_string("):]) if c.startswith("is_string(") and not c.endswith(EMPTY_TEXT_OK)
             else c for c in a] for a in alts]


def hollow(x):
    """the same shape with every string blanked and every list emptied"""
    if isinstance(x, dict):
        return {k: hollow(v) for k, v in x.items()}
    if isinstance(x, list):
        return []
    if isinstance(x, str):
        return ""
    return x


def emit_control(rid, rec, spec):
    k = idkey(rid)
    out = [f"# {rid} ({rec['assessment']}; {profiles_txt(rec)})", f"# {rec['title']}"]
    if spec.manual:
        out += rule_block(f"p_{k}", [[f'input.attest.{spec.name} == true']])
        msg = f"{rid}: manual review not attested — set attest.{spec.name} to true once reviewed ({rec['title']})"
        out += [f"violations contains {J(msg)} if not p_{k}", ""]
        return out
    alts = getattr(spec, "alts", [])
    ev, passes = list(spec.ev_alts), list(spec.passes)
    if alts:  # a v2.0.0 source decides only when the primary evidence is absent
        out += rule_block(f"ep_{k}", evidence_alts(spec.ev_alts))
        ev += [a_ev for a_ev, _ in alts]
        passes += [[f"not ep_{k}"] + a_p for _, a_p in alts]
    out += rule_block(f"e_{k}", evidence_alts(ev))
    out += rule_block(f"p_{k}", passes)
    miss = f"{rid}: evidence missing — {spec.field} not collected"
    fail = f"{rid}: not met — {rec['title']}"
    out += [f"violations contains {J(miss)} if not e_{k}", "",
            f"violations contains {J(fail)} if {{", f"\te_{k}", f"\tnot p_{k}", "}", ""]
    return out


def emit_module(name, ids, recs):
    f, key, title, _ = SEC[name]
    lines = [header(f"{BENCH} — {title}", len(ids)).rstrip("\n"), f"package {PKG}.{name}", "",
             f"# {BENCH}, {title}.",
             f"# Generated by {GEN} from the CIS workbook; do not edit by hand.",
             "# Each recommendation: e_<id> (facts collected), p_<id> (requirement met). Automated checks fail",
             "# closed with '<id>: evidence missing — <field> not collected'; Manual ones need input.attest.<name>.",
             f"# Roll-up: package {PKG} (section key {key}); profile scoping happens there, this package reports",
             "# every recommendation of the section.", "",
             "import rego.v1", "",
             "default compliant := false", "",
             "compliant if count(violations) == 0", "",
             f"controls := {json.dumps(sorted(ids, key=sort_key))}", "",
             "# v2.0.0 report shape (callers read report.compliant / total_violations / violations)",
             "report := {",
             '\t"compliant": compliant,',
             '\t"total_violations": count(violations),',
             '\t"violations": violations,',
             '\t"controls": controls,',
             '\t"compliance_summary": {"overall_compliant": compliant, "controls_evaluated": count(controls)},',
             '\t"collection_timestamp": object.get(input, ["collection_timestamp"], null),',
             "}", ""]
    lines += HELPERS.strip("\n").split("\n") + [""]
    if name in EXTRA:
        lines += EXTRA[name].strip("\n").split("\n") + [""]
    for rid in ids:
        lines += emit_control(rid, recs[rid], L[rid])
    return "\n".join(lines).rstrip("\n") + "\n"


def merge_good(a, b, path=()):
    """merge two passing inputs: objects recursively, lists concatenated (records of the same user merged),
    configuration texts joined line-wise; a scalar conflict aborts"""
    if isinstance(a, dict) and isinstance(b, dict):
        out = dict(a)
        for k, v in b.items():
            out[k] = merge_good(out[k], v, path + (k,)) if k in out else v
        return out
    if isinstance(a, list) and isinstance(b, list):
        out = list(a)
        for x in b:
            if isinstance(x, dict) and "user" in x and "has_password" in x:
                same = [i for i, y in enumerate(out) if isinstance(y, dict) and y.get("user") == x["user"]]
                if same:
                    out[same[0]] = dict(out[same[0]], **x)
                    continue
            if x not in out:
                out.append(x)
        return out
    if a == b:
        return a
    if path in GOOD_OVERRIDE:
        return GOOD_OVERRIDE[path]
    if isinstance(a, str) and isinstance(b, str):
        return a + "\n" + b
    raise SystemExit(f"good-input conflict at {'.'.join(path)}: {a!r} vs {b!r}")


GOOD_OVERRIDE = {("initial_setup", "crypto_policy"): "DEFAULT:PQ"}


def good_input(ids):
    g = {}
    for rid in ids:
        s = L[rid]
        g = merge_good(g, {"attest": {s.name: True}} if s.manual else s.tp)
    return g


def emit_tests(name, ids):
    lines = [f"package {PKG}.{name}_test", "", "import rego.v1", "",
             f"# Generated by {GEN}: a pass and a fail case for every recommendation, plus an evidence-missing",
             "# case for every Automated one. Assertions key on '<id>: ' (colon included).", ""]
    mod = f"data.{PKG}.{name}.violations"
    for rid in ids:
        s, k = L[rid], idkey(rid)
        pre = json.dumps(rid + ": ")
        if s.manual:
            tp, tf, miss = {"attest": {s.name: True}}, {"attest": {s.name: False}}, None
        else:
            tp, tf = s.tp, s.tf
            miss = J(f"{rid}: evidence missing — {s.field} not collected")
        lines += [f"test_{k}_pass if {{", f"\tvs := {mod} with input as {J(tp)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) == 0", "}", ""]
        lines += [f"test_{k}_fail if {{", f"\tvs := {mod} with input as {J(tf)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0"]
        if miss:
            lines.append(f"\tnot {miss} in vs")
        lines += ["}", ""]
        if miss:
            lines += [f"test_{k}_evidence_missing if {{", f"\tvs := {mod} with input as {J(s.tm)}",
                      f"\t{miss} in vs", "}", ""]
        if miss and not s.hollow_ok and hollow(s.tp) != s.tp:
            lines += [f"test_{k}_hollow_evidence_not_a_pass if {{", f"\tvs := {mod} with input as {J(hollow(s.tp))}",
                      f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0", "}", ""]
    g = good_input(ids)
    lines += ["# every recommendation of the section satisfied at once (the inputs above merged)",
              f"test_v3_{name}_section_compliant_on_merged_good_input if {{",
              f"\tvs := {mod} with input as {J(g)}", "\tcount(vs) == 0", "}", "",
              "# every violation names a recommendation of this section",
              f"test_v3_{name}_violations_carry_section_ids if {{",
              f"\tvs := {mod} with input as {{}}",
              f"\tevery v in vs {{ split(v, \":\")[0] in {json.dumps(sorted(ids, key=sort_key))} }}", "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


PROFILES = {"level1_server": ("Server", 1), "level2_server": ("Server", 2),
            "level1_workstation": ("Workstation", 1), "level2_workstation": ("Workstation", 2)}


def controls_table(recs):
    controls = {}
    for rid in sorted(recs, key=sort_key):
        r = recs[rid]
        lv = {}
        for kind in ("Server", "Workstation"):
            x = level(r["profiles"], kind)
            if x is not None:
                lv[kind.lower()] = x
        mod = section_of(rid)
        controls[rid] = {"title": r["title"], "assessment": r["assessment"], "profiles": r["profiles"],
                         "levels": lv, "section": SEC[mod][1], "package": mod}
    return controls


def emit_rollup(recs, by_mod):
    controls = controls_table(recs)
    imports = "\n".join(f"import data.{PKG}.{n}" for n, *_ in SECTIONS)
    modv = ",\n".join(f'\t"{n}": {n}.violations' for n, *_ in SECTIONS)
    sec_inputs = {SEC[n][1]: SECTION_INPUTS[n] for n, *_ in SECTIONS}
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    prio = {SEC[n][1]: p for n, (p, _, _) in RECOMMEND.items()}
    recs_rules = []
    for n, (p, sname, action) in RECOMMEND.items():
        key = SEC[n][1]
        cis = ", ".join(x + "x" for x in SEC[n][3])
        recs_rules.append(f'''generate_recommendations contains recommendation if {{
	section_compliance[{J(key)}].compliant == false
	recommendation := {{
		"priority": {J(p)},
		"section": {J(sname)},
		"issue": sprintf("%d {n} violations found", [section_compliance[{J(key)}].violation_count]),
		"action": {J(action)},
		"controls": {J("CIS " + cis)},
	}}
}}
''')
    return header(f"{BENCH} — roll-up", len(recs), density="orchestrator", basis="role") + f'''package {PKG}

# {BENCH}: roll-up of the {len(SECTIONS)} section packages.
# Generated by {GEN} from the CIS workbook; do not edit by hand.
# {len(recs)} recommendations ({a} Automated, {len(recs) - a} Manual), every one in the benchmark. Ids, titles,
# Automated/Manual status and profiles are the workbook's. Levels differ by platform (e.g. 1.1.1.9 is Level 1
# Server but Level 2 Workstation), so each control carries its own server/workstation level.
#
# PROFILE: input.profile selects one of level1_server (default), level2_server, level1_workstation,
# level2_workstation ("level1" / "level2", the v2.0.0 values, mean the server profiles). A Level 2 profile
# includes the Level 1 recommendations of the same platform.
#
# FAIL CLOSED: every Automated recommendation emits its own evidence-missing violation when its facts were
# not collected, and every Manual one fails until attested, so a control is counted as passed only when
# real facts satisfy it. A section none of whose input keys is present is also NOT ASSESSED (listed in
# not_assessed_sections, never compliant). The section keys are the v2.0.0 ones; each now holds the v3.0.0
# sections listed in the package title of its module.

import rego.v1

{imports}

# Extended hardening modules (STIG / NIST / EDA drift detection), not CIS RHEL 9 recommendations:
# evaluated separately so the CIS result is not affected.
import data.{PKG}.authorized_keys
import data.{PKG}.certificate_validation
import data.{PKG}.storage_encryption

standard_total := {len(recs)}

controls := {json.dumps(controls, indent=1, ensure_ascii=False)}

profiles := {json.dumps({k: {"platform": v[0].lower(), "level": v[1]} for k, v in PROFILES.items()})}

profile_aliases := {{"level1": "level1_server", "level2": "level2_server"}}

default profile := "level1_server"

profile := input.profile if input.profile in object.keys(profiles)

profile := profile_aliases[input.profile] if not input.profile in object.keys(profiles)

applies(c) if c.levels[profiles[profile].platform] <= profiles[profile].level

applicable_ids := {{id | some id, c in controls; applies(c)}}

module_violations := {{
{modv},
}}

all_violations := sort([v |
	some _, vs in module_violations
	some v in vs
	split(v, ":")[0] in applicable_ids
])

failing_ids := {{split(v, ":")[0] | some v in all_violations}}

# =============================================================================
# ASSESSMENT COVERAGE — which sections received any of the input keys they read
# =============================================================================

section_inputs := {json.dumps(sec_inputs, indent=1)}

section_assessed(name) if {{
	some k in section_inputs[name]
	input[k]
}}

missing_inputs := {{name: [k | some k in section_inputs[name]; not input[k]] | some name, _ in section_inputs}}

not_assessed_sections := sort([name | some name, _ in section_inputs; not section_assessed(name)])

assessed_sections := count(section_inputs) - count(not_assessed_sections)

assessment_coverage_percentage := round((assessed_sections * 100) / count(section_inputs))

section_assessment := {{name: {{"assessed": section_assessed(name) == true, "missing_inputs": missing_inputs[name]}} |
	some name, _ in section_inputs
}}

# =============================================================================
# SECTIONS (controls applicable to the selected profile)
# =============================================================================

section_ids(name) := {{id | some id in applicable_ids; controls[id].section == name}}

section_compliance := {{name: {{
	"compliant": count(section_ids(name) & failing_ids) == 0,
	"controls_covered": count(section_ids(name)),
	"violation_count": count([v | some v in all_violations; split(v, ":")[0] in section_ids(name)]),
}} |
	some name, _ in section_inputs
	count(section_ids(name)) > 0
}}

total_sections := count(section_compliance)

# A section counts as compliant only if it was actually assessed (fail-closed).
compliant_sections := count([name | some name, s in section_compliance; s.compliant == true; section_assessed(name)])

default compliance_percentage := 0

compliance_percentage := (compliant_sections * 100) / total_sections if total_sections > 0

total_controls_covered := count(applicable_ids)

default compliant := false

compliant if {{
	count(object.keys(input)) > 0
	count(applicable_ids) > 0
	count(failing_ids) == 0
	count(not_assessed_sections) == 0
}}

# =============================================================================
# RISK ASSESSMENT AND RECOMMENDATIONS (v2.0.0 shapes)
# =============================================================================

section_priority := {json.dumps(prio, indent=1)}

critical_risks := [v | some v in all_violations; section_priority[controls[split(v, ":")[0]].section] == "critical"]

high_risks := [v | some v in all_violations; section_priority[controls[split(v, ":")[0]].section] == "high"]

overall_risk_level := "critical" if {{
	count(critical_risks) > 0
}} else := "high" if {{
	count(high_risks) > 5
}} else := "medium" if {{
	count(all_violations) > 10
}} else := "low"

{chr(10).join(recs_rules)}
overall_status := "compliant" if {{
	compliant
}} else := "non_compliant"

level2_ids := {{id | some id, c in controls; object.get(c.levels, profiles[profile].platform, 0) == 2}}

l2_violations_active := [v | some v in all_violations; split(v, ":")[0] in level2_ids] if {{
	profiles[profile].level == 2
}} else := []

default l2_compliant := false

l2_compliant if profiles[profile].level == 1

l2_compliant if {{
	profiles[profile].level == 2
	count(l2_violations_active) == 0
}}

executive_summary := {{
	"overall_status": overall_status,
	"profile": profile,
	"compliance_score": compliance_percentage,
	"sections_evaluated": total_sections,
	"sections_compliant": compliant_sections,
	"total_violations": count(all_violations),
	"l1_violations": count(all_violations) - count(l2_violations_active),
	"l2_violations": count(l2_violations_active),
	"l2_compliant": l2_compliant,
	"critical_issues": count(critical_risks),
	"top_priorities": [rec | some rec in generate_recommendations; rec.priority == "critical"],
}}

# Consumer endpoint — /v1/data/{PKG}/compliance_assessment (customer profiles query it)
compliance_assessment := {{
	"framework": {J(BENCH)},
	"compliant": compliant,
	"score": compliance_percentage,
	"total_controls": count(applicable_ids),
	"sections_evaluated": total_sections,
	"sections_compliant": compliant_sections,
	"violations": all_violations,
	"violation_count": count(all_violations),
	"summary": executive_summary,
}}

# =============================================================================
# EXTENDED HARDENING (STIG / NIST / EDA drift detection) — separate from the CIS score
# =============================================================================

extended_violations := array.concat(
	array.concat(
		[v | some v in storage_encryption.violations],
		[v | some v in certificate_validation.violations],
	),
	[v | some v in authorized_keys.violations],
)

default hardening_compliant := false

hardening_compliant if {{
	storage_encryption.compliant
	certificate_validation.compliant
	authorized_keys.compliant
}}

default fully_hardened_compliant := false

fully_hardened_compliant if {{
	compliant
	hardening_compliant
}}

extended_hardening_summary := {{
	"hardening_compliant": hardening_compliant,
	"fully_hardened_compliant": fully_hardened_compliant,
	"extended_violation_count": count(extended_violations),
	"modules": {{
		"storage_encryption": {{
			"compliant": storage_encryption.compliant,
			"violations": count(storage_encryption.violations),
			"standards": ["STIG RHEL-09-231xxx", "STIG RHEL-09-672xxx", "NIST SC-28"],
		}},
		"certificate_validation": {{
			"compliant": certificate_validation.compliant,
			"violations": count(certificate_validation.violations),
			"standards": ["NIST SC-17", "NIST IA-5"],
		}},
		"authorized_keys": {{
			"compliant": authorized_keys.compliant,
			"violations": count(authorized_keys.violations),
			"standards": ["STIG RHEL-09-255xxx", "NIST IA-3"],
		}},
	}},
}}

generate_recommendations contains recommendation if {{
	not storage_encryption.compliant
	recommendation := {{
		"priority": "critical",
		"section": "Storage Encryption",
		"issue": sprintf("%d storage encryption violations found", [count(storage_encryption.violations)]),
		"action": "Enable LUKS encryption on data partitions, encrypt swap, verify FIPS mode",
		"controls": "STIG RHEL-09-231xxx / NIST SC-28",
	}}
}}

generate_recommendations contains recommendation if {{
	not certificate_validation.compliant
	recommendation := {{
		"priority": "high",
		"section": "Certificate and PKI",
		"issue": sprintf("%d certificate violations found", [count(certificate_validation.violations)]),
		"action": "Renew expired/expiring certificates, remove weak keys, audit trust store",
		"controls": "NIST SC-17 / IA-5",
	}}
}}

generate_recommendations contains recommendation if {{
	not authorized_keys.compliant
	recommendation := {{
		"priority": "critical",
		"section": "SSH Authorized Keys",
		"issue": sprintf("%d authorized_keys violations found", [count(authorized_keys.violations)]),
		"action": "Remove unauthorized keys, fix permissions, enforce approved-key baseline",
		"controls": "STIG RHEL-09-255xxx / NIST IA-3",
	}}
}}
'''


def emit_main(recs):
    return header(f"Bridge: expose {PKG} under the /v1/data/<framework>/main/compliance_report", 0,
                  entry="compliance_report", density="orchestrator", basis="role") + f'''package {PKG}.main

# Bridge: expose {PKG} under the /v1/data/<framework>/main/compliance_report entrypoint convention, so a caller
# can evaluate this framework by key without knowing its internal package layout.
# Generated by {GEN}; do not edit by hand.
#
# 1. Every field is sourced through a DEFAULTED helper: an undefined field inside an object literal makes the
#    whole object undefined, which OPA returns as {{}} — a silently empty result instead of an error.
# 2. total_controls is the number of recommendations applicable to the selected profile (counted from the
#    controls table, never hard-coded); standard_total is the benchmark's {len(recs)}. A control passes only
#    when real facts (or an attestation) satisfy it, so failed_controls counts distinct failing
#    recommendation ids, including those whose evidence is missing.
#
# FAIL-CLOSED GATE: with no facts at all every control counts as failed and an explicit FAIL-CLOSED
# violation is reported. PARTIAL INPUT: a section that received none of its input keys is NOT ASSESSED
# ({PKG}.section_inputs): it is listed, it adds an explicit violation, and its controls fail on their own
# evidence-missing violations, never as passed.

import data.{PKG} as bench
import rego.v1

default _compliant := false

_compliant := bench.compliant

default _facts_supplied := false

_facts_supplied if count(object.keys(input)) > 0

_no_facts_msg := sprintf(
	"FAIL-CLOSED: no facts supplied for %s — the assessment could not be evaluated. This is NOT a passing result; check that fact collection ran and produced input.",
	["{PKG}"],
)

default _bench_violations := []

_bench_violations := bench.all_violations

default _not_assessed := []

_not_assessed := bench.not_assessed_sections

default _missing_inputs := {{}}

_missing_inputs := bench.missing_inputs

_not_assessed_msgs := [sprintf("NOT ASSESSED: %s — required input not collected: %s", [name, concat(", ", _missing_inputs[name])]) | some name in _not_assessed]

default _sections := {{}}

_sections := bench.section_compliance

default _section_percentage := 0

_section_percentage := bench.compliance_percentage

default _coverage := 0

_coverage := bench.assessment_coverage_percentage

default _profile := "level1_server"

_profile := bench.profile

default _total_controls := 0

_total_controls := count(bench.applicable_ids)

default _failing := 0

_failing := count(bench.failing_ids)

_violations := array.concat(_bench_violations, [_no_facts_msg]) if not _facts_supplied

_violations := array.concat(_bench_violations, _not_assessed_msgs) if _facts_supplied

# No facts => every control is unverified, which counts as failed, not passed.
_failed := _total_controls if not _facts_supplied

_failed := min([_total_controls, _failing]) if _facts_supplied

# Clamp: more failures than known controls must not yield a negative count.
_passed := max([0, _total_controls - _failed])

default _percentage := 0

_percentage := round((_passed * 100000) / _total_controls) / 1000 if _total_controls > 0

compliance_report := {{
	"framework": "{PKG}",
	"benchmark": "{BENCH}",
	"version": "{VERSION}",
	"profile": _profile,
	"total_controls": _total_controls,
	"standard_total": {len(recs)},
	"controls_basis": "CIS recommendations applicable to the selected profile (every one evaluated)",
	"passed_controls": _passed,
	"failed_controls": _failed,
	"compliance_percentage": _percentage,
	"section_compliance_percentage": _section_percentage,
	"compliant": _compliant,
	"violations": _violations,
	"section_results": _sections,
	"sections_not_assessed": _not_assessed,
	"missing_inputs": _missing_inputs,
	"assessment_coverage_percentage": _coverage,
}}
'''


def emit_l2(recs):
    n = sum(1 for r in recs.values() if level(r["profiles"], "Server") == 2)
    return header(f"{BENCH} — Level 2 (Server) recommendations", n, entry="compliance_report",
                  density="orchestrator", basis="role") + f'''package {PKG}.l2

# {BENCH}: the Level 2 - Server recommendations, i.e. what the level2_server profile adds to level1_server.
# Generated by {GEN}; do not edit by hand. A view of the same section packages the roll-up reads (no logic
# of its own), so the Level 2 result can never disagree with the roll-up. Kept as cis_rhel9.l2 for callers
# of the v2.0.0 package; the roll-up selects Level 2 with input.profile = "level2_server" (or "level2").

import data.{PKG} as bench
import rego.v1

l2_ids := {{id | some id, c in bench.controls; object.get(c.levels, "server", 0) == 2}}

l2_total_controls := count(l2_ids)

violations contains v if {{
	some _, vs in bench.module_violations
	some v in vs
	split(v, ":")[0] in l2_ids
}}

failing_ids := {{split(v, ":")[0] | some v in violations}}

default compliant := false

compliant if {{
	count(object.keys(input)) > 0
	count(violations) == 0
}}

compliance_report := {{
	"framework": "{PKG}_l2",
	"benchmark": "{BENCH} (Level 2 - Server recommendations)",
	"profile": "level2_server",
	"total_controls": l2_total_controls,
	"failed_controls": count(failing_ids),
	"passed_controls": l2_total_controls - count(failing_ids),
	"compliant": compliant,
	"violations": [v | some v in violations],
}}

# v2.0.0 compatibility: sshd directive -> value(s); a directive `sshd -T` repeats is comma-joined
_active_lines(text) := [l |
	some raw in split(text, "\\n")
	l := trim_space(raw)
	l != ""
	not startswith(l, "#")
]

_sshd_pairs := [[lower(m[1]), trim_space(m[2])] |
	some l in _active_lines(input.ssh.sshd_config_raw)
	some m in regex.find_all_string_submatch_n(`^([A-Za-z0-9]+)(?:\\s*=\\s*|\\s+)(.*)$`, l, 1)
]

sshd_config := {{k: concat(",", [p[1] | some p in _sshd_pairs; p[0] == k]) | some q in _sshd_pairs; k := q[0]}}
'''


def emit_rollup_tests(recs):
    counts = {p: sum(1 for r in recs.values() if (level(r["profiles"], kind) or 99) <= lv)
              for p, (kind, lv) in PROFILES.items()}
    keys = ["framework", "benchmark", "version", "profile", "total_controls", "standard_total", "controls_basis",
            "passed_controls", "failed_controls", "compliance_percentage", "section_compliance_percentage", "compliant",
            "violations", "section_results", "sections_not_assessed", "missing_inputs", "assessment_coverage_percentage"]
    m = f"data.{PKG}.main.compliance_report"
    lines = [f"package {PKG}.main_test", "", "import rego.v1", "",
             f"# Generated by {GEN}: roll-up and bridge tests for {BENCH}.", "",
             "test_v3_every_recommendation_in_table if {",
             f"\tcount(data.{PKG}.controls) == {len(recs)}",
             f"\tcount([c | some c in data.{PKG}.controls; c.assessment == \"Manual\"]) == {EXPECTED_MANUAL}",
             f"\tdata.{PKG}.standard_total == {len(recs)}", "}", "",
             "test_v3_empty_input_fails_closed if {", f"\tr := {m} with input as {{}}",
             f"\tcount(object.keys(r)) == {len(keys)}",
             f"\tevery k in {json.dumps(keys)} {{ k in object.keys(r) }}",
             "\tr.compliant == false", f"\tr.total_controls == {counts['level1_server']}",
             "\tr.failed_controls == r.total_controls", "\tr.passed_controls == 0",
             "\tsome v in r.violations", "\tstartswith(v, \"FAIL-CLOSED: no facts supplied\")", "}", ""]
    for p, n in counts.items():
        lines += [f"test_v3_profile_{p}_count if {{", f"\tr := {m} with input as {{\"profile\": \"{p}\"}}",
                  f"\tr.profile == \"{p}\"", f"\tr.total_controls == {n}", f"\tr.standard_total == {len(recs)}",
                  "\tr.failed_controls == r.total_controls", "}", ""]
    lines += ["test_v3_legacy_profile_values_map_to_server_profiles if {",
              f"\tr1 := {m} with input as {{\"profile\": \"level1\"}}", "\tr1.profile == \"level1_server\"",
              f"\tr1.total_controls == {counts['level1_server']}",
              f"\tr2 := {m} with input as {{\"profile\": \"level2\"}}", "\tr2.profile == \"level2_server\"",
              f"\tr2.total_controls == {counts['level2_server']}", "}", "",
              "test_v3_unknown_profile_defaults_to_level1_server if {",
              f"\tr := {m} with input as {{\"profile\": \"bogus\"}}",
              "\tr.profile == \"level1_server\"", f"\tr.total_controls == {counts['level1_server']}", "}", "",
              "test_v3_one_attested_control_passes if {",
              f"\tr := {m} with input as {{\"attest\": {{\"gpg_keys_reviewed\": true}}}}",
              "\tr.passed_controls == 1", "\tr.failed_controls == r.total_controls - 1", "\tr.compliant == false", "}", "",
              "test_v3_level2_violation_hidden_in_level1_profile if {",
              f"\tvs := data.{PKG}.all_violations with input as {{\"profile\": \"level1_server\"}}",
              "\tcount([v | some v in vs; startswith(v, \"6.2.1.1: \")]) == 0",
              f"\tvs2 := data.{PKG}.all_violations with input as {{\"profile\": \"level2_server\"}}",
              "\tcount([v | some v in vs2; startswith(v, \"6.2.1.1: \")]) > 0", "}", "",
              "test_v3_every_violation_parses_to_a_control if {",
              f"\tvs := data.{PKG}.all_violations with input as {{\"profile\": \"level2_server\"}}",
              f"\tevery v in vs {{ split(v, \":\")[0] in object.keys(data.{PKG}.controls) }}",
              f"\tmvs := data.{PKG}.module_violations with input as {{}}",
              "\tevery _, mv in mvs {",
              f"\t\tevery v in mv {{ split(v, \":\")[0] in object.keys(data.{PKG}.controls) }}", "\t}", "}", "",
              "test_v3_l2_view_is_level2_server_only if {",
              f"\tr := data.{PKG}.l2.compliance_report with input as {{}}",
              f"\tr.total_controls == {sum(1 for x in recs.values() if level(x['profiles'], 'Server') == 2)}",
              "\tr.compliant == false",
              f"\tl2v := data.{PKG}.l2.violations with input as {{}}",
              f"\tevery v in l2v {{ data.{PKG}.controls[split(v, \":\")[0]].levels.server == 2 }}",
              "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    recs = read_workbook(sys.argv[1])
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    m = sum(1 for r in recs.values() if r["assessment"] == "Manual")
    if (len(recs), a, m) != (EXPECTED_TOTAL, EXPECTED_AUTOMATED, EXPECTED_MANUAL):
        raise SystemExit(f"workbook has {len(recs)} recommendations ({a} Automated, {m} Manual)")
    if set(L) != set(recs):
        raise SystemExit(f"logic/workbook mismatch: no logic {sorted(set(recs) - set(L), key=sort_key)}; "
                         f"no recommendation {sorted(set(L) - set(recs), key=sort_key)}")
    for rid, r in recs.items():
        if (r["assessment"] == "Manual") != L[rid].manual:
            raise SystemExit(f"{rid}: workbook says {r['assessment']}, logic table disagrees")
        title = r["title"].lower() + " "
        kw = re.sub(r"\s+", " ", KW[rid]).lower().lstrip()  # a trailing space anchors the keyword at the title's end
        if kw not in title:
            raise SystemExit(f"{rid}: logic keyword {KW[rid]!r} not in workbook title {r['title']!r}")
    by_mod = {n: [] for n, *_ in SECTIONS}
    for rid in sorted(recs, key=sort_key):
        by_mod[section_of(rid)].append(rid)
    (OUT / "tests").mkdir(parents=True, exist_ok=True)
    for name, fname, *_ in SECTIONS:
        (OUT / fname).write_text(emit_module(name, by_mod[name], recs))
        (OUT / "tests" / f"test_{name}_v3_controls.rego").write_text(emit_tests(name, by_mod[name]))
    (OUT / "cis_rhel9_complete.rego").write_text(emit_rollup(recs, by_mod))
    (OUT / "cis_rhel9_main.rego").write_text(emit_main(recs))
    (OUT / "l2_validation.rego").write_text(emit_l2(recs))
    (OUT / "tests" / "test_cis_rhel9_v3_rollup.rego").write_text(emit_rollup_tests(recs))
    if len(sys.argv) > 2:  # input paths each control reads + merged passing inputs (gap analysis, compat tests)
        pathlib.Path(sys.argv[2]).write_text(json.dumps(
            {"reads": {rid: [".".join(p) for p in L[rid].reads] for rid in sorted(recs, key=sort_key)},
             "good": {n: good_input(by_mod[n]) for n, *_ in SECTIONS}}, indent=1))
    print(f"{len(recs)} recommendations ({a} Automated, {m} Manual) -> {OUT}")


if __name__ == "__main__":
    main()
