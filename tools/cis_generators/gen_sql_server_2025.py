#!/usr/bin/env python3
"""CIS Microsoft SQL Server 2025 Benchmark v1.0.0: every recommendation, from the benchmark's own workbook.

    python3 gen_sql_server_2025.py CIS_Microsoft_SQL_Server_2025_Benchmark_v1.0.0.xlsx <repo root>

Needs openpyxl. Reads id, title, Assessment Status and profiles out of the CIS workbook (CIS-licensed, never
committed), checks the ids against the logic table below, and emits benchmarks/cis/sql_server_2025/
cis_sql_server_2025.rego, cis_sql_server_2025_main.rego and tests/test_cis_sql_server_2025.rego.

Adapted from tools/cis_generators/gen_oracle_db_26ai.py:
  * every Automated check fails closed with "<id>: evidence missing — <field> not collected";
  * every Manual review is an attestation boolean (not attested => violation);
  * the tests carry a pass, a fail and (Automated) an evidence-missing case for every recommendation;
  * unlike Oracle, the profiles really differ (Database Engine L1/L2, AWS RDS L1), so input.cis_profile selects
    which recommendations are in scope; it defaults to the strictest selection (all 49).
Input contract roots at input.sql, as benchmarks/stig/ms_sql_2016 does (input.sql.sa_account_disabled is shared).
"""
import json
import pathlib
import sys

import openpyxl

BENCH = "CIS Microsoft SQL Server 2025 Benchmark v1.0.0"
PKG = "cis.sql_server_2025"
EXPECTED_TOTAL = 49
SKIP_SHEETS = {"License", "Overview", "Combined Profiles"}
DE1, RDS1, DE2 = "Level 1 - Database Engine", "Level 1 - AWS RDS", "Level 2 - Database Engine"


def ptr(path):
    parts = path if isinstance(path, list) else path.split(".")
    return "/" + "/".join(p.replace("~", "~0").replace("/", "~1") for p in parts)


def rep(path, value):
    return {"op": "replace", "path": ptr(path), "value": value}


def add(path, value):
    return {"op": "add", "path": ptr(path), "value": value}


def rm(path):
    return {"op": "remove", "path": ptr(path)}


def spec(body, fx, ev, field, fails, miss, kind="automated"):
    """body / ev: list of alternatives, each a list of rego lines (any alternative satisfies the rule);
    fx: green fixture entries [(path, value)]; fails: json.patch op lists that must each make the check fail;
    miss: json.patch ops that remove the evidence."""
    return {"body": body, "fx": fx, "ev": ev, "field": field, "fails": fails, "miss": miss, "kind": kind}


def cfg(name, want):
    """sys.configurations row: both value_configured and value_in_use must equal `want`"""
    p = ["sql", "configurations", name]
    return spec([[f'_cfg_is("{name}", {want})']], [(p, {"value_configured": want, "value_in_use": want})],
                [[f'count(_cfg_rows("{name}")) > 0']], f"sql.configurations['{name}']",
                [[rep(p + ["value_in_use"], 1 - want)], [rep(p + ["value_configured"], 1 - want)]], [rm(p)])


def rows_empty(key, row, allow=None):
    """the audit query's offending rows; [] is 'no rows', a missing list is no evidence"""
    path = f"sql.{key}"
    body = f"count(input.{path}) == 0" if allow is None else \
        f"count([r | some r in input.{path}; not {allow}(r)]) == 0"
    return spec([[body]], [(path, [])], [[f"is_array(input.{path})"]], path, [[rep(path, [row])]], [rm(path)])


def fact(key, body, fx, ev, bad):
    path = f"sql.{key}"
    return spec([[b] for b in body] if isinstance(body, tuple) else [[body]], [(path, fx)], [[f"{ev}(input.{path})"]],
                path, [[rep(path, b)] for b in bad], [rm(path)])


def attest(name):
    path = f"sql.attest.{name}"
    return spec([[f"input.{path} == true"]], [(path, True)], None, path, [[rep(path, False)], [rm(path)]], None, kind="attest")


AUDIT_GROUPS = ["AUDIT_CHANGE_GROUP", "FAILED_LOGIN_GROUP", "SUCCESSFUL_LOGIN_GROUP", "DATABASE_ROLE_MEMBER_CHANGE_GROUP",
                "SERVER_ROLE_MEMBER_CHANGE_GROUP", "SERVER_OPERATION_GROUP"]
CONTAINED_AUDIT_GROUPS = ["SUCCESSFUL_DATABASE_AUTHENTICATION_GROUP", "FAILED_DATABASE_AUTHENTICATION_GROUP"]


def _audit_row(g, **kw):
    return {"audit_action_name": g, "audit_enabled": True, "specification_enabled": True,
            "audited_result": "SUCCESS AND FAILURE", **kw}


GREEN_AUDIT = [_audit_row(g) for g in AUDIT_GROUPS]

HELPERS = """
# ── helpers ──────────────────────────────────────────────────────────────────
# sys.configurations rows, matched by lowercased name (the workbook spells names with mixed case)
_cfg_rows(name) := [v | some k, v in input.sql.configurations; lower(k) == name]

# both value_configured and value_in_use must equal want, in every row the collector supplied for that name
_cfg_is(name, want) if {
	rows := _cfg_rows(name)
	count(rows) > 0
	every v in rows {
		to_number(v.value_configured) == want
		to_number(v.value_in_use) == want
	}
}

# AWS RDS allowances from the audit procedures; they apply only when the profile says the instance is RDS
_rds_rdsadmin_db(r) if {
	_rds
	lower(r.database) == "rdsadmin"
}

_rds_builtin_admins(r) if {
	_rds
	upper(replace(replace(r.name, "[", ""), "]", "")) == "BUILTIN\\\\ADMINISTRATORS"
}

_rds_master_rdsa(r) if {
	_rds
	lower(r.name) == "rdsa"
}

# 5.4: the six server audit action groups, plus the two contained-database authentication groups when the
# instance has contained databases; each needs an enabled audit, an enabled specification, success and failure.
_audit_groups_required contains g if some g in AUDIT_GROUPS

_audit_groups_required contains g if {
	input.sql.has_contained_databases == true
	some g in CONTAINED_AUDIT_GROUPS
}

_audit_group_ok(g) if {
	some r in input.sql.server_audit_actions
	upper(r.audit_action_name) == g
	r.audit_enabled == true
	r.specification_enabled == true
	contains(upper(r.audited_result), "SUCCESS")
	contains(upper(r.audited_result), "FAILURE")
}
""".replace("AUDIT_GROUPS\n", json.dumps(AUDIT_GROUPS) + "\n", 1).replace(
    "some g in CONTAINED_AUDIT_GROUPS", "some g in " + json.dumps(CONTAINED_AUDIT_GROUPS))

RULES = {
    "1.1": attest("latest_cumulative_update_installed"),
    "1.2": attest("single_function_member_server"),
    "2.1": cfg("ad hoc distributed queries", 0),
    "2.2": spec([['_cfg_is("clr strict security", 1)'], ['_cfg_is("clr enabled", 0)']],
                [(["sql", "configurations", "clr enabled"], {"value_configured": 0, "value_in_use": 0})],
                [['_cfg_is("clr strict security", 1)'], ['count(_cfg_rows("clr enabled")) > 0']],
                "sql.configurations['clr enabled']",
                [[rep(["sql", "configurations", "clr strict security"], {"value_configured": 0, "value_in_use": 0}),
                  rep(["sql", "configurations", "clr enabled"], {"value_configured": 1, "value_in_use": 1})]],
                [rep(["sql", "configurations", "clr strict security"], {"value_configured": 0, "value_in_use": 0}),
                 rm(["sql", "configurations", "clr enabled"])]),
    "2.3": cfg("cross db ownership chaining", 0),
    "2.4": cfg("database mail xps", 0),
    "2.5": cfg("ole automation procedures", 0),
    "2.6": cfg("remote access", 0),
    "2.7": spec([["input.sql.is_clustered == true"], ['_cfg_is("remote admin connections", 0)']],
                [("sql.is_clustered", False), (["sql", "configurations", "remote admin connections"], {"value_configured": 0, "value_in_use": 0})],
                [["input.sql.is_clustered == true"], ["is_boolean(input.sql.is_clustered)", 'count(_cfg_rows("remote admin connections")) > 0']],
                "sql.is_clustered / sql.configurations['remote admin connections']",
                [[rep(["sql", "configurations", "remote admin connections"], {"value_configured": 1, "value_in_use": 0})]],
                [rm("sql.is_clustered")]),
    "2.8": cfg("scan for startup procs", 0),
    "2.9": rows_empty("databases_trustworthy_on", {"database": "AppDB"}),
    "2.10": attest("unnecessary_protocols_disabled"),
    "2.11": fact("tcp_port_1433_count", "input.sql.tcp_port_1433_count == 0", 0, "is_number", [1]),
    "2.12": fact("hide_instance", "input.sql.hide_instance == 1", 1, "is_number", [0]),
    "2.13": fact("sa_account_disabled", "input.sql.sa_account_disabled == true", True, "is_boolean", [False]),
    "2.14": fact("sa_login_name", 'lower(input.sql.sa_login_name) != "sa"', "dbadmin_x", "is_string", ["sa", "SA"]),
    "2.15": rows_empty("contained_databases_auto_close_on", {"database": "ContainedDB"}),
    "2.16": rows_empty("logins_named_sa", {"principal_id": 270, "name": "sa"}),
    "2.17": cfg("clr strict security", 1),
    "3.1": fact("is_integrated_security_only", "input.sql.is_integrated_security_only == 1", 1, "is_number", [0]),
    "3.2": rows_empty("guest_connect_grants", {"database": "AppDB"}, "_rds_rdsadmin_db"),
    "3.3": rows_empty("orphaned_users", {"database": "AppDB", "name": "old_user"}),
    "3.4": rows_empty("contained_db_sql_auth_users", {"database": "ContainedDB", "name": "app_sql_user"}),
    "3.5": attest("mssql_service_account_not_admin"),
    "3.6": attest("sqlagent_service_account_not_admin"),
    "3.7": attest("fulltext_service_account_not_admin"),
    "3.8": rows_empty("public_role_extra_permissions", {"permission_name": "VIEW SERVER STATE", "class_desc": "SERVER"}),
    "3.9": rows_empty("builtin_logins", {"name": "BUILTIN\\Users", "permission_name": "CONNECT SQL"}, "_rds_builtin_admins"),
    "3.10": rows_empty("local_group_logins", {"name": "HOST01\\SQLAdmins", "permission_name": "CONNECT SQL"}),
    "3.11": rows_empty("msdb_public_proxy_grants", {"proxyname": "ssis_proxy"}),
    "3.12": attest("sysadmin_membership_reviewed"),
    "3.13": rows_empty("msdb_admin_role_members", {"role": "db_owner", "member": "app_user"}),
    "3.14": attest("control_server_grants_reviewed"),
    "3.15": attest("external_rest_endpoint_access_reviewed"),
    "4.1": attest("must_change_reviewed"),
    "4.2": rows_empty("sysadmin_logins_without_expiration_check", {"name": "app_admin"}, "_rds_master_rdsa"),
    "4.3": rows_empty("sql_logins_without_check_policy", {"name": "app_user", "is_disabled": 0}),
    "5.1": fact("num_error_logs", ("input.sql.num_error_logs == -1", "input.sql.num_error_logs >= 12"), 12, "is_number", [6, 0]),
    "5.2": cfg("default trace enabled", 1),
    "5.3": fact("login_audit_level", 'lower(input.sql.login_audit_level) in {"failure", "all"}', "failure", "is_string", ["none", "success"]),
    "5.4": spec([["every g in _audit_groups_required { _audit_group_ok(g) }"]],
                [("sql.server_audit_actions", GREEN_AUDIT), ("sql.has_contained_databases", False)],
                [["is_array(input.sql.server_audit_actions)", "is_boolean(input.sql.has_contained_databases)"]],
                "sql.server_audit_actions / sql.has_contained_databases",
                [[rep("sql.server_audit_actions", [r for r in GREEN_AUDIT if r["audit_action_name"] != "FAILED_LOGIN_GROUP"])],
                 [rep("sql.server_audit_actions", [_audit_row(g, audit_enabled=(g != "AUDIT_CHANGE_GROUP")) for g in AUDIT_GROUPS])],
                 [rep("sql.server_audit_actions", [_audit_row(g, audited_result="SUCCESS" if g == "SUCCESSFUL_LOGIN_GROUP" else "SUCCESS AND FAILURE") for g in AUDIT_GROUPS])],
                 [rep("sql.has_contained_databases", True)]],
                [rm("sql.server_audit_actions")]),
    "6.1": attest("input_sanitization_reviewed"),
    "6.2": spec([['every a in input.sql.user_assemblies { upper(a.permission_set_desc) == "SAFE_ACCESS" }']],
                [("sql.user_assemblies", [{"database": "AppDB", "name": "SafeLib", "permission_set_desc": "SAFE_ACCESS"}])],
                [["is_array(input.sql.user_assemblies)"]], "sql.user_assemblies",
                [[rep("sql.user_assemblies", [{"database": "AppDB", "name": "SafeLib", "permission_set_desc": "UNSAFE_ACCESS"}])],
                 [rep("sql.user_assemblies", [{"database": "AppDB", "name": "ExtLib", "permission_set_desc": "EXTERNAL_ACCESS"}])]],
                [rm("sql.user_assemblies")]),
    "7.1": spec([['every k in input.sql.symmetric_keys { upper(k.algorithm_desc) in {"AES_128", "AES_192", "AES_256"} }']],
                [("sql.symmetric_keys", [{"database": "AppDB", "name": "AppKey", "algorithm_desc": "AES_256"}])],
                [["is_array(input.sql.symmetric_keys)"]], "sql.symmetric_keys",
                [[rep("sql.symmetric_keys", [{"database": "AppDB", "name": "OldKey", "algorithm_desc": "TRIPLE_DES"}])]],
                [rm("sql.symmetric_keys")]),
    "7.2": spec([["every k in input.sql.asymmetric_keys { to_number(k.key_length) >= 2048 }"]],
                [("sql.asymmetric_keys", [{"database": "AppDB", "name": "AppAsym", "key_length": 2048}])],
                [["is_array(input.sql.asymmetric_keys)"]], "sql.asymmetric_keys",
                [[rep("sql.asymmetric_keys", [{"database": "AppDB", "name": "WeakAsym", "key_length": 1024}])]],
                [rm("sql.asymmetric_keys")]),
    "7.3": rows_empty("unencrypted_backups", {"database_name": "AppDB", "backup_finish_date": "2026-10-01T00:00:00"}),
    # 7.4: every non-shared-memory connection (outside encrypted mirroring endpoints) must report encrypt_option TRUE.
    # An empty list means the query saw no such connection; it is not treated as proof, so it fails.
    "7.4": spec([["count(input.sql.connection_encrypt_options) > 0",
                  'every o in input.sql.connection_encrypt_options { upper(sprintf("%v", [o])) == "TRUE" }']],
                [("sql.connection_encrypt_options", ["TRUE"])],
                [["is_array(input.sql.connection_encrypt_options)"]], "sql.connection_encrypt_options",
                [[rep("sql.connection_encrypt_options", ["TRUE", "FALSE"])], [rep("sql.connection_encrypt_options", [])]],
                [rm("sql.connection_encrypt_options")]),
    "7.5": rows_empty("databases_without_tde", {"database_id": 5, "name": "AppDB"}),
    "8.1": attest("browser_service_reviewed"),
}


def key(rid):
    return [int(p) for p in rid.split(".")]


def ident(rid):
    return rid.replace(".", "_")


def rego_str(s):
    return json.dumps(s, ensure_ascii=False)


def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    recs, sections, combined = {}, {}, set()
    for sheet in wb.sheetnames:
        if sheet in ("License", "Overview"):
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        h = rows[0]
        ci = {c: h.index(c) for c in ("Section #", "Recommendation #", "Title", "Assessment Status")}
        for r in rows[1:]:
            rid, title, status = r[ci["Recommendation #"]], r[ci["Title"]], r[ci["Assessment Status"]]
            if not rid:
                sec = str(r[ci["Section #"]] or "")
                if sec and "." not in sec:
                    sections[sec] = " ".join(str(title).split())
                continue
            rid = str(rid).strip()
            if sheet == "Combined Profiles":
                combined.add(rid)
                continue
            title = " ".join(str(title).split()).replace('"', "'")
            level = int(sheet.split()[1])
            if rid in recs:
                if recs[rid]["title"] != title or recs[rid]["status"] != status:
                    raise SystemExit(f"{rid}: title/status differ between profile sheets")
                recs[rid]["profiles"].append(sheet)
                recs[rid]["level"] = min(recs[rid]["level"], level)
            else:
                recs[rid] = {"title": title, "status": status, "level": level, "profiles": [sheet]}
    if combined != set(recs):
        raise SystemExit(f"Combined Profiles disagrees with the profile sheets: {sorted(combined ^ set(recs))}")
    return recs, sections


def rule_lines(name, alts):
    out = []
    for alt in alts:
        out += [f"{name} if {{"] + ["\t" + b for b in alt] + ["}", ""]
    return out


def main(xlsx, root):
    out = pathlib.Path(root) / "benchmarks" / "cis" / "sql_server_2025"
    recs, all_sections = read_workbook(xlsx)
    missing = sorted(set(recs) - set(RULES), key=key)
    extra = sorted(set(RULES) - set(recs), key=key)
    if missing or extra:
        raise SystemExit(f"logic table out of step with the workbook: missing {missing} extra {extra}")
    if len(recs) != EXPECTED_TOTAL:
        raise SystemExit(f"workbook has {len(recs)} recommendations, expected {EXPECTED_TOTAL}")
    for rid, r in recs.items():
        if (RULES[rid]["kind"] == "attest") != (r["status"] == "Manual"):
            raise SystemExit(f"{rid} is {r['status']} in the workbook but the logic table says {RULES[rid]['kind']}")
    ids = sorted(recs, key=key)
    tops = sorted({rid.split(".")[0] for rid in ids}, key=int)
    sections = {s: all_sections[s] for s in tops}
    n_auto = sum(1 for r in recs.values() if r["status"] == "Automated")
    n_manual = len(ids) - n_auto
    by_prof = {p: [r for r in ids if p in recs[r]["profiles"]] for p in (DE1, RDS1, DE2)}

    L = ["# METADATA", f'# title: "{BENCH}"', "# custom:", "#   class: security", "#   framework: cis_sql_server_2025",
         "#   source: cis", "#   domains: [database, mssql]", f"package {PKG}", "",
         f"# {BENCH}.",
         f"# {len(ids)} recommendations ({n_auto} Automated, {n_manual} Manual), every one in the benchmark; ids, titles,",
         "# status and profiles read from the CIS workbook and the logic written from each audit procedure",
         "# (generated by gen_sql_server_2025.py, adapted from tools/cis_generators/gen_oracle_db_26ai.py).",
         f"# Profiles: '{DE1}' ({len(by_prof[DE1])}), '{RDS1}' ({len(by_prof[RDS1])}, a subset),",
         f"# '{DE2}' ({len(by_prof[DE2])}: {', '.join(by_prof[DE2])}).",
         "# input.cis_profile selects the recommendations in scope:",
         "#   platform: \"database_engine\" (default) | \"aws_rds\"; level: 1 | 2 (default 2; Level 2 includes Level 1).",
         "#   The default selection is the strictest one: all 49. Out-of-scope recommendations raise nothing.",
         "#   aws_rds also turns on the RDS allowances the audit procedures name (3.2 rdsadmin, 3.9 BUILTIN\\Administrators,",
         "#   4.2 rdsa).",
         f"# {n_manual} Manual reviews carry an attestation boolean (not attested => violation).",
         "# Automated checks fail closed: facts not collected => '<id>: evidence missing — <field> not collected'.",
         "#",
         "# Input contract (root input.sql, as benchmarks/stig/ms_sql_2016; sql.sa_account_disabled is shared with it):",
         "#   sql.configurations[<name>] = {value_configured, value_in_use}  : sys.configurations rows (name any case)",
         "#   sql.is_clustered (bool)                     : SERVERPROPERTY('IsClustered') = 1 (2.7 is N/A on a cluster)",
         "#   sql.tcp_port_1433_count (number)            : 2.11 query result",
         "#   sql.hide_instance (number)                  : SuperSocketNetLib HideInstance",
         "#   sql.sa_account_disabled (bool)              : sid 0x01 login is_disabled = 1",
         "#   sql.sa_login_name (string)                  : name of the sid 0x01 login",
         "#   sql.is_integrated_security_only (number)    : SERVERPROPERTY('IsIntegratedSecurityOnly')",
         "#   sql.num_error_logs (number)                 : NumErrorLogs, -1 when no limit is configured",
         "#   sql.login_audit_level (string)              : xp_loginconfig 'audit level' config_value",
         "#   sql.server_audit_actions[] = {audit_action_name, audit_enabled, specification_enabled, audited_result}",
         "#   sql.has_contained_databases (bool)",
         "#   sql.user_assemblies[] / symmetric_keys[] / asymmetric_keys[] : user-defined assemblies and non-system-db keys",
         "#   sql.connection_encrypt_options[]            : 7.4 query result (distinct encrypt_option values)",
         "#   sql.<offender list>[]                       : rows the audit query must not return ([] = none; missing = not collected):",
         "#     databases_trustworthy_on, contained_databases_auto_close_on, logins_named_sa, guest_connect_grants,",
         "#     orphaned_users, contained_db_sql_auth_users, public_role_extra_permissions, builtin_logins,",
         "#     local_group_logins, msdb_public_proxy_grants, msdb_admin_role_members,",
         "#     sysadmin_logins_without_expiration_check, sql_logins_without_check_policy, unencrypted_backups,",
         "#     databases_without_tde",
         "#   sql.attest.*                                : attestation booleans for the Manual recommendations",
         "", "import rego.v1", "",
         "# ── profile ──────────────────────────────────────────────────────────────────",
         "default _cis_profile := {}", "", "_cis_profile := p if {", "\tp := input.cis_profile", "\tis_object(p)", "}", "",
         'default platform := "database_engine"', "", 'platform := "aws_rds" if _cis_profile.platform == "aws_rds"', "",
         "default level := 2", "", "level := 1 if _cis_profile.level == 1", "",
         "default _rds := false", "", '_rds if platform == "aws_rds"', "",
         f'_selected_profiles := {{"{RDS1}"}} if _rds', "",
         f'_selected_profiles := {{"{DE1}"}} if {{', "\tnot _rds", "\tlevel == 1", "}", "",
         f'_selected_profiles := {{"{DE1}", "{DE2}"}} if {{', "\tnot _rds", "\tlevel == 2", "}", "",
         "in_scope contains id if {", "\tsome id, r in recommendations", "\tsome p in r.profiles", "\tp in _selected_profiles", "}", "",
         "default compliant := false", "", "compliant if {", "\tinput.sql", "\tcount(violations) == 0", "}",
         HELPERS.rstrip(), ""]
    fx = {}
    for rid in ids:
        r, s, i = recs[rid], RULES[rid], ident(rid)
        L += [f"# {rid} ({r['status']}, {'; '.join(r['profiles'])})", f"default r_{i} := false"]
        L += rule_lines(f"r_{i}", s["body"])
        if s["kind"] == "attest":
            L += [f"violations contains {rego_str(rid + ': ' + r['title'] + ' (manual review not attested)')} if {{",
                  f'\t"{rid}" in in_scope', f"\tnot r_{i}", "}", ""]
        else:
            L += [f"default e_{i} := false"] + rule_lines(f"e_{i}", s["ev"])
            L += [f"violations contains {rego_str(rid + ': evidence missing — ' + s['field'] + ' not collected')} if {{",
                  f'\t"{rid}" in in_scope', f"\tnot e_{i}", "}", "",
                  f"violations contains {rego_str(rid + ': ' + r['title'])} if {{", f'\t"{rid}" in in_scope', f"\te_{i}", f"\tnot r_{i}", "}", ""]
        for path, val in s["fx"]:
            d = fx
            parts = path if isinstance(path, list) else path.split(".")
            for p in parts[:-1]:
                d = d.setdefault(p, {})
            d[parts[-1]] = val
    recmap = {rid: {"title": recs[rid]["title"], "level": recs[rid]["level"], "assessment": recs[rid]["status"],
                    "profiles": recs[rid]["profiles"]} for rid in ids}
    L += ['violations contains "CIS SQL Server 2025: FAIL-CLOSED - facts \'input.sql\' were not supplied; the instance was not evaluated" if {',
          "\tnot input.sql", "}", "",
          "sections := " + json.dumps(sections, indent=1, ensure_ascii=False).replace("\n", "\n"), "",
          "# id -> CIS title, level, Automated/Manual status and the workbook profiles it belongs to",
          "recommendations := " + json.dumps(recmap, indent=1, ensure_ascii=False), "",
          "compliance_report := {", f'\t"benchmark": "{BENCH}",', '\t"profile": {"platform": platform, "level": level},',
          '\t"total_controls": count(in_scope),', f'\t"standard_total": {len(ids)},',
          f'\t"standard_total_source": "{BENCH}, recommendations counted from the CIS workbook",',
          '\t"sections": sections,', '\t"compliant": compliant,', '\t"violation_count": count(violations),', '\t"violations": violations,', "}", ""]
    out.mkdir(parents=True, exist_ok=True)
    (out / "cis_sql_server_2025.rego").write_text("\n".join(L))

    M = ["# METADATA", f'# title: "{BENCH} - compliance_report entrypoint"', "# custom:", "#   class: security",
         "#   framework: cis_sql_server_2025", "#   source: cis", "#   domains: [database, mssql]", f"package {PKG}.main", "",
         f"# OPA query path: /v1/data/{PKG.replace('.', '/')}/main/compliance_report",
         f"# Thin alias of data.{PKG}; that package already fails closed when input.sql is absent.", "",
         f"import data.{PKG} as bench", "import rego.v1", "",
         "default _facts_supplied := false", "", "_facts_supplied if count(object.keys(input)) > 0", "",
         'compliance_report := object.union(bench.compliance_report, {"facts_supplied": _facts_supplied})', "",
         "compliant := bench.compliant", "", "violations := bench.violations", ""]
    (out / "cis_sql_server_2025_main.rego").write_text("\n".join(M))

    fxs = json.dumps(fx, indent=1, ensure_ascii=False).replace("\n", "\n\t")
    p_rds = {"cis_profile": {"platform": "aws_rds", "level": 1}}
    T = [f"package {PKG}_test", "", "import rego.v1", "", f"import data.{PKG} as b", f"import data.{PKG}.main as m", "",
         f"green := {fxs}", "",
         "test_green_instance_is_compliant if {", "\tr := b.compliance_report with input as green", "\tr.compliant == true",
         "\tcount(r.violations) == 0", f"\tr.total_controls == {len(ids)}", f"\tr.standard_total == {len(ids)}", "}", "",
         "test_main_alias_matches if {", "\tr := m.compliance_report with input as green", "\tr.compliant == true",
         "\tr.facts_supplied == true", "\tm.compliant with input as green", "}", "",
         "test_fail_closed_on_empty_input if {", "\tr := b.compliance_report with input as {}", "\tr.compliant == false",
         "\tcount(r.violations) > 0", "\tm.compliance_report.facts_supplied == false with input as {}", "}", "",
         "test_every_recommendation_fires_without_facts if {", '\tv := b.violations with input as {"sql": {}}',
         '\tids := {id | some msg in v; id := split(msg, ":")[0]}', f"\tcount(ids) == {len(ids)}", "}", "",
         "test_recommendations_cover_the_workbook if {", f"\tcount(b.recommendations) == {len(ids)}",
         f'\tcount([r | some r in b.recommendations; r.assessment == "Manual"]) == {n_manual}',
         f'\tcount([r | some r in b.recommendations; r.level == 2]) == {len(by_prof[DE2])}', "}", "",
         "test_profile_scopes if {", f"\tcount(b.in_scope) == {len(ids)} with input as green",
         f'\tcount(b.in_scope) == {len(by_prof[DE1])} with input as object.union(green, {{"cis_profile": {{"level": 1}}}})',
         f"\tcount(b.in_scope) == {len(by_prof[RDS1])} with input as object.union(green, {json.dumps(p_rds)})", "}", "",
         "test_out_of_scope_recommendation_is_silent if {",
         '\tbad := json.patch(green, [{"op": "remove", "path": "/sql/databases_without_tde"}])',
         '\tv := b.violations with input as object.union(bad, {"cis_profile": {"level": 1}})',
         '\tcount([msg | some msg in v; startswith(msg, "7.5:")]) == 0', "}", "",
         "test_rds_allowances_only_apply_on_rds if {",
         '\tg := json.patch(green, [{"op": "replace", "path": "/sql/guest_connect_grants", "value": [{"database": "rdsadmin"}]}, '
         '{"op": "replace", "path": "/sql/builtin_logins", "value": [{"name": "[BUILTIN]\\\\Administrators"}]}, '
         '{"op": "replace", "path": "/sql/sysadmin_logins_without_expiration_check", "value": [{"name": "rdsa"}]}])',
         f"\tb.compliant with input as object.union(g, {json.dumps(p_rds)})",
         "\tv := b.violations with input as g",
         '\tcount([msg | some msg in v; regex.match(`^(3\\.2|3\\.9|4\\.2):`, msg)]) == 3', "}", "",
         "test_2_2_not_applicable_when_clr_strict_security_on if {",
         '\tok := json.patch(green, [{"op": "add", "path": "/sql/configurations/clr enabled", "value": {"value_configured": 1, "value_in_use": 1}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_2_7_not_applicable_on_cluster if {",
         '\tok := json.patch(green, [{"op": "replace", "path": "/sql/is_clustered", "value": true}, {"op": "remove", "path": "/sql/configurations/remote admin connections"}])',
         "\tb.compliant with input as ok", "}", "",
         "test_config_names_match_case_insensitively if {",
         '\tok := json.patch(green, [{"op": "remove", "path": "/sql/configurations/database mail xps"}, '
         '{"op": "add", "path": "/sql/configurations/Database Mail XPs", "value": {"value_configured": "0", "value_in_use": "0"}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_5_1_unlimited_error_logs_passes if {",
         '\tok := json.patch(green, [{"op": "replace", "path": "/sql/num_error_logs", "value": -1}])', "\tb.compliant with input as ok", "}", "",
         "test_5_3_all_logins_passes if {",
         '\tok := json.patch(green, [{"op": "replace", "path": "/sql/login_audit_level", "value": "ALL"}])', "\tb.compliant with input as ok", "}", "",
         "test_5_4_contained_groups_satisfy_contained_requirement if {",
         f'\trows := array.concat(green.sql.server_audit_actions, [{{"audit_action_name": g, "audit_enabled": true, "specification_enabled": true, "audited_result": "SUCCESS AND FAILURE"}} | some g in {json.dumps(CONTAINED_AUDIT_GROUPS)}])',
         '\tok := json.patch(green, [{"op": "replace", "path": "/sql/has_contained_databases", "value": true}, {"op": "replace", "path": "/sql/server_audit_actions", "value": rows}])',
         "\tb.compliant with input as ok", "}", ""]
    for rid in ids:
        r, s, i = recs[rid], RULES[rid], ident(rid)
        T += [f"# {rid} ({r['status']})", f"test_{i}_pass if {{", "\tv := b.violations with input as green",
              f'\tcount([msg | some msg in v; startswith(msg, "{rid}:")]) == 0', "}", ""]
        fail_msg = rid + ": " + r["title"] + (" (manual review not attested)" if s["kind"] == "attest" else "")
        for n, ops in enumerate(s["fails"], 1):
            suffix = "" if len(s["fails"]) == 1 else f"_{n}"
            T += [f"test_{i}_fail{suffix} if {{", f"\tbad := json.patch(green, {json.dumps(ops, ensure_ascii=False)})",
                  f"\t{rego_str(fail_msg)} in b.violations with input as bad", "}", ""]
        if s["kind"] != "attest":
            T += [f"test_{i}_evidence_missing if {{", f"\tbad := json.patch(green, {json.dumps(s['miss'], ensure_ascii=False)})",
                  f"\t{rego_str(rid + ': evidence missing — ' + s['field'] + ' not collected')} in b.violations with input as bad", "}", ""]
    (out / "tests").mkdir(exist_ok=True)
    (out / "tests" / "test_cis_sql_server_2025.rego").write_text("\n".join(T))
    print(f"cis_sql_server_2025: {len(ids)} recommendations ({n_auto} Automated, {n_manual} Manual); "
          f"profiles {', '.join(f'{p}={len(v)}' for p, v in by_prof.items())}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
