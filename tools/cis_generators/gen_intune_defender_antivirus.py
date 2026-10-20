#!/usr/bin/env python3
"""CIS Microsoft Intune for Microsoft Defender Antivirus Benchmark v1.0.0, from the benchmark's own workbook.

    python3 gen_intune_defender_antivirus.py CIS_Microsoft_Intune_for_Microsoft_Defender_Antivirus_Benchmark_v1.0.0.xlsx <repo root>

Needs openpyxl. Reads id, title, Assessment Status and profile out of the CIS workbook (CIS-licensed, never
committed), checks the ids against the setting table below and emits benchmarks/cis/intune_defender_antivirus/
cis_intune_defender_antivirus.rego, cis_intune_defender_antivirus_main.rego and tests/test_cis_intune_defender_antivirus.rego.

Same shape as gen_sql_server_2025.py / tools/cis_generators/gen_oracle_db_26ai.py (fail closed per recommendation,
pass/fail/evidence-missing tests). The evidence is the effective Intune policy value of each setting, keyed by its
CSP node name (the windows_server_2022_modular idea of one value per key, with the Intune key instead of the registry
path the workbook's audit reads). The table maps each recommendation to its CSP node; the OMA-URI goes into
`recommendations` so a collector author can map Graph / Settings Catalog output onto it.
"""
import json
import pathlib
import sys

import openpyxl

BENCH = "CIS Microsoft Intune for Microsoft Defender Antivirus Benchmark v1.0.0"
PKG = "cis.intune_defender_antivirus"
EXPECTED_TOTAL = 55
L1, L2 = "Level 1 - Workstation", "Level 2 - Workstation"
POL = "./Device/Vendor/MSFT/Policy/Config/Defender/"
CFG = "./Device/Vendor/MSFT/Defender/Configuration/"
ADMX = "./Device/Vendor/MSFT/Policy/Config/ADMX_MicrosoftDefenderAntivirus/"
ASR = "AttackSurfaceReductionRules"
TSDA = "ThreatSeverityDefaultAction"


def ptr(parts):
    return "/" + "/".join(p.replace("~", "~0").replace("/", "~1") for p in parts)


def sp(name):
    return ptr(["intune_defender", "settings", name])


def rep(name, value):
    return {"op": "replace", "path": sp(name), "value": value}


def rm(name):
    return {"op": "remove", "path": sp(name)}


def s(uri, check, arg, green, bads):
    """uri: OMA-URI (its last segment is the setting key); check: rego call template with {n} for the key;
    green: compliant value; bads: values that must each fail"""
    name = uri.rsplit("/", 1)[1]
    return {"uri": uri, "name": name, "body": [check.format(n=name, a=arg)], "ev": [f'_set("{name}")'],
            "fx": green, "fails": [[rep(name, b)] for b in bads], "miss": [rm(name)]}


def int_in(uri, allowed, green, bads):
    return s(uri, '_int_in("{n}", {a})', "{" + ", ".join(str(a) for a in allowed) + "}", green, bads)


def between(uri, lo, hi, green, bads):
    return s(uri, '_int_between("{n}", {a})', f"{lo}, {hi}", green, bads)


def at_least(uri, lo, green, bads):
    return s(uri, '_int_at_least("{n}", {a})', str(lo), green, bads)


def admx(name, state):
    other = "enabled" if state == "disabled" else "disabled"
    # a registry-style DWORD is not a policy state; it must not be read as one (the orientation differs per policy)
    return s(ADMX + name, '_admx_state("{n}", "{a}")', state, state, [other, 0, 1])


def asr(guid, allowed):
    guid = guid.lower()
    bad = {(1,): 2, (1, 2): 0}[tuple(allowed)]
    r = {"uri": POL + ASR, "name": ASR, "guid": guid, "allowed": allowed,
         "body": [f'_asr_ok("{guid}", {{{", ".join(map(str, allowed))}}})'], "ev": [f'_set("{ASR}")'],
         "fails": [("asr_value", bad), ("asr_absent", None), ("asr_conflict", 6)], "miss": [rm(ASR)]}
    return r


def tsda(level):
    return {"uri": POL + TSDA, "name": TSDA, "sev": str(level), "body": [f'_tsda_ok("{level}")'], "ev": [f'_set("{TSDA}")'],
            "fails": [("tsda_value", "6"), ("tsda_absent", None)], "miss": [rm(TSDA)]}


AUDIT_OR_BLOCK = (1, 2)
BLOCK = (1,)
RULES = {
    "1.14": admx("DisableRoutinelyTakingAction", "disabled"),
    "1.3.1": admx("Spynet_LocalSettingOverrideSpynetReporting", "disabled"),
    "1.3.2": admx("DisableBlockAtFirstSeen", "enabled"),
    "1.8.1": admx("RealtimeProtection_DisableScanOnRealtimeEnable", "enabled"),
    "1.10.1": admx("Reporting_DisablegenericrePorts", "disabled"),
    "2.1": int_in(POL + "AllowArchiveScanning", [1], 1, [0]),
    "2.2": int_in(POL + "AllowBehaviorMonitoring", [1], 1, [0]),
    "2.3": None,  # absent-or-1, built below
    "2.4": int_in(POL + "AllowEmailScanning", [1], 1, [0]),
    "2.5": int_in(POL + "AllowFullScanRemovableDriveScanning", [1], 1, [0]),
    "2.6": int_in(POL + "AllowOnAccessProtection", [1], 1, [0]),
    "2.7": int_in(POL + "AllowRealtimeMonitoring", [1], 1, [0]),
    "2.8": int_in(POL + "AllowIOAVProtection", [1], 1, [0]),
    "2.9": int_in(POL + "AllowScriptScanning", [1], 1, [0]),
    "2.10": asr("56a863a9-875e-4185-98a7-b882c64b5ce5", BLOCK),
    "2.11": asr("7674ba52-37eb-4a4f-a9a1-f0f9a1619a2c", BLOCK),
    "2.12": asr("d4f940ab-401b-4efc-aadc-ad5f3c50688a", AUDIT_OR_BLOCK),
    "2.13": asr("9e6c4e1f-7d60-472f-ba1a-a39ef669e4b2", BLOCK),
    "2.14": asr("be9ba2d9-53ea-4cdc-84e5-9b1eeee46550", BLOCK),
    "2.15": asr("01443614-cd74-433a-b99e-2ecdc07bfc25", AUDIT_OR_BLOCK),
    "2.16": asr("5beb7efe-fd9a-4556-801d-275e5ffc04cc", AUDIT_OR_BLOCK),
    "2.17": asr("d3e037e1-3eb8-44c8-a917-57927947596d", BLOCK),
    "2.18": asr("3b576869-a4ec-4529-8536-b80a7769e899", BLOCK),
    "2.19": asr("75668c1f-73b5-4cf0-bb93-3ecf5cb7cc84", BLOCK),
    "2.20": asr("26190899-1602-49e8-8b27-eb1d0a1ce869", AUDIT_OR_BLOCK),
    "2.21": asr("e6db77e5-3df2-4cf1-b95a-636979351e5b", BLOCK),
    "2.22": asr("d1e49aac-8f56-4280-b9ba-993a6d77406c", AUDIT_OR_BLOCK),
    "2.23": asr("33ddedf1-c6e0-47cb-833e-de6133960387", BLOCK),
    "2.24": asr("b2b3f03d-6a65-4f7b-a9c7-1c7ef74a9ba4", BLOCK),
    "2.25": asr("92e97fa1-2edf-4476-bdd6-9dd0b4dddc7b", BLOCK),
    "2.26": asr("c1db55ab-c21a-4637-bb3f-a12568109d35", AUDIT_OR_BLOCK),
    "2.27": int_in(POL + "CheckForSignaturesBeforeRunningScan", [1], 1, [0]),
    "2.28": int_in(POL + "CloudBlockLevel", [2, 4, 6], 2, [0, 1]),
    "2.29": at_least(POL + "CloudExtendedTimeout", 1, 50, [0]),
    "2.30": int_in(CFG + "DaysUntilAggressiveCatchupQuickScan", [7], 7, [0, 30]),
    "2.31": int_in(CFG + "EnableConvertWarnToBlock", [1], 1, [0]),
    "2.32": int_in(CFG + "Reporting/EnableDynamicSignatureDroppedEventReporting", [1], 1, [0]),
    "2.33": int_in(CFG + "EnableFileHashComputation", [1], 1, [0]),
    "2.34": int_in(POL + "EnableNetworkProtection", [1], 1, [0, 2]),
    "2.35": int_in(CFG + "HideExclusionsFromLocalUsers", [1], 1, [0]),
    "2.36": int_in(CFG + "OobeEnableRtpAndSigUpdate", [1], 1, [0]),
    "2.37": int_in(POL + "PUAProtection", [1], 1, [0, 2]),
    "2.38": int_in(CFG + "QuickScanIncludeExclusions", [1], 1, [0]),
    "2.39": int_in(POL + "RealTimeScanDirection", [0], 0, [1, 2]),
    "2.40": int_in(CFG + "BehavioralNetworkBlocks/RemoteEncryptionProtection/RemoteEncryptionProtectionAggressiveness", [1, 2], 1, [0]),
    "2.41": int_in(CFG + "BehavioralNetworkBlocks/RemoteEncryptionProtection/RemoteEncryptionProtectionConfiguredState", [1, 2], 1, [0, 4]),
    "2.42": int_in(POL + "ScanParameter", [1, 2], 1, [0]),
    "2.43": at_least(POL + "ScheduleQuickScanTime", 1, 120, [0]),
    "2.44": at_least(POL + "ScheduleScanTime", 1, 120, [0]),
    "2.45": between(POL + "ScheduleScanDay", 0, 7, 0, [8, -1]),
    "2.46": between(POL + "SignatureUpdateInterval", 1, 4, 4, [0, 5, 24]),
    "2.47": int_in(POL + "SubmitSamplesConsent", [1, 3], 1, [0, 2]),
    "2.48": tsda(2),
    "2.49": tsda(4),
    "2.50": tsda(5),
}
_cp = POL + "AllowCloudProtection"
RULES["2.3"] = {"uri": _cp, "name": "AllowCloudProtection",
                "body": ['_absent_or_int_in("AllowCloudProtection", {1})'], "ev": ["count(_settings) > 0"],
                "fx": 1, "fails": [[rep("AllowCloudProtection", 0)]],
                "miss": [{"op": "replace", "path": "/intune_defender/settings", "value": {}}],
                "field": "intune_defender.settings"}

HELPERS = r"""
# ── helpers ──────────────────────────────────────────────────────────────────
default _settings := {}

_settings := st if {
	st := input.intune_defender.settings
	is_object(st)
}

# every value the collector supplied for a CSP node name (matched case-insensitively)
_vals(name) := [v | some k, v in _settings; lower(k) == lower(name)]

_set(name) if count(_vals(name)) > 0

# CSP integers arrive as numbers or as numeric strings
_num(v) := v if is_number(v)

_num(v) := to_number(trim_space(v)) if {
	is_string(v)
	regex.match(`^\s*-?[0-9]+\s*$`, v)
}

_int_in(name, allowed) if {
	_set(name)
	every v in _vals(name) {
		_num(v) in allowed
	}
}

_int_between(name, lo, hi) if {
	_set(name)
	every v in _vals(name) {
		_num(v) >= lo
		_num(v) <= hi
	}
}

_int_at_least(name, lo) if {
	_set(name)
	every v in _vals(name) {
		_num(v) >= lo
	}
}

# 2.3: in effect when the setting is not configured or is 1; "not configured" only counts when settings were collected
_absent_or_int_in(name, _) if {
	count(_settings) > 0
	not _set(name)
}

_absent_or_int_in(name, allowed) if _int_in(name, allowed)

# ADMX-backed policies (section 1): the value is the policy STATE, "enabled" or "disabled" (the SyncML forms
# "<enabled/>" / "<disabled/>" are accepted too) -- never the registry DWORD the workbook's audit reads (that DWORD
# is inverted for some of them). Anything else fails.
_admx_norm(v) := trim_space(trim_suffix(trim_prefix(lower(trim_space(v)), "<"), "/>")) if is_string(v)

_admx_state(name, state) if {
	_set(name)
	every v in _vals(name) {
		_admx_norm(v) == state
	}
}

# "k=v|k=v" CSP strings (AttackSurfaceReductionRules, ThreatSeverityDefaultAction) -> set of [lower(k), v]
_pairs(name) := {[lower(trim_space(p[0])), trim_space(p[1])] |
	some v in _vals(name)
	is_string(v)
	some e in split(v, "|")
	p := split(e, "=")
	count(p) == 2
}

# the key is present with an allowed value, and no entry for the same key carries another value
_pair_ok(name, k, allowed) if {
	some p in _pairs(name)
	p[0] == k
	_num(p[1]) in allowed
	not _pair_conflict(name, k, allowed)
}

_pair_conflict(name, k, allowed) if {
	some p in _pairs(name)
	p[0] == k
	not _num(p[1]) in allowed
}

# ASR rule GUID (lowercase) -> 1 = Block, 2 = Audit, 6 = Warn, 0 = Off
_asr_ok(guid, allowed) if _pair_ok("AttackSurfaceReductionRules", guid, allowed)

# threat alert level (2 = Moderate/Medium, 4 = High, 5 = Severe) -> action 2 = Quarantine, 3 = Remove
_tsda_ok(sev) if _pair_ok("ThreatSeverityDefaultAction", sev, {2, 3})
"""


def key(rid):
    return [int(p) for p in rid.split(".")]


def ident(rid):
    return rid.replace(".", "_")


def rego_str(x):
    return json.dumps(x, ensure_ascii=False)


def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    recs, sections, combined = {}, {}, set()
    for sheet in wb.sheetnames:
        if sheet in ("License", "Overview"):
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        h = rows[0]
        ci = {c: h.index(c) for c in ("Section #", "Recommendation #", "Title", "Assessment Status")}
        for r in rows[1:]:
            rid, title, status = r[ci["Recommendation #"]], r[ci["Title"]], r[ci["Assessment Status"]]
            if not rid:
                sec = str(r[ci["Section #"]] or "")
                if sec and "." not in sec:
                    sections[sec] = " ".join(str(title).split())
                continue
            rid = str(rid).strip()
            if sheet == "Combined Profiles":
                combined.add(rid)
                continue
            title = " ".join(str(title).split()).replace('"', "'")
            level = int(sheet.split()[1])
            if rid in recs:
                raise SystemExit(f"{rid} appears in more than one profile sheet")
            recs[rid] = {"title": title, "status": status, "level": level, "profiles": [sheet]}
    if combined != set(recs):
        raise SystemExit(f"Combined Profiles disagrees with the profile sheets: {sorted(combined ^ set(recs))}")
    return recs, sections


def main(xlsx, root):
    out = pathlib.Path(root) / "benchmarks" / "cis" / "intune_defender_antivirus"
    recs, all_sections = read_workbook(xlsx)
    missing = sorted(set(recs) - set(RULES), key=key)
    extra = sorted(set(RULES) - set(recs), key=key)
    if missing or extra:
        raise SystemExit(f"setting table out of step with the workbook: missing {missing} extra {extra}")
    if len(recs) != EXPECTED_TOTAL:
        raise SystemExit(f"workbook has {len(recs)} recommendations, expected {EXPECTED_TOTAL}")
    if any(r["status"] != "Automated" for r in recs.values()):
        raise SystemExit("a Manual recommendation appeared; this generator has no attestation path")
    ids = sorted(recs, key=key)
    sections = {x: all_sections[x] for x in sorted({rid.split(".")[0] for rid in ids}, key=int)}
    l2 = [r for r in ids if recs[r]["level"] == 2]

    L = ["# METADATA", f'# title: "{BENCH}"', "# custom:", "#   class: security", "#   framework: cis_intune_defender_antivirus",
         "#   source: cis", "#   domains: [endpoint, windows, intune, antivirus]", f"package {PKG}", "",
         f"# {BENCH}.",
         f"# {len(ids)} recommendations, all Automated, every one in the benchmark; ids, titles and profiles read from the",
         "# CIS workbook (generated by gen_intune_defender_antivirus.py). Profiles: 'Level 1 - Workstation'",
         f"# ({len(ids) - len(l2)}) and 'Level 2 - Workstation' ({len(l2)}: {', '.join(l2)}). input.cis_profile.level (1 | 2,",
         "# default 2, which includes Level 1) selects the recommendations in scope.",
         "# Every check fails closed: setting not collected => '<id>: evidence missing — <field> not collected'.",
         "#",
         "# Input contract: the EFFECTIVE Intune policy value of each setting for the assessed device group, keyed by CSP",
         "# node name (the last segment of its OMA-URI; `recommendations[id].oma_uri` gives the full path):",
         "#   input.intune_defender.settings[<CSP node name>] = value     (names match case-insensitively)",
         "#     Policy/Config/Defender/* and Defender/Configuration/* integers : number or numeric string",
         "#     ADMX_MicrosoftDefenderAntivirus/* (section 1)                   : policy state \"enabled\" | \"disabled\" (or <enabled/> / <disabled/>)",
         "#     AttackSurfaceReductionRules                                     : CSP string \"<rule guid>=<0|1|2|6>|...\"",
         "#     ThreatSeverityDefaultAction                                     : CSP string \"<1|2|4|5>=<2|3|6>|...\"",
         "#   input.cis_profile.level (optional)",
         "# A setting with no value is not collected (fail closed), except 2.3, which the benchmark treats as in effect when",
         "# not configured; that holds only when the settings object itself was collected and is non-empty.",
         "", "import rego.v1", "",
         "default _cis_profile := {}", "", "_cis_profile := p if {", "\tp := input.cis_profile", "\tis_object(p)", "}", "",
         "default level := 2", "", "level := 1 if _cis_profile.level == 1", "",
         "in_scope contains id if {", "\tsome id, r in recommendations", "\tr.level <= level", "}", "",
         "default compliant := false", "", "compliant if {", "\tinput.intune_defender", "\tcount(violations) == 0", "}",
         HELPERS.rstrip(), ""]
    settings = {}
    asr_green, tsda_green = {}, {"2": "2", "4": "2", "5": "3"}
    for rid in ids:
        r, x, i = recs[rid], RULES[rid], ident(rid)
        field = x.get("field", f"intune_defender.settings.{x['name']}")
        L += [f"# {rid} ({r['status']}, {r['profiles'][0]}) {x['uri']}", f"default r_{i} := false", f"r_{i} if {{"]
        L += ["\t" + b for b in x["body"]] + ["}", "", f"default e_{i} := false", f"e_{i} if {{"] + ["\t" + b for b in x["ev"]] + ["}", ""]
        L += [f"violations contains {rego_str(rid + ': evidence missing — ' + field + ' not collected')} if {{",
              f'\t"{rid}" in in_scope', f"\tnot e_{i}", "}", "",
              f"violations contains {rego_str(rid + ': ' + r['title'])} if {{", f'\t"{rid}" in in_scope', f"\te_{i}", f"\tnot r_{i}", "}", ""]
        if "guid" in x:
            asr_green[x["guid"]] = 1
        elif "sev" not in x:
            settings[x["name"]] = x["fx"]
    # the green ASR string mixes case on purpose (the workbook prints the GUIDs both ways)
    settings[ASR] = "|".join(f"{(g.upper() if n % 2 else g)}={v}" for n, (g, v) in enumerate(asr_green.items()))
    settings[TSDA] = "|".join(f"{k}={v}" for k, v in tsda_green.items())
    recmap = {rid: {"title": recs[rid]["title"], "level": recs[rid]["level"], "assessment": recs[rid]["status"],
                    "profiles": recs[rid]["profiles"], "setting": RULES[rid]["name"], "oma_uri": RULES[rid]["uri"]}
              for rid in ids}
    L += ['violations contains "CIS Intune Defender Antivirus: FAIL-CLOSED - facts \'input.intune_defender\' were not supplied; the policy was not evaluated" if {',
          "\tnot input.intune_defender", "}", "",
          "sections := " + json.dumps(sections, indent=1, ensure_ascii=False), "",
          "# id -> CIS title, level, status, profile, and the CSP setting it reads (name and full OMA-URI)",
          "recommendations := " + json.dumps(recmap, indent=1, ensure_ascii=False), "",
          "compliance_report := {", f'\t"benchmark": "{BENCH}",', '\t"profile": {"level": level},',
          '\t"total_controls": count(in_scope),', f'\t"standard_total": {len(ids)},',
          f'\t"standard_total_source": "{BENCH}, recommendations counted from the CIS workbook",',
          '\t"sections": sections,', '\t"compliant": compliant,', '\t"violation_count": count(violations),', '\t"violations": violations,', "}", ""]
    out.mkdir(parents=True, exist_ok=True)
    (out / "cis_intune_defender_antivirus.rego").write_text("\n".join(L))

    M = ["# METADATA", f'# title: "{BENCH} - compliance_report entrypoint"', "# custom:", "#   class: security",
         "#   framework: cis_intune_defender_antivirus", "#   source: cis", "#   domains: [endpoint, windows, intune, antivirus]",
         f"package {PKG}.main", "", f"# OPA query path: /v1/data/{PKG.replace('.', '/')}/main/compliance_report",
         f"# Thin alias of data.{PKG}; that package already fails closed when input.intune_defender is absent.", "",
         f"import data.{PKG} as bench", "import rego.v1", "",
         "default _facts_supplied := false", "", "_facts_supplied if count(object.keys(input)) > 0", "",
         'compliance_report := object.union(bench.compliance_report, {"facts_supplied": _facts_supplied})', "",
         "compliant := bench.compliant", "", "violations := bench.violations", ""]
    (out / "cis_intune_defender_antivirus_main.rego").write_text("\n".join(M))

    green = {"intune_defender": {"settings": settings}}
    gs = json.dumps(green, indent=1, ensure_ascii=False).replace("\n", "\n\t")

    def asr_str(guid, val):
        items = dict(asr_green)
        if val is None:
            items.pop(guid)
        else:
            items[guid] = val
        return "|".join(f"{g}={v}" for g, v in items.items())

    T = [f"package {PKG}_test", "", "import rego.v1", "", f"import data.{PKG} as b", f"import data.{PKG}.main as m", "",
         f"green := {gs}", "",
         "test_green_policy_is_compliant if {", "\tr := b.compliance_report with input as green", "\tr.compliant == true",
         "\tcount(r.violations) == 0", f"\tr.total_controls == {len(ids)}", f"\tr.standard_total == {len(ids)}", "}", "",
         "test_main_alias_matches if {", "\tr := m.compliance_report with input as green", "\tr.compliant == true",
         "\tr.facts_supplied == true", "\tm.compliant with input as green", "}", "",
         "test_fail_closed_on_empty_input if {", "\tr := b.compliance_report with input as {}", "\tr.compliant == false",
         "\tcount(r.violations) > 0", "\tm.compliance_report.facts_supplied == false with input as {}", "}", "",
         "test_every_recommendation_fires_without_facts if {", '\tv := b.violations with input as {"intune_defender": {"settings": {}}}',
         '\tids := {id | some msg in v; id := split(msg, ":")[0]}', f"\tcount(ids) == {len(ids)}", "}", "",
         "test_recommendations_cover_the_workbook if {", f"\tcount(b.recommendations) == {len(ids)}",
         f"\tcount([r | some r in b.recommendations; r.level == 2]) == {len(l2)}", "}", "",
         "test_level_1_scope if {", f'\tcount(b.in_scope) == {len(ids) - len(l2)} with input as object.union(green, {{"cis_profile": {{"level": 1}}}})',
         f'\tb.compliant with input as object.union(json.patch(green, [{{"op": "remove", "path": "{sp("EnableConvertWarnToBlock")}"}}]), {{"cis_profile": {{"level": 1}}}})', "}", "",
         "test_numeric_strings_and_key_case_are_accepted if {",
         f'\tok := json.patch(green, [{{"op": "remove", "path": "{sp("CloudBlockLevel")}"}}, {{"op": "add", "path": "{sp("cloudblocklevel")}", "value": " 4 "}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_non_numeric_value_fails if {",
         f'\tbad := json.patch(green, [{{"op": "replace", "path": "{sp("AllowArchiveScanning")}", "value": "allowed"}}])',
         f'\t{rego_str("2.1: " + recs["2.1"]["title"])} in b.violations with input as bad', "}", "",
         "test_audit_or_higher_rule_accepts_audit if {",
         f'\tok := json.patch(green, [{{"op": "replace", "path": "{sp(ASR)}", "value": {rego_str(asr_str(RULES["2.12"]["guid"], 2))}}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_2_3_absent_is_in_effect if {",
         f'\tok := json.patch(green, [{{"op": "remove", "path": "{sp("AllowCloudProtection")}"}}])', "\tb.compliant with input as ok", "}", "",
         "test_admx_state_is_case_insensitive if {",
         f'\tok := json.patch(green, [{{"op": "replace", "path": "{sp("DisableBlockAtFirstSeen")}", "value": "Enabled"}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_admx_syncml_state_is_accepted if {",
         f'\tok := json.patch(green, [{{"op": "replace", "path": "{sp("DisableBlockAtFirstSeen")}", "value": "<enabled/>"}}, {{"op": "replace", "path": "{sp("DisableRoutinelyTakingAction")}", "value": "<disabled/>"}}])',
         "\tb.compliant with input as ok", "}", "",
         "test_admx_syncml_wrong_state_fails if {",
         f'\tbad := json.patch(green, [{{"op": "replace", "path": "{sp("DisableBlockAtFirstSeen")}", "value": "<disabled/>"}}])',
         f'\t{rego_str("1.3.2: " + recs["1.3.2"]["title"])} in b.violations with input as bad', "}", ""]
    for rid in ids:
        r, x, i = recs[rid], RULES[rid], ident(rid)
        field = x.get("field", f"intune_defender.settings.{x['name']}")
        T += [f"# {rid}", f"test_{i}_pass if {{", "\tv := b.violations with input as green",
              f'\tcount([msg | some msg in v; startswith(msg, "{rid}:")]) == 0', "}", ""]
        fails = []
        for f in x["fails"]:
            if isinstance(f, tuple):
                kind, val = f
                if kind.startswith("asr"):
                    if kind == "asr_conflict":  # the GUID listed twice, once allowed and once not
                        v = settings[ASR] + f"|{x['guid']}={val}"
                    else:
                        v = asr_str(x["guid"], val)
                    fails.append([rep(ASR, v)])
                else:
                    items = dict(tsda_green)
                    if val is None:
                        items.pop(x["sev"])
                    else:
                        items[x["sev"]] = val
                    fails.append([rep(TSDA, "|".join(f"{k}={v}" for k, v in items.items()))])
            else:
                fails.append(f)
        for n, ops in enumerate(fails, 1):
            suffix = "" if len(fails) == 1 else f"_{n}"
            T += [f"test_{i}_fail{suffix} if {{", f"\tbad := json.patch(green, {json.dumps(ops, ensure_ascii=False)})",
                  f"\t{rego_str(rid + ': ' + r['title'])} in b.violations with input as bad", "}", ""]
        T += [f"test_{i}_evidence_missing if {{", f"\tbad := json.patch(green, {json.dumps(x['miss'], ensure_ascii=False)})",
              f"\t{rego_str(rid + ': evidence missing — ' + field + ' not collected')} in b.violations with input as bad", "}", ""]
    (out / "tests").mkdir(exist_ok=True)
    (out / "tests" / "test_cis_intune_defender_antivirus.rego").write_text("\n".join(T))
    print(f"cis_intune_defender_antivirus: {len(ids)} recommendations (all Automated; Level 2: {', '.join(l2)})")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
