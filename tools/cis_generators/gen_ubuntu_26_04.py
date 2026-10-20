#!/usr/bin/env python3
"""CIS Ubuntu Linux 26.04 LTS Benchmark v1.0.0: every recommendation, from the benchmark's own workbook.

    python3 tools/cis_generators/gen_ubuntu_26_04.py CIS_Ubuntu_Linux_26.04_LTS_Benchmark_v1.0.0.xlsx [fields.json]

Needs openpyxl. Reads id, title, Assessment Status and profiles (Level 1/2 - Server/Workstation sheets)
out of the CIS workbook (CIS-licensed, never committed), checks the ids against the hand-authored logic
table below (a recommendation without logic, logic without a recommendation, or a logic keyword missing
from the workbook title aborts) and emits benchmarks/cis/ubuntu_26_04/: one package per benchmark section
(cis_ubuntu_26_04.<section>), the roll-up package cis_ubuntu_26_04 (controls table, profile selector,
section coverage) and the bridge cis_ubuntu_26_04.main (compliance_report, same keys as cis_rhel9.main),
plus tests.

Adapted from gen_suse_16.py. Layout and input contract follow benchmarks/cis/rhel_9 (the modules the
shared Linux mapper ansible/tasks/cis_linux_build_opa_input.yml feeds): wherever the control is the same,
the rhel_9 field name is reused. Ubuntu-specific facts (dpkg packages, apt, AppArmor, ufw, Ubuntu unit
names such as ssh.service / cron.service / chrony.service, pam-auth-update common-* files) use new fields;
the optional fields.json output lists the input paths each control reads.

Rules of the generated code:
  * every Automated recommendation evaluates real facts; facts not collected =>
    "<id>: evidence missing — <field> not collected" (fail closed, never a pass);
  * every Manual recommendation is an attestation boolean input.attest.<name> (not attested => violation);
  * messages are our words; only the id and the short CIS title are carried from the workbook;
  * the tests give every recommendation a pass and a fail case, plus an evidence-missing case for
    every Automated one.
"""
import json
import pathlib
import re
import sys

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE.parent.parent / "benchmarks" / "cis" / "ubuntu_26_04"
PKG = "cis_ubuntu_26_04"
FILE_STEM = "cis_ubuntu_26_04"
GEN = "tools/cis_generators/gen_ubuntu_26_04.py"
BENCH = "CIS Ubuntu Linux 26.04 LTS Benchmark v1.0.0"
EXPECTED_TOTAL = 350
EXPECTED_AUTOMATED = 333
EXPECTED_MANUAL = 17
AS_OF = "v2.4.2"
SKIP_SHEETS = {"License", "Overview", "Combined Profiles"}

# module name, title, id prefixes (Ubuntu 26.04 section numbering)
SECTIONS = [
    ("filesystem", "Section 1.1: Filesystem", ["1.1."]),
    ("package_management", "Section 1.2: Package Management (apt)", ["1.2."]),
    ("apparmor", "Section 1.3: Mandatory Access Control (AppArmor)", ["1.3."]),
    ("bootloader", "Section 1.4: Bootloader", ["1.4."]),
    ("process_hardening", "Section 1.5: Additional Process Hardening", ["1.5."]),
    ("banners", "Section 1.6: Command Line Warning Banners", ["1.6."]),
    ("gdm", "Section 1.7: GNOME Display Manager", ["1.7."]),
    ("services", "Sections 2.1-2.2: Server and Client Services", ["2.1.", "2.2."]),
    ("time_sync", "Section 2.3: Time Synchronization", ["2.3."]),
    ("job_schedulers", "Section 2.4: Job Schedulers", ["2.4."]),
    ("network", "Section 3: Network", ["3."]),
    ("firewall", "Section 4: Host Based Firewall (ufw)", ["4."]),
    ("ssh", "Section 5.1: SSH Server", ["5.1."]),
    ("sudo", "Section 5.2: Privilege Escalation", ["5.2."]),
    ("pam", "Section 5.3: Pluggable Authentication Modules", ["5.3."]),
    ("accounts", "Section 5.4: User Accounts and Environment", ["5.4."]),
    ("logging", "Section 6.1: System Logging", ["6.1."]),
    ("auditd", "Section 6.2: System Auditing (auditd)", ["6.2."]),
    ("integrity", "Section 6.3: Integrity Checking (AIDE)", ["6.3."]),
    ("file_permissions", "Section 7.1: System File and Directory Access", ["7.1."]),
    ("user_group", "Section 7.2: Local User and Group Settings", ["7.2."]),
]


def section_of(rid):
    for name, _, prefixes in SECTIONS:
        if any(rid.startswith(p) for p in prefixes):
            return name
    raise SystemExit(f"no section for {rid}")


# ── input paths ────────────────────────────────────────────────────────────────
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def ref(path):
    """['a', 'b-c'] -> input.a["b-c"]"""
    out = "input"
    for p in path:
        out += f".{p}" if IDENT.match(p) else f"[{json.dumps(p)}]"
    return out


def dotted(path):
    return ".".join(path)


def nest(path, value):
    d = value
    for p in reversed(path):
        d = {p: d}
    return d


def merge(*ds):
    out = {}
    for d in ds:
        for k, v in d.items():
            if k in out and isinstance(out[k], dict) and isinstance(v, dict):
                out[k] = merge(out[k], v)
            else:
                out[k] = v
    return out


def P(s):
    """'a.b.c' -> ['a','b','c'] (keys containing dots are passed as lists by callers)"""
    return s.split(".")


# ── a control's logic ──────────────────────────────────────────────────────────
class Spec:
    """ev: rego lines (AND) proving the facts were collected (alternatives: list of lists in ev_alts);
    passes: list of alternatives, each a list of rego lines (AND); field: evidence-missing message field;
    tp/tf/tm: test inputs that pass / fail / lack evidence; reads: input paths read (for the report)."""

    def __init__(self, ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
        self.ev_alts = ev_alts if ev_alts is not None else [ev]
        self.passes = passes
        self.field = field
        self.tp, self.tf = tp, tf
        self.tm = {} if tm is None else tm
        self.reads = reads or []
        self.manual = False
        # True when an all-empty input (strings "", lists []) is a genuine pass: offender lists, where an
        # empty list collected by the collector means "none found". Otherwise a hollow-input test is generated.
        self.hollow_ok = False


class Attest:
    def __init__(self, name):
        self.name = name
        self.manual = True
        self.reads = [["attest", name]]


# ── builders ───────────────────────────────────────────────────────────────────
def kmod(listkey, name):
    lst = ref([listkey, "modules"])
    ev = [f'some e in {lst}', f'e.module == {json.dumps(name)}', 'kmod_evidence(e)']
    passes = [[f'count([e | some e in {lst}; e.module == {json.dumps(name)}]) > 0',
               f'every e in [e | some e in {lst}; e.module == {json.dumps(name)}] {{ kmod_unavailable(e) }}']]
    ok = {"module": name, "disabled": True, "loaded": False}
    return Spec(ev, passes, f"{listkey}.modules[{name}]",
                nest([listkey, "modules"], [ok]),
                nest([listkey, "modules"], [{"module": name, "disabled": False, "loaded": False}]),
                nest([listkey, "modules"], [{"module": "other", "disabled": True, "loaded": False}]),
                reads=[[listkey, "modules"]])


MOUNTS = ["ansible_facts", "ansible_mounts"]


def mnt_sep(path, extra_tmp=False):
    passes = [[f'mounted({json.dumps(path)})']]
    if extra_tmp:
        passes = [[f'mounted({json.dumps(path)})', 'not tmp_mount_unit_off']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "tmpfs", "options": "rw,nodev,nosuid,noexec"}]),
                nest(MOUNTS, [other]), {}, reads=[MOUNTS] + ([["services", "tmp.mount"]] if extra_tmp else []))


def mnt_opt(path, opt):
    # The option applies to the partition when it exists; the mount table is the collected fact
    # (the separate-partition recommendation fails on its own when the partition is missing).
    passes = [[f'not mounted({json.dumps(path)})'],
              [f'every m in mounts_at({json.dumps(path)}) {{ {json.dumps(opt)} in mount_opts(m) }}']]
    other = {"mount": "/", "fstype": "xfs", "options": "rw"}
    good = "rw,nodev,nosuid,noexec"
    bad = ",".join(o for o in ["rw", "nodev", "nosuid", "noexec"] if o != opt)
    return Spec(['mounts_collected'], passes, "ansible_facts.ansible_mounts",
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": good}]),
                nest(MOUNTS, [other, {"mount": path, "fstype": "xfs", "options": bad}]),
                {}, reads=[MOUNTS])


def sysctl(param, allowed, ipv6=False):
    vals = json.dumps(allowed)
    run, per = ref(["sysctl", param]), ref(["sysctl_persisted", param])
    ev = [f'{run} != null', f'{per} != null']
    passes = [[f'sprintf("%v", [{run}]) in {vals}', f'sprintf("%v", [{per}]) in {vals}']]
    good = {"sysctl": {param: allowed[0]}, "sysctl_persisted": {param: allowed[0]}}
    bad_v = "99"
    tf = {"sysctl": {param: allowed[0]}, "sysctl_persisted": {param: bad_v}}
    ev_alts = [ev]
    reads = [["sysctl", param], ["sysctl_persisted", param]]
    if ipv6:  # "IF IPv6 is enabled": a host with IPv6 disabled passes on that collected fact
        ev_alts = [ev, ['input.ipv6.disabled == true']]
        passes = passes + [['input.ipv6.disabled == true']]
        reads.append(["ipv6", "disabled"])
    return Spec(None, passes, f"sysctl/sysctl_persisted[{param}]", good, tf, {}, reads=reads, ev_alts=ev_alts)


def pkg_present(pkg, path=None):
    p = path or ["packages", pkg]
    return Spec([f'is_boolean({ref(p)})'], [[f'{ref(p)} == true']], dotted(p),
                nest(p, True), nest(p, False), {}, reads=[p])


def pkg_absent(pkgs, path=None):
    if path:  # a single rhel_9-style boolean
        return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == false']], dotted(path),
                    nest(path, False), nest(path, True), {}, reads=[path])
    ev = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    passes = [[f'{ref(["packages", p])} == false' for p in pkgs]]
    tp = {"packages": {p: False for p in pkgs}}
    tf = {"packages": {p: (i == 0) for i, p in enumerate(pkgs)}}
    return Spec(ev, passes, "packages[" + ",".join(pkgs) + "]", tp, tf, {"packages": {"audit": True}},
                reads=[["packages", p] for p in pkgs])


def not_in_use(units_by_pkg):
    """CIS 'not in use': every package not installed, OR (installed for a dependency) each unit the audit
    names for it neither enabled nor active. Ansible's service_facts lists .service units only, so a
    unit missing from input.services proves nothing: for an installed package every named unit must be
    reported (LoadState not-found is fine), or the evidence is missing."""
    pkgs = list(units_by_pkg)
    pk = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    none_installed = [f'{ref(["packages", p])} == false' for p in pkgs]
    m = json.dumps(units_by_pkg)
    reported = f'every p, us in {m} {{ unit_set_reported(p, us) }}'
    ev_alts = [pk + none_installed, pk + ['services_collected', reported]]
    passes = [none_installed,
              ['services_collected', reported, f'every p, us in {m} {{ unit_set_off(p, us) }}']]
    tp = {"packages": {p: False for p in pkgs}}
    allu = {u: {"state": "inactive", "status": "disabled"} for us in units_by_pkg.values() for u in us}
    first = units_by_pkg[pkgs[0]][0]
    tf = {"packages": {p: True for p in pkgs},
          "services": dict(allu, **{first: {"state": "running", "status": "enabled"}},
                           **{"ssh.service": {"state": "running", "status": "enabled"}})}
    tm = {"packages": {p: True for p in pkgs},
          "services": {"ssh.service": {"state": "running", "status": "enabled"}}}
    return Spec(None, passes, "packages[" + ",".join(pkgs) + "] / services[" + ",".join(
        u for us in units_by_pkg.values() for u in us) + "]", tp, tf, tm,
                reads=[["packages", p] for p in pkgs] + [["services"]], ev_alts=ev_alts)


def units_off(units):
    """The named units must each be reported (not-found is fine) and be neither enabled nor active."""
    passes = [[f'every u in {json.dumps(units)} {{ not unit_on(u) }}']]
    other = {"ssh.service": {"state": "running", "status": "enabled"}}
    nf = {u: {"state": "inactive", "status": "not-found"} for u in units}
    return Spec(['services_collected', f'every u in {json.dumps(units)} {{ unit_reported(u) }}'], passes,
                "services[" + ",".join(units) + "]",
                {"services": merge(other, nf)},
                {"services": merge(other, nf, {units[0]: {"state": "running", "status": "enabled"}})},
                {"services": other}, reads=[["services"]])


def unit_running(unit):
    passes = [[f'unit_enabled_active({json.dumps(unit)})']]
    other = {"ssh.service": {"state": "running", "status": "enabled"}}
    return Spec(['services_collected'], passes, "services",
                {"services": merge(other, {unit: {"state": "running", "status": "enabled"}})},
                {"services": merge(other, {unit: {"state": "stopped", "status": "disabled"}})},
                {}, reads=[["services"]])


def boolean(path, want):
    return Spec([f'is_boolean({ref(path)})'], [[f'{ref(path)} == {json.dumps(want)}']], dotted(path),
                nest(path, want), nest(path, not want), {}, reads=[path])


def access(mode, owner, group, mx, owners, groups, exists=None, optional=False):
    """One file: mode within mx, owner in owners, group in groups. exists: path of a boolean; when
    optional, a collected 'does not exist' passes; when required, it fails."""
    ev = [f'{ref(mode)} != null', f'{ref(owner)} != null', f'{ref(group)} != null']
    ok = [f'mode_within({ref(mode)}, {json.dumps(mx)})', f'{ref(owner)} in {json.dumps(owners)}',
          f'{ref(group)} in {json.dumps(groups)}']
    tp = merge(nest(mode, mx), nest(owner, owners[0]), nest(group, groups[0]))
    tf = merge(nest(mode, "0777"), nest(owner, owners[0]), nest(group, groups[0]))
    ev_alts, passes, reads = [ev], [ok], [mode, owner, group]
    if exists:
        reads.append(exists)
        ev_alts = [[f'{ref(exists)} == true'] + ev, [f'{ref(exists)} == false']]
        passes = [[f'{ref(exists)} == true'] + ok]
        tp = merge(tp, nest(exists, True))
        tf = merge(tf, nest(exists, True))
        if optional:
            passes.append([f'{ref(exists)} == false'])
        else:
            tf = nest(exists, False)
    return Spec(None, passes, "/".join(dotted(p) for p in reads), tp, tf, {}, reads=reads, ev_alts=ev_alts)


def access_list(path, mx, owners, groups, mode="mode", owner="owner", group="group", sample="/x", nonempty=True):
    """An inventory of files that always exist on a host must be non-empty to count as evidence."""
    lst = ref(path)
    passes = [[f'every f in {lst} {{ list_item_ok(f, {json.dumps(mode)}, {json.dumps(owner)}, '
               f'{json.dumps(group)}, {json.dumps(mx)}, {json.dumps(owners)}, {json.dumps(groups)}) }}']]
    good = {"path": sample, mode: mx, owner: owners[0], group: groups[0]}
    bad = dict(good, **{mode: "0777"})
    ev = [f'is_array({lst})'] + ([f'count({lst}) > 0'] if nonempty else [])
    return Spec(ev, passes, dotted(path), nest(path, [good]), nest(path, [good, bad]), {}, reads=[path])


def crit(path, mx, groups, optional=False):
    """critical_files.analysis entry {file, exists, actual_mode, actual_owner, actual_group} (rhel_9 shape);
    the expected_* / *_compliant fields the rhel_9 collector pre-computes are not trusted here."""
    p = json.dumps(path)
    ent = f'[f | some f in object.get(input, ["critical_files", "analysis"], []); f.file == {p}]'
    ok = (f'every f in {ent} {{ f.exists == true; mode_within(f.actual_mode, {json.dumps(mx)}); '
          f'f.actual_owner == "root"; f.actual_group in {json.dumps(groups)} }}')
    ev = [f'count({ent}) > 0', f'every f in {ent} {{ is_boolean(f.exists) }}']
    passes = [[f'count({ent}) > 0', ok]]
    if optional:
        passes.append([f'count({ent}) > 0', f'every f in {ent} {{ f.exists == false }}'])
    good = {"file": path, "exists": True, "actual_mode": mx, "actual_owner": "root", "actual_group": groups[0]}
    bad = dict(good, actual_mode="0666")
    other = {"file": "/etc/other", "exists": True, "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"}
    return Spec(ev, passes, f"critical_files.analysis[{path}]",
                {"critical_files": {"analysis": [other, good]}},
                {"critical_files": {"analysis": [other, bad]}},
                {"critical_files": {"analysis": [other]}}, reads=[["critical_files", "analysis"]])


def attest(name):
    return Attest(name)


def custom(ev, passes, field, tp, tf, tm=None, reads=None, ev_alts=None):
    return Spec(ev, passes, field, tp, tf, tm, reads=reads, ev_alts=ev_alts)


# ── sshd (effective config; `sshd -T` output preferred, raw sshd_config accepted) ──────────────
SSH_RAW = ["ssh", "sshd_config_raw"]


def ssh_in(lines):
    return {"ssh": {"sshd_config_raw": "\n".join(lines)}}


def sshd(key, cond, good, bad, explicit=False):
    """cond: rego expression over `v` (the effective lowercase value). explicit: the directive must be
    present (no reliable compiled default), otherwise its absence is missing evidence."""
    ev = ['ssh_collected', f'sshd_first({json.dumps(key)})']
    passes = [[f'v := lower(sshd_value({json.dumps(key)}))', cond]]
    tm = ssh_in(["port 22"])
    return Spec(ev, passes, f"ssh.sshd_config_raw (effective {key})",
                ssh_in(["port 22", f"{key} {good}"]), ssh_in(["port 22", f"{key} {bad}"]), tm, reads=[SSH_RAW])


def text_has(path, regex, want, good, bad):
    """want=True: some active line of the text matches regex; False: none does."""
    t = ref(path)
    cond = f'regex_in_text({t}, {json.dumps(regex)})'
    return Spec([f'is_string({t})'], [[cond if want else f'not {cond}']], dotted(path),
                nest(path, good), nest(path, bad), {}, reads=[path])


def conf(path, key_re, cond, good, bad, default=None):
    """key = value configuration text; last occurrence wins. cond is an expression over `v`.
    default: the documented compiled-in value used when the key is not set (None: unset fails)."""
    t = ref(path)
    passes = [[f'v := conf_last({t}, {json.dumps(key_re)})', cond]]
    if default is not None:
        passes.append([f'not conf_last({t}, {json.dumps(key_re)})', f'v := {json.dumps(default)}', cond])
    return Spec([f'is_string({t})'], passes, dotted(path), nest(path, good), nest(path, bad), {}, reads=[path])


# ── Rego helpers copied into every module (no file may call a function defined in another) ────
HELPERS = r'''
# Parse an octal permission string ("644", "0644", "0o644") to an integer; undefined if not octal.
octal(s) := n if {
	t := trim_prefix(lower(trim_space(sprintf("%v", [s]))), "0o")
	regex.match(`^[0-7]{1,4}$`, t)
	p := concat("", [substring("0000", 0, 4 - count(t)), t])
	ds := [to_number(substring(p, i, 1)) | some i in numbers.range(0, 3)]
	n := ((ds[0] * 512) + (ds[1] * 64)) + ((ds[2] * 8) + ds[3])
}

# true when mode grants nothing beyond max (both octal strings)
mode_within(mode, max) if bits.and(octal(mode), bits.and(bits.negate(octal(max)), 4095)) == 0

# Active (non-blank, non-comment) lines of a configuration text.
active_lines(text) := [l |
	some raw in split(text, "\n")
	l := trim_space(raw)
	l != ""
	not startswith(l, "#")
]

# a collected text: a non-blank string ("" is not evidence)
text_collected(t) if {
	is_string(t)
	trim_space(t) != ""
}

regex_in_text(text, re) if {
	some l in active_lines(text)
	regex.match(re, l)
}

# Values of a `key = value` / `key value` directive (key is a case-insensitive regex); last one wins.
conf_values(text, key_re) := [trim(m[1], "\"' ") |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(concat("", ["(?i)^", key_re, `(?:\s*=\s*|\s+)(.*?)\s*$`]), l, 1)
]

conf_last(text, key_re) := vs[count(vs) - 1] if {
	vs := conf_values(text, key_re)
	count(vs) > 0
}

num(v) := to_number(trim_space(sprintf("%v", [v]))) if regex.match(`^\s*-?[0-9]+\s*$`, sprintf("%v", [v]))

# ── packages / services ──
services_collected if {
	is_object(input.services)
	count(input.services) > 0
}

unit_on(u) if input.services[u].status == "enabled"

unit_on(u) if lower(input.services[u].state) in {"running", "active"}

# every unit named for an installed package is reported in the inventory
unit_set_reported(p, _) if input.packages[p] == false

unit_set_reported(p, us) if {
	input.packages[p] == true
	every u in us { unit_reported(u) }
}

unit_reported(u) if {
	is_string(input.services[u].status)
	input.services[u].status != ""
}

unit_set_off(p, _) if input.packages[p] == false

unit_set_off(p, us) if {
	input.packages[p] == true
	every u in us { not unit_on(u) }
}

unit_enabled_active(u) if {
	input.services[u].status == "enabled"
	lower(input.services[u].state) in {"running", "active"}
}

# tmp.mount masked or disabled means systemd will not mount /tmp at boot
tmp_mount_unit_off if input.services["tmp.mount"].status in {"masked", "disabled"}

# ── kernel modules: {module, disabled, loaded, available?}; "unknown" proves nothing ──
kmod_evidence(e) if is_boolean(e.available)

kmod_evidence(e) if {
	is_boolean(e.disabled)
	is_boolean(e.loaded)
}

kmod_unavailable(e) if e.available == false

kmod_unavailable(e) if {
	e.disabled == true
	e.loaded == false
}

# ── mounts ──
mounts_collected if {
	is_array(input.ansible_facts.ansible_mounts)
	count(input.ansible_facts.ansible_mounts) > 0
}

mounts_at(path) := [m | some m in input.ansible_facts.ansible_mounts; m.mount == path]

mounted(path) if count(mounts_at(path)) > 0

mount_opts(m) := split(m.options, ",") if is_string(m.options)

mount_opts(m) := m.options if is_array(m.options)

# ── file lists ──
list_item_ok(f, mk, ok, gk, max, owners, groups) if {
	mode_within(f[mk], max)
	f[ok] in owners
	f[gk] in groups
}

# ── sshd: first occurrence wins (sshd semantics); `sshd -T` prints effective values ──
ssh_collected if {
	is_string(input.ssh.sshd_config_raw)
	trim_space(input.ssh.sshd_config_raw) != ""
}

sshd_pairs := [[lower(m[1]), trim_space(m[2])] |
	some l in active_lines(input.ssh.sshd_config_raw)
	some m in regex.find_all_string_submatch_n(`^([A-Za-z0-9]+)(?:\s*=\s*|\s+)(.*)$`, l, 1)
]

sshd_first(key) := vs[0] if {
	vs := [p[1] | some p in sshd_pairs; p[0] == key]
	count(vs) > 0
}

# effective value: `sshd -T` prints every directive, so a missing directive means the effective
# configuration was not collected (a raw sshd_config misses /usr/etc/ssh and sshd_config.d)
sshd_value(key) := sshd_first(key)

sshd_list(key) := [trim_space(x) | some x in split(lower(sshd_value(key)), ",")]

# seconds from an sshd time value ("60", "60s", "1m")
sshd_seconds(v) := to_number(v) if regex.match(`^[0-9]+$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "s")) if regex.match(`^[0-9]+s$`, v)

sshd_seconds(v) := to_number(trim_suffix(v, "m")) * 60 if regex.match(`^[0-9]+m$`, v)

# ── /etc/passwd, /etc/group (rhel_9 fields user_group.passwd_content / group_content) ──
rows(text) := [f |
	some l in active_lines(text)
	f := split(l, ":")
]

passwd_rows := [f | some f in rows(input.user_group.passwd_content); count(f) >= 7]

group_rows := [f | some f in rows(input.user_group.group_content); count(f) >= 3]

valid_shells := {s |
	some s in active_lines(input.user_group.shells_content)
	not regex.match(`/nologin$`, s)
}
'''


# ════════════════════════════════════════════════════════════════════════════════════════════
# LOGIC TABLE — one entry per recommendation, written from each audit procedure of the 26.04
# workbook. put(id, keyword, spec): the keyword must occur in the workbook title of that id
# (case-insensitive), so a logic entry attached to a renumbered id aborts the run.
# ════════════════════════════════════════════════════════════════════════════════════════════
L = {}
KW = {}


def put(rid, kw, spec):
    if rid in L:
        raise SystemExit(f"{rid} defined twice")
    L[rid] = spec
    KW[rid] = kw
    return spec


def gate_pkg(spec, pkg):
    """'IF <pkg> is installed': a collected 'not installed' passes (and counts as evidence)."""
    g = f'{ref(["packages", pkg])} == false'
    spec.ev_alts = spec.ev_alts + [[g]]
    spec.passes = spec.passes + [[g]]
    spec.reads = spec.reads + [["packages", pkg]]
    spec.field = f"{spec.field} (or packages.{pkg})"
    return spec


def pkgs_present(pkgs):
    ev = [f'is_boolean({ref(["packages", p])})' for p in pkgs]
    ok = [f'{ref(["packages", p])} == true' for p in pkgs]
    return Spec(ev, [ok], "packages[" + ",".join(pkgs) + "]", {"packages": {p: True for p in pkgs}},
                {"packages": {p: (i != 0) for i, p in enumerate(pkgs)}}, {}, reads=[["packages", p] for p in pkgs])


def dir_access(base, mx, owners=("root",), groups=("root",), optional=False):
    b = list(base)
    return access(b + ["mode"], b + ["owner"], b + ["group"], mx, list(owners), list(groups),
                  exists=(b + ["exists"]) if optional else None, optional=optional)


FSM = "disabled_filesystem_modules"   # mapper field (rhel_9 shape)
NETM = "disabled_network_modules"     # same shape, network modules (new)

# 1.1.1 filesystem kernel modules
for rid, mod in [("1.1.1.1", "cramfs"), ("1.1.1.2", "freevxfs"), ("1.1.1.3", "hfs"), ("1.1.1.4", "hfsplus"),
                 ("1.1.1.5", "jffs2"), ("1.1.1.6", "overlay"), ("1.1.1.7", "squashfs"), ("1.1.1.8", "udf"),
                 ("1.1.1.9", "firewire-core"), ("1.1.1.10", "usb-storage")]:
    put(rid, mod + " kernel module", kmod(FSM, mod))
put("1.1.1.11", "unused filesystems", attest("unused_filesystem_modules_reviewed"))

# 1.1.2 partitions
put("1.1.2.1.1", "/tmp is tmpfs", mnt_sep("/tmp", extra_tmp=True))
put("1.1.2.2.1", "/dev/shm is tmpfs", mnt_sep("/dev/shm"))
for base, part in [("1.1.2.3", "/home"), ("1.1.2.4", "/var"), ("1.1.2.5", "/var/tmp"), ("1.1.2.6", "/var/log"),
                   ("1.1.2.7", "/var/log/audit")]:
    put(base + ".1", "separate partition exists for " + part, mnt_sep(part))
OPTS = {"1.1.2.1": ("/tmp", ["nodev", "nosuid", "noexec"]), "1.1.2.2": ("/dev/shm", ["nodev", "nosuid", "noexec"]),
        "1.1.2.3": ("/home", ["nodev", "nosuid"]), "1.1.2.4": ("/var", ["nodev", "nosuid"]),
        "1.1.2.5": ("/var/tmp", ["nodev", "nosuid", "noexec"]), "1.1.2.6": ("/var/log", ["nodev", "nosuid", "noexec"]),
        "1.1.2.7": ("/var/log/audit", ["nodev", "nosuid", "noexec"])}
for base, (part, opts) in OPTS.items():
    for i, o in enumerate(opts, start=2):
        put(f"{base}.{i}", f"{o} option set on {part} partition", mnt_opt(part, o))

# 1.2 package management (apt). apt.config_dump is the text of `apt-config dump`.
APTD = ["apt", "config_dump"]


def apt_opt(key, allowed):
    return custom([f'is_string({ref(APTD)})'], [[f'apt_last({json.dumps(key)}) in {json.dumps(allowed)}']], dotted(APTD),
                  nest(APTD, f'APT::Architecture "amd64";\n{key} "{allowed[0]}";'),
                  nest(APTD, f'APT::Architecture "amd64";\n{key} "1";' if allowed[0] != "true" else f'{key} "false";'),
                  reads=[APTD])


put("1.2.1.1", "Signed-By", attest("apt_sources_signed_by_reviewed"))
put("1.2.1.2", "weak dependencies", custom(
    [f'is_string({ref(APTD)})'],
    [['apt_last("APT::Install-Recommends") in {"0", "false"}', 'apt_last("APT::Install-Suggests") in {"0", "false"}']],
    dotted(APTD), nest(APTD, 'APT::Install-Recommends "0";\nAPT::Install-Suggests "0";'),
    nest(APTD, 'APT::Install-Recommends "1";\nAPT::Install-Suggests "0";'), reads=[APTD]))
put("1.2.1.3", "gpg key files", access_list(["apt", "key_files"], "0644", ["root"], ["root"],
                                            sample="/usr/share/keyrings/ubuntu-archive-keyring.gpg"))
put("1.2.1.4", "/etc/apt/trusted.gpg.d", dir_access(["apt", "trusted_gpg_d"], "0755"))
put("1.2.1.5", "/etc/apt/auth.conf.d directory", dir_access(["apt", "auth_conf_d"], "0755"))
put("1.2.1.6", "files in the /etc/apt/auth.conf.d", access_list(["apt", "auth_conf_files"], "0640", ["root"], ["root"],
                                                                sample="/etc/apt/auth.conf.d/x.conf", nonempty=False))
put("1.2.1.7", "/usr/share/keyrings", dir_access(["apt", "usr_share_keyrings"], "0755"))
put("1.2.1.8", "/etc/apt/sources.list.d directory", dir_access(["apt", "sources_list_d"], "0755"))
put("1.2.1.9", "files in /etc/apt/sources.list.d", access_list(["apt", "sources_list_d_files"], "0644", ["root"], ["root"],
                                                               sample="/etc/apt/sources.list.d/ubuntu.sources", nonempty=False))
SL = ["apt", "sources_list_raw"]
_s = put("1.2.1.10", "https in /etc/apt/sources.list", text_has(SL, "(?i)^deb\\s+http://", False,
                                                                "deb https://archive.example.com/ubuntu resolute main",
                                                                "deb http://archive.example.com/ubuntu resolute main"))
_s.ev_alts.append(['input.apt.sources_list_exists == false'])
_s.passes.append(['input.apt.sources_list_exists == false'])
_s.reads.append(["apt", "sources_list_exists"])
_s.field = "apt.sources_list_raw/sources_list_exists"
put("1.2.1.11", "https in /etc/apt/sources.list.d", text_has(
    ["apt", "sources_list_d_raw"], "(?i)^(deb\\s+http://|URIs:\\s+http://)", False,
    "Types: deb\nURIs: https://archive.example.com/ubuntu\nSigned-By: /usr/share/keyrings/ubuntu-archive-keyring.gpg",
    "Types: deb\nURIs: http://archive.example.com/ubuntu"))
put("1.2.1.12", "AllowInsecureRepositories", apt_opt("Acquire::AllowInsecureRepositories", ["0", "false"]))
put("1.2.1.13", "AllowWeakRepositories", apt_opt("Acquire::AllowWeakRepositories", ["0", "false"]))
put("1.2.1.14", "AllowDowngradeToInsecureRepositories", apt_opt("Acquire::AllowDowngradeToInsecureRepositories", ["0", "false"]))
put("1.2.1.15", "Check-Date", apt_opt("Acquire::Check-Date", ["true", "1"]))
put("1.2.2.1", "updates, patches", attest("updates_installed"))

# 1.3 AppArmor (Ubuntu uses AppArmor, not SELinux)
GL = ["boot_security", "grub_cfg_linux_lines"]   # `linux` lines of /boot/grub/grub.cfg
LINUX = "linux /boot/vmlinuz root=/dev/mapper/root ro quiet splash"
put("1.3.1.1", "apparmor packages", pkgs_present(["apparmor", "apparmor-utils"]))
put("1.3.1.2", "AppArmor is enabled", custom(
    [f'is_array({ref(GL)})', 'count(grub_linux_lines) > 0', 'services_collected', 'input.apparmor.profiles.loaded != null'],
    [['every l in grub_linux_lines { not regex.match(`(^|\\s)apparmor=0(\\s|$)`, l) }', 'not apparmor_grub_d_disabled',
      'unit_enabled_active("apparmor.service")', 'num(input.apparmor.profiles.loaded) > 0']],
    "boot_security.grub_cfg_linux_lines/services[apparmor.service]/apparmor.profiles.loaded",
    {"boot_security": {"grub_cfg_linux_lines": [LINUX]}, "apparmor": {"profiles": {"loaded": 34}},
     "services": {"apparmor.service": {"state": "running", "status": "enabled"}}},
    {"boot_security": {"grub_cfg_linux_lines": [LINUX + " apparmor=0"]}, "apparmor": {"profiles": {"loaded": 34}},
     "services": {"apparmor.service": {"state": "running", "status": "enabled"}}},
    reads=[GL, ["apparmor", "profiles"], ["apparmor", "grub_d_cfg_raw"], ["services"]]))
_prof_ok = {"loaded": 34, "enforce": 34, "complain": 0, "prompt": 0, "kill": 0, "unconfined": 0}
put("1.3.1.3", "AppArmor Profiles are enforcing", custom(
    ['input.apparmor.profiles.loaded != null', 'input.apparmor.profiles.enforce != null',
     'input.apparmor.processes.unconfined != null'],
    [['num(input.apparmor.profiles.loaded) > 0', 'num(input.apparmor.profiles.enforce) == num(input.apparmor.profiles.loaded)',
      'every k in ["complain", "prompt", "kill", "unconfined"] { num(object.get(input.apparmor.profiles, k, -1)) == 0 }',
      'num(input.apparmor.processes.unconfined) == 0']],
    "apparmor.profiles/apparmor.processes",
    {"apparmor": {"profiles": _prof_ok, "processes": {"unconfined": 0}}},
    {"apparmor": {"profiles": dict(_prof_ok, enforce=32, complain=2), "processes": {"unconfined": 0}}},
    reads=[["apparmor", "profiles"], ["apparmor", "processes"]]))
put("1.3.1.4", "apparmor_restrict_unprivileged_unconfined",
    sysctl("kernel.apparmor_restrict_unprivileged_unconfined", ["1"]))

# 1.4 bootloader (/boot/grub/grub.cfg)
put("1.4.1", "bootloader password", custom(
    ['is_boolean(input.boot_security.grub_password_set)', 'is_boolean(input.boot_security.grub_password_hashed)'],
    [['input.boot_security.grub_password_set == true', 'input.boot_security.grub_password_hashed == true']],
    "boot_security.grub_password_set/grub_password_hashed",
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": True}},
    {"boot_security": {"grub_password_set": True, "grub_password_hashed": False}},
    reads=[["boot_security", "grub_password_set"], ["boot_security", "grub_password_hashed"]]))
put("1.4.2", "bootloader config", access(["boot_security", "grub_cfg_mode"], ["boot_security", "grub_cfg_owner"],
                                         ["boot_security", "grub_cfg_group"], "0600", ["root"], ["root"]))

# 1.5 process hardening
for rid, p, vals in [("1.5.1", "fs.protected_hardlinks", ["1"]), ("1.5.2", "fs.protected_symlinks", ["1"]),
                     ("1.5.3", "kernel.yama.ptrace_scope", ["1", "2", "3"]), ("1.5.4", "fs.suid_dumpable", ["0"]),
                     ("1.5.5", "kernel.dmesg_restrict", ["1"]), ("1.5.8", "kernel.kptr_restrict", ["1", "2"]),
                     ("1.5.9", "kernel.randomize_va_space", ["2"])]:
    put(rid, p, sysctl(p, vals))
put("1.5.6", "prelink", pkg_absent(["prelink"]))
APD = ["apport", "default_raw"]   # /etc/default/apport
put("1.5.7", "Automatic Error Reporting", custom(
    None,
    [['input.packages.apport == false'],
     ['input.packages.apport == true', f'not regex_in_text({ref(APD)}, `(?i)^enabled\\s*=\\s*[^0]`)',
      'not lower(input.services["apport.service"].state) in {"running", "active"}']],
    "packages.apport/apport.default_raw/services[apport.service]",
    {"packages": {"apport": True}, "apport": {"default_raw": "enabled=0"},
     "services": {"apport.service": {"state": "stopped", "status": "disabled"}}},
    {"packages": {"apport": True}, "apport": {"default_raw": "enabled=1"},
     "services": {"apport.service": {"state": "running", "status": "enabled"}}},
    {"packages": {"apport": True}}, reads=[["packages", "apport"], APD, ["services"]],
    ev_alts=[['input.packages.apport == false'],
             ['input.packages.apport == true', f'is_string({ref(APD)})', 'services_collected', 'unit_reported("apport.service")']]))
LIM = ["limits", "limits_raw"]   # /etc/security/limits.conf + limits.d/*.conf in parse order
put("1.5.10", "core file size", custom(
    [f'is_string({ref(LIM)})'], [['count(core_hard_limits) > 0', 'every v in core_hard_limits { v == "0" }']], dotted(LIM),
    nest(LIM, "* hard core 0"), nest(LIM, "* hard core 0\n* hard core unlimited"), reads=[LIM]))
put("1.5.11", "ProcessSizeMax", custom(['input.boot_security.systemd_coredump_process_size_max != null'],
    [['sprintf("%v", [input.boot_security.systemd_coredump_process_size_max]) == "0"']],
    "boot_security.systemd_coredump_process_size_max",
    {"boot_security": {"systemd_coredump_process_size_max": "0"}},
    {"boot_security": {"systemd_coredump_process_size_max": "2G"}},
    reads=[["boot_security", "systemd_coredump_process_size_max"]]))
put("1.5.12", "coredump Storage", custom(['is_string(input.boot_security.systemd_coredump_storage)'],
    [['lower(input.boot_security.systemd_coredump_storage) == "none"']], "boot_security.systemd_coredump_storage",
    {"boot_security": {"systemd_coredump_storage": "none"}}, {"boot_security": {"systemd_coredump_storage": "external"}},
    reads=[["boot_security", "systemd_coredump_storage"]]))

# 1.6 banners — no OS/version escapes (\m \r \s \v) or the os-release ID ("ubuntu")
SYSINFO = "(?i)(\\\\[mrsv]|\\bubuntu\\b)"
for rid, f in [("1.6.1", "motd"), ("1.6.2", "issue"), ("1.6.3", "issue_net")]:
    ex, ct = ["initial_setup", f"{f}_exists"], ["initial_setup", f"{f}_content"]
    alts = [[f'{ref(ex)} == false'], [f'{ref(ex)} == true', f'is_string({ref(ct)})']]
    ok = [f'{ref(ex)} == true', f'not regex.match({json.dumps(SYSINFO)}, {ref(ct)})']
    passes = [ok] + ([[f'{ref(ex)} == false']] if f == "motd" else [])
    put(rid, "/etc/" + f.replace("_", ".") + " is configured", custom(
        None, passes, f"initial_setup.{f}_exists/{f}_content",
        merge(nest(ex, True), nest(ct, "Authorized uses only. All activity may be monitored.")),
        merge(nest(ex, True), nest(ct, "Ubuntu 26.04 LTS \\n \\l")), {}, reads=[ex, ct], ev_alts=alts))
PMF = ["initial_setup", "pam_motd_files"]   # files named by motd= on pam_motd.so (sshd, login, su, gdm-password)
put("1.6.4", "pam_motd", custom(
    [f'is_array({ref(PMF)})'],
    [[f'every f in {ref(PMF)} {{ is_string(f.content); not regex.match({json.dumps(SYSINFO)}, f.content) }}']], dotted(PMF),
    nest(PMF, [{"path": "/run/motd.dynamic", "content": "Authorized uses only."}]),
    nest(PMF, [{"path": "/run/motd.dynamic", "content": "Welcome to Ubuntu 26.04 LTS"}]), reads=[PMF]))
put("1.6.5", "sshd warning Banner", custom(
    None,
    [['v := sshd_value("banner")', 'lower(v) != "none"', 'is_string(input.ssh.banner_content)',
      f'not regex.match({json.dumps(SYSINFO)}, input.ssh.banner_content)']],
    "ssh.sshd_config_raw/ssh.banner_content",
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_content": "Authorized uses only."}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_content": "Ubuntu 26.04 \\r"}},
    {}, reads=[SSH_RAW, ["ssh", "banner_content"]],
    ev_alts=[['ssh_collected', 'lower(sshd_value("banner")) == "none"'],
             ['ssh_collected', 'sshd_first("banner")', 'is_string(input.ssh.banner_content)']]))
put("1.6.6", "access to /etc/motd", access(["initial_setup", "motd_mode"], ["initial_setup", "motd_owner"],
    ["initial_setup", "motd_group"], "0644", ["root"], ["root"], exists=["initial_setup", "motd_exists"], optional=True))
put("1.6.7", "access to /etc/issue ", access(["initial_setup", "issue_mode"], ["initial_setup", "issue_owner"],
    ["initial_setup", "issue_group"], "0644", ["root"], ["root"], exists=["initial_setup", "issue_exists"]))
put("1.6.8", "access to /etc/issue.net", access(["initial_setup", "issue_net_mode"], ["initial_setup", "issue_net_owner"],
    ["initial_setup", "issue_net_group"], "0644", ["root"], ["root"], exists=["initial_setup", "issue_net_exists"]))
put("1.6.9", "pam_motd file", access_list(["initial_setup", "pam_motd_file_permissions"], "0644", ["root"], ["root"],
                                          sample="/etc/motd", nonempty=False))
BP = ["ssh", "banner_permissions"]
put("1.6.10", "access to sshd warning banner", custom(
    None,
    [['lower(sshd_value("banner")) == "none"'],
     ['lower(sshd_value("banner")) != "none"', 'mode_within(input.ssh.banner_permissions.mode, "0644")',
      'input.ssh.banner_permissions.pw_name == "root"', 'input.ssh.banner_permissions.gr_name == "root"']],
    "ssh.banner_permissions",
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_permissions": {"mode": "0644", "pw_name": "root", "gr_name": "root"}}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net", "banner_permissions": {"mode": "0666", "pw_name": "root", "gr_name": "root"}}},
    {"ssh": {"sshd_config_raw": "banner /etc/issue.net"}}, reads=[SSH_RAW, BP],
    ev_alts=[['ssh_collected', 'lower(sshd_value("banner")) == "none"'],
             ['ssh_collected', 'sshd_first("banner")', 'input.ssh.banner_permissions.mode != null',
              'input.ssh.banner_permissions.pw_name != null', 'input.ssh.banner_permissions.gr_name != null']]))
put("1.6.11", "update-notifier-motd.service", not_in_use({"update-notifier-common": ["update-notifier-motd.service"]}))
put("1.6.12", "update-notifier-motd.timer", not_in_use({"update-notifier-common": ["update-notifier-motd.timer"]}))


# 1.7 GDM — not applicable (passes) when GDM is not installed, a collected boolean
def gdm(fields, conds, good, bad):
    inst = ["initial_setup", "gdm_installed"]
    fp = [["initial_setup", f] for f in fields]
    have = [f'{ref(p)} != null' for p in fp]
    return custom(None, [['input.initial_setup.gdm_installed == false'],
                         ['input.initial_setup.gdm_installed == true'] + conds],
                  "initial_setup." + "/".join(["gdm_installed"] + fields),
                  {"initial_setup": dict(good, gdm_installed=True)},
                  {"initial_setup": dict(bad, gdm_installed=True)},
                  {"initial_setup": {"gdm_installed": True}}, reads=[inst] + fp,
                  ev_alts=[['input.initial_setup.gdm_installed == false'],
                           ['input.initial_setup.gdm_installed == true'] + have])


I = "input.initial_setup."
put("1.7.1", "GDM login banner", gdm(["gdm_banner_configured", "gdm_banner_override_allowed"],
    [I + "gdm_banner_configured == true", I + "gdm_banner_override_allowed == false"],
    {"gdm_banner_configured": True, "gdm_banner_override_allowed": False},
    {"gdm_banner_configured": True, "gdm_banner_override_allowed": True}))
put("1.7.2", "disable-user-list", gdm(["gdm_disable_user_list", "gdm_disable_user_list_override_allowed"],
    [I + "gdm_disable_user_list == true", I + "gdm_disable_user_list_override_allowed == false"],
    {"gdm_disable_user_list": True, "gdm_disable_user_list_override_allowed": False},
    {"gdm_disable_user_list": False, "gdm_disable_user_list_override_allowed": False}))
put("1.7.3", "screen lock", gdm(["gdm_idle_delay", "gdm_lock_delay", "gdm_lock_override_allowed"],
    [f"num({I}gdm_idle_delay) > 0", f"num({I}gdm_idle_delay) <= 900",
     f"num({I}gdm_lock_delay) >= 0", f"num({I}gdm_lock_delay) <= 5", I + "gdm_lock_override_allowed == false"],
    {"gdm_idle_delay": 900, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False},
    {"gdm_idle_delay": 0, "gdm_lock_delay": 5, "gdm_lock_override_allowed": False}))
put("1.7.4", "automount", gdm(["gdm_automount_disabled", "gdm_automount_open_disabled", "gdm_automount_override_allowed"],
    [I + "gdm_automount_disabled == true", I + "gdm_automount_open_disabled == true", I + "gdm_automount_override_allowed == false"],
    {"gdm_automount_disabled": True, "gdm_automount_open_disabled": True, "gdm_automount_override_allowed": False},
    {"gdm_automount_disabled": True, "gdm_automount_open_disabled": False, "gdm_automount_override_allowed": False}))
put("1.7.5", "autorun-never", gdm(["gdm_autorun_never_enabled", "gdm_autorun_override_allowed"],
    [I + "gdm_autorun_never_enabled == true", I + "gdm_autorun_override_allowed == false"],
    {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": False},
    {"gdm_autorun_never_enabled": True, "gdm_autorun_override_allowed": True}))
put("1.7.6", "XDMCP", gdm(["gdm_xdmcp_enabled"], [I + "gdm_xdmcp_enabled == false"],
    {"gdm_xdmcp_enabled": False}, {"gdm_xdmcp_enabled": True}))
# the 26.04 audit checks WaylandEnable=false in the gdm3/gdm custom.conf/daemon.conf
put("1.7.7", "Xwayland", gdm(["gdm_wayland_enabled"], [I + "gdm_wayland_enabled == false"],
    {"gdm_wayland_enabled": False}, {"gdm_wayland_enabled": True}))

# 2.1 server services: package not installed (dpkg), or every unit neither enabled nor active
for rid, kw, units_by_pkg in [
    ("2.1.1", "autofs", {"autofs": ["autofs.service"]}),
    ("2.1.2", "avahi", {"avahi-daemon": ["avahi-daemon.socket", "avahi-daemon.service"]}),
    # the audit greps dpkg for any package named kea*: packages.kea = "some kea package is installed"
    ("2.1.3", "dhcp server", {"kea": ["kea-dhcp-ddns-server.service", "kea-dhcp4-server.service", "kea-dhcp6-server.service"]}),
    ("2.1.4", "dns server", {"bind9": ["named.service"]}),
    ("2.1.5", "dnsmasq", {"dnsmasq": ["dnsmasq.service"]}),
    ("2.1.6", "ftp server", {"vsftpd": ["vsftpd.service"]}),
    ("2.1.7", "ldap server", {"slapd": ["slapd.service"]}),
    ("2.1.8", "message access", {"dovecot-imapd": ["dovecot.socket", "dovecot.service"],
                                 "dovecot-pop3d": ["dovecot.socket", "dovecot.service"]}),
    ("2.1.9", "network file system", {"nfs-kernel-server": ["nfs-server.service"]}),
    ("2.1.10", "nis server", {"ypserv": ["ypserv.service"]}),
    ("2.1.11", "print server", {"cups": ["cups.socket", "cups.service"]}),
    ("2.1.12", "rpcbind", {"rpcbind": ["rpcbind.socket", "rpcbind.service"]}),
    ("2.1.13", "rsync", {"rsync": ["rsync.service"]}),
    ("2.1.14", "samba", {"samba": ["smbd.service"]}),
    ("2.1.15", "snmp", {"snmpd": ["snmpd.service"]}),
    ("2.1.16", "telnet server", {"telnetd": ["inetutils-inetd.service"], "telnetd-ssl": ["inetutils-inetd.service"]}),
    ("2.1.17", "tftp server", {"tftpd-hpa": ["tftpd-hpa.service"]}),
    ("2.1.18", "web proxy", {"squid": ["squid.service"]}),
    ("2.1.19", "web server", {"apache2": ["apache2.socket", "apache2.service"], "nginx": ["nginx.service"]}),
    ("2.1.20", "xinetd", {"xinetd": ["xinetd.service"]}),
]:
    put(rid, kw, not_in_use(units_by_pkg))
put("2.1.21", "X window server", pkg_absent(["xserver-common"]))
put("2.1.22", "mail transfer agents", custom(
    ['is_array(input.mta.non_loopback_listeners)'],
    [['count(input.mta.non_loopback_listeners) == 0', 'not mta_interfaces_open']],
    "mta.non_loopback_listeners",
    {"mta": {"non_loopback_listeners": [], "inet_interfaces": "loopback-only"}},
    {"mta": {"non_loopback_listeners": ["0.0.0.0:25"], "inet_interfaces": "all"}},
    reads=[["mta", "non_loopback_listeners"], ["mta", "inet_interfaces"]]))
put("2.1.23", "approved services", attest("listening_services_approved"))
# 2.2 client packages
put("2.2.1", "nis client", pkg_absent(["nis"]))
put("2.2.2", "rsh client", pkg_absent(["rsh-client"]))
put("2.2.3", "talk client", pkg_absent(["talk"]))
put("2.2.4", "telnet client", pkg_absent(["telnet", "inetutils-telnet", "telnet-ssl"]))
put("2.2.5", "ldap client", pkg_absent(["ldap-utils"]))
put("2.2.6", "ftp client", pkg_absent(["ftp", "tnftp", "atftp", "inetutils-ftp", "ftp-ssl"]))

# 2.3 time synchronization: exactly one of systemd-timesyncd / chrony in use; the other daemon's
# recommendations are not applicable (pass) only on collected unit state showing the other one alone in use.
TS, CH = "systemd-timesyncd.service", "chrony.service"
SSHU = {"ssh.service": {"state": "running", "status": "enabled"}}
ON = {"state": "running", "status": "enabled"}
OFF = {"state": "stopped", "status": "disabled"}
put("2.3.1.1", "single time synchronization", custom(
    ['services_collected', 'time_unit_reported'], [['timesyncd_only'], ['chrony_only']], f"services[{TS},{CH}]",
    {"services": dict(SSHU, **{TS: ON})}, {"services": dict(SSHU, **{TS: ON, CH: ON})}, {"services": SSHU},
    reads=[["services"]]))
TSC = ["timesyncd", "config_raw"]   # systemd-analyze cat-config systemd/timesyncd.conf
put("2.3.2.1", "authorized timeserver", custom(
    None, [['chrony_only'], [f'regex_in_text({ref(TSC)}, `^(NTP|FallbackNTP)\\s*=\\s*\\S+`)']], f"{dotted(TSC)}",
    nest(TSC, "[Time]\nNTP=time.example.com"), nest(TSC, "[Time]\n#NTP="), {}, reads=[TSC, ["services"]],
    ev_alts=[['chrony_only'], [f'is_string({ref(TSC)})']]))
put("2.3.2.2", "systemd-timesyncd is enabled", custom(
    ['services_collected'], [['chrony_only'], [f'unit_enabled_active("{TS}")']], f"services[{TS}]",
    {"services": dict(SSHU, **{TS: ON})}, {"services": dict(SSHU, **{TS: OFF})}, {}, reads=[["services"]]))
put("2.3.3.1", "chrony is configured", custom(
    None, [['timesyncd_only'], ['num(input.chronyd.servers_configured) > 0'], ['num(input.chronyd.pools_configured) > 0']],
    "chronyd.servers_configured/pools_configured",
    {"chronyd": {"servers_configured": 0, "pools_configured": 2}},
    {"chronyd": {"servers_configured": 0, "pools_configured": 0}}, {},
    reads=[["chronyd", "servers_configured"], ["chronyd", "pools_configured"], ["services"]],
    ev_alts=[['timesyncd_only'], ['input.chronyd.servers_configured != null', 'input.chronyd.pools_configured != null']]))
put("2.3.3.2", "_chrony", custom(
    None, [['timesyncd_only'], ['every u in input.chronyd.process_users { u == "_chrony" }']], "chronyd.process_users",
    {"chronyd": {"process_users": ["_chrony"]}}, {"chronyd": {"process_users": ["root"]}}, {},
    reads=[["chronyd", "process_users"], ["services"]],
    ev_alts=[['timesyncd_only'], ['is_array(input.chronyd.process_users)', 'count(input.chronyd.process_users) > 0']]))
put("2.3.3.3", "chrony is enabled", custom(
    ['services_collected'], [['timesyncd_only'], [f'unit_enabled_active("{CH}")']], f"services[{CH}]",
    {"services": dict(SSHU, **{CH: ON})}, {"services": dict(SSHU, **{CH: OFF})}, {}, reads=[["services"]]))

# 2.4 job schedulers — "IF cron (at) is installed": packages.cron / packages.at false passes
put("2.4.1.1", "cron daemon", gate_pkg(custom(
    ['services_collected'], [['unit_enabled_active("cron.service")']], "services[cron.service]",
    {"services": dict(SSHU, **{"cron.service": ON})}, {"services": dict(SSHU, **{"cron.service": OFF})}, {},
    reads=[["services"]]), "cron"))
put("2.4.1.2", "/etc/crontab", gate_pkg(access(["cron", "crontab_mode"], ["cron", "crontab_owner"], ["cron", "crontab_group"],
                                               "0600", ["root"], ["root"]), "cron"))
for rid, d in [("2.4.1.3", "cron_hourly"), ("2.4.1.4", "cron_daily"), ("2.4.1.5", "cron_weekly"),
               ("2.4.1.6", "cron_monthly"), ("2.4.1.7", "cron_yearly"), ("2.4.1.8", "cron_d")]:
    put(rid, "/etc/" + d.replace("_", "."), gate_pkg(access(["cron", f"{d}_mode"], ["cron", f"{d}_owner"],
                                                            ["cron", f"{d}_group"], "0700", ["root"], ["root"]), "cron"))


def allow_deny(svc, groups):
    """<svc>.allow must exist (mode 0640, root, groups); <svc>.deny, if present, the same."""
    c = "input.cron."
    a = [f"{svc}_allow_{k}" for k in ("exists", "mode", "owner", "group")]
    d = [f"{svc}_deny_{k}" for k in ("exists", "mode", "owner", "group")]
    g = json.dumps(groups)
    allow_ok = [f'{c}{a[0]} == true', f'mode_within({c}{a[1]}, "0640")', f'{c}{a[2]} == "root"', f'{c}{a[3]} in {g}']
    deny_ok_absent = [f'{c}{d[0]} == false']
    deny_ok_present = [f'{c}{d[0]} == true', f'mode_within({c}{d[1]}, "0640")', f'{c}{d[2]} == "root"', f'{c}{d[3]} in {g}']
    good = {a[0]: True, a[1]: "0640", a[2]: "root", a[3]: groups[0], d[0]: False}
    bad = {a[0]: False, d[0]: False}
    ev_allow = [[f'{c}{a[0]} == false'], [f'{c}{a[0]} == true', f'{c}{a[1]} != null', f'{c}{a[2]} != null', f'{c}{a[3]} != null']]
    ev_deny = [[f'{c}{d[0]} == false'], [f'{c}{d[0]} == true', f'{c}{d[1]} != null', f'{c}{d[2]} != null', f'{c}{d[3]} != null']]
    return custom(None, [allow_ok + deny_ok_absent, allow_ok + deny_ok_present],
                  f"cron.{svc}_allow_*/{svc}_deny_*", {"cron": good}, {"cron": bad}, {},
                  reads=[["cron", k] for k in a + d],
                  ev_alts=[x + y for x in ev_allow for y in ev_deny])


put("2.4.1.9", "access to crontab", gate_pkg(allow_deny("cron", ["root", "crontab"]), "cron"))
put("2.4.2.1", "access to at", gate_pkg(allow_deny("at", ["root", "daemon"]), "at"))

# 3 network
put("3.1.1", "IPv6 status", attest("ipv6_status_identified"))
put("3.1.2", "wireless interfaces", custom(
    None,
    [['input.wireless_interfaces.has_wireless == false'],
     ['input.wireless_interfaces.has_wireless == true', 'is_array(input.wireless_interfaces.modules)',
      'count(input.wireless_interfaces.modules) > 0',
      'every e in input.wireless_interfaces.modules { kmod_unavailable(e) }']],
    "wireless_interfaces.has_wireless/modules",
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": True, "loaded": False}]}},
    {"wireless_interfaces": {"has_wireless": True, "modules": [{"module": "iwlwifi", "disabled": False, "loaded": True}]}},
    {"wireless_interfaces": {"has_wireless": True}},
    reads=[["wireless_interfaces", "has_wireless"], ["wireless_interfaces", "modules"]],
    ev_alts=[['input.wireless_interfaces.has_wireless == false'],
             ['input.wireless_interfaces.has_wireless == true', 'is_array(input.wireless_interfaces.modules)']]))
put("3.1.3", "bluetooth", not_in_use({"bluez": ["bluetooth.service"]}))
for rid, mod in [("3.2.1", "atm"), ("3.2.2", "can"), ("3.2.3", "dccp"), ("3.2.4", "rds"), ("3.2.5", "sctp"), ("3.2.6", "tipc")]:
    put(rid, mod + " kernel module", kmod(NETM, mod))
put("3.2.7", "unneeded network protocol", attest("unneeded_network_modules_reviewed"))
for i, (p, v) in enumerate([
        ("net.ipv4.ip_forward", "0"), ("net.ipv4.conf.all.forwarding", "0"), ("net.ipv4.conf.default.forwarding", "0"),
        ("net.ipv4.conf.all.send_redirects", "0"), ("net.ipv4.conf.default.send_redirects", "0"),
        ("net.ipv4.icmp_ignore_bogus_error_responses", "1"), ("net.ipv4.icmp_echo_ignore_broadcasts", "1"),
        ("net.ipv4.conf.all.accept_redirects", "0"), ("net.ipv4.conf.default.accept_redirects", "0"),
        ("net.ipv4.conf.all.secure_redirects", "0"), ("net.ipv4.conf.default.secure_redirects", "0"),
        ("net.ipv4.conf.all.rp_filter", "1"), ("net.ipv4.conf.default.rp_filter", "1"),
        ("net.ipv4.conf.all.accept_source_route", "0"), ("net.ipv4.conf.default.accept_source_route", "0"),
        ("net.ipv4.conf.all.log_martians", "1"), ("net.ipv4.conf.default.log_martians", "1"),
        ("net.ipv4.tcp_syncookies", "1"), ("net.ipv4.conf.all.route_localnet", "0")], start=1):
    put(f"3.3.1.{i}", p + " is", sysctl(p, [v]))
for i, p in enumerate(["net.ipv6.conf.all.forwarding", "net.ipv6.conf.default.forwarding",
                       "net.ipv6.conf.all.accept_redirects", "net.ipv6.conf.default.accept_redirects",
                       "net.ipv6.conf.all.accept_source_route", "net.ipv6.conf.default.accept_source_route",
                       "net.ipv6.conf.all.accept_ra", "net.ipv6.conf.default.accept_ra"], start=1):
    put(f"3.3.2.{i}", p + " is", sysctl(p, ["0"], ipv6=True))

# 4 host firewall — the 26.04 benchmark covers ufw only (no nftables / iptables sections).
# ufw.status_verbose is the text of `ufw status verbose`; ufw.default_raw is /etc/default/ufw.
UV = ["ufw", "status_verbose"]
UVOK = "Status: active\nLogging: on (low)\nDefault: deny (incoming), deny (outgoing), deny (routed)\nNew profiles: skip"
put("4.1.1", "ufw is installed", pkg_present("ufw"))
put("4.1.2", "ufw service", custom(
    ['services_collected', f'is_string({ref(UV)})'],
    [['unit_enabled_active("ufw.service")', f'regex_in_text({ref(UV)}, `(?i)^Status:\\s*active\\b`)']],
    "services[ufw.service]/ufw.status_verbose",
    {"services": dict(SSHU, **{"ufw.service": ON}), "ufw": {"status_verbose": UVOK}},
    {"services": dict(SSHU, **{"ufw.service": ON}), "ufw": {"status_verbose": "Status: inactive"}},
    reads=[["services"], UV]))
for rid, kw, d, ok, bad in [("4.1.3", "incoming default", "incoming", '{"deny", "reject"}', "allow (incoming)"),
                            ("4.1.4", "outgoing default", "outgoing", '{"deny", "reject"}', "allow (outgoing)")]:
    put(rid, kw, custom([f'is_string({ref(UV)})'], [[f'ufw_default("{d}") in {ok}']], dotted(UV), nest(UV, UVOK),
                        nest(UV, UVOK.replace(f"deny ({d})", bad)), reads=[UV]))
put("4.1.5", "routed default", custom(
    [f'is_string({ref(UV)})'],
    [['ufw_default("routed") == "deny"'], ['ufw_default("routed") == "disabled"', 'ufw_forward_policy_drop']],
    "ufw.status_verbose/ufw.default_raw", nest(UV, UVOK), nest(UV, UVOK.replace("deny (routed)", "allow (routed)")),
    reads=[UV, ["ufw", "default_raw"]]))

# 5.1 SSH server (`sshd -T` effective configuration; Ubuntu unit ssh.service)
put("5.1.1", "/etc/ssh/sshd_config", access_list(["ssh", "config_files"], "0600", ["root"], ["root"], mode="mode",
                                                 owner="pw_name", group="gr_name", sample="/etc/ssh/sshd_config"))
put("5.1.2", "sshd_config.d/ directory", dir_access(["ssh", "config_dir_permissions"], "0700", optional=True))
put("5.1.3", "private host key", custom(
    ['is_array(input.ssh.private_key_permissions)', 'count(input.ssh.private_key_permissions) > 0'],
    [['every k in input.ssh.private_key_permissions { mode_within(k.stat.mode, "0600"); k.stat.pw_name == "root"; k.stat.gr_name == "root" }']],
    "ssh.private_key_permissions",
    {"ssh": {"private_key_permissions": [{"stat": {"path": "/etc/ssh/ssh_host_ed25519_key", "mode": "0600", "pw_name": "root", "gr_name": "root"}}]}},
    {"ssh": {"private_key_permissions": [{"stat": {"path": "/etc/ssh/ssh_host_ed25519_key", "mode": "0640", "pw_name": "root", "gr_name": "root"}}]}},
    reads=[["ssh", "private_key_permissions"]]))
put("5.1.4", "public host key", access_list(["ssh", "public_key_permissions"], "0644", ["root"], ["root"], mode="mode",
                                            owner="pw_name", group="gr_name", sample="/etc/ssh/ssh_host_ed25519_key.pub"))
put("5.1.5", "sshd access", custom(['ssh_collected'],
    [['some k in ["allowusers", "allowgroups", "denyusers", "denygroups"]', 'sshd_first(k) != ""']],
    "ssh.sshd_config_raw", ssh_in(["port 22", "allowgroups sshusers"]), ssh_in(["port 22"]), {}, reads=[SSH_RAW]))
put("5.1.6", "sshd Banner", sshd("banner", 'startswith(v, "/")', "/etc/issue.net", "none"))
WEAK_CIPH = ["3des-cbc", "aes128-cbc", "aes192-cbc", "aes256-cbc", "arcfour", "arcfour128", "arcfour256",
             "blowfish-cbc", "cast128-cbc", "rijndael-cbc@lysator.liu.se"]
WEAK_MAC = ["hmac-md5", "hmac-md5-96", "hmac-ripemd160", "hmac-sha1-96", "umac-64@openssh.com",
            "hmac-md5-etm@openssh.com", "hmac-md5-96-etm@openssh.com", "hmac-ripemd160-etm@openssh.com",
            "hmac-sha1-96-etm@openssh.com", "umac-64-etm@openssh.com", "umac-128-etm@openssh.com"]
WEAK_KEX = ["diffie-hellman-group1-sha1", "diffie-hellman-group14-sha1", "diffie-hellman-group-exchange-sha1"]


def sshd_list_none(key, weak, good, bad):
    s = sshd(key, f'count([x | some x in sshd_list({json.dumps(key)}); x in {json.dumps(weak)}]) == 0', good, bad,
             explicit=True)
    s.passes = [[f'count([x | some x in sshd_list({json.dumps(key)}); x in {json.dumps(weak)}]) == 0']]
    return s


put("5.1.7", "Ciphers", sshd_list_none("ciphers", WEAK_CIPH, "aes256-gcm@openssh.com,aes256-ctr", "aes256-ctr,aes128-cbc"))
put("5.1.8", "ClientAliveInterval", custom(
    ['ssh_collected', 'sshd_first("clientaliveinterval")', 'sshd_first("clientalivecountmax")'],
    [['num(sshd_value("clientaliveinterval")) > 0', 'num(sshd_value("clientalivecountmax")) > 0']],
    "ssh.sshd_config_raw (effective clientaliveinterval/clientalivecountmax)",
    ssh_in(["clientaliveinterval 15", "clientalivecountmax 3"]),
    ssh_in(["clientaliveinterval 0", "clientalivecountmax 3"]), {}, reads=[SSH_RAW]))
put("5.1.9", "DisableForwarding", sshd("disableforwarding", 'v == "yes"', "yes", "no"))
put("5.1.10", "GSSAPIAuthentication", sshd("gssapiauthentication", 'v == "no"', "no", "yes"))
put("5.1.11", "HostbasedAuthentication", sshd("hostbasedauthentication", 'v == "no"', "no", "yes"))
put("5.1.12", "IgnoreRhosts", sshd("ignorerhosts", 'v == "yes"', "yes", "no"))
put("5.1.13", "KexAlgorithms", sshd_list_none("kexalgorithms", WEAK_KEX, "curve25519-sha256,mlkem768x25519-sha256",
                                              "curve25519-sha256,diffie-hellman-group14-sha1"))
put("5.1.14", "post-quantum", attest("ssh_post_quantum_kex_reviewed"))
_s = put("5.1.15", "LoginGraceTime", sshd("logingracetime", 'sshd_seconds(v) >= 1', "60", "120"))
_s.passes = [['v := lower(sshd_value("logingracetime"))', 'sshd_seconds(v) >= 1', 'sshd_seconds(v) <= 60']]
put("5.1.16", "LogLevel", sshd("loglevel", 'v in {"info", "verbose"}', "VERBOSE", "QUIET"))
put("5.1.17", "MACs", sshd_list_none("macs", WEAK_MAC, "hmac-sha2-512-etm@openssh.com,hmac-sha2-256",
                                     "hmac-sha2-256,umac-64-etm@openssh.com"))
put("5.1.18", "MaxAuthTries", sshd("maxauthtries", 'num(v) <= 4', "4", "6"))
put("5.1.19", "MaxSessions", sshd("maxsessions", 'num(v) <= 10', "10", "20"))
put("5.1.20", "MaxStartups", sshd("maxstartups", 'maxstartups_ok(v)', "10:30:60", "10:30:100"))
put("5.1.21", "PermitEmptyPasswords", sshd("permitemptypasswords", 'v == "no"', "no", "yes"))
put("5.1.22", "PermitRootLogin", sshd("permitrootlogin", 'v == "no"', "no", "prohibit-password"))
put("5.1.23", "PermitUserEnvironment", sshd("permituserenvironment", 'v == "no"', "no", "yes"))
put("5.1.24", "UsePAM", sshd("usepam", 'v == "yes"', "yes", "no"))

# 5.2 privilege escalation (sudo or sudo-rs; /etc/sudoers and /etc/sudoers.d)
SU = ["sudo", "sudoers_content"]
put("5.2.1", "sudo is installed", custom(
    ['is_boolean(input.packages.sudo)', 'is_boolean(input.packages["sudo-rs"])'],
    [['input.packages.sudo == true'], ['input.packages["sudo-rs"] == true'],
     ['input.packages["libsss-sudo"] == true', 'input.packages.sssd == true']],
    "packages[sudo,sudo-rs]", {"packages": {"sudo": False, "sudo-rs": True}}, {"packages": {"sudo": False, "sudo-rs": False}},
    reads=[["packages", "sudo"], ["packages", "sudo-rs"], ["packages", "libsss-sudo"], ["packages", "sssd"]]))
put("5.2.2", "use pty", custom([f'is_string({ref(SU)})'],
    [['sudo_defaults_has(`(^|[\\s,])use_pty\\b`)', 'not sudo_defaults_has(`!\\s*use_pty\\b`)']],
    dotted(SU), nest(SU, "Defaults use_pty\nroot ALL=(ALL:ALL) ALL"),
    nest(SU, "Defaults !use_pty\nroot ALL=(ALL:ALL) ALL"), reads=[SU]))
# Method A: sudo events reach journald/rsyslog (sudo.events_logged); Method B: classic sudo with Defaults logfile=
put("5.2.3", "sudo log file", custom(
    None, [['input.sudo.events_logged == true'], ['lower(input.sudo.implementation) == "sudo"', 'sudo_logfile']],
    "sudo.events_logged/sudo.implementation/sudo.sudoers_content",
    {"sudo": {"events_logged": True}},
    {"sudo": {"events_logged": False, "implementation": "sudo-rs", "sudoers_content": "Defaults use_pty"}}, {},
    reads=[["sudo", "events_logged"], ["sudo", "implementation"], SU],
    ev_alts=[['is_boolean(input.sudo.events_logged)'], ['is_string(input.sudo.implementation)', f'is_string({ref(SU)})']]))
put("5.2.4", "password for escalation", text_has(SU, "\\bNOPASSWD\\s*:", False, "%sudo ALL=(ALL:ALL) ALL",
                                                 "%sudo ALL=(ALL:ALL) NOPASSWD: ALL"))
put("5.2.5", "re-authentication", text_has(SU, "!authenticate\\b", False, "Defaults use_pty", "Defaults !authenticate"))
put("5.2.6", "timestamp_timeout", custom(
    [f'is_string({ref(SU)})'], [['count(sudo_timeouts) > 0', 'every t in sudo_timeouts { t >= 0; t <= 15 }']],
    dotted(SU), nest(SU, "Defaults env_reset, timestamp_timeout=15"), nest(SU, "Defaults timestamp_timeout=-1"),
    nest(SU, "Defaults use_pty"), reads=[SU]))
L["5.2.6"].tm = {}
put("5.2.7", "su command", custom(
    ['is_string(input.pam.su)', 'is_string(input.user_group.group_content)'],
    [['some g in su_wheel_groups', 'su_group_empty(g)']],
    "pam.su/user_group.group_content",
    {"pam": {"su": "auth required pam_wheel.so use_uid group=sugroup"},
     "user_group": {"group_content": "root:x:0:\nsugroup:x:1001:"}},
    {"pam": {"su": "auth required pam_wheel.so use_uid group=sugroup"},
     "user_group": {"group_content": "root:x:0:\nsugroup:x:1001:alice"}},
    reads=[["pam", "su"], ["user_group", "group_content"]]))

# 5.3 PAM (pam-auth-update managed /etc/pam.d/common-*: pam.common_auth, common_account, common_password,
# common_session, common_session_noninteractive; /etc/security/{faillock,pwquality,pwhistory}.conf)
for rid, pkg in [("5.3.1.1", "libpam-runtime"), ("5.3.1.2", "libpam-modules"), ("5.3.1.3", "libpam-pwquality"),
                 ("5.3.1.4", "cracklib-runtime")]:
    pk = ref(["packages", pkg])
    put(rid, "latest version of " + ("pam" if pkg == "libpam-runtime" else pkg), custom(
        [f'is_boolean({pk})', 'is_array(input.pending_updates)'],
        [[f'{pk} == true', f'not {json.dumps(pkg)} in input.pending_updates']], f"packages.{pkg}/pending_updates",
        {"packages": {pkg: True}, "pending_updates": ["vim"]}, {"packages": {pkg: True}, "pending_updates": [pkg]},
        reads=[["packages", pkg], ["pending_updates"]]))
PCA, PCC, PCP = ["pam", "common_auth"], ["pam", "common_account"], ["pam", "common_password"]
COMMON_OK = {
    "common_auth": "auth requisite pam_faillock.so preauth\nauth [success=2 default=ignore] pam_unix.so try_first_pass\n"
                   "auth [default=die] pam_faillock.so authfail",
    "common_account": "account [success=1 new_authtok_reqd=done default=ignore] pam_unix.so\naccount required pam_faillock.so",
    "common_password": "password requisite pam_pwquality.so retry=3\n"
                       "password requisite pam_pwhistory.so remember=24 enforce_for_root use_authtok\n"
                       "password [success=1 default=ignore] pam_unix.so obscure use_authtok try_first_pass yescrypt",
    "common_session": "session required pam_unix.so",
    "common_session_noninteractive": "session required pam_unix.so",
}
ALL5 = [["pam", k] for k in COMMON_OK]


def pam_texts(paths, cond, field, good, bad, reads=None):
    return custom([f'is_string({ref(p)})' for p in paths], [cond], field, {"pam": good}, {"pam": bad},
                  reads=reads or paths)


put("5.3.2.1", "pam_unix module", pam_texts(ALL5, ['every k in pam_common_files { regex_in_text(input.pam[k], `\\bpam_unix\\.so\\b`) }'],
    "pam.common_*", COMMON_OK, dict(COMMON_OK, common_session="session optional pam_systemd.so")))
put("5.3.2.2", "pam_faillock module", pam_texts([PCA, PCC], [
    'regex_in_text(input.pam.common_auth, `^auth\\s+(\\[[^\\]]*\\]|\\S+)\\s+pam_faillock\\.so\\b.*\\spreauth\\b`)',
    'regex_in_text(input.pam.common_auth, `^auth\\s+(\\[[^\\]]*\\]|\\S+)\\s+pam_faillock\\.so\\b.*\\sauthfail\\b`)',
    'regex_in_text(input.pam.common_account, `^account\\s+(\\[[^\\]]*\\]|\\S+)\\s+pam_faillock\\.so\\b`)'],
    "pam.common_auth/common_account", COMMON_OK,
    dict(COMMON_OK, common_auth="auth [success=1 default=ignore] pam_unix.so try_first_pass")))
put("5.3.2.3", "pam_pwquality module", pam_texts([PCP], ['count(pwquality_lines) > 0'], "pam.common_password", COMMON_OK,
    dict(COMMON_OK, common_password="password [success=1 default=ignore] pam_unix.so obscure use_authtok yescrypt")))
put("5.3.2.4", "pam_pwhistory module", pam_texts([PCP], ['count(pwhistory_lines) > 0'], "pam.common_password", COMMON_OK,
    dict(COMMON_OK, common_password="password [success=1 default=ignore] pam_unix.so obscure use_authtok yescrypt")))
FL = ["pam", "faillock_conf_raw"]
FA = COMMON_OK["common_auth"]


def faillock(cond, good, bad):
    return custom([f'is_string({ref(FL)})', f'is_string({ref(PCA)})'], [cond], "pam.faillock_conf_raw/pam.common_auth",
                  {"pam": {"faillock_conf_raw": good, "common_auth": FA}}, {"pam": {"faillock_conf_raw": bad, "common_auth": FA}},
                  reads=[FL, PCA])


put("5.3.3.1.1", "failed attempts lockout is", faillock(
    ['v := num(conf_last(input.pam.faillock_conf_raw, "deny"))', 'v >= 1', 'v <= 5', 'not faillock_arg_out_of_range("deny", 1, 5)'],
    "deny = 5", "deny = 10"))
put("5.3.3.1.2", "unlock time", faillock(
    ['unlock_ok(conf_last(input.pam.faillock_conf_raw, "unlock_time"))', 'not faillock_unlock_arg_short'],
    "unlock_time = 900", "unlock_time = 600"))
put("5.3.3.1.3", "includes root", faillock(
    ['root_lockout_set', 'not root_unlock_short'], "deny = 3\neven_deny_root\nroot_unlock_time = 60",
    "deny = 3\nroot_unlock_time = 30"))
PQ = ["pam", "pwquality_conf_raw"]   # /etc/security/pwquality.conf + pwquality.conf.d/*.conf
PW_OK = COMMON_OK["common_password"]


def pwq(key, required, good, bad):
    cond = [f'every v in conf_values(input.pam.pwquality_conf_raw, {json.dumps(key)}) {{ pwq_ok({json.dumps(key)}, num(v)) }}',
            f'count([l | some l in pwquality_lines; n := pam_arg_num(l, {json.dumps(key)}); not pwq_ok({json.dumps(key)}, n)]) == 0']
    if required:
        cond = [f'count(conf_values(input.pam.pwquality_conf_raw, {json.dumps(key)})) > 0'] + cond
    return custom([f'is_string({ref(PQ)})', f'is_string({ref(PCP)})'], [cond], "pam.pwquality_conf_raw/pam.common_password",
                  {"pam": {"pwquality_conf_raw": good, "common_password": PW_OK}},
                  {"pam": {"pwquality_conf_raw": bad, "common_password": PW_OK}}, reads=[PQ, PCP])


put("5.3.3.2.1", "changed characters", pwq("difok", True, "difok = 2", "difok = 1"))
put("5.3.3.2.2", "password length", pwq("minlen", True, "minlen = 14", "minlen = 8"))
put("5.3.3.2.3", "password complexity", attest("password_complexity_reviewed"))
put("5.3.3.2.4", "same consecutive", pwq("maxrepeat", True, "maxrepeat = 3", "maxrepeat = 0"))
put("5.3.3.2.5", "maximum sequential", pwq("maxsequence", True, "maxsequence = 3", "maxsequence = 4"))
put("5.3.3.2.6", "dictionary check", pwq("dictcheck", False, "minlen = 14", "dictcheck = 0"))
put("5.3.3.2.7", "quality checking is enforced", pwq("enforcing", False, "minlen = 14", "enforcing = 0"))
put("5.3.3.2.8", "quality is enforced for the root", text_has(PQ, "(?i)^enforce_for_root\\b", True,
                                                               "minlen = 14\nenforce_for_root", "minlen = 14"))
PH = ["pam", "pwhistory_conf_raw"]   # /etc/security/pwhistory.conf (module arguments override it)
PWH = "password requisite pam_pwhistory.so remember=24 enforce_for_root use_authtok"
PWH_BARE = "password requisite pam_pwhistory.so"


def pwh(cond, field, good_line, bad_line, good_conf=None):
    good = {"common_password": good_line}
    if good_conf is not None:
        good["pwhistory_conf_raw"] = good_conf
    return custom([f'is_string({ref(PCP)})'], [['count(pwhistory_lines) > 0', cond]], f"pam.common_password/{field}",
                  {"pam": good}, {"pam": {"common_password": bad_line}}, reads=[PCP, PH])


put("5.3.3.3.1", "history remember", pwh('every l in pwhistory_lines { pwh_remember_ok(l) }', "pwhistory_conf_raw",
                                         PWH, PWH.replace("24", "5")))
put("5.3.3.3.2", "history is enforced for the root", pwh('every l in pwhistory_lines { pwh_flag(l, "enforce_for_root") }',
                                                         "pwhistory_conf_raw", PWH_BARE, PWH.replace(" enforce_for_root", ""),
                                                         good_conf="enforce_for_root"))
put("5.3.3.3.3", "pam_pwhistory includes use_authtok", pwh('every l in pwhistory_lines { pwh_flag(l, "use_authtok") }',
                                                           "pwhistory_conf_raw", PWH, PWH.replace(" use_authtok", "")))
UNIX = "password [success=1 default=ignore] pam_unix.so obscure use_authtok try_first_pass yescrypt"
put("5.3.3.4.1", "nullok", pam_texts(ALL5, ['count([l | some l in pam_unix_lines; regex.match(`\\snullok\\b`, l)]) == 0'],
    "pam.common_*", COMMON_OK, dict(COMMON_OK, common_auth="auth [success=2 default=ignore] pam_unix.so nullok")))
put("5.3.3.4.2", "does not include remember", pam_texts(ALL5,
    ['count([l | some l in pam_unix_lines; regex.match(`\\sremember=[0-9]+`, l)]) == 0'],
    "pam.common_*", COMMON_OK, dict(COMMON_OK, common_password=UNIX + " remember=5")))
put("5.3.3.4.3", "strong password hashing", pam_texts([PCP],
    ['count(pam_unix_password_lines) > 0', 'every l in pam_unix_password_lines { regex.match(`\\s(sha512|yescrypt)\\b`, l) }'],
    "pam.common_password", {"common_password": UNIX}, {"common_password": UNIX.replace("yescrypt", "md5")}))
put("5.3.3.4.4", "pam_unix includes use_authtok", pam_texts([PCP],
    ['count(pam_unix_password_lines) > 0', 'every l in pam_unix_password_lines { regex.match(`\\suse_authtok\\b`, l) }'],
    "pam.common_password", {"common_password": UNIX}, {"common_password": UNIX.replace(" use_authtok", "")}))

# 5.4 accounts — login.defs values (mapper fields pam.pass_*), per-user shadow facts (no hashes)
SH = ["user_group", "shadow_entries"]


def shadow_rule(defs_path, defs_cond, user_cond, good_def, bad_def, good_user):
    """login.defs default AND every user that has a password."""
    dp = ref(defs_path)
    return custom(
        [f'{dp} != null', f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
        [[f'v := num({dp})'] + defs_cond + [
          f'every u in [u | some u in {ref(SH)}; u.has_password == true] {{ {user_cond} }}']],
        f"{dotted(defs_path)}/user_group.shadow_entries",
        merge(nest(defs_path, good_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        merge(nest(defs_path, bad_def), nest(SH, [dict({"user": "alice", "has_password": True}, **good_user)])),
        nest(defs_path, good_def), reads=[defs_path, SH])


put("5.4.1.1", "password expiration is", shadow_rule(["pam", "pass_max_days"], ['v >= 1', 'v <= 365'],
    'u.max_days >= 1; u.max_days <= 365', 365, 99999, {"max_days": 365}))
put("5.4.1.2", "minimum password days", attest("minimum_password_days_reviewed"))
put("5.4.1.3", "warning days", shadow_rule(["pam", "pass_warn_age"], ['v >= 7'], 'u.warn_days >= 7', 7, 3, {"warn_days": 7}))
put("5.4.1.4", "strong password hashing algorithm is configured", custom(['is_string(input.pam.encrypt_method)'],
    [['upper(input.pam.encrypt_method) in {"SHA512", "YESCRYPT"}']], "pam.encrypt_method",
    {"pam": {"encrypt_method": "YESCRYPT"}}, {"pam": {"encrypt_method": "MD5"}}, reads=[["pam", "encrypt_method"]]))
put("5.4.1.5", "inactive password lock", shadow_rule(["pam", "inactive_days"], ['v >= 0', 'v <= 45'],
    'u.inactive >= 0; u.inactive <= 45', 45, -1, {"inactive": 30}))
put("5.4.1.6", "last password change", custom(
    [f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [[f'every u in {ref(SH)} {{ u.last_change <= today_days }}']], "user_group.shadow_entries",
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 19000}]),
    nest(SH, [{"user": "alice", "has_password": True, "last_change": 9999999}]), reads=[SH]))
PW = ["user_group", "passwd_content"]
GR = ["user_group", "group_content"]
PASSWD = "root:x:0:0:root:/root:/bin/bash\nbin:x:2:2:bin:/bin:/usr/sbin/nologin\nalice:x:1000:1000::/home/alice:/bin/bash"
GROUP = "root:x:0:\nbin:x:2:\nshadow:x:42:\nalice:x:1000:"
put("5.4.2.1", "only UID 0", custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'],
    [['count([f | some f in passwd_rows; f[2] == "0"; f[0] != "root"]) == 0']], dotted(PW),
    nest(PW, PASSWD), nest(PW, PASSWD + "\ntoor:x:0:0::/root:/bin/bash"), reads=[PW]))
put("5.4.2.2", "only GID 0 account", custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'],
    [['count([f | some f in passwd_rows; f[3] == "0"; not f[0] in {"root", "sync", "shutdown", "halt", "operator"}]) == 0']],
    dotted(PW), nest(PW, PASSWD), nest(PW, PASSWD + "\nbob:x:1001:0::/home/bob:/bin/bash"), reads=[PW]))
put("5.4.2.3", "only GID 0 group", custom([f'is_string({ref(GR)})', 'count(group_rows) > 0'],
    [['count([f | some f in group_rows; f[2] == "0"; f[0] != "root"]) == 0']], dotted(GR),
    nest(GR, GROUP), nest(GR, GROUP + "\nwheel2:x:0:"), reads=[GR]))
put("5.4.2.4", "root account access", custom(['is_string(input.user_group.root_password_status)'],
    [['regex.match(`^(P|L)`, input.user_group.root_password_status)']], "user_group.root_password_status",
    {"user_group": {"root_password_status": "L"}}, {"user_group": {"root_password_status": "NP"}},
    reads=[["user_group", "root_password_status"]]))
put("5.4.2.5", "root path integrity", custom(
    ['is_string(input.user_group.root_path)', 'is_array(input.user_group.root_path_entries)', 'count(input.user_group.root_path_entries) > 0'],
    [['root_path_text_ok', 'every d in input.user_group.root_path_entries { d.is_dir == true; d.owner == "root"; mode_within(d.mode, "0755") }']],
    "user_group.root_path/root_path_entries",
    {"user_group": {"root_path": "/usr/sbin:/usr/bin", "root_path_entries": [{"path": "/usr/bin", "is_dir": True, "owner": "root", "mode": "0755"}]}},
    {"user_group": {"root_path": "/usr/sbin:/usr/bin:.", "root_path_entries": [{"path": "/usr/bin", "is_dir": True, "owner": "root", "mode": "0755"}]}},
    reads=[["user_group", "root_path"], ["user_group", "root_path_entries"]]))
# /root/.profile and /root/.bashrc; neither present => no umask set there (root_shell_init_exists false)
_s = put("5.4.2.6", "root user umask", custom(['is_string(input.user_group.root_shell_init_raw)'],
    [['every u in umask_values(input.user_group.root_shell_init_raw) { umask_ok(u) }']],
    "user_group.root_shell_init_raw/root_shell_init_exists", {"user_group": {"root_shell_init_raw": "umask 0027"}},
    {"user_group": {"root_shell_init_raw": "umask 0022"}}, reads=[["user_group", "root_shell_init_raw"]]))
_s.ev_alts.append(['input.user_group.root_shell_init_exists == false'])
_s.passes.append(['input.user_group.root_shell_init_exists == false'])
_s.reads.append(["user_group", "root_shell_init_exists"])
SHELLS = ["user_group", "shells_content"]
UMIN = ["user_group", "uid_min"]
SHELLS_TXT = "/bin/bash\n/bin/sh\n/usr/sbin/nologin"
put("5.4.2.7", "valid login shell", custom(
    [f'is_string({ref(PW)})', 'count(passwd_rows) > 0', f'is_string({ref(SHELLS)})', f'is_number({ref(UMIN)})'],
    [['count(system_accounts_with_shell) == 0']], "user_group.passwd_content/shells_content/uid_min",
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "uid_min": 1000}},
    {"user_group": {"passwd_content": PASSWD + "\ndaemon:x:1:1::/:/bin/bash", "shells_content": SHELLS_TXT, "uid_min": 1000}},
    reads=[PW, SHELLS, UMIN]))
put("5.4.2.8", "are locked", custom(
    [f'is_string({ref(PW)})', 'count(passwd_rows) > 0', f'is_string({ref(SHELLS)})', f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [['count(unlocked_no_shell_accounts) == 0']], "user_group.passwd_content/shells_content/shadow_entries",
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "shadow_entries": [{"user": "bin", "locked": True}]}},
    {"user_group": {"passwd_content": PASSWD, "shells_content": SHELLS_TXT, "shadow_entries": [{"user": "bin", "locked": False}]}},
    reads=[PW, SHELLS, SH]))
put("5.4.3.1", "nologin", text_has(SHELLS, "/nologin\\b", False, "/bin/bash\n/bin/sh", SHELLS_TXT))
put("5.4.3.2", "shell timeout", custom(
    ['input.pam.tmout != null', 'is_boolean(input.pam.tmout_readonly)', 'is_boolean(input.pam.tmout_exported)'],
    [['num(input.pam.tmout) > 0', 'num(input.pam.tmout) <= 900', 'input.pam.tmout_readonly == true',
      'input.pam.tmout_exported == true']], "pam.tmout/tmout_readonly/tmout_exported",
    {"pam": {"tmout": 900, "tmout_readonly": True, "tmout_exported": True}},
    {"pam": {"tmout": 900, "tmout_readonly": False, "tmout_exported": True}},
    reads=[["pam", "tmout"], ["pam", "tmout_readonly"], ["pam", "tmout_exported"]]))
PDR = ["pam", "profile_d_raw"]   # umask lines of /etc/profile and /etc/profile.d/*
put("5.4.3.3", "default user umask", custom(
    [f'is_string({ref(PDR)})', 'input.pam.default_umask != null'],
    [[f'every u in umask_values({ref(PDR)}) {{ umask_ok(u) }}', 'umask_ok(sprintf("%v", [input.pam.default_umask]))']],
    "pam.profile_d_raw/pam.default_umask", {"pam": {"profile_d_raw": "umask 027", "default_umask": "027"}},
    {"pam": {"profile_d_raw": "umask 022", "default_umask": "027"}}, reads=[PDR, ["pam", "default_umask"]]))

# 6.1 logging (mapper journald / rsyslog fields)
put("6.1.1.1.1", "journald service is active", boolean(["journald", "service_active"], True))
put("6.1.1.1.2", "systemd-journal-remote", not_in_use(
    {"systemd-journal-remote": ["systemd-journal-remote.socket", "systemd-journal-remote.service"]}))
put("6.1.1.1.3", "send logs to rsyslog", custom(
    ['is_string(input.journald.forward_to_syslog)', 'is_boolean(input.rsyslog.service_active)'],
    [['lower(input.journald.forward_to_syslog) == "yes"', 'input.rsyslog.service_active == true']],
    "journald.forward_to_syslog/rsyslog.service_active",
    {"journald": {"forward_to_syslog": "yes"}, "rsyslog": {"service_active": True}},
    {"journald": {"forward_to_syslog": "no"}, "rsyslog": {"service_active": True}},
    reads=[["journald", "forward_to_syslog"], ["rsyslog", "service_active"]]))
put("6.1.1.1.4", "journald log file access", attest("journald_file_access_reviewed"))
put("6.1.1.1.5", "journald log file rotation", attest("journald_rotation_reviewed"))
put("6.1.1.1.6", "journald Storage", custom(['is_string(input.journald.storage_mode)'],
    [['lower(input.journald.storage_mode) == "persistent"']], "journald.storage_mode",
    {"journald": {"storage_mode": "persistent"}}, {"journald": {"storage_mode": "auto"}}, reads=[["journald", "storage_mode"]]))
put("6.1.1.1.7", "journald Compress", custom(['is_string(input.journald.compress_enabled)'],
    [['lower(input.journald.compress_enabled) == "yes"']], "journald.compress_enabled",
    {"journald": {"compress_enabled": "yes"}}, {"journald": {"compress_enabled": "no"}},
    reads=[["journald", "compress_enabled"]]))
put("6.1.2.1", "rsyslog is installed", boolean(["rsyslog", "installed"], True))
put("6.1.2.2", "rsyslog service is enabled", custom(
    ['is_boolean(input.rsyslog.service_enabled)', 'is_boolean(input.rsyslog.service_active)'],
    [['input.rsyslog.service_enabled == true', 'input.rsyslog.service_active == true']],
    "rsyslog.service_enabled/service_active", {"rsyslog": {"service_enabled": True, "service_active": True}},
    {"rsyslog": {"service_enabled": False, "service_active": True}},
    reads=[["rsyslog", "service_enabled"], ["rsyslog", "service_active"]]))
# $FileCreateMode unset => rsyslog's 0644 default, which fails; the mapper omits the key when unset
put("6.1.2.3", "file creation mode", custom(['is_boolean(input.rsyslog.installed)'],
    [['mode_within(input.rsyslog.file_create_mode, "0640")']], "rsyslog",
    {"rsyslog": {"installed": True, "file_create_mode": "0640"}},
    {"rsyslog": {"installed": True, "file_create_mode": "0644"}},
    reads=[["rsyslog", "installed"], ["rsyslog", "file_create_mode"]]))
put("6.1.2.4", "rsyslog logging is configured", attest("rsyslog_logging_reviewed"))
put("6.1.2.5", "remote log host", attest("rsyslog_remote_host_configured"))
RC = ["rsyslog", "config_raw"]
put("6.1.2.6", "receive logs", text_has(RC,
    "(?i)(^module\\(load=\"?im(tcp|udp)\"?\\)|^input\\(type=\"?im(tcp|udp)\"?|^\\$ModLoad\\s+im(tcp|udp)\\b|^\\$(InputTCPServerRun|UDPServerRun)\\b)",
    False, "*.* /var/log/syslog", 'module(load="imudp")'))
put("6.1.2.7", "logrotate", attest("logrotate_reviewed"))
put("6.1.2.8", "rsyslog-gnutls", pkg_present("rsyslog-gnutls"))
put("6.1.2.9", "gtls", text_has(RC, "(?i)StreamDriver\\s*=\\s*\"?gtls\"?", True,
    '*.* action(type="omfwd" target="loghost" StreamDriver="gtls" StreamDriverMode="1")',
    '*.* action(type="omfwd" target="loghost" protocol="tcp")'))
put("6.1.2.10", "CA certificates", attest("rsyslog_ca_certificates_configured"))
LF = ["log_file_permissions", "analysis"]
put("6.1.3.1", "all logfiles", custom(
    [f'is_array({ref(LF)})', f'count({ref(LF)}) > 0'], [[f'every f in {ref(LF)} {{ log_file_ok(f) }}']], dotted(LF),
    nest(LF, [{"file": "/var/log/syslog", "actual_mode": "0640", "actual_owner": "syslog", "actual_group": "adm"},
              {"file": "/var/log/apt/history.log", "actual_mode": "0644", "actual_owner": "root", "actual_group": "root"},
              {"file": "/var/log/wtmp", "actual_mode": "0664", "actual_owner": "root", "actual_group": "utmp"}]),
    nest(LF, [{"file": "/var/log/syslog", "actual_mode": "0644", "actual_owner": "syslog", "actual_group": "adm"}]),
    reads=[LF]))

# 6.2 auditd
put("6.2.1.1", "auditd packages", pkgs_present(["auditd", "audispd-plugins"]))
put("6.2.1.2", "auditd service", unit_running("auditd.service"))
put("6.2.1.3", "prior to auditd", custom(
    [f'is_array({ref(GL)})', 'count(grub_linux_lines) > 0'],
    [['every l in grub_linux_lines { regex.match(`(^|\\s)audit=1(\\s|$)`, l) }']], dotted(GL),
    nest(GL, [LINUX + " audit=1 audit_backlog_limit=8192"]), nest(GL, [LINUX + " audit=1", LINUX]), reads=[GL]))
put("6.2.1.4", "audit_backlog_limit", custom(
    [f'is_array({ref(GL)})', 'count(grub_linux_lines) > 0'],
    [['every l in grub_linux_lines { regex.match(`(^|\\s)audit_backlog_limit=[0-9]+(\\s|$)`, l) }']], dotted(GL),
    nest(GL, [LINUX + " audit=1 audit_backlog_limit=8192"]), nest(GL, [LINUX + " audit=1"]), reads=[GL]))
AC = ["auditd", "config"]
put("6.2.2.1", "storage size", conf(AC, "max_log_file", 'num(v) > 0', "max_log_file = 8", "max_log_file = 0"))
put("6.2.2.2", "not automatically deleted", conf(AC, "max_log_file_action", 'lower(v) == "keep_logs"',
                                                 "max_log_file_action = keep_logs", "max_log_file_action = rotate"))
put("6.2.2.3", "disabled when audit logs are full", custom([f'is_string({ref(AC)})'],
    [['lower(conf_last(input.auditd.config, "disk_full_action")) in {"halt", "single"}',
      'lower(conf_last(input.auditd.config, "disk_error_action")) in {"syslog", "single", "halt"}']],
    dotted(AC), nest(AC, "disk_full_action = halt\ndisk_error_action = syslog"),
    nest(AC, "disk_full_action = suspend\ndisk_error_action = syslog"), reads=[AC]))
put("6.2.2.4", "low on space", custom([f'is_string({ref(AC)})'],
    [['lower(conf_last(input.auditd.config, "space_left_action")) in {"email", "exec", "single", "halt"}',
      'lower(conf_last(input.auditd.config, "admin_space_left_action")) in {"single", "halt"}']],
    dotted(AC), nest(AC, "space_left_action = email\nadmin_space_left_action = single"),
    nest(AC, "space_left_action = ignore\nadmin_space_left_action = single"), reads=[AC]))
AR = ["auditd", "active_rules"]
W = "-a always,exit -F arch=b64 -F path={p} -F perm={perm} -F key=k"
S64 = "-a always,exit -F arch={arch} -S {s} {extra}-F key=k"
AUID = "-F auid>=1000 -F auid!=unset "


def audit(watches=(), syscalls=(), arches=False):
    """watches: [(path, perm)]; syscalls: [(names, [field names])]. Rules are matched on tokens.
    arches: every syscall rule is required for both arch=b32 and arch=b64."""
    conds, good = [], []
    for p, perm in watches:
        conds.append(f'audit_watch({json.dumps(p)}, {json.dumps(perm)})')
        good.append(W.format(p=p, perm=perm))
    for names, fields in syscalls:
        extra = ""
        for f in fields:
            extra += {"auid": AUID, "euid": "-C euid!=uid -F auid!=unset ", "eacces": "-F exit=-EACCES " + AUID,
                      "eperm": "-F exit=-EPERM " + AUID}[f]
        if arches:
            conds.append(f'every s in {json.dumps(names)} {{ every a in ["b32", "b64"] {{ audit_syscall_arch(s, {json.dumps(fields)}, a) }} }}')
            good += [S64.format(arch=a, s=",".join(names), extra=extra) for a in ("b32", "b64")]
        else:
            conds.append(f'every s in {json.dumps(names)} {{ audit_syscall(s, {json.dumps(fields)}) }}')
            good.append(S64.format(arch="b64", s=",".join(names), extra=extra))
    bad = good[:-1]
    return custom([f'is_array({ref(AR)})'], [conds], dotted(AR), nest(AR, good), nest(AR, bad), reads=[AR])


put("6.2.3.1", "/etc/sudoers", audit([("/etc/sudoers", "wa"), ("/etc/sudoers.d", "wa")]))
put("6.2.3.2", "another user", audit(syscalls=[(["execve"], ["euid"])]))
put("6.2.3.3", "sudo log file", custom(
    [f'is_array({ref(AR)})', f'is_string({ref(SU)})'], [['some f in sudo_logfiles', 'audit_watch(f, "wa")']],
    "auditd.active_rules/sudo.sudoers_content",
    merge(nest(AR, [W.format(p="/var/log/sudo.log", perm="wa")]), nest(SU, 'Defaults logfile="/var/log/sudo.log"')),
    merge(nest(AR, [W.format(p="/var/log/other.log", perm="wa")]), nest(SU, 'Defaults logfile="/var/log/sudo.log"')),
    reads=[AR, SU]))
put("6.2.3.4", "date and time", audit([("/etc/localtime", "wa")], [(["adjtimex", "settimeofday", "clock_settime"], [])]))
put("6.2.3.5", "sethostname", audit(syscalls=[(["sethostname", "setdomainname"], [])]))
put("6.2.3.6", "/etc/issue", audit([("/etc/issue", "wa"), ("/etc/issue.net", "wa")]))
put("6.2.3.7", "/etc/hosts", audit([("/etc/hosts", "wa"), ("/etc/hostname", "wa")]))
put("6.2.3.8", "network environment", audit([("/etc/network/interfaces", "wa"), ("/etc/network/interfaces.d", "wa"),
                                             ("/etc/netplan", "wa")]))
put("6.2.3.9", "/etc/NetworkManager", gate_pkg(audit([("/etc/NetworkManager", "wa")]), "network-manager"))
put("6.2.3.10", "privileged commands", custom(
    [f'is_array({ref(AR)})', 'is_array(input.auditd.privileged_commands)', 'count(input.auditd.privileged_commands) > 0'],
    [['every c in input.auditd.privileged_commands { audit_watch(c, "x") }']], "auditd.active_rules/privileged_commands",
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su"]}},
    {"auditd": {"active_rules": [W.format(p="/usr/bin/su", perm="x")], "privileged_commands": ["/usr/bin/su", "/usr/bin/passwd"]}},
    reads=[AR, ["auditd", "privileged_commands"]]))
put("6.2.3.11", "unsuccessful file access", audit(syscalls=[(["creat", "open", "openat", "truncate", "ftruncate"], ["eacces"]),
                                                            (["creat", "open", "openat", "truncate", "ftruncate"], ["eperm"])]))
put("6.2.3.12", "/etc/group", audit([("/etc/group", "wa")]))
put("6.2.3.13", "/etc/passwd", audit([("/etc/passwd", "wa")]))
put("6.2.3.14", "/etc/shadow", audit([("/etc/gshadow", "wa"), ("/etc/shadow", "wa")]))
put("6.2.3.15", "/etc/security/opasswd", audit([("/etc/security/opasswd", "wa")]))
put("6.2.3.16", "/etc/nsswitch.conf", audit([("/etc/nsswitch.conf", "wa")]))
put("6.2.3.17", "/etc/pam.conf", audit([("/etc/pam.conf", "wa"), ("/etc/pam.d", "wa")]))
put("6.2.3.18", "chmod", audit(syscalls=[(["chmod", "fchmod", "fchmodat", "fchmodat2"], ["auid"])]))
put("6.2.3.19", "chown", audit(syscalls=[(["chown", "fchown", "lchown", "fchownat"], ["auid"])]))
put("6.2.3.20", "setxattr", audit(syscalls=[(["setxattr", "lsetxattr", "fsetxattr", "removexattr", "lremovexattr",
                                              "fremovexattr"], ["auid"])]))
put("6.2.3.21", "mounts", audit(syscalls=[(["mount"], ["auid"])]))
put("6.2.3.22", "session initiation", audit([("/var/run/utmp", "wa"), ("/var/log/wtmp", "wa"), ("/var/log/btmp", "wa")]))
put("6.2.3.23", "login and logout", audit([("/var/log/lastlog", "wa"), ("/var/run/faillock", "wa")]))
put("6.2.3.24", "unlink", audit(syscalls=[(["unlink", "unlinkat"], ["auid"])]))
put("6.2.3.25", "rename", audit(syscalls=[(["rename", "renameat", "renameat2"], ["auid"])]))
put("6.2.3.26", "Mandatory Access Controls", audit([("/etc/apparmor", "wa"), ("/etc/apparmor.d", "wa")]))
put("6.2.3.27", "chcon", audit([("/usr/bin/chcon", "x")]))
put("6.2.3.28", "setfacl", audit([("/usr/bin/setfacl", "x")]))
put("6.2.3.29", "chacl", audit([("/usr/bin/chacl", "x")]))
put("6.2.3.30", "usermod", audit([("/usr/sbin/usermod", "x")]))
put("6.2.3.31", "kernel module loading", audit([("/usr/bin/kmod", "x")]))
put("6.2.3.32", "init_module", audit(syscalls=[(["init_module", "finit_module"], ["auid"])], arches=True))
put("6.2.3.33", "delete_module", audit(syscalls=[(["delete_module"], ["auid"])]))
ARD = ["auditd", "rules"]   # on-disk rules (/etc/audit/rules.d)
put("6.2.3.34", "regardless of errors", custom([f'is_array({ref(ARD)})'],
    [['some l in input.auditd.rules', 'regex.match(`^\\s*-c\\s*$`, l)']],
    dotted(ARD), nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 2"]), reads=[ARD]))
put("6.2.3.35", "immutable", custom([f'is_array({ref(ARD)})'], [['audit_last_e == "2"']], dotted(ARD),
    nest(ARD, ["-D", "-c", "-e 2"]), nest(ARD, ["-D", "-e 1"]), reads=[ARD]))
put("6.2.3.36", "running and on disk", attest("audit_running_and_disk_rules_match"))
ALF = ["auditd", "log_files"]
ALF_GOOD = [{"path": "/var/log/audit/audit.log", "mode": "0600", "owner": "root", "group": "adm"}]
put("6.2.4.1", "log files mode", custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0'],
    [[f'every f in {ref(ALF)} {{ mode_within(f.mode, "0640") }}']], dotted(ALF),
    nest(ALF, ALF_GOOD), nest(ALF, [dict(ALF_GOOD[0], mode="0644")]), reads=[ALF]))
put("6.2.4.2", "log files owner", custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0'],
    [[f'every f in {ref(ALF)} {{ f.owner == "root" }}']], dotted(ALF),
    nest(ALF, ALF_GOOD), nest(ALF, [dict(ALF_GOOD[0], owner="syslog")]), reads=[ALF]))
put("6.2.4.3", "log files group owner", custom([f'is_array({ref(ALF)})', f'count({ref(ALF)}) > 0', f'is_string({ref(AC)})'],
    [['not audit_log_group_bad', f'every f in {ref(ALF)} {{ f.group in {{"root", "adm"}} }}']],
    "auditd.log_files/auditd.config",
    merge(nest(ALF, ALF_GOOD), nest(AC, "log_group = adm")),
    merge(nest(ALF, ALF_GOOD), nest(AC, "log_group = users")), reads=[ALF, AC]))
put("6.2.4.4", "log file directory", custom(['input.auditd.log_dir_permissions.mode != null'],
    [['mode_within(input.auditd.log_dir_permissions.mode, "0750")']], "auditd.log_dir_permissions.mode",
    {"auditd": {"log_dir_permissions": {"mode": "0750"}}}, {"auditd": {"log_dir_permissions": {"mode": "0755"}}},
    reads=[["auditd", "log_dir_permissions", "mode"]]))
for rid, kw, path, key, cond, badv in [
        ("6.2.4.5", "configuration files mode", ["auditd", "config_files"], "mode", 'mode_within(f.mode, "0640")', "0644"),
        ("6.2.4.6", "configuration files owner", ["auditd", "config_files"], "owner", 'f.owner == "root"', "adm"),
        ("6.2.4.7", "configuration files group", ["auditd", "config_files"], "group", 'f.group == "root"', "adm"),
        ("6.2.4.8", "tools mode", ["auditd", "tools"], "mode", 'mode_within(f.mode, "0755")', "0777"),
        ("6.2.4.9", "tools owner", ["auditd", "tools"], "owner", 'f.owner == "root"', "adm"),
        ("6.2.4.10", "tools group", ["auditd", "tools"], "group", 'f.group == "root"', "adm")]:
    g = {"path": "/etc/audit/auditd.conf", "mode": "0640", "owner": "root", "group": "root"}
    put(rid, kw, custom([f'is_array({ref(path)})', f'count({ref(path)}) > 0'], [[f'every f in {ref(path)} {{ {cond} }}']],
                        dotted(path), nest(path, [g]), nest(path, [dict(g, **{key: badv})]), reads=[path]))

# 6.3 integrity (AIDE; Ubuntu schedules dailyaidecheck.timer)
put("6.3.1", "AIDE is installed", pkgs_present(["aide", "aide-common"]))
AT = ["aide", "dailyaidecheck"]   # {timer_unit_file_state, timer_active_state, service_unit_file_state}
put("6.3.2", "regularly checked", custom(
    ['is_string(input.aide.dailyaidecheck.timer_unit_file_state)', 'is_string(input.aide.dailyaidecheck.timer_active_state)',
     'is_string(input.aide.dailyaidecheck.service_unit_file_state)'],
    [['input.aide.dailyaidecheck.timer_unit_file_state == "enabled"', 'input.aide.dailyaidecheck.timer_active_state == "active"',
      'input.aide.dailyaidecheck.service_unit_file_state in {"static", "enabled"}']],
    dotted(AT),
    nest(AT, {"timer_unit_file_state": "enabled", "timer_active_state": "active", "service_unit_file_state": "static"}),
    nest(AT, {"timer_unit_file_state": "disabled", "timer_active_state": "inactive", "service_unit_file_state": "static"}),
    reads=[AT]))
AIDE_OPTS = "p+i+n+u+g+s+b+acl+xattrs+sha512"
TOOLS = ["auditctl", "auditd", "ausearch", "aureport", "augenrules"]
put("6.3.3", "integrity of audit tools", custom(
    ['is_object(input.aide.audit_tool_options)'],
    [[f'every t in {json.dumps(TOOLS)} {{ aide_opts_ok(input.aide.audit_tool_options[t]) }}']],
    "aide.audit_tool_options",
    {"aide": {"audit_tool_options": {t: AIDE_OPTS for t in TOOLS}}},
    {"aide": {"audit_tool_options": dict({t: AIDE_OPTS for t in TOOLS[:-1]}, augenrules="p+i+n")}},
    reads=[["aide", "audit_tool_options"]]))

# 7.1 system files (rhel_9 critical_files.analysis entries, judged here from actual_* values)
put("7.1.1", "/etc/passwd ", crit("/etc/passwd", "0644", ["root"]))
put("7.1.2", "/etc/passwd-", crit("/etc/passwd-", "0644", ["root"]))
put("7.1.3", "/etc/group ", crit("/etc/group", "0644", ["root"]))
put("7.1.4", "/etc/group-", crit("/etc/group-", "0644", ["root"]))
put("7.1.5", "/etc/shadow ", crit("/etc/shadow", "0640", ["root", "shadow"]))
put("7.1.6", "/etc/shadow-", crit("/etc/shadow-", "0640", ["root", "shadow"]))
put("7.1.7", "/etc/gshadow ", crit("/etc/gshadow", "0640", ["root", "shadow"]))
put("7.1.8", "/etc/gshadow-", crit("/etc/gshadow-", "0640", ["root", "shadow"], optional=True))
put("7.1.9", "/etc/shells", crit("/etc/shells", "0644", ["root"]))
_op = crit("/etc/security/opasswd", "0600", ["root"], optional=True)
_op2 = crit("/etc/security/opasswd.old", "0600", ["root"], optional=True)
_op.ev_alts = [_op.ev_alts[0] + _op2.ev_alts[0]]
_op.passes = [a + b for a in _op.passes for b in _op2.passes]
_op.field = "critical_files.analysis[/etc/security/opasswd, /etc/security/opasswd.old]"
_old = {"file": "/etc/security/opasswd.old", "exists": False}
_op.tp = {"critical_files": {"analysis": _op.tp["critical_files"]["analysis"] + [_old]}}
_op.tf = {"critical_files": {"analysis": _op.tf["critical_files"]["analysis"] + [_old]}}
put("7.1.10", "/etc/security/opasswd", _op)
put("7.1.11", "world writable", custom(
    ['is_array(input.world_writable.files)', 'is_array(input.sticky_bit.world_writable_without_sticky)'],
    [['count(input.world_writable.files) == 0', 'count(input.sticky_bit.world_writable_without_sticky) == 0']],
    "world_writable.files/sticky_bit.world_writable_without_sticky",
    {"world_writable": {"files": []}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": ["/opt/app/data.txt"]}, "sticky_bit": {"world_writable_without_sticky": []}},
    {"world_writable": {"files": []}},
    reads=[["world_writable", "files"], ["sticky_bit", "world_writable_without_sticky"]]))
put("7.1.12", "without an owner", custom(
    ['is_array(input.unowned_files.files)', 'is_array(input.ungrouped_files.files)'],
    [['count(input.unowned_files.files) == 0', 'count(input.ungrouped_files.files) == 0']],
    "unowned_files.files/ungrouped_files.files",
    {"unowned_files": {"files": []}, "ungrouped_files": {"files": []}},
    {"unowned_files": {"files": []}, "ungrouped_files": {"files": ["/srv/orphan"]}},
    reads=[["unowned_files", "files"], ["ungrouped_files", "files"]]))
put("7.1.13", "SUID and SGID", attest("suid_sgid_files_reviewed"))

# 7.2 local users and groups
put("7.2.1", "shadowed passwords", custom([f'is_string({ref(PW)})', 'count(passwd_rows) > 0'],
    [['count([f | some f in passwd_rows; f[1] != "x"]) == 0']],
    dotted(PW), nest(PW, PASSWD), nest(PW, PASSWD + "\nbob::1001:1000::/home/bob:/bin/bash"), reads=[PW]))
put("7.2.2", "password fields are not empty", custom([f'is_array({ref(SH)})', f'count({ref(SH)}) > 0'],
    [[f'count([u | some u in {ref(SH)}; u.password_empty != false]) == 0']], dotted(SH),
    nest(SH, [{"user": "root", "password_empty": False}]), nest(SH, [{"user": "bob", "password_empty": True}]), reads=[SH]))
put("7.2.3", "all groups in /etc/passwd", custom(
    [f'is_string({ref(PW)})', f'is_string({ref(GR)})', 'count(passwd_rows) > 0', 'count(group_rows) > 0'],
    [['count([f | some f in passwd_rows; not f[3] in {g[2] | some g in group_rows}]) == 0']],
    "user_group.passwd_content/group_content", {"user_group": {"passwd_content": PASSWD, "group_content": GROUP}},
    {"user_group": {"passwd_content": PASSWD + "\nbob:x:1001:4242::/home/bob:/bin/bash", "group_content": GROUP}},
    reads=[PW, GR]))
put("7.2.4", "shadow group is empty", custom(
    [f'is_string({ref(PW)})', f'is_string({ref(GR)})', 'count(passwd_rows) > 0', 'count(group_rows) > 0'],
    [['count([f | some f in group_rows; f[0] == "shadow"; count(f) >= 4; trim_space(f[3]) != ""]) == 0',
      'count([p | some p in passwd_rows; some g in group_rows; g[0] == "shadow"; p[3] == g[2]]) == 0']],
    "user_group.passwd_content/group_content", {"user_group": {"passwd_content": PASSWD, "group_content": GROUP}},
    {"user_group": {"passwd_content": PASSWD, "group_content": GROUP.replace("shadow:x:42:", "shadow:x:42:alice")}},
    reads=[PW, GR]))
for rid, kw, rows_name, col, path, txt, dup in [
        ("7.2.5", "duplicate UIDs", "passwd_rows", 2, PW, PASSWD, "\ntoor:x:1000:1000::/home/t:/bin/bash"),
        ("7.2.6", "duplicate GIDs", "group_rows", 2, GR, GROUP, "\nstaff:x:1000:"),
        ("7.2.7", "duplicate user names", "passwd_rows", 0, PW, PASSWD, "\nalice:x:1005:1000::/home/a2:/bin/bash"),
        ("7.2.8", "duplicate group names", "group_rows", 0, GR, GROUP, "\nalice:x:1200:")]:
    put(rid, kw, custom([f'is_string({ref(path)})', f'count({rows_name}) > 0'],
                        [[f'count([f | some f in {rows_name}]) == count({{f[{col}] | some f in {rows_name}}})']],
                        dotted(path), nest(path, txt), nest(path, txt + dup), reads=[path]))
HOMES = ["user_group", "interactive_homes"]
put("7.2.9", "home directories", custom(
    [f'is_array({ref(HOMES)})'],
    [[f'every h in {ref(HOMES)} {{ h.exists == true; h.owner == h.user; mode_within(h.mode, "0750") }}']], dotted(HOMES),
    nest(HOMES, [{"user": "alice", "home": "/home/alice", "exists": True, "owner": "alice", "mode": "0750"}]),
    nest(HOMES, [{"user": "alice", "home": "/home/alice", "exists": True, "owner": "alice", "mode": "0755"}]), reads=[HOMES]))
DOT = ["user_group", "dotfiles"]
put("7.2.10", "dot files", custom(
    [f'is_array({ref(DOT)})'], [[f'every d in {ref(DOT)} {{ dotfile_ok(d) }}']], dotted(DOT),
    nest(DOT, [{"user": "alice", "path": "/home/alice/.bashrc", "mode": "0644", "owner": "alice", "group": "alice", "primary_group": "alice"},
               {"user": "alice", "path": "/home/alice/.bash_history", "mode": "0600", "owner": "alice", "group": "alice", "primary_group": "alice"}]),
    nest(DOT, [{"user": "alice", "path": "/home/alice/.rhosts", "mode": "0600", "owner": "alice", "group": "alice", "primary_group": "alice"}]),
    reads=[DOT]))
DDIR = ["user_group", "dot_dirs"]
put("7.2.11", "dot directories", custom(
    [f'is_array({ref(DDIR)})'], [[f'every d in {ref(DDIR)} {{ dotdir_ok(d) }}']], dotted(DDIR),
    nest(DDIR, [{"user": "alice", "path": "/home/alice/.ssh", "mode": "0700", "owner": "alice", "group": "alice", "primary_group": "alice"},
                {"user": "alice", "path": "/home/alice/.config", "mode": "0750", "owner": "alice", "group": "alice", "primary_group": "alice"}]),
    nest(DDIR, [{"user": "alice", "path": "/home/alice/.ssh", "mode": "0750", "owner": "alice", "group": "alice", "primary_group": "alice"}]),
    reads=[DDIR]))


# An empty collected list or blank body is a genuine pass here (offender lists; empty banners; no interactive
# users / dotfiles / dot directories; empty apt drop-in directories; no pam_motd file). Collector must always emit
# these keys.
for _rid in ["1.2.1.6", "1.2.1.9", "1.6.1", "1.6.2", "1.6.3", "1.6.4", "1.6.9", "5.3.1.1", "5.3.1.2", "5.3.1.3",
             "5.3.1.4", "7.1.11", "7.1.12", "7.2.9", "7.2.10", "7.2.11"]:
    L[_rid].hollow_ok = True

# ── helpers used by one module only ──────────────────────────────────────────────────────────
EXTRA = {
    "package_management": r'''
# values of an option in `apt-config dump` output (`Key "value";`); the last one wins
apt_values(key) := [m[1] |
	some l in active_lines(input.apt.config_dump)
	some m in regex.find_all_string_submatch_n(concat("", ["(?i)^", key, `\s+"([^"]*)";\s*$`]), l, 1)
]

apt_last(key) := lower(vs[count(vs) - 1]) if {
	vs := apt_values(key)
	count(vs) > 0
}
''',
    "apparmor": r'''
# `linux` lines of /boot/grub/grub.cfg
grub_linux_lines := [l |
	some l in input.boot_security.grub_cfg_linux_lines
	is_string(l)
	regex.match(`^\s*linux`, l)
]

# /etc/default/grub.d/apparmor.cfg is optional; when collected it must not set apparmor=0
apparmor_grub_d_disabled if regex_in_text(input.apparmor.grub_d_cfg_raw, `(^|[\s"])apparmor=0\b`)
''',
    "process_hardening": r'''
# `* hard core <n>` entries of limits.conf and limits.d, in parse order
core_hard_limits := [m[1] |
	some l in active_lines(input.limits.limits_raw)
	some m in regex.find_all_string_submatch_n(`^\*\s+hard\s+core\s+(\S+)`, l, 1)
]
''',
    "time_sync": r'''
time_unit_reported if unit_reported("systemd-timesyncd.service")

time_unit_reported if unit_reported("chrony.service")

# exactly one of the two daemons enabled or active (collected unit state)
timesyncd_only if {
	services_collected
	unit_on("systemd-timesyncd.service")
	not unit_on("chrony.service")
}

chrony_only if {
	services_collected
	unit_on("chrony.service")
	not unit_on("systemd-timesyncd.service")
}
''',
    "services": r'''
# postfix inet_interfaces (when collected): "all" or a non-loopback value listens off-host
mta_interfaces_open if regex.match(`\ball\b`, lower(input.mta.inet_interfaces))

mta_interfaces_open if {
	is_string(input.mta.inet_interfaces)
	not regex.match(`(0\.0\.0\.0|::1|loopback-only|localhost)`, lower(input.mta.inet_interfaces))
}
''',
    "firewall": r'''
# `ufw status verbose`: "Default: deny (incoming), allow (outgoing), disabled (routed)"
ufw_default(dir) := lower(m[1]) if {
	some l in active_lines(input.ufw.status_verbose)
	startswith(l, "Default:")
	some m in regex.find_all_string_submatch_n(concat("", [`([A-Za-z]+) \(`, dir, `\)`]), l, 1)
}

# /etc/default/ufw DEFAULT_FORWARD_POLICY (used when ufw shows "disabled (routed)" because forwarding is off)
ufw_forward_policy_drop if upper(conf_last(input.ufw.default_raw, "DEFAULT_FORWARD_POLICY")) == "DROP"
''',
    "ssh": r'''
maxstartups_ok(v) if {
	parts := split(v, ":")
	num(parts[0]) <= 10
	ms_le(parts, 1, 30)
	ms_le(parts, 2, 60)
}

ms_le(parts, i, _) if count(parts) <= i

ms_le(parts, i, lim) if num(parts[i]) <= lim
''',
    "sudo": r'''
sudo_defaults_has(re) if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	regex.match(re, l)
}

sudo_logfile if sudo_defaults_has(`(^|[\s,])logfile\s*=`)

sudo_timeouts := [num(m[1]) |
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`timestamp_timeout\s*=\s*(-?[0-9]+)`, l, -1)
]

# groups named by `auth required|requisite pam_wheel.so use_uid group=<g>`
su_wheel_groups contains g if {
	some l in active_lines(input.pam.su)
	regex.match(`^auth\s+(required|requisite)\s+pam_wheel\.so\b`, l)
	regex.match(`\suse_uid\b`, l)
	some m in regex.find_all_string_submatch_n(`\sgroup=(\S+)`, l, 1)
	g := m[1]
}

su_group_empty(g) if {
	some f in group_rows
	f[0] == g
	count(f) >= 4
	trim_space(f[3]) == ""
}
''',
    "pam": r'''
pam_common_files := ["common_auth", "common_account", "common_password", "common_session", "common_session_noninteractive"]

pam_module_lines(text, type, mod) := [l |
	some l in active_lines(text)
	regex.match(concat("", ["^", type, `\s+(\[[^\]]*\]|\S+)\s+`, mod, `\b`]), l)
]

pwhistory_lines := pam_module_lines(input.pam.common_password, "password", `pam_pwhistory\.so`)

pwquality_lines := pam_module_lines(input.pam.common_password, "password", `pam_pwquality\.so`)

faillock_auth_lines := pam_module_lines(input.pam.common_auth, "auth", `pam_faillock\.so`)

pam_unix_lines := [l |
	some k in pam_common_files
	some l in active_lines(input.pam[k])
	regex.match(`\spam_unix\.so\b`, l)
]

pam_unix_password_lines := pam_module_lines(input.pam.common_password, "password", `pam_unix\.so`)

pam_arg_num(l, k) := num(m[1]) if {
	some m in regex.find_all_string_submatch_n(concat("", [`\s`, k, `=(-?[0-9]+)`]), l, 1)
}

pam_has_arg(l, k) if regex.match(concat("", [`\s`, k, `(=|\s|$)`]), l)

# pam_faillock module arguments override faillock.conf
faillock_arg_out_of_range(k, lo, hi) if {
	some l in faillock_auth_lines
	pam_has_arg(l, k)
	not in_range(pam_arg_num(l, k), lo, hi)
}

in_range(n, lo, hi) if {
	n >= lo
	n <= hi
}

unlock_ok(v) if num(v) == 0

unlock_ok(v) if num(v) >= 900

faillock_unlock_arg_short if {
	some l in faillock_auth_lines
	n := pam_arg_num(l, "unlock_time")
	n > 0
	n < 900
}

root_lockout_set if regex_in_text(input.pam.faillock_conf_raw, `(?i)^even_deny_root\b`)

root_lockout_set if conf_last(input.pam.faillock_conf_raw, "root_unlock_time")

root_unlock_short if num(conf_last(input.pam.faillock_conf_raw, "root_unlock_time")) < 60

root_unlock_short if {
	some l in faillock_auth_lines
	pam_arg_num(l, "root_unlock_time") < 60
}

# pwquality: every value in pwquality.conf / pwquality.conf.d and every pam_pwquality argument must comply
pwq_ok("difok", n) if n >= 2

pwq_ok("minlen", n) if n >= 14

pwq_ok("maxrepeat", n) if in_range(n, 1, 3)

pwq_ok("maxsequence", n) if in_range(n, 1, 3)

pwq_ok("dictcheck", n) if n != 0

pwq_ok("enforcing", n) if n != 0

# pam_pwhistory: a module argument overrides /etc/security/pwhistory.conf
pwh_remember_ok(l) if pam_arg_num(l, "remember") >= 24

pwh_remember_ok(l) if {
	not regex.match(`\sremember=`, l)
	num(conf_last(input.pam.pwhistory_conf_raw, "remember")) >= 24
}

pwh_flag(l, f) if regex.match(concat("", [`\s`, f, `\b`]), l)

pwh_flag(_, f) if regex_in_text(input.pam.pwhistory_conf_raw, concat("", [`(?i)^`, f, `\b`]))
''',
    "accounts": r'''
today_days := time.now_ns() / 86400000000000

root_path_text_ok if {
	p := input.user_group.root_path
	not contains(p, "::")
	not endswith(trim_space(p), ":")
	not startswith(trim_space(p), ":")
	not "." in split(p, ":")
}

umask_values(text) := [m[1] |
	some l in active_lines(text)
	some m in regex.find_all_string_submatch_n(`(?i)^umask\s+(\S+)`, l, 1)
]

# at least as restrictive as 027 (027 octal = 23)
umask_ok(u) if bits.and(octal(u), 23) == 23

umask_ok(u) if {
	regex.match(`^([ugoa]+[=+-][rwx]*,?)+$`, u)
	not regex.match(`g=[rx]*w`, u)
	not regex.match(`o=[rwx]+`, u)
}

system_uid(uid) if uid < input.user_group.uid_min

system_uid(uid) if uid == 65534

system_accounts_with_shell := [f[0] |
	some f in passwd_rows
	not f[0] in {"root", "halt", "sync", "shutdown", "nfsnobody"}
	system_uid(to_number(f[2]))
	f[6] in valid_shells
]

shadow_locked(u) if {
	some e in input.user_group.shadow_entries
	e.user == u
	e.locked == true
}

unlocked_no_shell_accounts := [f[0] |
	some f in passwd_rows
	f[0] != "root"
	not f[6] in valid_shells
	not shadow_locked(f[0])
]
''',
    "logging": r'''
# files under /var/log/apt: 0644, root, root|adm (checked before the per-name rules)
log_rule_for(path) := ["0644", {"root"}, {"root", "adm"}] if {
	regex.match(`^/var/log/apt/[^/]+$`, path)
} else := log_rule(regex.replace(path, `^.*/`, ""))

# 6.1.3.1: per-file expectations by file name, from the benchmark's audit script (max mode, owners, groups)
log_rule(b) := ["0664", {"root"}, {"root", "utmp"}] if {
	regex.match(`^((lastlog|wtmp|btmp)([.-].*)?|README)$`, b)
} else := ["0644", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(cloud-init\.log|localmessages|waagent\.log)`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}] if {
	regex.match(`^(secure([.-].*)?|auth\.log|syslog|messages)$`, b)
} else := ["0660", {"root", "SSSD"}, {"root", "SSSD"}] if {
	b in {"SSSD", "sssd"}
} else := ["0660", {"root"}, {"root", "gdm", "gdm3"}] if {
	b in {"gdm", "gdm3"}
} else := ["0640", {"root"}, {"root", "systemd-journal"}] if {
	regex.match(`\.journal~?$`, b)
} else := ["0640", {"root", "syslog"}, {"root", "adm"}]

log_file_ok(f) if {
	r := log_rule_for(f.file)
	mode_within(f.actual_mode, r[0])
	f.actual_owner in r[1]
	f.actual_group in r[2]
}
''',
    "auditd": r'''
audit_rule_lines := [l |
	some r in input.auditd.active_rules
	l := trim_space(r)
	l != ""
	not startswith(l, "#")
]

audit_toks(l) := regex.split(`\s+`, l)

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some t in toks
	t in {concat("", ["path=", p]), concat("", ["dir=", p]), concat("", ["path=", p, "/"]), concat("", ["dir=", p, "/"])}
}

audit_path_tok(l, p) if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-w"
	toks[i + 1] in {p, concat("", [p, "/"])}
}

audit_perm(l) := substring(t, 5, -1) if {
	some t in audit_toks(l)
	startswith(t, "perm=")
}

audit_perm(l) := toks[i + 1] if {
	toks := audit_toks(l)
	some i, t in toks
	t == "-p"
}

audit_watch(p, perm) if {
	some l in audit_rule_lines
	audit_path_tok(l, p)
	pv := audit_perm(l)
	every i in numbers.range(0, count(perm) - 1) { contains(pv, substring(perm, i, 1)) }
}

line_syscalls(l) := {x |
	toks := audit_toks(l)
	some i, t in toks
	t == "-S"
	some x in split(toks[i + 1], ",")
}

audit_unset := `auid!=(unset|-1|4294967295)(\s|$)`

audit_field(l, "auid") if {
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "euid") if {
	regex.match(`(euid!=uid|uid!=euid)(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eacces") if {
	regex.match(`exit=-EACCES(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_field(l, "eperm") if {
	regex.match(`exit=-EPERM(\s|$)`, l)
	regex.match(`auid>=1000(\s|$)`, l)
	regex.match(audit_unset, l)
}

audit_syscall(s, fields) if {
	some l in audit_rule_lines
	regex.match(`^-a\s+(always,exit|exit,always)(\s|$)`, l)
	s in line_syscalls(l)
	every f in fields { audit_field(l, f) }
}

# the last `-e` line of the on-disk rules decides the enabled flag
audit_e_values := [m[1] |
	some r in input.auditd.rules
	some m in regex.find_all_string_submatch_n(`^\s*-e\s+([0-9])\b`, r, 1)
]

audit_last_e := audit_e_values[count(audit_e_values) - 1] if count(audit_e_values) > 0

sudo_logfiles contains f if {
	some l in active_lines(input.sudo.sudoers_content)
	regex.match(`^Defaults\b`, l)
	some m in regex.find_all_string_submatch_n(`logfile\s*=\s*("[^"]+"|'[^']+'|[^,\s]+)`, l, -1)
	f := trim(m[1], "\"'")
}

# 6.2.4.3: log_group unset (root default) or root/adm
audit_log_group_bad if {
	v := conf_last(input.auditd.config, "log_group")
	not lower(v) in {"root", "adm"}
}

audit_syscall_arch(s, fields, arch) if {
	some l in audit_rule_lines
	regex.match(`^-a\s+(always,exit|exit,always)(\s|$)`, l)
	concat("", ["arch=", arch]) in audit_toks(l)
	s in line_syscalls(l)
	every f in fields { audit_field(l, f) }
}

# `linux` lines of /boot/grub/grub.cfg
grub_linux_lines := [l |
	some l in input.boot_security.grub_cfg_linux_lines
	is_string(l)
	regex.match(`^\s*linux`, l)
]
''',
    "integrity": r'''
aide_opts_ok(s) if {
	opts := {trim_space(o) | some o in split(s, "+")}
	every want in ["p", "i", "n", "u", "g", "s", "b", "acl", "xattrs", "sha512"] { want in opts }
}
''',
    "user_group": r'''
dotdir_max(b) := "0700" if b == ".ssh" else := "0750"

dotdir_ok(d) if {
	mode_within(d.mode, dotdir_max(regex.replace(d.path, `^.*/`, "")))
	d.owner == d.user
	d.group == d.primary_group
}

dot_max(b) := "0600" if b in {".netrc", ".bash_history"} else := "0644"

dotfile_ok(d) if {
	b := regex.replace(d.path, `^.*/`, "")
	not b in {".forward", ".rhosts"}
	mode_within(d.mode, dot_max(b))
	d.owner == d.user
	d.group == d.primary_group
}
''',
}


# ════════════════════════════════════════════════════════════════════════════════════════════
# workbook
# ════════════════════════════════════════════════════════════════════════════════════════════
def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    recs, order = {}, []
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        rows = ws.iter_rows(values_only=True)
        head = next(rows)
        for r in rows:
            d = dict(zip(head, r))
            rid = d.get("Recommendation #")
            if not rid:
                continue
            rid = str(rid).strip()
            if rid not in recs:
                recs[rid] = {"title": str(d["Title"]).strip(), "assessment": str(d["Assessment Status"]).strip(),
                             "profiles": []}
                order.append(rid)
            recs[rid]["profiles"].append(ws.title)
    comb = wb["Combined Profiles"]
    rows = comb.iter_rows(values_only=True)
    head = next(rows)
    combined = {str(dict(zip(head, r))["Recommendation #"]).strip() for r in rows if dict(zip(head, r))["Recommendation #"]}
    if combined != set(recs):
        raise SystemExit(f"Combined Profiles disagrees with the level sheets: {sorted(combined ^ set(recs))}")
    return recs, order


def level(profiles, kind):
    lv = [int(m.group(1)) for p in profiles for m in [re.match(r"Level (\d) - (\w+)", p)] if m and m.group(2) == kind]
    return min(lv) if lv else None


def J(x):
    return json.dumps(x, ensure_ascii=False)


def idkey(rid):
    return rid.replace(".", "_")


def sort_key(rid):
    return [int(x) for x in rid.split(".")]


# ════════════════════════════════════════════════════════════════════════════════════════════
# emit
# ════════════════════════════════════════════════════════════════════════════════════════════
def header(title, checks, entry="none", density="sized", basis="size"):
    return f'''# METADATA
# title: {J(title)}
# custom:
#   class: security
#   framework: cis_ubuntu_26_04
#   source: cis
#   domains: [linux, ubuntu]
#   coverage:
#     checks: {checks}
#     cited: {checks}
#     density: {density}
#     basis: {basis}
#     tested: true
#     entrypoint: {entry}
#     guard: self
#     as_of: {AS_OF}
'''


def rule_block(name, alts):
    out = [f"default {name} := false", ""]
    for alt in alts:
        out.append(f"{name} if {{")
        out += [f"\t{c}" for c in alt]
        out += ["}", ""]
    return out


def profiles_txt(rec):
    return ", ".join(rec["profiles"])


EMPTY_TEXT_OK = ("motd_content)", "issue_content)", "issue_net_content)", "banner_content)")  # motd/issue/issue.net/sshd banner bodies: an empty banner is a real value


def evidence_alts(spec):
    out = []
    for alt in spec.ev_alts:
        out.append([("text_collected(" + c[len("is_string("):]) if c.startswith("is_string(") and not c.endswith(EMPTY_TEXT_OK)
                    else c for c in alt])
    return out


def hollow(x):
    """the same shape with every string blanked and every list emptied"""
    if isinstance(x, dict):
        return {k: hollow(v) for k, v in x.items()}
    if isinstance(x, list):
        return []
    if isinstance(x, str):
        return ""
    return x


def emit_control(rid, rec, spec):
    k = idkey(rid)
    out = [f"# {rid} ({rec['assessment']}; {profiles_txt(rec)})", f"# {rec['title']}"]
    if spec.manual:
        out += rule_block(f"p_{k}", [[f'input.attest.{spec.name} == true']])
        msg = f"{rid}: manual review not attested — set attest.{spec.name} to true once reviewed ({rec['title']})"
        out += [f"violations contains {J(msg)} if not p_{k}", ""]
        return out
    out += rule_block(f"e_{k}", evidence_alts(spec))
    out += rule_block(f"p_{k}", spec.passes)
    miss = f"{rid}: evidence missing — {spec.field} not collected"
    fail = f"{rid}: not met — {rec['title']}"
    out += [f"violations contains {J(miss)} if not e_{k}", "",
            f"violations contains {J(fail)} if {{", f"\te_{k}", f"\tnot p_{k}", "}", ""]
    return out


def emit_module(name, title, ids, recs):
    lines = [header(f"{BENCH} — {title}", len(ids)).rstrip("\n"), f"package cis_ubuntu_26_04.{name}", "",
             f"# {BENCH}, {title}.",
             "# Generated by tools/cis_generators/gen_ubuntu_26_04.py from the CIS workbook; do not edit by hand.",
             "# Each recommendation: e_<id> (facts collected), p_<id> (requirement met). Automated checks fail",
             "# closed with '<id>: evidence missing — <field> not collected'; Manual ones need input.attest.<name>.",
             "", "import rego.v1", "",
             "default compliant := false", "",
             "compliant if count(violations) == 0", "",
             f"controls := {json.dumps(sorted(ids, key=sort_key))}", ""]
    lines += HELPERS.strip("\n").split("\n") + [""]
    if name in EXTRA:
        lines += EXTRA[name].strip("\n").split("\n") + [""]
    for rid in ids:
        lines += emit_control(rid, recs[rid], L[rid])
    return "\n".join(lines).rstrip("\n") + "\n"


def emit_tests(name, ids, recs):
    lines = [f"package cis_ubuntu_26_04.{name}_test", "", "import rego.v1", "",
             "# Generated by tools/cis_generators/gen_ubuntu_26_04.py: a pass and a fail case for every recommendation,",
             "# plus an evidence-missing case for every Automated one. Assertions key on '<id>: ' (colon included).", ""]
    mod = f"data.cis_ubuntu_26_04.{name}.violations"
    for rid in ids:
        s, k = L[rid], idkey(rid)
        pre = json.dumps(rid + ": ")
        if s.manual:
            tp, tf = {"attest": {s.name: True}}, {"attest": {s.name: False}}
            miss = None
        else:
            tp, tf = s.tp, s.tf
            miss = J(f"{rid}: evidence missing — {s.field} not collected")
        lines += [f"test_{k}_pass if {{", f"\tvs := {mod} with input as {J(tp)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) == 0", "}", ""]
        lines += [f"test_{k}_fail if {{", f"\tvs := {mod} with input as {J(tf)}",
                  f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0"]
        if miss:
            lines.append(f"\tnot {miss} in vs")
        lines += ["}", ""]
        if miss:
            lines += [f"test_{k}_evidence_missing if {{", f"\tvs := {mod} with input as {J(s.tm)}",
                      f"\t{miss} in vs", "}", ""]
        if miss and not s.hollow_ok and hollow(s.tp) != s.tp:  # present-but-hollow facts (blank strings, empty lists) never pass
            lines += [f"test_{k}_hollow_evidence_not_a_pass if {{", f"\tvs := {mod} with input as {J(hollow(s.tp))}",
                      f"\tcount([v | some v in vs; startswith(v, {pre})]) > 0", "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


PROFILES = {"level1_server": ("Server", 1), "level2_server": ("Server", 2),
            "level1_workstation": ("Workstation", 1), "level2_workstation": ("Workstation", 2)}


def emit_rollup(recs, by_mod):
    controls = {}
    for rid in sorted(recs, key=sort_key):
        r = recs[rid]
        lv = {}
        for kind in ("Server", "Workstation"):
            x = level(r["profiles"], kind)
            if x is not None:
                lv[kind.lower()] = x
        controls[rid] = {"title": r["title"], "assessment": r["assessment"], "profiles": r["profiles"],
                         "levels": lv, "section": section_of(rid)}
    sec_inputs = {}
    for name, _, _ in SECTIONS:
        keys = sorted({L[rid].reads[i][0] for rid in by_mod[name] for i in range(len(L[rid].reads))} - {"attest"})
        sec_inputs[name] = keys if keys else ["attest"]
    imports = "\n".join(f"import data.cis_ubuntu_26_04.{n}" for n, _, _ in SECTIONS)
    modv = ",\n".join(f'\t"{n}": {n}.violations' for n, _, _ in SECTIONS)
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    return header(f"{BENCH} — roll-up", len(recs), density="orchestrator", basis="role") + f'''package cis_ubuntu_26_04

# {BENCH}: roll-up of the {len(SECTIONS)} section packages.
# Generated by tools/cis_generators/gen_ubuntu_26_04.py from the CIS workbook; do not edit by hand.
# {len(recs)} recommendations ({a} Automated, {len(recs) - a} Manual), every one in the benchmark. Ids, titles,
# Automated/Manual status and profiles are the workbook's. Levels differ by platform (e.g. 1.1.1.9 is
# Level 1 Server but Level 2 Workstation), so each control carries its own server/workstation level, and three
# recommendations (1.6.12, 2.1.21, 3.1.2) have no Workstation profile at all.
#
# PROFILE: input.profile selects one of level1_server (default), level2_server, level1_workstation,
# level2_workstation. A Level 2 profile includes the Level 1 recommendations of the same platform.
#
# FAIL CLOSED: every Automated recommendation emits its own evidence-missing violation when its facts were
# not collected, and every Manual one fails until attested, so a control is counted as passed only when
# real facts satisfy it. sections_not_assessed is informational (no input key a section reads is present).

import rego.v1

{imports}

standard_total := {len(recs)}

controls := {json.dumps(controls, indent=1, ensure_ascii=False)}

profiles := {json.dumps({k: {"platform": v[0].lower(), "level": v[1]} for k, v in PROFILES.items()})}

default profile := "level1_server"

profile := input.profile if input.profile in object.keys(profiles)

applies(c) if c.levels[profiles[profile].platform] <= profiles[profile].level

applicable_ids := {{id | some id, c in controls; applies(c)}}

module_violations := {{
{modv},
}}

all_violations := sort([v |
	some _, vs in module_violations
	some v in vs
	split(v, ":")[0] in applicable_ids
])

failing_ids := {{split(v, ":")[0] | some v in all_violations}}

# =============================================================================
# ASSESSMENT COVERAGE — which sections received any of the input keys they read
# =============================================================================

section_inputs := {json.dumps(sec_inputs, indent=1)}

section_assessed(name) if {{
	some k in section_inputs[name]
	input[k]
}}

missing_inputs := {{name: [k | some k in section_inputs[name]; not input[k]] | some name, _ in section_inputs}}

not_assessed_sections := sort([name | some name, _ in section_inputs; not section_assessed(name)])

assessment_coverage_percentage := round(((count(section_inputs) - count(not_assessed_sections)) * 100) / count(section_inputs))

# =============================================================================
# SECTIONS (controls applicable to the selected profile)
# =============================================================================

section_ids(name) := {{id | some id in applicable_ids; controls[id].section == name}}

section_compliance := {{name: {{
	"compliant": count(section_ids(name) & failing_ids) == 0,
	"controls_covered": count(section_ids(name)),
	"violation_count": count([v | some v in all_violations; split(v, ":")[0] in section_ids(name)]),
}} |
	some name, _ in section_inputs
	count(section_ids(name)) > 0
}}

compliant_sections := count([name | some name, s in section_compliance; s.compliant == true; section_assessed(name)])

default compliance_percentage := 0

compliance_percentage := (compliant_sections * 100) / count(section_compliance) if count(section_compliance) > 0

default compliant := false

compliant if {{
	count(object.keys(input)) > 0
	count(applicable_ids) > 0
	count(failing_ids) == 0
}}
'''


def emit_main(recs):
    return header("Bridge: expose cis_ubuntu_26_04 under the /v1/data/<framework>/main/compliance_report", 0,
                  entry="compliance_report", density="orchestrator", basis="role") + f'''package cis_ubuntu_26_04.main

# Bridge: expose cis_ubuntu_26_04 under the /v1/data/<framework>/main/compliance_report convention, with the
# same keys as cis_rhel9.main (plus profile and standard_total). Every field is sourced through a DEFAULTED
# helper: an undefined field inside an object literal would collapse the whole report to {{}}.
#
# total_controls is the number of recommendations applicable to the selected profile; standard_total is
# the benchmark's {len(recs)}. A control passes only when real facts (or an attestation) satisfy it, so
# failed_controls counts distinct failing recommendation ids, including those whose evidence is missing.

import data.cis_ubuntu_26_04 as bench
import rego.v1

default _compliant := false

_compliant := bench.compliant

default _facts_supplied := false

_facts_supplied if count(object.keys(input)) > 0

_no_facts_msg := "FAIL-CLOSED: no facts supplied for cis_ubuntu_26_04 — the assessment could not be evaluated. This is NOT a passing result; check that fact collection ran and produced input."

default _bench_violations := []

_bench_violations := bench.all_violations

default _not_assessed := []

_not_assessed := bench.not_assessed_sections

default _missing_inputs := {{}}

_missing_inputs := bench.missing_inputs

_not_assessed_msgs := [sprintf("NOT ASSESSED: %s — required input not collected: %s", [name, concat(", ", _missing_inputs[name])]) | some name in _not_assessed]

default _sections := {{}}

_sections := bench.section_compliance

default _section_percentage := 0

_section_percentage := bench.compliance_percentage

default _coverage := 0

_coverage := bench.assessment_coverage_percentage

default _profile := "level1_server"

_profile := bench.profile

default _total_controls := 0

_total_controls := count(bench.applicable_ids)

default _failing := 0

_failing := count(bench.failing_ids)

_violations := array.concat(_bench_violations, [_no_facts_msg]) if not _facts_supplied

_violations := array.concat(_bench_violations, _not_assessed_msgs) if _facts_supplied

# No facts => every control is unverified, which counts as failed, not passed.
_failed := _total_controls if not _facts_supplied

_failed := min([_total_controls, _failing]) if _facts_supplied

_passed := max([0, _total_controls - _failed])

default _percentage := 0

_percentage := round((_passed * 100000) / _total_controls) / 1000 if _total_controls > 0

compliance_report := {{
	"framework": "cis_ubuntu_26_04",
	"benchmark": "{BENCH}",
	"version": "v1.0.0",
	"profile": _profile,
	"total_controls": _total_controls,
	"standard_total": {len(recs)},
	"controls_basis": "CIS recommendations applicable to the selected profile (every one evaluated)",
	"passed_controls": _passed,
	"failed_controls": _failed,
	"compliance_percentage": _percentage,
	"section_compliance_percentage": _section_percentage,
	"compliant": _compliant,
	"violations": _violations,
	"section_results": _sections,
	"sections_not_assessed": _not_assessed,
	"missing_inputs": _missing_inputs,
	"assessment_coverage_percentage": _coverage,
}}
'''


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    recs, _ = read_workbook(sys.argv[1])
    a = sum(1 for r in recs.values() if r["assessment"] == "Automated")
    m = sum(1 for r in recs.values() if r["assessment"] == "Manual")
    if (len(recs), a, m) != (EXPECTED_TOTAL, EXPECTED_AUTOMATED, EXPECTED_MANUAL):
        raise SystemExit(f"workbook has {len(recs)} recommendations ({a} Automated, {m} Manual)")
    if set(L) != set(recs):
        raise SystemExit(f"logic/workbook mismatch: no logic {sorted(set(recs) - set(L), key=sort_key)}; "
                         f"no recommendation {sorted(set(L) - set(recs), key=sort_key)}")
    for rid, r in recs.items():
        if (r["assessment"] == "Manual") != L[rid].manual:
            raise SystemExit(f"{rid}: workbook says {r['assessment']}, logic table disagrees")
        title = " ".join(r["title"].split()).lower() + " "
        if " ".join(KW[rid].split()).lower() not in title and KW[rid].lower() not in title:
            raise SystemExit(f"{rid}: logic keyword {KW[rid]!r} not in workbook title {r['title']!r}")
    by_mod = {n: [] for n, _, _ in SECTIONS}
    for rid in sorted(recs, key=sort_key):
        by_mod[section_of(rid)].append(rid)
    (OUT / "tests").mkdir(parents=True, exist_ok=True)
    for name, title, _ in SECTIONS:
        (OUT / f"{name}.rego").write_text(emit_module(name, title, by_mod[name], recs))
        (OUT / "tests" / f"test_{name}.rego").write_text(emit_tests(name, by_mod[name], recs))
    (OUT / "cis_ubuntu_26_04.rego").write_text(emit_rollup(recs, by_mod))
    (OUT / "cis_ubuntu_26_04_main.rego").write_text(emit_main(recs))
    (OUT / "tests" / "test_cis_ubuntu_26_04_main.rego").write_text(emit_rollup_tests(recs))
    if len(sys.argv) > 2:  # input paths each control reads (for collector/mapper gap analysis)
        pathlib.Path(sys.argv[2]).write_text(json.dumps(
            {rid: [".".join(p) for p in L[rid].reads] for rid in sorted(recs, key=sort_key)}, indent=1))
    print(f"{len(recs)} recommendations ({a} Automated, {m} Manual) -> {OUT}")



def emit_rollup_tests(recs):
    counts = {}
    for p, (kind, lv) in PROFILES.items():
        counts[p] = sum(1 for r in recs.values() if (level(r["profiles"], kind) or 99) <= lv)
    keys = ["framework", "benchmark", "version", "profile", "total_controls", "standard_total", "controls_basis",
            "passed_controls", "failed_controls", "compliance_percentage", "section_compliance_percentage", "compliant",
            "violations", "section_results", "sections_not_assessed", "missing_inputs", "assessment_coverage_percentage"]
    lines = ["package cis_ubuntu_26_04.main_test", "", "import rego.v1", "",
             "# Generated by tools/cis_generators/gen_ubuntu_26_04.py: roll-up and bridge smoke tests.", "",
             "test_every_recommendation_in_table if {",
             f"\tcount(data.cis_ubuntu_26_04.controls) == {len(recs)}",
             f"\tcount([c | some c in data.cis_ubuntu_26_04.controls; c.assessment == \"Manual\"]) == {EXPECTED_MANUAL}",
             "}", "",
             "test_empty_input_fails_closed if {", "\tr := data.cis_ubuntu_26_04.main.compliance_report with input as {}",
             f"\tcount(object.keys(r)) == {len(keys)}",
             f"\tevery k in {json.dumps(keys)} {{ k in object.keys(r) }}",
             "\tr.compliant == false", f"\tr.total_controls == {counts['level1_server']}",
             "\tr.failed_controls == r.total_controls", "\tr.passed_controls == 0",
             "\tsome v in r.violations", "\tstartswith(v, \"FAIL-CLOSED\")", "}", ""]
    for p, n in counts.items():
        lines += [f"test_profile_{p}_count if {{", f"\tr := data.cis_ubuntu_26_04.main.compliance_report with input as {{\"profile\": \"{p}\"}}",
                  f"\tr.profile == \"{p}\"", f"\tr.total_controls == {n}", f"\tr.standard_total == {len(recs)}",
                  "\tr.failed_controls == r.total_controls", "}", ""]
    lines += ["test_unknown_profile_defaults_to_level1_server if {", "\tr := data.cis_ubuntu_26_04.main.compliance_report with input as {\"profile\": \"bogus\"}",
              "\tr.profile == \"level1_server\"", f"\tr.total_controls == {counts['level1_server']}", "}", "",
              "test_one_attested_control_passes if {",
              "\tr := data.cis_ubuntu_26_04.main.compliance_report with input as {\"attest\": {\"apt_sources_signed_by_reviewed\": true}}",
              "\tr.passed_controls == 1", "\tr.failed_controls == r.total_controls - 1", "\tr.compliant == false", "}", "",
              "test_unassessed_sections_listed if {", "\tr := data.cis_ubuntu_26_04.main.compliance_report with input as {\"ssh\": {\"sshd_config_raw\": \"permitrootlogin no\"}}",
              "\tnot \"ssh\" in r.sections_not_assessed", "\t\"auditd\" in r.sections_not_assessed",
              "\tr.assessment_coverage_percentage < 100", "}", "",
              "test_level2_violation_hidden_in_level1_profile if {",
              "\tvs := data.cis_ubuntu_26_04.all_violations with input as {\"profile\": \"level1_server\"}",
              "\tcount([v | some v in vs; startswith(v, \"6.2.1.1: \")]) == 0",
              "\tvs2 := data.cis_ubuntu_26_04.all_violations with input as {\"profile\": \"level2_server\"}",
              "\tcount([v | some v in vs2; startswith(v, \"6.2.1.1: \")]) > 0", "}", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


if __name__ == "__main__":
    main()
